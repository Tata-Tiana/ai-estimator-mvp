from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
EXTRACTED_DIR = DATA_DIR / "extracted"
MAPPED_DIR = DATA_DIR / "mapped"
REVIEW_DIR = DATA_DIR / "review"
REPORTS_DIR = DATA_DIR / "reports"
MAPPED_PARAMETERS_PATH = MAPPED_DIR / "mapped_parameters.json"
NORMALIZED_PARAMETERS_PATH = MAPPED_DIR / "normalized_parameters.json"
CANDIDATES_PATH = EXTRACTED_DIR / "candidates.json"
INTEGRITY_REPORT_PATH = REPORTS_DIR / "integrity_report.md"
REVIEW_PACK_PATH = REVIEW_DIR / "review_pack.xlsx"


FILLS = {
    "found_from_pdf": PatternFill("solid", fgColor="D9EAD3"),
    "low_confidence": PatternFill("solid", fgColor="FFF2CC"),
    "true_missing_in_project": PatternFill("solid", fgColor="F4CCCC"),
    "supplier_required": PatternFill("solid", fgColor="CFE2F3"),
    "manual_required": PatternFill("solid", fgColor="EADCF8"),
    "mapping_gap": PatternFill("solid", fgColor="F9CB9C"),
    "header": PatternFill("solid", fgColor="D9EAD3"),
}


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def action_needed(row: dict[str, Any]) -> str:
    status = row.get("found_status")
    if status == "found_from_pdf":
        return "Проверить и подтвердить."
    if status == "low_confidence":
        return "Проверить единицу/контекст."
    if status == "true_missing_in_project":
        return "Запросить у проектировщика или заполнить после review."
    if status == "supplier_required":
        return "Нужен поставщик/раскладка."
    if status == "manual_required":
        return "Нужно ручное решение сметчика."
    if status == "mapping_gap":
        return "Доработать mapping/parser; похожие данные есть."
    return ""


def setup_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = FILLS["header"]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def append_table(ws, headers: list[str], rows: list[dict[str, Any]], keys: list[str]) -> None:
    ws.append(headers)
    setup_header(ws)
    for row in rows:
        ws.append([row.get(key, "") for key in keys])
        fill = FILLS.get(row.get("found_status"))
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 24


def build_main_review_sheet(wb: Workbook, mapped: list[dict[str, Any]]) -> None:
    headers = [
        "section_name",
        "label",
        "calculator_input_key",
        "found_status",
        "value_from_pdf",
        "unit",
        "source_pdf",
        "page",
        "source_fragment",
        "confidence",
        "needs_elena_review",
        "elena_status",
        "elena_value",
        "elena_comment",
        "final_value",
        "action_needed",
        "candidate_id",
    ]
    ws = wb.create_sheet("01_Проверка_Елены")
    ws.append(headers)
    setup_header(ws)
    include = {"found_from_pdf", "low_confidence", "true_missing_in_project", "supplier_required", "manual_required"}
    for row in [item for item in mapped if item.get("found_status") in include]:
        excel_row = ws.max_row + 1
        ws.append(
            [
                row.get("section_name"),
                row.get("target_label"),
                row.get("calculator_input_key"),
                row.get("found_status"),
                row.get("value"),
                row.get("unit"),
                row.get("source_pdf"),
                row.get("page"),
                row.get("source_fragment"),
                row.get("confidence"),
                "да" if row.get("needs_elena_review") else "нет",
                "",
                "",
                "",
                f'=IF(M{excel_row}<>"",M{excel_row},E{excel_row})',
                action_needed(row),
                row.get("candidate_id"),
            ]
        )
        fill = FILLS.get(row.get("found_status"))
        if fill:
            for cell in ws[excel_row]:
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    validation = DataValidation(
        type="list",
        formula1='"ok,исправить,заполнить,не используется,спросить проектировщика,нужен поставщик"',
        allow_blank=True,
    )
    ws.add_data_validation(validation)
    if ws.max_row > 1:
        validation.add(f"L2:L{ws.max_row}")
    widths = [24, 34, 34, 22, 16, 10, 20, 8, 70, 14, 16, 18, 18, 28, 18, 34, 24]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = width


def build_review_pack() -> Path:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    mapped = load_json(MAPPED_PARAMETERS_PATH)
    candidates = load_json(CANDIDATES_PATH)
    normalized = load_json(NORMALIZED_PARAMETERS_PATH)
    wb = Workbook()
    wb.remove(wb.active)

    readme = wb.create_sheet("00_README")
    readme.append(["Strict parser v2"])
    readme.append(["Значения на листе 01 пришли только из PDF candidates/spec rows."])
    readme.append(["Если значение не найдено, оно не подставляется руками."])
    readme.append(["mapping_gap вынесен отдельно: это проблема parser/mapping, а не обязательно проблема проекта."])
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 100

    build_main_review_sheet(wb, mapped)
    append_table(
        wb.create_sheet("02_Все_candidates"),
        ["candidate_id", "source_pdf", "page", "page_title", "raw_label", "raw_value", "unit", "candidate_type", "confidence", "raw_context"],
        candidates,
        ["candidate_id", "source_pdf", "page", "page_title", "raw_label", "raw_value", "unit", "candidate_type", "confidence", "raw_context"],
    )
    normalized_rows = []
    for key, value in normalized.items():
        normalized_rows.append({"structure": key, "json": json.dumps(value, ensure_ascii=False)})
    append_table(wb.create_sheet("03_Normalized_Data"), ["structure", "json"], normalized_rows, ["structure", "json"])
    append_table(
        wb.create_sheet("04_Реально_не_найдено"),
        ["section_name", "calculator_input_key", "target_label", "mapping_reason"],
        [row for row in mapped if row.get("found_status") == "true_missing_in_project"],
        ["section_name", "calculator_input_key", "target_label", "mapping_reason"],
    )
    append_table(
        wb.create_sheet("05_Low_Confidence"),
        ["section_name", "calculator_input_key", "target_label", "value", "unit", "source_pdf", "page", "source_fragment", "mapping_reason"],
        [row for row in mapped if row.get("found_status") == "low_confidence"],
        ["section_name", "calculator_input_key", "target_label", "value", "unit", "source_pdf", "page", "source_fragment", "mapping_reason"],
    )
    append_table(
        wb.create_sheet("06_Mapping_Gaps"),
        ["section_name", "calculator_input_key", "target_label", "value", "unit", "source_pdf", "page", "source_fragment", "mapping_reason"],
        [row for row in mapped if row.get("found_status") == "mapping_gap"],
        ["section_name", "calculator_input_key", "target_label", "value", "unit", "source_pdf", "page", "source_fragment", "mapping_reason"],
    )
    append_table(
        wb.create_sheet("07_SUPPLIER_INPUT"),
        ["section_name", "calculator_input_key", "target_label", "value", "source_fragment", "action_needed"],
        [dict(row, action_needed=action_needed(row)) for row in mapped if row.get("found_status") == "supplier_required"],
        ["section_name", "calculator_input_key", "target_label", "value", "source_fragment", "action_needed"],
    )
    append_table(
        wb.create_sheet("08_Готово_к_draft"),
        ["section_name", "calculator_input_key", "target_label", "value", "unit", "source_pdf", "page", "candidate_id"],
        [row for row in mapped if row.get("found_status") == "found_from_pdf" and row.get("confidence") in {"high", "medium"}],
        ["section_name", "calculator_input_key", "target_label", "value", "unit", "source_pdf", "page", "candidate_id"],
    )
    integrity_text = INTEGRITY_REPORT_PATH.read_text(encoding="utf-8") if INTEGRITY_REPORT_PATH.exists() else ""
    integrity_rows = [{"line": line} for line in integrity_text.splitlines()]
    append_table(wb.create_sheet("09_Integrity_Check"), ["line"], integrity_rows, ["line"])

    wb.save(REVIEW_PACK_PATH)
    load_workbook(REVIEW_PACK_PATH, read_only=True).close()
    return REVIEW_PACK_PATH


def main() -> int:
    path = build_review_pack()
    print(f"review_pack: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
