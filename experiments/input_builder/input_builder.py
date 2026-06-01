"""Build calculator input.json files from reviewed_parameters.xlsx.

This is an experimental MVP bridge:
reviewed PDF parameters -> calculator input.json.
It does not parse PDF and does not calculate estimates.
"""

from __future__ import annotations

import copy
import json
import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from section_input_registry import SECTION_INPUT_REGISTRY, enabled_sections


REPO_ROOT = Path(__file__).resolve().parents[2]
READY_STATUSES = {"reviewed", "corrected", "manual", "manual_entered", "default_used"}
WARNING_STATUSES = {"needs_review"}
BLOCKING_STATUSES = {"missing", "manual_required", "ignored"}


@dataclass
class ParameterRow:
    raw: dict[str, Any]
    effective_value: Any = None
    value_missing: bool = False
    warnings: list[str] = field(default_factory=list)

    @property
    def section_code(self) -> str:
        return str(self.raw.get("section_code") or "")

    @property
    def section_name(self) -> str:
        return str(self.raw.get("section_name") or self.section_code)

    @property
    def parameter_code(self) -> str:
        return str(self.raw.get("parameter_code") or "")

    @property
    def calculator_input_key(self) -> str:
        return str(self.raw.get("calculator_input_key") or "")

    @property
    def label(self) -> str:
        return str(self.raw.get("label") or self.parameter_code)

    @property
    def elena_status(self) -> str:
        return str(self.raw.get("elena_status") or "")

    @property
    def required(self) -> bool:
        return parse_bool(self.raw.get("required"))

    @property
    def use_for_calculation(self) -> bool:
        return parse_bool(self.raw.get("use_for_calculation"))

    @property
    def input_type(self) -> str:
        return str(self.raw.get("input_type") or "")


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in {"true", "1", "yes", "да"}


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def normalize_scalar(value: Any) -> Any:
    if is_blank(value):
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()
    lowered = text.lower()
    if lowered in {"true", "false"}:
        return lowered == "true"

    if text.startswith("[") and text.endswith("]"):
        try:
            parsed = json.loads(text)
            return normalize_value(parsed)
        except json.JSONDecodeError:
            pass

    if ";" in text:
        return [normalize_scalar(part) for part in text.split(";")]

    compact = text.replace("\u00a0", " ").replace(" ", "")
    if re.fullmatch(r"[-+]?\d+(,\d+)?", compact):
        compact = compact.replace(",", ".")
    if re.fullmatch(r"[-+]?\d+(\.\d+)?", compact):
        try:
            number = Decimal(compact)
        except InvalidOperation:
            return text
        if number == number.to_integral_value():
            return int(number)
        return float(number)

    return text


def normalize_value(value: Any) -> Any:
    if isinstance(value, list):
        return [normalize_value(item) for item in value]
    if isinstance(value, dict):
        return {key: normalize_value(child) for key, child in value.items()}
    return normalize_scalar(value)


def effective_value_for(row: dict[str, Any]) -> tuple[Any, bool]:
    for column in ("corrected_value", "final_value", "extracted_value"):
        value = normalize_value(row.get(column))
        if value is not None:
            return value, False
    return None, True


def read_reviewed_parameters(path: Path) -> list[ParameterRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "parameters" not in workbook.sheetnames:
        raise ValueError(f"Sheet 'parameters' not found in {path}")
    sheet = workbook["parameters"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[ParameterRow] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        raw = {header: values[index] for index, header in enumerate(headers)}
        effective, missing = effective_value_for(raw)
        parameter = ParameterRow(raw=raw, effective_value=effective, value_missing=missing)
        status = parameter.elena_status
        if status in WARNING_STATUSES and parameter.use_for_calculation:
            parameter.warnings.append("needs_review value used; Elena has not confirmed this parameter")
        if status not in READY_STATUSES | WARNING_STATUSES | BLOCKING_STATUSES | {"control_only", ""}:
            parameter.warnings.append(f"unknown elena_status: {status}")
        rows.append(parameter)
    return rows


PATH_TOKEN_RE = re.compile(r"([^\.\[\]]+)(?:\[(\d+)\])?")


def parse_path(path: str) -> list[tuple[str, int | None]]:
    parts: list[tuple[str, int | None]] = []
    for raw_part in path.split("."):
        pos = 0
        while pos < len(raw_part):
            match = PATH_TOKEN_RE.match(raw_part, pos)
            if not match:
                raise ValueError(f"Unsupported calculator_input_key segment: {raw_part}")
            key = match.group(1)
            index = int(match.group(2)) if match.group(2) is not None else None
            parts.append((key, index))
            pos = match.end()
    return parts


def ensure_list_size(value: list[Any], index: int) -> None:
    while len(value) <= index:
        value.append({})


def set_nested_value(target: dict[str, Any], path: str, value: Any) -> None:
    current: Any = target
    parts = parse_path(path)
    for part_index, (key, list_index) in enumerate(parts):
        is_last = part_index == len(parts) - 1
        if list_index is None:
            if is_last:
                current[key] = value
                return
            if key not in current or not isinstance(current[key], (dict, list)):
                next_key, next_list_index = parts[part_index + 1]
                current[key] = [] if next_list_index is not None else {}
            current = current[key]
            continue

        if key not in current or not isinstance(current[key], list):
            current[key] = []
        ensure_list_size(current[key], list_index)
        if is_last:
            current[key][list_index] = value
            return
        if not isinstance(current[key][list_index], dict):
            current[key][list_index] = {}
        current = current[key][list_index]


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as file:
        return json.load(file)


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
        file.write("\n")


def template_value_for(template: dict[str, Any], path: str) -> Any:
    current: Any = template
    for key, index in parse_path(path):
        if not isinstance(current, dict) or key not in current:
            return None
        current = current[key]
        if index is not None:
            if not isinstance(current, list) or index >= len(current):
                return None
            current = current[index]
    return copy.deepcopy(current)


def clean_generated_inputs(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    for child in path.iterdir():
        if child.is_file() and child.suffix == ".json":
            child.unlink()


def build_inputs(case_dir: Path, mode_override: str | None = None) -> dict[str, Any]:
    config_path = case_dir / "project_config.json"
    config = load_json(config_path)
    builder_mode = mode_override or config.get("builder_mode") or "strict"
    if builder_mode not in {"strict", "demo_with_template_fallback"}:
        raise ValueError(f"Unsupported builder_mode: {builder_mode}")

    reviewed_path = resolve_repo_path(config["reviewed_parameters_path"])
    rows = read_reviewed_parameters(reviewed_path)
    rows_by_section: dict[str, list[ParameterRow]] = {}
    for row in rows:
        rows_by_section.setdefault(row.section_code, []).append(row)

    generated_dir = case_dir / "generated_inputs"
    output_dir = REPO_ROOT / "experiments/input_builder/output" / case_dir.name
    output_generated_dir = output_dir / "generated_inputs"
    clean_generated_inputs(generated_dir)
    clean_generated_inputs(output_generated_dir)

    sections_result: list[dict[str, Any]] = []
    all_warnings: list[str] = []
    missing_rows: list[ParameterRow] = []
    needs_review_rows: list[ParameterRow] = []

    for section_code in SECTION_INPUT_REGISTRY:
        enabled = section_code in enabled_sections(config.get("sections"))
        registry = SECTION_INPUT_REGISTRY[section_code]
        section_rows = rows_by_section.get(section_code, [])
        if not enabled:
            sections_result.append(section_result(section_code, registry["section_name"], False, "skipped", section_rows))
            continue

        template_path = resolve_repo_path(registry["calculator_case_template_input_path"])
        template_input = load_json(template_path)
        generated_input = copy.deepcopy(template_input)
        warnings: list[str] = []
        section_missing: list[ParameterRow] = []
        section_needs_review: list[ParameterRow] = []
        used_parameters = 0

        for parameter in section_rows:
            if not parameter.use_for_calculation:
                continue

            if parameter.elena_status in WARNING_STATUSES:
                section_needs_review.append(parameter)
                needs_review_rows.append(parameter)
                warnings.extend(format_parameter_warnings(parameter))

            value = parameter.effective_value
            missing = parameter.value_missing

            if missing and builder_mode == "demo_with_template_fallback":
                value = template_value_for(template_input, parameter.calculator_input_key)
                if value is not None:
                    missing = False
                    warning = (
                        f"template fallback used for {parameter.parameter_code} "
                        f"({parameter.calculator_input_key})"
                    )
                    parameter.warnings.append(warning)
                    warnings.append(warning)

            if parameter.required and parameter.use_for_calculation and missing:
                section_missing.append(parameter)
                missing_rows.append(parameter)
                continue

            if not parameter.calculator_input_key:
                warnings.append(f"empty calculator_input_key for {parameter.parameter_code}")
                continue

            if value is not None:
                set_nested_value(generated_input, parameter.calculator_input_key, value)
                used_parameters += 1

        status = "generated"
        generated_path = ""
        if section_missing and builder_mode == "strict":
            status = "blocked_missing_required_params"
        else:
            generated_input["pricing"] = config.get("pricing", {})
            generated_input["generated_from"] = {
                "source": "reviewed_parameters.xlsx",
                "source_path": str(reviewed_path.relative_to(REPO_ROOT)),
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "builder_mode": builder_mode,
                "section_code": section_code,
            }
            generated_input["warnings"] = warnings
            target_path = generated_dir / registry["generated_input_filename"]
            write_json(target_path, generated_input)
            shutil.copy2(target_path, output_generated_dir / registry["generated_input_filename"])
            generated_path = str(target_path.relative_to(REPO_ROOT))

        sections_result.append({
            "section_code": section_code,
            "section_name": registry["section_name"],
            "enabled": enabled,
            "status": status,
            "parameters_total": len(section_rows),
            "used_parameters": used_parameters,
            "missing_required": len(section_missing),
            "manual_required": sum(1 for item in section_missing if item.elena_status == "manual_required"),
            "needs_review": len(section_needs_review),
            "generated_input_path": generated_path,
            "warnings": warnings,
        })
        all_warnings.extend(warnings)

    result = {
        "project_name": config.get("project_name"),
        "reviewed_parameters_path": config.get("reviewed_parameters_path"),
        "builder_mode": builder_mode,
        "sections": sections_result,
        "totals": {
            "sections_enabled": sum(1 for item in sections_result if item["enabled"]),
            "sections_generated": sum(1 for item in sections_result if item["status"] == "generated"),
            "sections_blocked": sum(1 for item in sections_result if item["status"] == "blocked_missing_required_params"),
            "missing_required_total": len(missing_rows),
            "manual_required_total": sum(1 for item in missing_rows if item.elena_status == "manual_required"),
        },
        "warnings": all_warnings,
    }

    write_json(case_dir / "result.json", result)
    write_json(output_dir / "result.json", result)
    report_config = dict(config)
    report_config["builder_mode"] = builder_mode
    write_missing_report(case_dir / "missing_parameters_report.md", report_config, sections_result, missing_rows, needs_review_rows)
    write_missing_report(output_dir / "missing_parameters_report.md", report_config, sections_result, missing_rows, needs_review_rows)
    write_result_md(case_dir / "result.md", result)
    write_result_md(output_dir / "result.md", result)
    return result


def resolve_repo_path(path: str | Path) -> Path:
    path = Path(path)
    return path if path.is_absolute() else REPO_ROOT / path


def format_parameter_warnings(parameter: ParameterRow) -> list[str]:
    if parameter.warnings:
        return [
            f"{parameter.section_code}:{parameter.parameter_code}: {warning}"
            for warning in parameter.warnings
        ]
    return []


def section_result(section_code: str, section_name: str, enabled: bool, status: str, rows: list[ParameterRow]) -> dict[str, Any]:
    return {
        "section_code": section_code,
        "section_name": section_name,
        "enabled": enabled,
        "status": status,
        "parameters_total": len(rows),
        "used_parameters": 0,
        "missing_required": 0,
        "manual_required": 0,
        "needs_review": 0,
        "generated_input_path": "",
        "warnings": [],
    }


def write_missing_report(
    path: Path,
    config: dict[str, Any],
    sections_result: list[dict[str, Any]],
    missing_rows: list[ParameterRow],
    needs_review_rows: list[ParameterRow],
) -> None:
    lines = [
        "# Missing parameters report",
        "",
        "## Project",
        "",
        f"- project_name: {config.get('project_name')}",
        f"- reviewed_parameters_path: {config.get('reviewed_parameters_path')}",
        f"- builder_mode: {config.get('builder_mode', 'strict')}",
        "",
        "## Blocked sections",
        "",
        "| section_code | section_name | missing required count | manual_required count | can_generate_input |",
        "|---|---|---:|---:|---|",
    ]
    for section in sections_result:
        can_generate = "yes" if section["status"] == "generated" else "no"
        lines.append(
            f"| {section['section_code']} | {section['section_name']} | "
            f"{section['missing_required']} | {section['manual_required']} | {can_generate} |"
        )

    lines.extend([
        "",
        "## Missing parameters",
        "",
        "| section_code | parameter_code | calculator_input_key | label | input_type | unit | elena_status | reason |",
        "|---|---|---|---|---|---|---|---|",
    ])
    for row in missing_rows:
        reason = row.raw.get("missing_reason") or "required parameter has no effective value"
        lines.append(
            f"| {row.section_code} | {row.parameter_code} | `{row.calculator_input_key}` | "
            f"{row.label} | {row.input_type} | {row.raw.get('unit') or ''} | {row.elena_status} | {reason} |"
        )

    lines.extend([
        "",
        "## Needs review warnings",
        "",
        "| section_code | parameter_code | label | effective_value | source_file | page | warning |",
        "|---|---|---|---|---|---|---|",
    ])
    for row in needs_review_rows:
        warning = "; ".join(row.warnings) or "needs_review value used"
        lines.append(
            f"| {row.section_code} | {row.parameter_code} | {row.label} | {row.effective_value} | "
            f"{row.raw.get('source_file') or ''} | {row.raw.get('page') or ''} | {warning} |"
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_result_md(path: Path, result: dict[str, Any]) -> None:
    lines = [
        "# Input builder result",
        "",
        "## Что сделано",
        "",
        "- Прочитан `reviewed_parameters.xlsx`.",
        "- Собраны `input.json` для разделов, где хватает обязательных данных.",
        "- Сформирован `missing_parameters_report.md`.",
        "",
        "## Сводка по разделам",
        "",
        "| Раздел | Статус | Параметров использовано | Missing | Manual required | Needs review | Generated input |",
        "|---|---|---:|---:|---:|---:|---|",
    ]
    for section in result["sections"]:
        lines.append(
            f"| {section['section_name']} | {section['status']} | {section['used_parameters']} | "
            f"{section['missing_required']} | {section['manual_required']} | {section['needs_review']} | "
            f"{section['generated_input_path']} |"
        )

    lines.extend([
        "",
        "## Что делать Елене",
        "",
        "Если раздел `blocked_missing_required_params`:",
        "",
        "- открыть `reviewed_parameters.xlsx`;",
        "- заполнить `corrected_value`;",
        "- поставить `elena_status = manual / reviewed / corrected`;",
        "- повторно запустить `input_builder`.",
        "",
        "## Warnings",
        "",
    ])
    if result["warnings"]:
        lines.extend([f"- {warning}" for warning in result["warnings"]])
    else:
        lines.append("Warnings нет.")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
