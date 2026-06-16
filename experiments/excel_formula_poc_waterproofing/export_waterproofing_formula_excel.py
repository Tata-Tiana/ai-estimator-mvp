from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
FIXTURES_DIR = BASE_DIR / "fixtures"
OUTPUT_PATH = BASE_DIR / "output" / "waterproofing_formula_demo.xlsx"

SECTION_TITLE = "ГИДРОИЗОЛЯЦИЯ, УТЕПЛЕНИЕ БОРТОВ ПЛИТ"
ADDRESS = (
    'Российская федерация, Московская область, г. о. Домодедово, '
    'п. Государственного племенного завода "Константиново", '
    'тер. КП "Юсупово Виладж", участок с кад. номером 50:28:0050421:2962'
)

FORBIDDEN_FORMULA_TOKENS = [
    "slab_formwork_perimeter_m",
    "slab_edge_height_m",
    "eps100_wall_volume_m3",
    "waterproofing_work_unit_price",
    "primer_consumption_l_per_m2",
    "mastic_consumption_kg_per_m2_per_layer",
    "eps100_wall_insulation_work_unit_price",
]

RIGHT_FIELD_FORBIDDEN_WORDS = [
    "price_code",
    "Прайс",
    "registry",
]


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as file:
        return json.load(file)


def result_line_by_code(result: dict[str, Any], code: str) -> dict[str, Any]:
    for line in result["estimate_lines"]:
        if line["code"] == code:
            return line
    raise KeyError(code)


def line_price(line: dict[str, Any], price_kind: str) -> float:
    value = line[f"{price_kind}_unit_price"]
    return float(value or 0)


def load_poc_values() -> dict[str, Any]:
    input_data = load_json(FIXTURES_DIR / "input.json")
    result = load_json(FIXTURES_DIR / "waterproofing_result.json")
    non_insulated_lengths = input_data.get("non_insulated_edge_lengths_m", [])

    return {
        "perimeter_m": input_data["slab_formwork_perimeter_m"],
        "edge_height_m": input_data["slab_edge_height_m"],
        "primer_consumption_l_per_m2": input_data["primer_consumption_l_per_m2"],
        "primer_canister_volume_l": input_data["primer_canister_volume_l"],
        "mastic_consumption_kg_per_m2_per_layer": input_data[
            "mastic_consumption_kg_per_m2_per_layer"
        ],
        "mastic_layers": input_data["mastic_layers"],
        "mastic_bucket_weight_kg": input_data["mastic_bucket_weight_kg"],
        "eps100_wall_volume_m3": input_data["eps100_wall_volume_m3"],
        "eps100_wall_thickness_m": input_data["eps100_wall_thickness_m"],
        "non_insulated_edge_total_m": sum(non_insulated_lengths),
        "eps_waste_coeff": input_data["eps_waste_coeff"],
        "eps100_pack_volume_m3": input_data["eps100_pack_volume_m3"],
        "glue_foam_coverage_m2_per_can": input_data["glue_foam_coverage_m2_per_can"],
        "glue_foam_min_units": input_data["glue_foam_min_units"],
        "waterproofing_logistics_coeff": input_data["waterproofing_logistics_coeff"],
        "waterproofing_consumables_coeff": input_data["waterproofing_consumables_coeff"],
        "waterproofing_work_price": line_price(
            result_line_by_code(result, "waterproofing_bitumen_mastic_work"), "work"
        ),
        "primer_price": line_price(result_line_by_code(result, "bitumen_primer_aquamast_18l"), "material"),
        "mastic_price": line_price(result_line_by_code(result, "bitumen_mastic_aquamast_18kg"), "material"),
        "eps_work_price": line_price(result_line_by_code(result, "eps100_wall_insulation_work"), "work"),
        "eps100_price": line_price(
            result_line_by_code(result, "eps100_wall_penoplex_geo_material"), "material"
        ),
        "glue_foam_price": line_price(result_line_by_code(result, "eps_glue_foam"), "material"),
    }


def style_header_cell(cell: Any, fill: PatternFill, border: Border) -> None:
    cell.fill = fill
    cell.border = border
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup_top_header(sheet: Any) -> None:
    grey_fill = PatternFill("solid", fgColor="B7B7B7")
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    section_fill = PatternFill("solid", fgColor="EFEFEF")
    thin = Side(style="thin", color="303030")
    blue = Side(style="thick", color="1D4ED8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    split_border = Border(left=blue, right=thin, top=thin, bottom=thin)

    for row in range(1, 9):
        for col in range(10, 23):
            sheet.cell(row, col).fill = grey_fill

    sheet.merge_cells("A1:I4")
    logo_path = FIXTURES_DIR / "brick_house_logo.png"
    if logo_path.exists():
        logo = OpenpyxlImage(logo_path)
        logo.width = 300
        logo.height = 50
        sheet.add_image(logo, "C2")
    else:
        sheet["A1"] = "Brick House"
        sheet["A1"].font = Font(size=18, bold=True)
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells("A5:I5")
    sheet["A5"] = "СМЕТНЫЙ РАСЧЕТ НА СТРОИТЕЛЬСТВО ДОМА"
    sheet["A5"].font = Font(size=12, bold=True)
    sheet["A5"].alignment = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells("A6:I6")
    sheet["A6"] = ADDRESS
    sheet["A6"].font = Font(size=10, bold=False)
    sheet["A6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    sheet.row_dimensions[6].height = 24

    group_header_row = 9
    subheader_row = 10
    header_merges = [
        (group_header_row, 1, subheader_row, 1, "№"),
        (group_header_row, 2, subheader_row, 2, "Наименование работ"),
        (group_header_row, 3, subheader_row, 3, "Ед. изм."),
        (group_header_row, 4, subheader_row, 4, "Кол-во"),
        (group_header_row, 5, group_header_row, 6, "Стоимость материалов,\nмашин и механизмов, руб."),
        (group_header_row, 7, group_header_row, 8, "Стоимость работ, руб."),
        (group_header_row, 9, subheader_row, 9, "Итого, руб."),
        (group_header_row, 10, subheader_row, 10, "Кол-во"),
        (group_header_row, 11, group_header_row, 12, "Стоимость материалов,\nмашин и механизмов, руб."),
        (group_header_row, 13, group_header_row, 14, "Стоимость работ, руб."),
        (group_header_row, 15, subheader_row, 15, "Итого, руб."),
    ]
    for start_row, start_col, end_row, end_col, value in header_merges:
        sheet.merge_cells(
            start_row=start_row,
            start_column=start_col,
            end_row=end_row,
            end_column=end_col,
        )
        sheet.cell(start_row, start_col).value = value

    for col, value in {
        5: "За ед",
        6: "Итого",
        7: "За ед",
        8: "Итого",
        11: "За ед",
        12: "Итого",
        13: "За ед",
        14: "Итого",
    }.items():
        sheet.cell(subheader_row, col).value = value

    for row in [group_header_row, subheader_row]:
        for col in range(1, 16):
            fill = grey_fill if col >= 10 else white_fill
            style_header_cell(sheet.cell(row, col), fill, split_border if col == 10 else border)
    sheet.row_dimensions[group_header_row].height = 46
    sheet.row_dimensions[subheader_row].height = 22

    section_row = 11
    for col in range(1, 16):
        cell = sheet.cell(section_row, col)
        cell.fill = grey_fill if col >= 10 else section_fill
        cell.border = split_border if col == 10 else border
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.cell(section_row, 1).value = 4
    sheet.cell(section_row, 2).value = SECTION_TITLE
    sheet.cell(section_row, 2).alignment = Alignment(horizontal="left", vertical="center")


def estimate_rows(values: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        {
            "name": "Гидроизоляция фундаментной плиты битумной мастикой в 2 слоя",
            "unit": "м2",
            "qty": "=R12",
            "mat_price": 0,
            "work_price": values["waterproofing_work_price"],
        },
        {
            "name": "Праймер битумный AquaMast, 18 л",
            "unit": "шт",
            "qty": "=ROUNDUP(S13,0)",
            "mat_price": values["primer_price"],
            "work_price": 0,
        },
        {
            "name": "Мастика гидроизоляционная битумная для фундаментов AquaMast, 18 кг",
            "unit": "шт",
            "qty": "=ROUNDUP(T14,0)",
            "mat_price": values["mastic_price"],
            "work_price": 0,
        },
        {
            "name": "Утепление стен плиты ЭППС 100 мм",
            "unit": "м2",
            "qty": "=T15",
            "mat_price": 0,
            "work_price": values["eps_work_price"],
        },
        {
            "name": "Пеноплэкс ГЕО 100 мм",
            "unit": "м3",
            "qty": "=ROUNDUP(R16,0)*P16",
            "mat_price": values["eps100_price"],
            "work_price": 0,
            "qty_format": "0.00",
        },
        {
            "name": "Клей-пена для ЭППС",
            "unit": "баллон",
            "qty": "=MAX(S17,ROUNDUP(R17,0))",
            "mat_price": values["glue_foam_price"],
            "work_price": 0,
        },
        {
            "name": "Логистика и снабжение",
            "unit": "-",
            "qty": 1,
            "mat_price": "=SUM(O12:O17)*0.02",
            "work_price": 0,
        },
        {
            "name": "Расходные материалы, амортизация инструмента",
            "unit": "комплект",
            "qty": 1,
            "mat_price": "=SUM(O12:O17)*0.03",
            "work_price": 0,
        },
    ]


def write_estimate_table(sheet: Any, values: dict[str, Any]) -> None:
    grey_fill = PatternFill("solid", fgColor="B7B7B7")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="303030")
    blue = Side(style="thick", color="1D4ED8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    split_border = Border(left=blue, right=thin, top=thin, bottom=thin)

    first_row = 12
    for offset, row_data in enumerate(estimate_rows(values)):
        row = first_row + offset
        grey_values = {
            10: row_data["qty"],
            11: row_data["mat_price"],
            12: f"=J{row}*K{row}",
            13: row_data["work_price"],
            14: f"=J{row}*M{row}",
            15: f"=L{row}+N{row}",
        }
        white_values = {
            2: row_data["name"],
            3: row_data["unit"],
            4: f"=J{row}",
            5: f"=K{row}",
            6: f"=L{row}",
            7: f"=M{row}",
            8: f"=N{row}",
            9: f"=O{row}",
        }

        for col in range(1, 16):
            cell = sheet.cell(row, col)
            cell.border = split_border if col == 10 else border
            cell.fill = grey_fill if col >= 10 else PatternFill("solid", fgColor="FFFFFF")
            cell.alignment = Alignment(
                horizontal="center" if col in {1, 3, 4, 10} else None,
                vertical="top",
                wrap_text=col == 2,
            )
            if col in white_values:
                cell.value = white_values[col]
            if col in grey_values:
                cell.value = grey_values[col]

        for col in [4, 10]:
            sheet.cell(row, col).number_format = row_data.get("qty_format", "0.00")
        for col in [5, 6, 7, 8, 9, 11, 12, 13, 14, 15]:
            sheet.cell(row, col).number_format = "0"

    last_row = first_row + len(estimate_rows(values)) - 1
    sheet.merge_cells(start_row=first_row, start_column=1, end_row=last_row, end_column=1)
    section_cell = sheet.cell(first_row, 1)
    section_cell.value = SECTION_TITLE
    section_cell.font = Font(bold=True)
    section_cell.fill = PatternFill("solid", fgColor="FFFFFF")
    section_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        text_rotation=90,
        wrap_text=True,
    )
    section_cell.border = border
    for row in range(first_row + 1, last_row + 1):
        sheet.cell(row, 1).border = border
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor="FFFFFF")

    total_row = last_row + 1
    total_values = {
        2: "Итого по разделу:",
        6: f"=L{total_row}",
        8: f"=N{total_row}",
        9: f"=O{total_row}",
        12: f"=SUM(L{first_row}:L{last_row})",
        14: f"=SUM(N{first_row}:N{last_row})",
        15: f"=SUM(O{first_row}:O{last_row})",
    }
    for col in range(1, 16):
        cell = sheet.cell(total_row, col)
        cell.value = total_values.get(col, "")
        cell.font = Font(bold=True)
        cell.fill = grey_fill if col >= 10 else total_fill
        cell.border = split_border if col == 10 else border
        cell.number_format = "0"


def row_helper_values(values: dict[str, Any]) -> dict[int, dict[int, Any]]:
    return {
        12: {
            16: values["perimeter_m"],
            17: values["edge_height_m"],
            18: "=P12*Q12",
        },
        13: {
            16: "=J12",
            17: values["primer_consumption_l_per_m2"],
            18: values["primer_canister_volume_l"],
            19: "=P13*Q13/R13",
        },
        14: {
            16: "=J12",
            17: values["mastic_consumption_kg_per_m2_per_layer"],
            18: values["mastic_layers"],
            19: values["mastic_bucket_weight_kg"],
            20: "=P14*Q14*R14/S14",
        },
        15: {
            16: values["non_insulated_edge_total_m"],
            17: "=(P12-P15)*Q12",
            18: values["eps100_wall_volume_m3"],
            19: values["eps100_wall_thickness_m"],
            20: "=R15/S15",
        },
        16: {
            16: values["eps100_pack_volume_m3"],
            17: "=J15*0.1*1.05",
            18: "=Q16/P16",
            19: "=ROUNDUP(R16,0)",
        },
        17: {
            16: "=J15",
            17: values["glue_foam_coverage_m2_per_can"],
            18: "=P17/Q17",
            19: values["glue_foam_min_units"],
        },
    }


def write_working_field(sheet: Any, values: dict[str, Any]) -> None:
    grey_fill = PatternFill("solid", fgColor="B7B7B7")
    thin = Side(style="thin", color="303030")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col in range(16, 23):
        sheet.column_dimensions[get_column_letter(col)].width = 12

    for row in range(1, 24):
        for col in range(16, 23):
            cell = sheet.cell(row, col)
            cell.fill = grey_fill
            cell.border = border
            cell.font = Font(size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for row, values_by_col in row_helper_values(values).items():
        for col, value in values_by_col.items():
            cell = sheet.cell(row, col)
            cell.value = value
            cell.number_format = "0.0000" if row == 16 and col in {16, 17, 18} else "0.00"
            if row in {13, 14, 17} and col in {19, 20, 18}:
                cell.number_format = "0.000"
            if row in {18, 19, 20}:
                cell.number_format = "0"


def setup_dimensions(sheet: Any) -> None:
    widths = {
        1: 6,
        2: 58,
        3: 10,
        4: 12,
        5: 16,
        6: 16,
        7: 14,
        8: 14,
        9: 14,
        10: 12,
        11: 16,
        12: 16,
        13: 14,
        14: 14,
        15: 14,
    }
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "E12"


def build_workbook() -> Workbook:
    values = load_poc_values()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Смета"
    setup_top_header(sheet)
    write_estimate_table(sheet, values)
    write_working_field(sheet, values)
    setup_dimensions(sheet)
    return workbook


def formula_cells(sheet: Any, cell_range: str) -> list[str]:
    formulas = []
    for row in sheet[cell_range]:
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append(cell.value)
    return formulas


def text_values(sheet: Any, cell_range: str) -> list[str]:
    values = []
    for row in sheet[cell_range]:
        for cell in row:
            if cell.value is not None:
                values.append(str(cell.value))
    return values


def validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, data_only=False)
    if "Смета" not in workbook.sheetnames:
        raise AssertionError("Missing sheet: Смета")
    visible_sheets = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if visible_sheets != ["Смета"]:
        raise AssertionError(f"Expected only visible Смета sheet, got: {visible_sheets}")

    sheet = workbook["Смета"]
    grey_formulas = formula_cells(sheet, "J12:O19")
    white_formulas = formula_cells(sheet, "D12:I19")
    working_field_formulas = formula_cells(sheet, "Q1:V55")
    if not grey_formulas:
        raise AssertionError("Grey zone formulas not found")
    if not white_formulas:
        raise AssertionError("White zone formulas not found")
    if not working_field_formulas:
        raise AssertionError("Right working field formulas not found")
    if not all(formula.startswith("=") for formula in grey_formulas + white_formulas + working_field_formulas):
        raise AssertionError("Formula without leading = found")

    expected_white_links = {"=J12", "=K12", "=L12", "=M12", "=N12", "=O12"}
    first_white_row = {sheet.cell(12, col).value for col in range(4, 10)}
    if first_white_row != expected_white_links:
        raise AssertionError(f"White zone first row is not linked to grey zone: {first_white_row}")

    expected_row_helpers = {
        "P12": 81,
        "Q12": 0.3,
        "R12": "=P12*Q12",
        "S13": "=P13*Q13/R13",
        "T14": "=P14*Q14*R14/S14",
        "Q15": "=(P12-P15)*Q12",
        "T15": "=R15/S15",
        "P16": 0.2776,
        "Q16": "=J15*0.1*1.05",
        "R16": "=Q16/P16",
        "R17": "=P17/Q17",
    }
    for cell_ref, expected in expected_row_helpers.items():
        if sheet[cell_ref].value != expected:
            raise AssertionError(f"Unexpected helper value in {cell_ref}: {sheet[cell_ref].value}")
    for row in [18, 19, 20]:
        helper_values = [sheet.cell(row, col).value for col in range(16, 23)]
        if any(value is not None for value in helper_values):
            raise AssertionError(f"Helper field must be empty for row {row}: {helper_values}")

    if sheet["J16"].value != "=ROUNDUP(R16,0)*P16":
        raise AssertionError("Penoplex quantity must be raw ordered volume from row helper cells")
    if sheet["L16"].value != "=J16*K16":
        raise AssertionError("Penoplex material total must be calculated from raw J16")
    if sheet["J16"].number_format != "0.00":
        raise AssertionError("Penoplex quantity should display as 0.00 while storing raw formula")

    all_formulas = "\n".join(
        grey_formulas + white_formulas + working_field_formulas + formula_cells(sheet, "A1:V55")
    )
    for token in FORBIDDEN_FORMULA_TOKENS:
        if re.search(rf"\b{re.escape(token)}\b", all_formulas):
            raise AssertionError(f"Technical token found in formula: {token}")

    right_field_text = "\n".join(text_values(sheet, "P1:V55"))
    forbidden_panel_words = [
        "РАСЧЁТНОЕ ПОЛЕ",
        "Площадь гидроизоляции",
        "Праймер",
        "Мастика",
        "ЭППС 100 мм",
        "Клей-пена",
        "База / коэффициенты",
    ]
    for word in RIGHT_FIELD_FORBIDDEN_WORDS + forbidden_panel_words:
        if word in right_field_text:
            raise AssertionError(f"Forbidden word found in right working field: {word}")

    if not sheet["K13"].value or not sheet["M12"].value:
        raise AssertionError("Prices must stay in estimate columns K/M")
    if any(
        sheet.cell(row, col).value in {2770, 2780, 10000, 490, 350, 500}
        for row in range(1, 55)
        for col in range(16, 23)
    ):
        raise AssertionError("Price-looking constants found in right working field")

    has_logo_or_text = bool(getattr(sheet, "_images", [])) or sheet["A1"].value == "Brick House"
    if not has_logo_or_text:
        raise AssertionError("Logo or Brick House fallback text not found")


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook()
    workbook.save(OUTPUT_PATH)
    validate_workbook(OUTPUT_PATH)
    print(f"OK: generated {OUTPUT_PATH.relative_to(BASE_DIR.parents[1])}")
    print("Expected totals:")
    print("materials = 33961")
    print("works = 17255")
    print("section_total = 51216")


if __name__ == "__main__":
    main()
