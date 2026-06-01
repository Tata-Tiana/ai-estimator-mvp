"""Build reviewed_parameters.xlsx from section schema and section review cards."""

from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from section_schema import REPO_ROOT, get_section_schema


PARAMETERS_HEADERS = [
    "project",
    "section_code",
    "section_name",
    "parameter_code",
    "calculator_input_key",
    "label",
    "input_type",
    "required",
    "extracted_value",
    "corrected_value",
    "final_value",
    "unit",
    "source_file",
    "source_id",
    "page",
    "section_title",
    "source_text",
    "confidence",
    "parser_status",
    "elena_status",
    "use_for_calculation",
    "missing_reason",
    "comment",
]

MISSING_HEADERS = ["section_code", "section_name", "parameter_code", "calculator_input_key", "label", "unit", "input_type", "missing_reason", "comment"]
QUESTION_HEADERS = ["project", "section_code", "question", "answer", "status", "comment"]
SOURCE_HEADERS = ["source_id", "source_file", "title", "pages", "parsed_dir", "text_characters", "tables_found", "page_images_count", "artifacts"]
SUMMARY_HEADERS = ["project", "generated_at", "sources_count", "extraction_method", "demo_parameters_for_review", "ai_used", "notes"]
RAW_HEADERS = ["path", "value"]


def read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def style_sheet(ws, columns_count: int) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=2, column=1)
    ws.auto_filter.ref = ws.dimensions
    for idx in range(1, columns_count + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 18
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def append_rows(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws, len(headers))


def flatten_raw(value: Any, prefix: str = "") -> list[tuple[str, str]]:
    if isinstance(value, dict):
        rows: list[tuple[str, str]] = []
        for key, child in value.items():
            rows.extend(flatten_raw(child, f"{prefix}.{key}" if prefix else key))
        return rows
    if isinstance(value, list):
        rows = []
        for index, child in enumerate(value):
            rows.extend(flatten_raw(child, f"{prefix}[{index}]"))
        return rows
    return [(prefix, stringify(value))]


def load_review_cards(case_dir: Path) -> dict[str, dict[str, Any]]:
    cards: dict[str, dict[str, Any]] = {}
    for path in sorted((case_dir / "review_cards").glob("*_review_card.json")):
        card = read_json(path, {})
        section_code = card.get("section") or card.get("section_code")
        if section_code:
            cards[section_code] = card
    return cards


def card_parameters_by_code(card: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {item.get("parameter_code"): item for item in card.get("parameters", []) if item.get("parameter_code")}


def determine_elena_status(parameter: dict[str, Any], card_item: dict[str, Any] | None, final_value: Any) -> str:
    if parameter["input_type"] == "control_only":
        return "control_only"
    if card_item and card_item.get("value") not in (None, ""):
        return "needs_review"
    if final_value not in (None, ""):
        return "default_used"
    if parameter["input_type"] == "manual":
        return "manual_required"
    if parameter["required"]:
        return "missing"
    return "optional"


def build_review_sheet(case_dir: Path) -> dict[str, Any]:
    manifest = read_json(case_dir / "input/project_manifest.json", {})
    extraction_report = (case_dir / "input/extraction_report.md").read_text(encoding="utf-8") if (case_dir / "input/extraction_report.md").exists() else ""
    raw_extraction = read_json(case_dir / "input/extracted_parameters_for_review.json", {})
    project = manifest.get("project", "usv_yusupovo_village")
    sections = get_section_schema(REPO_ROOT)
    cards = load_review_cards(case_dir)

    parameter_rows: list[list[Any]] = []
    missing_rows: list[list[Any]] = []
    section_stats: dict[str, Counter] = defaultdict(Counter)

    for section in sections:
        card = cards.get(section["section_code"], {})
        card_params = card_parameters_by_code(card)
        for parameter in section["parameters"]:
            card_item = card_params.get(parameter["parameter_code"])
            extracted_value = card_item.get("value") if card_item else None
            default_value = parameter.get("default_value")
            final_value = extracted_value if extracted_value not in (None, "") else default_value
            elena_status = determine_elena_status(parameter, card_item, final_value)
            parser_status = card_item.get("review_status") if card_item else "not_found"
            use_for_calculation = bool(parameter["required"] and parameter["input_type"] != "control_only")
            missing_reason = "required parameter not found in parser/review cards" if parameter["required"] and final_value in (None, "") else ""

            stats = section_stats[section["section_code"]]
            stats["required_total"] += 1 if parameter["required"] else 0
            stats["extracted_found"] += 1 if extracted_value not in (None, "") else 0
            stats["manual_required"] += 1 if elena_status == "manual_required" else 0
            stats["missing"] += 1 if elena_status == "missing" else 0
            stats["control_only"] += 1 if parameter["input_type"] == "control_only" else 0
            stats["total"] += 1

            row = [
                project,
                section["section_code"],
                section["section_name"],
                parameter["parameter_code"],
                parameter["calculator_input_key"],
                parameter["label"],
                parameter["input_type"],
                parameter["required"],
                stringify(extracted_value),
                "",
                stringify(final_value),
                parameter.get("unit") or (card_item or {}).get("unit", ""),
                (card_item or {}).get("source_file", ""),
                (card_item or {}).get("source_id", ""),
                (card_item or {}).get("page", ""),
                (card_item or {}).get("section_title", ""),
                (card_item or {}).get("source_text", ""),
                (card_item or {}).get("confidence", ""),
                parser_status,
                elena_status,
                use_for_calculation,
                missing_reason,
                "",
            ]
            parameter_rows.append(row)

            if parameter["required"] and final_value in (None, ""):
                missing_rows.append([
                    section["section_code"],
                    section["section_name"],
                    parameter["parameter_code"],
                    parameter["calculator_input_key"],
                    parameter["label"],
                    parameter.get("unit", ""),
                    parameter["input_type"],
                    missing_reason,
                    "",
                ])

    question_rows = []
    for section_code, card in cards.items():
        for question in card.get("review_questions_for_elena", []):
            question_rows.append([project, section_code, question, "", "open", ""])

    source_rows = []
    for source in manifest.get("sources", []):
        source_rows.append([
            source.get("source_id", ""),
            source.get("source_file", ""),
            source.get("title", ""),
            source.get("pages", ""),
            source.get("parsed_dir", ""),
            source.get("text_characters", ""),
            source.get("tables_found", ""),
            source.get("page_images_count", ""),
            stringify(source.get("artifacts", [])),
        ])

    summary_rows = [[
        project,
        manifest.get("generated_at", datetime.now().isoformat(timespec="seconds")),
        len(manifest.get("sources", [])),
        "pdf_parser + section review cards + required parameter schema",
        raw_extraction.get("parameters") and len(raw_extraction.get("parameters", [])) or "",
        "false",
        extraction_report.splitlines()[0] if extraction_report else "",
    ]]

    wb = Workbook()
    ws = wb.active
    ws.title = "parameters"
    append_rows(ws, PARAMETERS_HEADERS, parameter_rows)
    append_rows(wb.create_sheet("missing_parameters"), MISSING_HEADERS, missing_rows)
    append_rows(wb.create_sheet("review_questions"), QUESTION_HEADERS, question_rows)
    append_rows(wb.create_sheet("sources"), SOURCE_HEADERS, source_rows)
    append_rows(wb.create_sheet("extraction_summary"), SUMMARY_HEADERS, summary_rows)
    append_rows(wb.create_sheet("raw_extraction_diagnostic"), RAW_HEADERS, [[path, value] for path, value in flatten_raw(raw_extraction)])

    xlsx_path = case_dir / "reviewed_parameters.xlsx"
    wb.save(xlsx_path)

    section_rows = []
    for section in sections:
        stats = section_stats[section["section_code"]]
        section_rows.append({
            "section_code": section["section_code"],
            "section_name": section["section_name"],
            "required_total": stats["required_total"],
            "extracted_found": stats["extracted_found"],
            "manual_required": stats["manual_required"],
            "missing": stats["missing"],
            "control_only": stats["control_only"],
            "review_card_json": str(case_dir / "review_cards" / f"{section['section_code']}_review_card.json"),
            "review_card_md": str(case_dir / "review_cards" / f"{section['section_code']}_review_card.md"),
        })

    return {
        "project": project,
        "sources": manifest.get("sources", []),
        "sections": section_rows,
        "parameters_total": len(parameter_rows),
        "parameters_for_calculation": sum(1 for row in parameter_rows if row[20] is True),
        "missing_total": len(missing_rows),
        "manual_required_total": sum(1 for row in parameter_rows if row[19] == "manual_required"),
        "output_xlsx": str(xlsx_path),
        "warnings": [],
    }
