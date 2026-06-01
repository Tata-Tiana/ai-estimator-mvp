from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]
CASE_DIR = REPO_ROOT / "experiments/pdf_parser_pipeline/cases/mvp_usv_demo"
OUTPUT_DIR = REPO_ROOT / "experiments/pdf_parser_pipeline/output/mvp_usv_demo"


SECTION_ORDER = [
    "earthworks",
    "foundation_slab",
    "waterproofing",
    "load_bearing_walls_lintels",
    "floor_slab_1",
    "floor_slab_2",
    "flat_roof",
    "schiedel_vent_channels",
]


TECHNICAL_TOKENS = [
    ".code",
    ".name",
    ".steel_class",
    ".diameter_mm",
    "unit_price",
    "work_unit_price",
    "material_unit_price",
    "rate_per",
    "_rate",
    "_coeff",
    "waste_coeff",
    "kg_per_meter",
    "rod_length",
    "pack_volume",
    "roll_area",
    "coverage",
    "round_step",
    "capacity",
    "min_",
    "display",
    "_raw",
    "calc_method",
    "strategy",
    "enabled",
    "box_total_metal_weight",
    "concrete_mixer",
    "planterband",
    "plywood",
    "sheet",
    "roll",
    "bucket",
    "canister",
    "pallet",
    "bag_weight",
    "truck",
    "delivery",
    "crane",
    "pump",
    "consumables",
    "logistics",
    "expected_material_total",
    "rail_piece",
    "source_weight",
    "weight_parts",
    "geotextile_laying_area",
    "sand_truck_step",
    "block_height_m",
    "gas_block_length_m",
    "mastic_",
    "primer_",
]


RU_LABELS = {
    "pit_area_m2": "Площадь котлована",
    "manual_refinement_depth_m": "Глубина ручной доработки",
    "trench_volume_m3": "Объем траншей",
    "communications_length_m": "Длина коммуникаций",
    "geotextile_laying_area_m2": "Площадь укладки геотекстиля",
    "sand_truck_step_m3": "Шаг заказа песка машиной",
    "slab_formwork_perimeter_m": "Периметр бортов/опалубки плиты",
    "slab_edge_height_m": "Высота борта плиты",
    "thermal_insert_length_m": "Длина термовставок",
    "thermal_insert_piece_length_m": "Длина элемента термовставки",
    "thermal_insert_piece_width_m": "Ширина элемента термовставки",
    "thermal_insert_piece_height_m": "Высота элемента термовставки",
    "eps100_wall_volume_m3": "Объем ЭППС 100 мм по стене/торцу",
    "non_insulated_edge_lengths_m": "Длины участков без утепления",
    "cutoff_waterproofing_wall_400_lengths_m": "Длины отсечной гидроизоляции стен 400 мм",
    "cutoff_waterproofing_wall_250_lengths_m": "Длины отсечной гидроизоляции стен 250 мм",
    "main_wall_external_length_m": "Длина наружных стен",
    "main_wall_internal_250_control_length_m": "Контрольная длина внутренних стен 250 мм",
    "lintel_lengths_m": "Длины и количества перемычек",
    "lintel_section_height_m": "Высота сечения перемычки",
    "lintel_section_width_m": "Ширина сечения перемычки",
    "main_wall_reinforcement_rows": "Количество рядов армирования кладки",
    "main_wall_250_reinforcement_threads": "Количество ниток армирования стены 250 мм",
    "main_wall_400_reinforcement_threads": "Количество ниток армирования стены 400 мм",
    "parapet_masonry_volume_m3": "Объем кладки парапета",
    "parapet_chasing_base_length_m": "Базовая длина штробления парапета",
    "parapet_rebar_base_length_m": "Базовая длина арматуры парапета",
    "second_light_masonry_volume_m3": "Объем кладки второго света",
    "second_light_chasing_base_length_m": "Базовая длина штробления второго света",
    "second_light_rebar_base_length_m": "Базовая длина арматуры второго света",
    "vent_chimney_segment_lengths_m": "Длины сегментов обкладки вентканалов",
    "total_concrete_volume_from_spec_m3": "Объем бетона по спецификации",
    "total_eps_volume_from_spec_m3": "Объем ЭППС по спецификации",
    "slab_edge_perimeter_m": "Периметр торца плиты",
    "edge_formwork_height_m": "Высота торцевой опалубки",
    "edge_insulation_height_m": "Высота утепления торца",
    "main_formwork_area_m2": "Основная площадь опалубки",
    "slab_area_m2": "Площадь плиты",
    "slab_length_m": "Длина плиты",
    "slab_width_m": "Ширина плиты",
    "length_m": "Длина",
    "width_m": "Ширина",
    "height_m": "Высота",
    "roof_area_level_1_m2": "Площадь кровли уровня 1",
    "roof_area_level_2_m2": "Площадь кровли уровня 2",
    "roof_area_total_m2": "Общая площадь кровли",
    "parapet_length_level_1_m": "Длина парапета уровня 1",
    "parapet_length_level_2_m": "Длина парапета уровня 2",
    "vent_wall_abutment_level_1_m": "Длина примыканий к вентшахтам/стенам уровня 1",
    "vent_wall_abutment_level_2_m": "Длина примыканий к вентшахтам/стенам уровня 2",
    "parapet_and_abutment_total_length_m": "Суммарная длина парапетов и примыканий",
    "eps50_supplier_required_volume_m3": "Объем ЭППС 50 мм по раскладке поставщика",
    "slope_plate_a_supplier_required_volume_m3": "Объем уклонных плит A по раскладке поставщика",
    "slope_plate_b_supplier_required_volume_m3": "Объем уклонных плит B по раскладке поставщика",
    "slope_plate_j_supplier_required_volume_m3": "Объем уклонных плит J по раскладке поставщика",
    "slope_plate_k_supplier_required_volume_m3": "Объем уклонных плит K по раскладке поставщика",
    "internal_drain_height_per_drain_m": "Высота внутреннего водостока на одну воронку",
    "vent_channel_1_height_m": "Высота вентканала 1",
    "vent_channel_2_height_m": "Высота вентканала 2",
    "vent_channel_2_count": "Количество вентканалов типа 2",
    "schiedel_masonry_total_length_m": "Общая длина кладки вентканалов Schiedel",
}


LEAF_LABELS = {
    "code": "код позиции",
    "name": "наименование позиции",
    "steel_class": "класс стали",
    "diameter_mm": "диаметр",
    "source_weight_kg": "вес по спецификации",
    "source_weight_parts_kg": "вес по спецификации",
    "weight_parts_kg": "вес по спецификации",
    "weight_kg": "вес по спецификации",
    "length_m": "длина",
    "width_m": "ширина",
    "height_m": "высота",
    "count": "количество",
}


def load_parameters() -> list[dict[str, Any]]:
    workbook = load_workbook(CASE_DIR / "reviewed_parameters.xlsx", read_only=True, data_only=True)
    sheet = workbook["parameters"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        rows.append({header: values[index] for index, header in enumerate(headers)})
    return rows


def key_leaf(calculator_input_key: str) -> str:
    leaf = calculator_input_key.split(".")[-1]
    if "[" in leaf:
        leaf = leaf.split("[")[0]
    return leaf


def array_index(calculator_input_key: str, array_name: str) -> int | None:
    import re

    match = re.search(rf"{array_name}\[(\d+)\]", calculator_input_key)
    return int(match.group(1)) + 1 if match else None


def contextual_label(calculator_input_key: str, leaf: str) -> str | None:
    context = ""
    if "beams.items" in calculator_input_key:
        index = array_index(calculator_input_key, "items")
        context = f"Балка Б-{index}" if index else "Балка"
    elif "lintel_lengths_m" in calculator_input_key:
        index = array_index(calculator_input_key, "lintel_lengths_m")
        context = f"Перемычка {index}" if index else "Перемычка"
    elif "lintel_rebar_items" in calculator_input_key:
        index = array_index(calculator_input_key, "lintel_rebar_items")
        context = f"Арматура перемычек {index}" if index else "Арматура перемычек"
    elif "rebar_items" in calculator_input_key:
        index = array_index(calculator_input_key, "rebar_items")
        context = f"Арматура {index}" if index else "Арматура"
    elif "vent_chimney_segment_lengths_m" in calculator_input_key:
        index = array_index(calculator_input_key, "vent_chimney_segment_lengths_m")
        context = f"Сегмент обкладки вентканалов {index}" if index else "Сегмент обкладки вентканалов"

    if not context or leaf not in LEAF_LABELS:
        return None

    label = LEAF_LABELS[leaf]
    if "source_weight_parts_kg" in calculator_input_key:
        part = array_index(calculator_input_key, "source_weight_parts_kg")
        if part:
            label = f"{label}, часть {part}"
    return f"{context}: {label}"


def human_label(row: dict[str, Any]) -> str:
    key = str(row.get("calculator_input_key") or "")
    leaf = key_leaf(key)
    contextual = contextual_label(key, leaf)
    if contextual:
        return contextual
    if leaf in RU_LABELS:
        return RU_LABELS[leaf]
    label = str(row.get("label") or "")
    return label.replace("_", " ")


def is_technical(row: dict[str, Any]) -> bool:
    key = str(row.get("calculator_input_key") or "").lower()
    code = str(row.get("parameter_code") or "").lower()
    return any(token in key or token in code for token in TECHNICAL_TOKENS)


def is_missing_for_calculation(row: dict[str, Any]) -> bool:
    return (
        row.get("use_for_calculation") is True
        and row.get("elena_status") in {"missing", "manual_required"}
        and not row.get("final_value")
    )


def split_rows(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    all_missing = [row for row in rows if is_missing_for_calculation(row)]
    designer_rows: list[dict[str, Any]] = []
    estimator_rows: list[dict[str, Any]] = []

    for row in all_missing:
        if row.get("input_type") == "parsed" and not is_technical(row):
            designer_rows.append(row)
        else:
            estimator_rows.append(row)
    return designer_rows, estimator_rows, all_missing


def normalize_row(row: dict[str, Any], responsibility: str) -> dict[str, Any]:
    return {
        "section_code": row.get("section_code"),
        "section_name": row.get("section_name"),
        "parameter_code": row.get("parameter_code"),
        "calculator_input_key": row.get("calculator_input_key"),
        "label": human_label(row),
        "unit": row.get("unit"),
        "input_type": row.get("input_type"),
        "elena_status": row.get("elena_status"),
        "responsibility": responsibility,
        "question": build_question(row, responsibility),
        "comment": "",
    }


def build_question(row: dict[str, Any], responsibility: str) -> str:
    label = human_label(row)
    unit = row.get("unit") or "-"
    if responsibility == "designer_or_project":
        return f"Уточнить по проекту/у проектировщика: {label}, ед. изм. {unit}."
    return f"Заполнить вручную в сметной логике/у Елены: {label}, ед. изм. {unit}."


def section_sort_key(row: dict[str, Any]) -> tuple[int, str]:
    section = str(row.get("section_code") or "")
    try:
        order = SECTION_ORDER.index(section)
    except ValueError:
        order = 999
    return order, str(row.get("parameter_code") or "")


def write_markdown(designer_rows: list[dict[str, Any]], estimator_rows: list[dict[str, Any]], all_missing: list[dict[str, Any]], path: Path) -> None:
    by_section: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in sorted(designer_rows, key=section_sort_key):
        by_section[str(row["section_code"])].append(row)

    estimator_by_section: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in sorted(estimator_rows, key=section_sort_key):
        estimator_by_section[str(row["section_code"])].append(row)

    lines = [
        "# Что parser не нашел по 8 разделам сметы",
        "",
        "Файл для Елены: список параметров, которые не удалось уверенно извлечь из PDF.",
        "",
        "Главный блок ниже — вопросы к проектировщикам/проекту. Отдельно вынесены параметры, которые скорее заполняет Елена или сметчик, а не проектировщик.",
        "",
        "## Сводка",
        "",
        "| Раздел | К проектировщикам/по проекту | Ручные сметные/технические | Всего missing/manual |",
        "|---|---:|---:|---:|",
    ]

    for section in SECTION_ORDER:
        designer_count = len(by_section.get(section, []))
        estimator_count = len(estimator_by_section.get(section, []))
        total_count = sum(1 for row in all_missing if row.get("section_code") == section)
        name = next((row.get("section_name") for row in all_missing if row.get("section_code") == section), section)
        lines.append(f"| {name} (`{section}`) | {designer_count} | {estimator_count} | {total_count} |")

    lines.extend(["", "## Вопросы к проектировщикам / по проекту", ""])
    for section in SECTION_ORDER:
        rows = by_section.get(section, [])
        name = rows[0]["section_name"] if rows else section
        lines.extend([f"### {name}", ""])
        if not rows:
            lines.extend(["Явных проектных missing-параметров нет.", ""])
            continue
        lines.extend(["| Что нужно уточнить | Ед. | calculator_input_key | Статус |", "|---|---|---|---|"])
        for row in rows:
            lines.append(f"| {row['label']} | {row['unit']} | `{row['calculator_input_key']}` | {row['elena_status']} |")
        lines.append("")

    lines.extend(["## Ручные сметные / технические параметры не для проектировщиков", ""])
    for section in SECTION_ORDER:
        rows = estimator_by_section.get(section, [])
        if not rows:
            continue
        name = rows[0]["section_name"]
        lines.extend([f"### {name}", "", "| Что заполнить | Ед. | calculator_input_key | Статус |", "|---|---|---|---|"])
        for row in rows:
            lines.append(f"| {row['label']} | {row['unit']} | `{row['calculator_input_key']}` | {row['elena_status']} |")
        lines.append("")

    lines.extend([
        "## Примечание",
        "",
        "- Parser/AI не считают смету.",
        "- Найденные значения всё равно требуют проверки Еленой.",
        "- Этот файл показывает только незаполненные параметры из текущего `reviewed_parameters.xlsx`.",
        "- Полный raw-список есть в Excel-версии на листе `all_missing_raw`.",
        "",
    ])
    path.write_text("\n".join(lines), encoding="utf-8")


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = sheet.cell(row=2, column=1)
    sheet.auto_filter.ref = sheet.dimensions
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def append_table(sheet, headers: list[str], rows: list[dict[str, Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    style_sheet(sheet)


def write_xlsx(designer_rows: list[dict[str, Any]], estimator_rows: list[dict[str, Any]], all_missing: list[dict[str, Any]], path: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    summary.append(["section_code", "section_name", "designer_or_project", "estimator_manual", "all_missing_manual"])
    for section in SECTION_ORDER:
        name = next((row.get("section_name") for row in all_missing if row.get("section_code") == section), section)
        summary.append([
            section,
            name,
            sum(1 for row in designer_rows if row.get("section_code") == section),
            sum(1 for row in estimator_rows if row.get("section_code") == section),
            sum(1 for row in all_missing if row.get("section_code") == section),
        ])
    style_sheet(summary)

    headers = [
        "section_code",
        "section_name",
        "parameter_code",
        "calculator_input_key",
        "label",
        "unit",
        "input_type",
        "elena_status",
        "responsibility",
        "question",
        "comment",
    ]
    append_table(wb.create_sheet("for_designers"), headers, designer_rows)
    append_table(wb.create_sheet("estimator_manual"), headers, estimator_rows)
    append_table(wb.create_sheet("all_missing_raw"), headers, all_missing)
    wb.save(path)


def main() -> int:
    rows = load_parameters()
    designer_raw, estimator_raw, all_raw = split_rows(rows)
    designer_rows = [normalize_row(row, "designer_or_project") for row in sorted(designer_raw, key=section_sort_key)]
    estimator_rows = [normalize_row(row, "estimator_or_internal") for row in sorted(estimator_raw, key=section_sort_key)]
    all_missing = [
        normalize_row(row, "designer_or_project" if row in designer_raw else "estimator_or_internal")
        for row in sorted(all_raw, key=section_sort_key)
    ]

    for target_dir in [CASE_DIR, OUTPUT_DIR]:
        target_dir.mkdir(parents=True, exist_ok=True)
        write_markdown(designer_rows, estimator_rows, all_missing, target_dir / "elena_missing_parameters_by_section.md")
        write_xlsx(designer_rows, estimator_rows, all_missing, target_dir / "elena_missing_parameters_by_section.xlsx")

    print(f"designer_or_project: {len(designer_rows)}")
    print(f"estimator_or_internal: {len(estimator_rows)}")
    print(f"all_missing_manual: {len(all_missing)}")
    print(f"output_md: {OUTPUT_DIR / 'elena_missing_parameters_by_section.md'}")
    print(f"output_xlsx: {OUTPUT_DIR / 'elena_missing_parameters_by_section.xlsx'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
