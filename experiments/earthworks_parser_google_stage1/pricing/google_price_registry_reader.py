from __future__ import annotations

import csv
import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook

from config import EXPERIMENT_DIR
from source_paths import LOCAL_PRICE_REGISTRY_PATH


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_price(value: Any) -> float | None:
    if value in {None, ""}:
        return None
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def read_local_registry_rows(path: Path = LOCAL_PRICE_REGISTRY_PATH) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["price_registry"] if "price_registry" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = [clean(cell) for cell in next(rows)]
    result = []
    for raw in rows:
        row = {header: raw[idx] if idx < len(raw) else None for idx, header in enumerate(headers)}
        price_code = clean(row.get("price_code"))
        if not price_code:
            continue
        result.append(
            {
                "Раздел": clean(row.get("Раздел")),
                "Наименование": clean(row.get("Наименование")),
                "Ед. изм.": clean(row.get("Ед. изм.")),
                "Цена": parse_price(row.get("Цена")),
                "Дата обновления": clean(row.get("Дата обновления")),
                "Комментарий": clean(row.get("Комментарий")),
                "price_code": price_code,
                "source": "local_price_registry_snapshot",
            }
        )
    return result


def read_google_registry_rows() -> tuple[list[dict[str, Any]], list[str]]:
    load_dotenv(EXPERIMENT_DIR / ".env")
    warnings: list[str] = []
    spreadsheet_id = os.getenv("GOOGLE_PRICE_REGISTRY_SPREADSHEET_ID", "").strip()
    if not spreadsheet_id:
        return [], ["GOOGLE_PRICE_REGISTRY_SPREADSHEET_ID не задан; используется локальный snapshot прайса."]
    try:
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request
        from google_auth_oauthlib.flow import InstalledAppFlow
        from googleapiclient.discovery import build
    except ImportError:
        return [], ["Google libraries не установлены; используется локальный snapshot прайса."]

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive.file",
    ]
    credentials_path = os.getenv("GOOGLE_OAUTH_CREDENTIALS_PATH", "credentials.json")
    token_path = os.getenv("GOOGLE_TOKEN_PATH", "token.json")
    creds = None
    if Path(token_path).exists():
        creds = Credentials.from_authorized_user_file(token_path, scopes)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(credentials_path, scopes)
            creds = flow.run_local_server(port=0)
        Path(token_path).write_text(creds.to_json(), encoding="utf-8")

    sheet_name = os.getenv("GOOGLE_PRICE_REGISTRY_SHEET_NAME", "price_registry")
    service = build("sheets", "v4", credentials=creds)
    values = (
        service.spreadsheets()
        .values()
        .get(spreadsheetId=spreadsheet_id, range=f"{sheet_name}!A:Z")
        .execute()
        .get("values", [])
    )
    if not values:
        return [], [f"Google price_registry `{sheet_name}` пустой; используется локальный snapshot."]
    headers = [clean(cell) for cell in values[0]]
    result = []
    for raw in values[1:]:
        row = {header: raw[idx] if idx < len(raw) else "" for idx, header in enumerate(headers)}
        price_code = clean(row.get("price_code"))
        if not price_code:
            continue
        result.append(
            {
                "Раздел": clean(row.get("Раздел")),
                "Наименование": clean(row.get("Наименование")),
                "Ед. изм.": clean(row.get("Ед. изм.")),
                "Цена": parse_price(row.get("Цена")),
                "Дата обновления": clean(row.get("Дата обновления")),
                "Комментарий": clean(row.get("Комментарий")),
                "price_code": price_code,
                "source": "google_price_registry",
            }
        )
    return result, warnings


def load_price_registry_snapshot(job_dir: Path) -> tuple[list[dict[str, Any]], list[str]]:
    pricing_dir = job_dir / "pricing"
    pricing_dir.mkdir(parents=True, exist_ok=True)
    rows, warnings = read_google_registry_rows()
    if not rows:
        rows = read_local_registry_rows()
    snapshot = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "rows_count": len(rows),
        "rows": rows,
        "warnings": warnings,
        "read_only": True,
    }
    (pricing_dir / "price_registry_snapshot.json").write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    with (pricing_dir / "price_registry_snapshot.csv").open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["price_code", "Раздел", "Наименование", "Ед. изм.", "Цена", "Дата обновления", "Комментарий", "source"],
        )
        writer.writeheader()
        writer.writerows(rows)
    return rows, warnings
