from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import DEFAULTS, SECTION_NAME_RU, SMOKE_TEST_KEY
from parameter_dictionary import EARTHWORKS_PARAMETER_DICTIONARY, SOURCE_STATUS_RU, parameter_meta
from source_paths import REVIEW_HINTS_PATH, REVIEW_TABLE_PATH


HEADERS = [
    "Раздел",
    "Группа параметра",
    "Параметр",
    "Технический ключ",
    "Как написано в проекте",
    "Значение parser",
    "Ед.",
    "Confidence",
    "Нужно проверить Елене",
    "Лист проекта",
    "Спецификация / блок проекта",
    "Файл PDF",
    "Страница PDF",
    "logical_sheet_type",
    "Фрагмент PDF",
    "Где используется в расчёте",
    "Статус источника",
    "Можно исправить",
    "Исправленное значение",
    "Ручное значение",
    "Итоговое значение",
    "Статус проверки",
    "Комментарий Елены",
    "Комментарий системы",
]

FILL_HEADER = PatternFill("solid", fgColor="D9E2F3")
FILL_PDF = PatternFill("solid", fgColor="D9EAD3")
FILL_AUTO = PatternFill("solid", fgColor="D9EAF7")
FILL_DEFAULT = PatternFill("solid", fgColor="FFF2CC")
FILL_REVIEW = PatternFill("solid", fgColor="FCE4D6")
FILL_MANUAL = PatternFill("solid", fgColor="EADCF8")


def yes_no(value: bool) -> str:
    return "да" if value else "нет"


def evidence_value(evidence: dict[str, Any] | None, key: str) -> Any:
    return (evidence or {}).get(key) or ""


def base_row(key: str, value: Any, unit: str, evidence: dict[str, Any] | None, system_comment: str) -> dict[str, Any]:
    meta = parameter_meta(key)
    needs_review = bool(meta["needs_review"])
    status_source = str(meta["source_status"])
    if status_source in {"DEFAULT_VALUE", "MATERIAL_CATALOG", "AUTO_CALCULATED"}:
        review_status = "не требуется"
        final_value = value
    elif value in {None, ""}:
        review_status = "не найдено"
        final_value = ""
    else:
        review_status = "ожидает проверки"
        final_value = value
    return {
        "Раздел": SECTION_NAME_RU,
        "Группа параметра": meta["source_group_ru"],
        "Параметр": meta["ru_label"],
        "Технический ключ": key,
        "Как написано в проекте": meta["project_wording"],
        "Значение parser": value if value is not None else "",
        "Ед.": unit,
        "Confidence": evidence_value(evidence, "confidence") or ("default" if status_source in {"DEFAULT_VALUE", "MATERIAL_CATALOG"} else ""),
        "Нужно проверить Елене": yes_no(needs_review),
        "Лист проекта": evidence_value(evidence, "logical_sheet_title"),
        "Спецификация / блок проекта": meta["project_wording"],
        "Файл PDF": evidence_value(evidence, "source_pdf"),
        "Страница PDF": evidence_value(evidence, "physical_page_number"),
        "logical_sheet_type": evidence_value(evidence, "logical_sheet_type"),
        "Фрагмент PDF": evidence_value(evidence, "raw_context"),
        "Где используется в расчёте": meta["used_in"],
        "Статус источника": status_source,
        "Можно исправить": "да" if meta["visible_to_elena"] else "нет",
        "Исправленное значение": "",
        "Ручное значение": "",
        "Итоговое значение": final_value,
        "Статус проверки": review_status,
        "Комментарий Елены": "",
        "Комментарий системы": system_comment,
    }


def review_rows(extracted: dict[str, Any]) -> list[dict[str, Any]]:
    params = extracted["parameters"]
    rows = []
    for key in [
        "pit_area_m2",
        "pit_excavation_depth_m",
        "sand_base_volume_m3",
        "trench_volume_m3",
        "geotextile_area_m2",
        "communications_length_m",
    ]:
        source = params[key]
        comment = "Найдено parser/PDF; ожидает проверки Елены."
        if key == "pit_excavation_depth_m":
            comment = "Интерпретация диапазона глубины: взят максимум для POC; обязательно проверить."
        if key == "communications_length_m":
            comment = "Рассчитано автоматически как сумма труб коммуникаций."
        rows.append(base_row(key, source.get("value"), source.get("unit", ""), source.get("evidence"), comment))

    geotextile = params["geotextile_area_m2"]
    rows.append(
        base_row(
            "geotextile_laying_area_m2",
            geotextile.get("value"),
            "м2",
            geotextile.get("evidence"),
            "POC assumption: площадь укладки принята равной площади геотекстиля; требует проверки, если в проекте есть отдельная площадь.",
        )
    )
    rows.extend(
        [
            base_row(
                "trench_routes",
                f"{len(params['trench_routes']['value'])} routes",
                "routes",
                params["trench_routes"].get("evidence"),
                "Структура маршрутов вынесена на лист 02_Траншеи.",
            ),
            base_row(
                "communications_pipe_items",
                f"{len(params['communications_pipe_items']['value'])} items",
                "items",
                None,
                "Структура труб вынесена на лист 03_Коммуникации; ф110 считается коммуникацией.",
            ),
        ]
    )
    for key, value in DEFAULTS.items():
        rows.append(base_row(key, value, "", None, "DEFAULT_VALUE / catalog value from local POC config."))
    smoke = base_row(
        SMOKE_TEST_KEY,
        "",
        "",
        None,
        "smoke-test gate; удалить перед production",
    )
    smoke["Группа параметра"] = "Ручной ввод"
    smoke["Статус источника"] = "MANUAL_REQUIRED"
    smoke["Нужно проверить Елене"] = "да"
    smoke["Можно исправить"] = "да"
    smoke["Итоговое значение"] = ""
    smoke["Статус проверки"] = "ожидает проверки"
    rows.append(smoke)
    return rows


def style_table(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        row_values = {ws.cell(1, col).value: row[col - 1].value for col in range(1, ws.max_column + 1)}
        fill = FILL_PDF
        if row_values.get("Статус источника") in {"DEFAULT_VALUE", "MATERIAL_CATALOG"}:
            fill = FILL_DEFAULT
        if row_values.get("Статус источника") == "AUTO_CALCULATED":
            fill = FILL_AUTO
        if row_values.get("Статус источника") == "MANUAL_REQUIRED":
            fill = FILL_MANUAL
        if row_values.get("Нужно проверить Елене") == "да" and row_values.get("Статус проверки") == "ожидает проверки":
            fill = FILL_REVIEW
        for cell in row:
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "Параметр": 32,
        "Технический ключ": 34,
        "Фрагмент PDF": 58,
        "Где используется в расчёте": 42,
        "Комментарий системы": 46,
    }
    for idx, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(header, min(max(len(header) + 2, 14), 28))


def write_rows_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    style_table(ws)


def build_review_workbook(extracted: dict[str, Any]) -> list[dict[str, Any]]:
    wb = Workbook()
    wb.remove(wb.active)
    rows = review_rows(extracted)
    write_rows_sheet(wb, "01_Проверка", rows, HEADERS)

    trench_rows = []
    for route in extracted["parameters"]["trench_routes"]["value"]:
        ev = route["evidence"]
        trench_rows.append(
            {
                "Маршрут": route["name"],
                "Длина, м": route["length_m"],
                "Глубина, м": route["depth_m"],
                "Ширина, м": route["width_m"],
                "Объём, м3": route["volume_m3"],
                "Файл PDF": ev.get("source_pdf"),
                "Страница PDF": ev.get("physical_page_number"),
                "Лист проекта": ev.get("logical_sheet_title"),
                "Фрагмент PDF": ev.get("raw_context"),
                "Confidence": ev.get("confidence"),
                "Комментарий": "Проверьте маршрут и объем по таблице траншей.",
            }
        )
    write_simple_sheet(wb, "02_Траншеи", trench_rows)

    pipe_rows = []
    for item in extracted["parameters"]["communications_pipe_items"]["value"]:
        ev = item["evidence"]
        pipe_rows.append(
            {
                "Код": item["code"],
                "Наименование из проекта": item["name"],
                "Длина одной трубы, м": item["pipe_length_m"],
                "Количество": item["quantity"],
                "Итоговая длина, м": item["total_length_m"],
                "Включено в расчёт": yes_no(item["include_in_communications"]),
                "Файл PDF": ev.get("source_pdf"),
                "Страница PDF": ev.get("physical_page_number"),
                "Лист проекта": ev.get("logical_sheet_title"),
                "Фрагмент PDF": ev.get("raw_context"),
                "Confidence": ev.get("confidence"),
                "Комментарий": "ф110 is communication pipe, not rebar",
            }
        )
    write_simple_sheet(wb, "03_Коммуникации", pipe_rows)
    REVIEW_TABLE_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(REVIEW_TABLE_PATH)
    return rows


def write_simple_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["Сообщение"])
        ws.append(["Нет данных"])
    else:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24


def build_hints_workbook() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    write_instruction_sheet(wb)
    write_parameter_hints_sheet(wb)
    write_status_sheet(wb)
    write_empty_help_sheet(wb)
    REVIEW_HINTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(REVIEW_HINTS_PATH)


def write_instruction_sheet(wb: Workbook) -> None:
    rows = [
        ["1", "Откройте earthworks_parameter_review.xlsx."],
        ["2", "Проверьте строки, где 'Нужно проверить Елене' = да."],
        ["3", "Если parser ошибся, внесите правильное значение в 'Исправленное значение'."],
        ["4", "Если параметр ручной, заполните 'Ручное значение'."],
        ["5", "Поставьте 'Статус проверки': проверено / исправлено / ручной ввод."],
        ["6", "После проверки запустите validate, затем calculate."],
    ]
    ws = wb.create_sheet("01_Как_проверять")
    ws.append(["Шаг", "Что сделать"])
    for row in rows:
        ws.append(row)
    style_hint_sheet(ws)


def write_parameter_hints_sheet(wb: Workbook) -> None:
    rows = []
    for key, meta in EARTHWORKS_PARAMETER_DICTIONARY.items():
        rows.append(
            [
                key,
                meta["ru_label"],
                meta["used_in"],
                "PDF / specification / review card" if meta["source_status"] == "AUTO_PROJECT" else meta["source_group_ru"],
                meta["project_wording"],
                meta["used_in"],
                "да" if meta["visible_to_elena"] else "нет",
                "Расчет блокируется, если обязательное поле не проверено или ручной ввод пуст.",
            ]
        )
    ws = wb.create_sheet("02_Параметры")
    ws.append(["Technical key", "Русское название", "Что это такое", "Где искать в проекте", "Как обычно написано", "Где используется", "Можно исправлять", "Что будет, если пусто"])
    for row in rows:
        ws.append(row)
    style_hint_sheet(ws)


def write_status_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("03_Статусы")
    ws.append(["Статус", "Что значит"])
    for status, ru in SOURCE_STATUS_RU.items():
        ws.append([status, ru])
    style_hint_sheet(ws)


def write_empty_help_sheet(wb: Workbook) -> None:
    rows = [
        ["pit_area_m2", "Найти на листе План котлована / если нет, заполнить вручную."],
        ["pit_excavation_depth_m", "Проверить глубину котлована; если диапазон, выбрать правильное значение."],
        ["sand_base_volume_m3", "Найти строку Песок в спецификации котлована."],
        ["trench_volume_m3", "Найти итог м3 в таблице траншей."],
        ["communications_pipe_items", "Проверить спецификацию труб на листе Схема коммуникаций."],
        ["manual_required_smoke_test", "Заполнить любое значение для проверки gate в POC."],
    ]
    ws = wb.create_sheet("04_Что_делать_если_пусто")
    ws.append(["Параметр", "Подсказка"])
    for row in rows:
        ws.append(row)
    style_hint_sheet(ws)


def style_hint_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 34


def dictionary_complete() -> bool:
    return all(meta.get("ru_label") and meta.get("used_in") for meta in EARTHWORKS_PARAMETER_DICTIONARY.values())
