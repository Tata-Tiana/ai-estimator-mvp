from __future__ import annotations

from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import PRICE_CODE_TO_INTERNAL_PRICE_KEY
from parameter_dictionary import EARTHWORKS_PARAMETER_DICTIONARY, parameter_meta
from source_paths import FINAL_EXCEL_PATH, REVIEW_TABLE_PATH


HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
PDF_FILL = PatternFill("solid", fgColor="D9EAD3")
AUTO_FILL = PatternFill("solid", fgColor="D9EAF7")
DEFAULT_FILL = PatternFill("solid", fgColor="FFF2CC")
WARNING_FILL = PatternFill("solid", fgColor="F4CCCC")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")


def setup_table(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for cell in ws.iter_rows(min_col=idx, max_col=idx, values_only=True):
            if cell[0] is not None:
                max_len = max(max_len, min(len(str(cell[0])), 60))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 48)


def copy_review_sheets(target_wb: Workbook) -> None:
    review_wb = load_workbook(REVIEW_TABLE_PATH, data_only=True)
    for sheet_name in ["01_Проверка", "02_Траншеи", "03_Коммуникации"]:
        source = review_wb[sheet_name]
        ws = target_wb.create_sheet(sheet_name)
        for row in source.iter_rows(values_only=True):
            ws.append(list(row))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22


def line_price_source(line: dict[str, Any], price_sources: dict[str, Any]) -> tuple[str, str]:
    price_code = line.get("price_code") or ""
    internal_key = PRICE_CODE_TO_INTERNAL_PRICE_KEY.get(price_code)
    keys = internal_key if isinstance(internal_key, tuple) else (internal_key,)
    sources = []
    warnings = []
    for key in keys:
        if not key:
            continue
        src = price_sources.get(key, {})
        if src.get("source"):
            sources.append(src["source"])
        if src.get("warning"):
            warnings.append(src["warning"])
    return ", ".join(sorted(set(sources))), "; ".join(warnings)


def build_sources_rows(extracted: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for key, payload in extracted["parameters"].items():
        if key in {"trench_routes", "communications_pipe_items"}:
            continue
        meta = parameter_meta(key) if key in EARTHWORKS_PARAMETER_DICTIONARY else {}
        ev = payload.get("evidence") or {}
        rows.append(
            [
                meta.get("ru_label", key),
                payload.get("value"),
                meta.get("source_group_ru", ""),
                ev.get("source_pdf"),
                ev.get("physical_page_number"),
                ev.get("logical_sheet_title"),
                ev.get("logical_sheet_type"),
                ev.get("raw_context"),
            ]
        )
    return rows


def export_excel(extracted: dict[str, Any], payload: dict[str, Any], result: dict[str, Any], output_path=FINAL_EXCEL_PATH) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    copy_review_sheets(wb)

    calc_rows = []
    for line in result["estimate_lines"]:
        price_source, warning = line_price_source(line, result.get("price_sources", {}))
        calc_rows.append(
            [
                line.get("code"),
                line.get("name"),
                line.get("unit"),
                line.get("quantity"),
                line.get("material_unit_price"),
                line.get("material_total"),
                line.get("work_unit_price"),
                line.get("work_total"),
                line.get("line_total"),
                line.get("price_code", ""),
                price_source,
                warning,
            ]
        )
    setup_table(
        wb.create_sheet("04_Расчет"),
        [
            "Код строки",
            "Наименование",
            "Ед.",
            "Кол-во",
            "Цена материала",
            "Сумма материала",
            "Цена работы",
            "Сумма работы",
            "Итого",
            "price_code",
            "Источник цены",
            "Предупреждение по цене",
        ],
        calc_rows,
    )

    totals = result["internal_totals"]
    setup_table(
        wb.create_sheet("05_Итоги"),
        ["Показатель", "Значение"],
        [
            ["Материалы", totals.get("internal_materials_total")],
            ["Работы", totals.get("internal_works_total")],
            ["Итого раздел", totals.get("internal_section_total")],
        ],
    )

    setup_table(
        wb.create_sheet("06_Источники"),
        ["Параметр", "Значение", "Источник", "Файл PDF", "Страница PDF", "Лист проекта", "logical_sheet_type", "Фрагмент PDF"],
        build_sources_rows(extracted),
    )

    log_rows = []
    for message in [
        "Прочитаны outputs strict parser v3 read-only.",
        "Собрана промежуточная таблица проверки параметров.",
        "Input собран из проверенной review table.",
        "Расчет выполнен через calculate_earthworks(), без старого runner.",
        "Excel статический, формул нет.",
    ]:
        log_rows.append(["Что сделано", message])
    for warning in result.get("price_warnings", []):
        log_rows.append(["Предупреждение цены", warning])
    for assumption in result.get("poc_assumptions", []):
        log_rows.append(["Assumption", assumption])
    log_rows.append(["Проверка Елены", "pit_excavation_depth_m интерпретирован как max depth из диапазона 300-250 мм."])
    setup_table(wb.create_sheet("07_Лог"), ["Тип", "Сообщение"], log_rows)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    raise ValueError(f"Formula detected in final Excel: {ws.title}!{cell.coordinate}")
        if ws.max_row > 1:
            for row in ws.iter_rows(min_row=2):
                fill = None
                values = [cell.value for cell in row]
                joined = " ".join(str(value) for value in values if value is not None)
                if "fallback" in joined or "warning" in joined.lower():
                    fill = WARNING_FILL
                elif "auto_reviewed_for_poc" in joined:
                    fill = PDF_FILL
                elif "DEFAULT_VALUE" in joined or "Стандартное" in joined:
                    fill = DEFAULT_FILL
                if fill:
                    for cell in row:
                        cell.fill = fill
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
