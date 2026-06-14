from __future__ import annotations

import json
from typing import Any

from openpyxl import load_workbook

from config import (
    DEFAULTS,
    EXPECTED_PRICE_REGISTRY_UNITS,
    PRICE_CODE_TO_INTERNAL_PRICE_KEY,
    PRODUCTION_METHODS,
    PROJECT_NAME,
    RISKY_REVIEW_REQUIRED_KEYS,
    SMOKE_TEST_KEY,
)
from parameter_dictionary import EARTHWORKS_PARAMETER_DICTIONARY, parameter_meta
from source_paths import FALLBACK_LIVE_INPUT_PATH, PRICE_REGISTRY_PATH, REVIEW_INPUT_PATH, REVIEW_TABLE_PATH


def load_json(path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def parse_cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            if "." in text or "," in text:
                return float(text.replace(",", "."))
            return int(text)
        except ValueError:
            return text
    return value


def review_rows() -> list[dict[str, Any]]:
    if not REVIEW_TABLE_PATH.exists():
        raise FileNotFoundError(f"Review table not found. Run prepare first: {REVIEW_TABLE_PATH}")
    wb = load_workbook(REVIEW_TABLE_PATH, data_only=True)
    ws = wb["01_Проверка"]
    headers = [cell.value for cell in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {header: value for header, value in zip(headers, values)}
        if row.get("Технический ключ"):
            rows.append(row)
    return rows


def row_by_key(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {str(row["Технический ключ"]): row for row in rows}


def preferred_value(row: dict[str, Any]) -> Any:
    corrected = parse_cell_value(row.get("Исправленное значение"))
    if corrected is not None:
        return corrected
    manual = parse_cell_value(row.get("Ручное значение"))
    if manual is not None:
        return manual
    final = parse_cell_value(row.get("Итоговое значение"))
    if final is not None:
        return final
    return parse_cell_value(row.get("Значение parser"))


def validate_review() -> tuple[list[str], list[str], list[dict[str, Any]]]:
    rows = review_rows()
    errors: list[str] = []
    warnings: list[str] = []
    for row in rows:
        key = str(row.get("Технический ключ"))
        status = str(row.get("Статус проверки") or "").strip()
        source_status = str(row.get("Статус источника") or "").strip()
        parser_value = parse_cell_value(row.get("Значение parser"))
        manual_value = parse_cell_value(row.get("Ручное значение"))
        if source_status == "MANUAL_REQUIRED":
            if manual_value is None or status not in {"ручной ввод", "проверено"}:
                errors.append(f"{key}: требуется ручное значение и статус `ручной ввод` или `проверено`.")
            continue
        if key in RISKY_REVIEW_REQUIRED_KEYS and status not in {"проверено", "исправлено"}:
            errors.append(f"{key}: рискованная интерпретация должна быть проверена или исправлена Еленой.")
            continue
        if source_status == "AUTO_PROJECT" and parser_value is not None and status == "ожидает проверки":
            warnings.append(f"{key}: найден parser value, но статус еще `ожидает проверки`.")
        if source_status == "AUTO_PROJECT" and parser_value is None and preferred_value(row) is None:
            errors.append(f"{key}: значение не найдено и не заполнено.")
    return errors, warnings, rows


def registry_prices() -> tuple[dict[str, float], dict[str, dict[str, Any]], list[str]]:
    warnings = []
    prices: dict[str, float] = {}
    meta: dict[str, dict[str, Any]] = {}
    if not PRICE_REGISTRY_PATH.exists():
        return prices, meta, [f"Price registry not found: {PRICE_REGISTRY_PATH}"]
    wb = load_workbook(PRICE_REGISTRY_PATH, read_only=True, data_only=True)
    ws = wb["price_registry"] if "price_registry" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = [str(cell) if cell is not None else "" for cell in next(rows)]
    index = {header: pos for pos, header in enumerate(headers)}
    for row in rows:
        price_code = row[index.get("price_code", -1)] if index.get("price_code", -1) >= 0 else None
        if not price_code or price_code not in EXPECTED_PRICE_REGISTRY_UNITS:
            continue
        unit = row[index.get("Ед. изм.", -1)] if index.get("Ед. изм.", -1) >= 0 else None
        price = row[index.get("Цена", -1)] if index.get("Цена", -1) >= 0 else None
        expected_unit = EXPECTED_PRICE_REGISTRY_UNITS[price_code]
        if unit != expected_unit:
            warnings.append(
                f"Registry price_code {price_code} skipped: unit `{unit}` does not match calculator unit `{expected_unit}`."
            )
            continue
        if price is None:
            continue
        internal_key = PRICE_CODE_TO_INTERNAL_PRICE_KEY[price_code]
        keys = internal_key if isinstance(internal_key, tuple) else (internal_key,)
        for key in keys:
            prices[key] = float(price)
            meta[key] = {"source": "price_registry", "price_code": price_code, "warning": ""}
    return prices, meta, warnings


def fallback_prices() -> dict[str, float]:
    data = load_json(FALLBACK_LIVE_INPUT_PATH, {})
    raw_prices = data.get("internal_prices") or {}
    return {key: float(val) for key, val in raw_prices.items()}


def build_internal_prices() -> tuple[dict[str, float], dict[str, dict[str, Any]], list[str]]:
    prices, sources, warnings = registry_prices()
    fallback = fallback_prices()
    for key, price in fallback.items():
        if key in prices:
            continue
        prices[key] = price
        sources[key] = {
            "source": "fallback_input_for_poc",
            "price_code": "",
            "warning": "Price not found in compatible registry rows; using fallback price from live input for POC only.",
        }
        warnings.append(f"Fallback price used for {key}.")
    return prices, sources, warnings


def read_routes() -> list[dict[str, Any]]:
    wb = load_workbook(REVIEW_TABLE_PATH, data_only=True)
    ws = wb["02_Траншеи"]
    headers = [cell.value for cell in ws[1]]
    routes = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {header: value for header, value in zip(headers, values)}
        if not row.get("Маршрут"):
            continue
        routes.append(
            {
                "route_code": {
                    "К1": "K1",
                    "К2": "K2",
                    "Вода": "water",
                    "Эл. кабель": "electric_cable",
                }.get(row["Маршрут"], row["Маршрут"]),
                "name": row["Маршрут"],
                "length_m": float(row["Длина, м"]),
                "depth_m": float(row["Глубина, м"]),
            }
        )
    return routes


def read_pipe_items() -> list[dict[str, Any]]:
    wb = load_workbook(REVIEW_TABLE_PATH, data_only=True)
    ws = wb["03_Коммуникации"]
    headers = [cell.value for cell in ws[1]]
    items = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {header: value for header, value in zip(headers, values)}
        if not row.get("Код"):
            continue
        items.append(
            {
                "code": row["Код"],
                "name": row["Наименование из проекта"],
                "pipe_length_m": row["Длина одной трубы, м"],
                "quantity": row["Количество"],
                "total_length_m": row["Итоговая длина, м"],
                "include_in_communications": row["Включено в расчёт"] == "да",
            }
        )
    return items


def build_input_from_review() -> dict[str, Any]:
    errors, warnings, rows = validate_review()
    if errors:
        raise ValueError("Review validation failed: " + "; ".join(errors))
    by_key = row_by_key(rows)
    prices, price_sources, price_warnings = build_internal_prices()

    def val(key: str) -> Any:
        if key in DEFAULTS:
            return DEFAULTS[key]
        return preferred_value(by_key[key])

    earthworks_input = {
        "project_name": PROJECT_NAME,
        "case_meta": {
            "experiment": "earthworks_mini_mvp_poc",
            "review_source": "earthworks_parameter_review.xlsx",
            "not_final_estimate": True,
        },
        "assumptions": {
            "manual_excavation_override": False,
            "sand_override": False,
            "geotextile_override": False,
        },
        **PRODUCTION_METHODS,
        "pit_area_m2": val("pit_area_m2"),
        "pit_excavation_depth_m": val("pit_excavation_depth_m"),
        "sand_base_volume_m3": val("sand_base_volume_m3"),
        "trench_volume_m3": val("trench_volume_m3"),
        "trench_routes": read_routes(),
        "geotextile_area_m2": val("geotextile_area_m2"),
        "geotextile_laying_area_m2": val("geotextile_laying_area_m2"),
        "communications_pipe_items": read_pipe_items(),
        "communications_length_m": val("communications_length_m"),
        **DEFAULTS,
        "internal_prices": prices,
    }
    payload = {
        "input": earthworks_input,
        "price_sources": price_sources,
        "price_warnings": price_warnings,
        "validation_warnings": warnings,
        "assumptions": [
            "geotextile_laying_area_m2 comes from reviewed table; initial prepare used geotextile_area_m2 as POC suggestion.",
            "manual_required_smoke_test is validated but not passed to EarthworksInput.",
        ],
    }
    REVIEW_INPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REVIEW_INPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload


def review_statistics(rows: list[dict[str, Any]]) -> dict[str, int]:
    total = len(rows)
    need_review = sum(1 for row in rows if row.get("Нужно проверить Елене") == "да")
    waiting = sum(1 for row in rows if row.get("Статус проверки") == "ожидает проверки")
    missing = sum(1 for row in rows if row.get("Статус проверки") == "не найдено")
    checked = sum(1 for row in rows if row.get("Статус проверки") in {"проверено", "исправлено", "ручной ввод", "не требуется"})
    return {"total": total, "need_review": need_review, "waiting": waiting, "missing": missing, "checked_or_not_required": checked}

