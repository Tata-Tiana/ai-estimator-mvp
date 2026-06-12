from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from mapping_rules import EXCLUDED_FROM_REVIEW_SHEET


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
MAPPED_PATH = DATA_DIR / "mapped_parameters.json"
CANDIDATES_PATH = DATA_DIR / "extracted_candidates.json"
REVIEW_PACK_PATH = DATA_DIR / "review_pack.xlsx"


FILLS = {
    "found": PatternFill("solid", fgColor="D9EAD3"),
    "low_confidence": PatternFill("solid", fgColor="FFF2CC"),
    "conflict": PatternFill("solid", fgColor="F9CB9C"),
    "missing": PatternFill("solid", fgColor="F4CCCC"),
    "supplier_required": PatternFill("solid", fgColor="CFE2F3"),
    "manual_required": PatternFill("solid", fgColor="EADCF8"),
    "header": PatternFill("solid", fgColor="D9EAD3"),
}


REVIEW_HEADERS = [
    "section_name",
    "label",
    "calculator_input_key",
    "Что это значит",
    "required_status",
    "value_from_pdf",
    "unit",
    "source_pdf",
    "page",
    "source_fragment",
    "confidence",
    "found_status",
    "needs_elena_review",
    "elena_status",
    "elena_value",
    "elena_comment",
    "final_value",
    "action_needed",
]


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def save_workbook(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    load_workbook(path, read_only=True).close()


def action_needed(row: dict[str, Any]) -> str:
    status = row.get("found_status")
    required = row.get("required_status")
    if status == "found":
        return "Проверить значение и подтвердить."
    if status == "low_confidence":
        return "Проверить единицу/контекст и подтвердить или исправить."
    if status == "conflict":
        return "Выбрать правильное значение."
    if status == "missing":
        return "Заполнить вручную или спросить проектировщика."
    if status == "supplier_required":
        return "Нужен поставщик/раскладка."
    if status == "manual_required" or required == "MANUAL_REQUIRED":
        return "Нужно ручное решение сметчика."
    return ""


def final_value_formula(row_number: int) -> str:
    return f'=IF(O{row_number}<>"",O{row_number},F{row_number})'


def include_in_review(row: dict[str, Any]) -> bool:
    if row.get("required_status") in EXCLUDED_FROM_REVIEW_SHEET:
        return False
    if row.get("found_status") in {"conflict", "low_confidence"}:
        return True
    if row.get("required_status") in {"SUPPLIER_INPUT", "MANUAL_REQUIRED"}:
        return True
    if row.get("required_status") == "OPTIONAL_CONTROL" and row.get("visible_to_elena") == "да":
        return True
    return row.get("required_status") == "AUTO_PROJECT"


def append_rows(ws, headers: list[str], rows: list[dict[str, Any]], keys: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = FILLS["header"]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row.get(key, "") for key in keys])
        fill = FILLS.get(row.get("found_status"))
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = 24


def build_review_sheet(wb: Workbook, mapped: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("01_Проверка_Елены")
    ws.append(REVIEW_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = FILLS["header"]
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in [item for item in mapped if include_in_review(item)]:
        excel_row = ws.max_row + 1
        ws.append(
            [
                row.get("section_name"),
                row.get("label"),
                row.get("calculator_input_key"),
                row.get("what_it_means") or row.get("reason"),
                row.get("required_status"),
                row.get("value_from_pdf"),
                row.get("unit"),
                row.get("source_pdf"),
                row.get("page"),
                row.get("source_fragment"),
                row.get("confidence"),
                row.get("found_status"),
                "да" if row.get("needs_elena_review") else "нет",
                "",
                "",
                "",
                final_value_formula(excel_row),
                action_needed(row),
            ]
        )
        fill = FILLS.get(row.get("found_status"))
        if fill:
            for cell in ws[excel_row]:
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    validation = DataValidation(
        type="list",
        formula1='"ok,исправить,заполнить,не используется,спросить проектировщика,нужен поставщик"',
        allow_blank=True,
    )
    ws.add_data_validation(validation)
    if ws.max_row > 1:
        validation.add(f"N2:N{ws.max_row}")
    ws.freeze_panes = "A2"
    widths = [22, 28, 36, 42, 18, 14, 10, 20, 8, 60, 14, 18, 18, 18, 18, 30, 18, 34]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width


def build_review_pack() -> Path:
    mapped = load_json(MAPPED_PATH)
    candidates = load_json(CANDIDATES_PATH)
    wb = Workbook()
    wb.remove(wb.active)
    build_review_sheet(wb, mapped)
    append_rows(
        wb.create_sheet("02_Все_извлеченные_кандидаты"),
        ["source_pdf", "page", "page_title", "raw_label", "raw_value", "unit", "confidence", "candidate_type", "raw_context", "notes"],
        candidates,
        ["source_pdf", "page", "page_title", "raw_label", "raw_value", "unit", "confidence", "candidate_type", "raw_context", "notes"],
    )
    append_rows(
        wb.create_sheet("03_Покрытие_AUTO_PROJECT"),
        ["section_name", "label", "calculator_input_key", "found_status", "value_from_pdf", "unit", "source_pdf", "page", "confidence", "notes"],
        [row for row in mapped if row.get("required_status") == "AUTO_PROJECT"],
        ["section_name", "label", "calculator_input_key", "found_status", "value_from_pdf", "unit", "source_pdf", "page", "confidence", "notes"],
    )
    append_rows(
        wb.create_sheet("04_Не_найдено"),
        ["section_name", "label", "calculator_input_key", "required_status", "reason", "action_needed"],
        [dict(row, action_needed=action_needed(row)) for row in mapped if row.get("found_status") == "missing"],
        ["section_name", "label", "calculator_input_key", "required_status", "reason", "action_needed"],
    )
    append_rows(
        wb.create_sheet("05_Конфликты"),
        ["section_name", "label", "calculator_input_key", "value_from_pdf", "unit", "source_fragment", "notes"],
        [row for row in mapped if row.get("found_status") == "conflict"],
        ["section_name", "label", "calculator_input_key", "value_from_pdf", "unit", "source_fragment", "notes"],
    )
    append_rows(
        wb.create_sheet("06_SUPPLIER_INPUT"),
        ["section_name", "label", "calculator_input_key", "value_from_pdf", "source_fragment", "action_needed"],
        [dict(row, action_needed=action_needed(row)) for row in mapped if row.get("found_status") == "supplier_required"],
        ["section_name", "label", "calculator_input_key", "value_from_pdf", "source_fragment", "action_needed"],
    )
    append_rows(
        wb.create_sheet("07_Готово_к_расчету"),
        ["section_name", "label", "calculator_input_key", "value_from_pdf", "unit", "source_pdf", "page", "confidence"],
        [
            row for row in mapped
            if row.get("found_status") == "found" and row.get("confidence") in {"high", "medium"} and row.get("required_status") == "AUTO_PROJECT"
        ],
        ["section_name", "label", "calculator_input_key", "value_from_pdf", "unit", "source_pdf", "page", "confidence"],
    )
    append_rows(
        wb.create_sheet("08_Отчет_по_проектировщику"),
        ["section_name", "calculator_input_key", "label", "found_status", "action_needed", "reason"],
        [dict(row, action_needed=action_needed(row)) for row in mapped if row.get("found_status") in {"missing", "low_confidence", "conflict"}],
        ["section_name", "calculator_input_key", "label", "found_status", "action_needed", "reason"],
    )
    save_workbook(wb, REVIEW_PACK_PATH)
    return REVIEW_PACK_PATH


def main() -> int:
    path = build_review_pack()
    print(f"review_pack: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
