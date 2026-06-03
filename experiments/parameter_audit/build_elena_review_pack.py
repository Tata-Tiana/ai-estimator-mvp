"""Build Elena review pack from parameter audit results.

This script is intentionally a reporting layer only. It does not update
section_schema.py, calculators, expected files, or reviewed_parameters.xlsx.
"""

from __future__ import annotations

import re
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]
AUDIT_XLSX = REPO_ROOT / "experiments/parameter_audit/output/mvp_usv_demo/parameter_audit_result.xlsx"
OUTPUT_DIR = REPO_ROOT / "experiments/parameter_audit/output/mvp_usv_demo"
PACK_XLSX = OUTPUT_DIR / "elena_parameter_review_pack.xlsx"
AGENDA_MD = OUTPUT_DIR / "elena_parameter_review_agenda.md"


DISCUSSION_STATUSES = {
    "MANUAL_REQUIRED",
    "REQUIRES_VALIDATION",
    "DEFAULT_VALUE",
    "AUTO_CALCULATED",
    "AUTO_PROJECT",
}


EXPLANATION_KEYS = [
    "used_in_estimate_lines",
    "calculator_formula",
    "usv_formula_example",
    "what_elena_should_check",
    "why_needed",
]

EXPLANATION_HEADERS = [
    "Где используется в смете",
    "Формула калькулятора",
    "Пример формулы ЮСВ",
    "Что проверить Елене",
    "Зачем нужен параметр",
]


LABEL_OVERRIDES = {
    "pit_area_m2": "Площадь котлована",
    "manual_refinement_depth_m": "Глубина ручной доработки котлована",
    "trench_volume_m3": "Объем траншей",
    "sand_truck_step_m3": "Шаг заказа песка машиной",
    "communications_length_m": "Длина коммуникаций",
    "geotextile_laying_area_m2": "Площадь укладки геотекстиля",
    "planterband_per_membrane_roll": "Количество ленты PLANTERBAND на один рулон мембраны",
    "slab_formwork_perimeter_m": "Периметр бортов/опалубки плиты",
    "slab_edge_height_m": "Высота борта плиты",
    "thermal_insert_length_m": "Длина термовставок",
    "thermal_insert_piece_length_m": "Термовставка: длина элемента",
    "thermal_insert_piece_width_m": "Термовставка: ширина элемента",
    "thermal_insert_piece_height_m": "Термовставка: высота элемента",
    "main_wall_external_length_m": "Длина наружных несущих стен",
    "main_wall_reinforcement_rows": "Количество рядов армирования кладки",
    "main_wall_400_reinforcement_threads": "Количество ниток армирования стены 400 мм",
    "main_wall_250_reinforcement_threads": "Количество ниток армирования стены 250 мм",
    "lintel_section_width_m": "Перемычки: ширина сечения",
    "lintel_section_height_m": "Перемычки: высота сечения",
    "parapet_masonry_volume_m3": "Объем кладки парапета",
    "second_light_masonry_volume_m3": "Объем кладки зоны второго света",
    "parapet_chasing_base_length_m": "Базовая длина штробления парапета",
    "second_light_chasing_base_length_m": "Базовая длина штробления зоны второго света",
    "parapet_rebar_base_length_m": "Базовая длина арматуры парапета",
    "second_light_rebar_base_length_m": "Базовая длина арматуры зоны второго света",
    "geometry.total_concrete_volume_from_spec_m3": "Объем бетона по спецификации",
    "insulation.total_eps_volume_from_spec_m3": "Объем ЭППС по спецификации",
    "slab_length_m": "Длина плиты",
    "slab_width_m": "Ширина плиты",
    "slab_area_m2": "Площадь плиты",
    "parapet_length_level_1_m": "Длина парапета уровня 1",
    "parapet_length_level_2_m": "Длина парапета уровня 2",
    "vent_wall_abutment_level_1_m": "Длина примыканий к вентшахтам/стенам уровня 1",
    "vent_wall_abutment_level_2_m": "Длина примыканий к вентшахтам/стенам уровня 2",
    "eps50_supplier_required_volume_m3": "Объем ЭППС 50 мм по раскладке поставщика",
    "slope_plate_a_supplier_required_volume_m3": "Объем уклонных плит A по раскладке поставщика",
    "slope_plate_b_supplier_required_volume_m3": "Объем уклонных плит B по раскладке поставщика",
    "slope_plate_j_supplier_required_volume_m3": "Объем уклонных плит J по раскладке поставщика",
    "slope_plate_k_supplier_required_volume_m3": "Объем уклонных плит K по раскладке поставщика",
    "vent_channel_1_height_m": "Высота вентканала 1",
    "vent_channel_2_height_m": "Высота вентканала 2",
    "vent_channel_2_count": "Количество вентканалов типа 2",
    "schiedel_masonry_total_length_m": "Общая длина кладки вентканалов Schiedel",
    "schiedel_delivery_trips": "Количество доставок вентканалов Schiedel",
}


TEMPLATE_INPUTS: dict[str, dict[str, Any]] = {}


TEMPLATE_INPUT_PATHS = {
    "earthworks": REPO_ROOT / "experiments/earthworks_calculator/cases/usv_yusupovo_village/input.json",
    "foundation_slab": REPO_ROOT / "experiments/foundation_slab_calculator/cases/test_foundation_slab/input.json",
    "waterproofing": REPO_ROOT / "experiments/waterproofing_calculator/cases/test_waterproofing_foundation_slab/input.json",
    "load_bearing_walls_lintels": REPO_ROOT / "experiments/load_bearing_walls_lintels_calculator/cases/test_load_bearing_walls_lintels/input.json",
    "floor_slab_1": REPO_ROOT / "experiments/floor_slab_1_calculator/cases/test_floor_slab_1/input.json",
    "floor_slab_2": REPO_ROOT / "experiments/floor_slab_2_calculator/cases/test_floor_slab_2/input.json",
    "flat_roof": REPO_ROOT / "experiments/flat_roof_calculator/cases/test_flat_roof_usv/input.json",
    "schiedel_vent_channels": REPO_ROOT / "experiments/schiedel_vent_channels_calculator/cases/test_schiedel_vent_channels_usv/input.json",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_template_inputs() -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for section, path in TEMPLATE_INPUT_PATHS.items():
        if not path.exists():
            result[section] = {}
            continue
        with path.open("r", encoding="utf-8") as f:
            result[section] = json.load(f)
    return result


def get_path(data: dict[str, Any], path: str) -> Any:
    current: Any = data
    for part in path.split("."):
        match = re.fullmatch(r"(.+)\[(\d+)\]", part)
        if match:
            key = match.group(1)
            index = int(match.group(2))
            if not isinstance(current, dict) or key not in current:
                return None
            current = current[key]
            if not isinstance(current, list) or index >= len(current):
                return None
            current = current[index]
        else:
            if not isinstance(current, dict) or part not in current:
                return None
            current = current[part]
    return current


def template_value(row: dict[str, Any]) -> Any:
    section = clean(row.get("section_code"))
    key = clean(row.get("calculator_input_key"))
    data = TEMPLATE_INPUTS.get(section, {})
    return get_path(data, key)


def fmt(value: Any) -> str:
    if value is None:
        return "..."
    if isinstance(value, float):
        return f"{value:g}"
    return clean(value)


def value_note(row: dict[str, Any]) -> str:
    value = template_value(row)
    return f"В эталонном кейсе ЮСВ сейчас: {fmt(value)}." if value is not None else "В текущей таблице значение не найдено; нужно подтвердить источник."


def context_pack(row: dict[str, Any]) -> dict[str, str]:
    section = clean(row.get("section_code"))
    key = clean(row.get("calculator_input_key"))
    label = display_label(row)
    status = clean(row.get("recommended_source_status"))
    lower = f"{key} {label}".lower()
    current = value_note(row)

    used = "Связь со строками сметы требует уточнения."
    formula = "Используется как входной параметр калькулятора."
    example = current
    check = f"Подтвердить значение параметра: {label}."
    why = "Без этого значения калькулятор не может корректно собрать количество или стоимость строки."

    if section == "earthworks":
        if "manual_excavation_override" in key:
            used = "Блок расчета строки `Разработка грунта вручную`."
            formula = "Если override включен, калькулятор может брать ручное количество `manual_excavation_quantity_for_estimate_m3`; иначе считает `pit_area_m2 * manual_refinement_depth_m + trench_volume_m3`."
            example = "ЮСВ: override включен; для строки сметы используется ручное количество 44.695 м3."
            check = "Подтвердить, нужно ли в этом проекте использовать ручное количество разработки грунта."
            why = "Это переключатель между формульным расчетом и ручным сметным количеством."
        elif "sand_override" in key:
            used = "Строки `Отсыпка песком`, `Песок строительный`, `Перемещение песка вручную`."
            formula = "Если override включен, количество песка можно задать вручную; иначе `sand_order_volume_m3 = ceil((sand_base_volume_m3 + trench_volume_m3) * sand_compaction_coeff / sand_truck_step_m3) * sand_truck_step_m3`."
            example = "ЮСВ: override выключен; расчетный заказ песка 160 м3."
            check = "Подтвердить, нужно ли считать песок формулой или задать вручную."
            why = "От этого зависит количество сразу в трех строках песка."
        elif "geotextile_override" in key:
            used = "Строка материала `Геотекстиль Дорнит 300`."
            formula = "Если override включен, количество геотекстиля можно задать вручную; иначе `ceil(geotextile_area_m2 * overlap_coeff / roll_area_m2) * roll_area_m2`."
            example = "ЮСВ: override выключен; 330 * 1.10 = 363 м2, заказ 400 м2."
            check = "Подтвердить, можно ли считать геотекстиль по площади и рулонам."
            why = "Это переключатель между формульным расчетом и ручным количеством материала."
        elif key == "pit_area_m2":
            used = "Строка `Разработка грунта вручную`."
            formula = "`manual_pit_volume_m3 = pit_area_m2 * manual_refinement_depth_m`; затем `manual_excavation_total = manual_pit_volume_m3 + trench_volume_m3`."
            example = "ЮСВ: 330 * 0.08 = 26.4 м3; 26.4 + 18.29 = 44.69 м3."
            check = "Подтвердить площадь котлована, от которой считается ручная доработка дна."
            why = "Площадь котлована участвует в количестве ручной разработки грунта."
        elif key == "manual_refinement_depth_m":
            used = "Строка `Разработка грунта вручную`."
            formula = "`manual_pit_volume_m3 = pit_area_m2 * manual_refinement_depth_m`."
            example = "ЮСВ: 330 * 0.08 = 26.4 м3 ручной доработки дна котлована."
            check = "Подтвердить глубину ручной доработки для текущего проекта."
            why = "Это сметное правило, которое превращает площадь котлована в объем ручных работ."
        elif key == "trench_volume_m3":
            used = "Строки `Разработка грунта вручную`, `Отсыпка песком`, `Песок строительный`."
            formula = "Если объем задан, берется `trench_volume_m3`; иначе `trench_length_m * trench_depth_m * trench_width_m`. Также песок по траншеям: `trench_volume_m3 * sand_compaction_coeff`."
            example = "ЮСВ: траншеи 18.29 м3; в ручной разработке 26.4 + 18.29 = 44.69 м3; в песке 18.29 * 1.3 = 23.777 м3."
            check = "Подтвердить объем траншей или дать длину/ширину/глубину."
            why = "Объем траншей влияет и на ручную разработку, и на песчаную подготовку."
        elif key == "sand_truck_step_m3":
            used = "Строки песка: работы, материал, перемещение."
            formula = "`sand_order_volume_m3 = ceil(sand_total_m3 / sand_truck_step_m3) * sand_truck_step_m3`."
            example = "ЮСВ: 149.357 м3 округляется вверх шагом 20 м3 до 160 м3."
            check = "Подтвердить стандартный шаг заказа песка машиной."
            why = "Это не проектный объем, а правило закупочного округления."
        elif key == "geotextile_roll_area_m2":
            used = "Строка `Геотекстиль Дорнит 300`."
            formula = "`rolls = ceil(geotextile_area_m2 * overlap_coeff / geotextile_roll_area_m2)`; материал = `rolls * geotextile_roll_area_m2`."
            example = "ЮСВ: 330 * 1.10 = 363 м2; ceil(363 / 100) = 4 рулона; заказ 400 м2."
            check = "Подтвердить площадь рулона геотекстиля как каталожное значение."
            why = "Нужно для закупочного количества геотекстиля."
        elif key == "communications_length_m":
            used = "Строки `Закладка технологических входов коммуникаций` и `Материалы для устройства входов коммуникаций`."
            formula = "Работы: `communications_length_m * communications_work_unit_price`; материалы: `communications_length_m * communications_material_unit_price`."
            example = "ЮСВ: 115 м * ставка работ; 115 м * ставка материалов."
            check = "Подтвердить длину технологических вводов до границы дома."
            why = "Одна длина задает количество и для работ, и для материалов коммуникаций."
        elif key == "axis_marking_shifts":
            used = "Строка `Вынос осей фундамента, котлована на участок`."
            formula = "`line_total = axis_marking_shifts * axis_marking_work_unit_price`."
            example = "ЮСВ: 1 смена * 20 000 = 20 000."
            check = "Подтвердить количество смен для разбивки осей."
            why = "Это организационная сметная позиция, не извлекается напрямую из проекта."
        elif key == "excavator_shifts":
            used = "Строка `Механизированная разработка грунта, Экскаватор JCB`."
            formula = "Материалы/механизмы: `excavator_shifts * excavator_material_unit_price`; работа: `excavator_shifts * excavator_work_unit_price`."
            example = "ЮСВ: 3 смены * 26 000 = 78 000; 3 * 3 500 = 10 500."
            check = "Подтвердить количество смен экскаватора."
            why = "Это решение по организации работ и технике."
        elif key == "geotextile_laying_area_m2":
            used = "Строка `Укладка геотекстиля`."
            formula = "`work_total = geotextile_laying_area_m2 * geotextile_laying_work_unit_price`."
            example = "ЮСВ: 340 м2 * 35 = 11 900."
            check = "Подтвердить площадь укладки геотекстиля для работ."
            why = "Площадь работ может отличаться от закупочного количества геотекстиля с нахлестом."
        elif key == "manual_excavation_quantity_for_estimate_m3":
            used = "Строка `Разработка грунта вручную`."
            formula = "Если заполнено, строка берет это количество напрямую; иначе считает `pit_area_m2 * manual_refinement_depth_m + trench_volume_m3`."
            example = "ЮСВ: напрямую используется 44.695 м3; сумма строки = 44.695 * 1 400 = 62 573."
            check = "Подтвердить ручной объем, если он должен переопределять формулу."
            why = "Это финальное количество для строки ручной разработки грунта при override."

    elif section == "foundation_slab":
        if "membrane_roll_area" in key:
            used = "Строки `Planter Standard Технониколь` и `PLANTERBAND`."
            formula = "`membrane_rolls = ceil(membrane_area_m2 / membrane_roll_area_m2)`; `planterband_quantity = membrane_rolls * planterband_per_membrane_roll`."
            example = f"{current} В ЮСВ рулоны мембраны считаются от площади мембраны и площади рулона."
            check = "Подтвердить площадь рулона мембраны как каталожное значение."
            why = "Определяет закупочное количество рулонов мембраны и ленты."
        elif "planterband_per_membrane_roll" in key:
            used = "Строка `PLANTERBAND 10м х 10см`."
            formula = "`planterband_quantity = membrane_rolls * planterband_per_membrane_roll`."
            example = f"{current} Количество ленты привязано к числу рулонов PLANTER."
            check = "Подтвердить норму ленты на один рулон мембраны."
            why = "Это закупочная норма, не проектный объем."
        elif key in {"slab_formwork_perimeter_m", "slab_edge_height_m"}:
            used = "Строки опалубки, фанеры и пиломатериала для отбортовки плиты."
            formula = "`formwork_area_m2 = slab_formwork_perimeter_m * slab_edge_height_m`; дальше площадь участвует в фанере и пиломатериале."
            example = f"{current} Формула ЮСВ: периметр борта * высота борта = площадь отбортовки."
            check = "Подтвердить периметр и высоту борта фундаментной плиты."
            why = "Эти параметры задают количество опалубки по торцам плиты."
        elif "plywood_sheet" in key:
            used = "Строки `Фанера ФК 1,52*1,52` и `Пиломатериал`."
            formula = "`plywood_sheets = ceil(formwork_area_m2 / plywood_sheet_working_area_m2 + reserve)`; листы и площадь влияют на закупку фанеры."
            example = f"{current} В ЮСВ рабочая площадь листа используется как делитель при расчете листов."
            check = "Подтвердить размер/рабочую площадь листа фанеры как каталог/default."
            why = "Это системная настройка материала, ее не нужно спрашивать у проектировщика."
        elif "thermal_insert" in key:
            used = "Строки `Устройство и монтаж термовкладыша` и материалы ЭППС 150 мм."
            formula = "Работы: `thermal_insert_length_m * work_rate`; объем ЭППС термовставок считается из длины и геометрии элемента."
            example = f"{current} В ЮСВ термовставки считаются по длине и размерам элемента."
            check = "Подтвердить геометрию и длину термовставок по проекту."
            why = "Без этих размеров нельзя посчитать работы и объем ЭППС термовставок."
        elif "rebar_items" in key:
            used = "Строки арматуры фундаментной плиты и нулевая строка `Изготовление и монтаж каркаса армирования`."
            formula = "Для каждой арматуры: `base_length = source_weight_kg / kg_per_meter`; `length_with_waste = base_length * waste_coeff`; `rods = ceil(length_with_waste / rod_length)`; `order_length = rods * rod_length`; материал = `order_length * price_per_m`."
            example = f"{current} Метаданные арматуры формируют price_code и название; веса по спецификации задают закупочную длину."
            check = "Подтвердить веса арматуры по спецификации; код/класс/диаметр лучше брать из каталога или спецификации."
            why = "Арматура считается по весу из спецификации, но закупается в погонных метрах с округлением до хлыстов."
        elif "rebar_crane_shifts" in key:
            used = "Строка `Подача арматуры автокраном`."
            formula = "`line_total = rebar_crane_shifts * crane_shift_price`."
            example = f"{current} В ЮСВ это фиксированное количество смен крана."
            check = "Подтвердить количество смен крана для подачи арматуры."
            why = "Это организационная строка, не геометрия проекта."
        elif "rebar_metal_delivery_trucks" in key or "box_total_metal_weight" in key:
            used = "Строка `Доставка арматуры, металла`."
            formula = "В будущем: `trucks = ceil(total_metal_weight_kg / max_weight_per_truck_kg)`; сейчас количество машин может быть ручным."
            example = f"{current} Для ЮСВ доставка металла должна контролироваться общим весом коробки."
            check = "Подтвердить правило доставки металла: ручное количество или расчет по общему весу."
            why = "Чтобы не задвоить доставку металла между разделами."
        elif "concrete_mixer_volume" in key:
            used = "Строка `Доставка бетона до объекта`."
            formula = "`concrete_delivery_trips = ceil(concrete_order_volume_m3 / concrete_mixer_volume_m3)`."
            example = f"{current} Объем миксера нужен для количества рейсов бетона."
            check = "Подтвердить стандартный объем миксера."
            why = "Это правило логистики бетона."
        elif "concrete_pump_shifts" in key:
            used = "Строка `Работа бетононасоса`."
            formula = "`line_total = concrete_pump_shifts * concrete_pump_shift_price`."
            example = f"{current} В ЮСВ бетононасос считается сменами."
            check = "Подтвердить количество смен бетононасоса."
            why = "Это ручная/организационная строка техники."
        elif "plywood_calc_method" in key or "slab_edge_height_strategy" in key:
            used = "Блок расчета опалубки фундаментной плиты."
            formula = "Метод/стратегия выбирает правило расчета, например какую высоту торца брать и как считать фанеру."
            example = f"{current} Это не проектное количество, а выбранное сметное правило."
            check = "Подтвердить, какое правило применять для текущего типа фундамента."
            why = "Разные правила дают разные закупочные количества фанеры/пиломатериала."

    elif section == "waterproofing":
        if key in {"slab_formwork_perimeter_m", "slab_edge_height_m"}:
            used = "Строки вертикальной гидроизоляции, праймера, мастики и утепления ЭППС 100 мм."
            formula = "`waterproofing_area = slab_formwork_perimeter_m * slab_edge_height_m` с корректировками по участкам без утепления."
            example = f"{current} В ЮСВ периметр и высота борта дают площадь вертикальных работ."
            check = "Подтвердить периметр и высоту вертикальной поверхности."
            why = "От этой площади зависят работы и материалы гидроизоляции."
        elif "primer" in key:
            used = "Строка праймера перед битумной мастикой."
            formula = "`primer_liters = waterproofing_area * primer_consumption_l_per_m2`; `canisters = ceil(primer_liters / primer_canister_volume_l)`."
            example = f"{current} В ЮСВ расход и объем канистры переводят площадь в закупку праймера."
            check = "Подтвердить расход праймера и объем канистры как каталог/default."
            why = "Это технологическая норма материала."
        elif "mastic" in key:
            used = "Строка битумной мастики."
            formula = "`mastic_required = waterproofing_area * consumption * mastic_layers`; `buckets = ceil(mastic_required / mastic_bucket_weight_kg)`."
            example = f"{current} В ЮСВ мастика считается по площади, слоям и весу ведра."
            check = "Подтвердить количество слоев и вес ведра мастики."
            why = "Эти параметры переводят площадь гидроизоляции в закупочные ведра."
        elif "non_insulated_edge_lengths" in key:
            used = "Блок вертикального утепления/гидроизоляции по участкам без утепления."
            formula = "`insulated_length = total_edge_length - sum(non_insulated_edge_lengths_m)`; дальше площадь = `insulated_length * slab_edge_height_m`."
            example = f"{current} В ЮСВ эти участки исключаются из утепляемой длины."
            check = "Подтвердить длину каждого участка, где утепление не выполняется."
            why = "Чтобы не завысить площадь ЭППС и клея."
        elif "glue_foam" in key:
            used = "Строка `Клей-пена для ЭППС`."
            formula = "`glue_units = max(glue_foam_min_units, ceil(eps_area / coverage_per_unit))`."
            example = f"{current} Минимум нужен, если расчетная площадь дает меньше одной единицы."
            check = "Подтвердить минимальный заказ клей-пены."
            why = "Это закупочное ограничение, не проектная геометрия."

    elif section == "load_bearing_walls_lintels":
        if "scaffolding" in key:
            used = "Строки устройства подмостей/лесов и пиломатериала для них."
            formula = "Работы/материал считаются напрямую от заданного количества: `quantity * unit_price`."
            example = f"{current} В ЮСВ это отдельное количество для организации кладки."
            check = "Подтвердить объем/количество подмостей для кладочных работ."
            why = "Это вспомогательная сметная позиция для выполнения стен."
        elif "cutoff_waterproofing" in key:
            used = "Строка отсечной гидроизоляции под стены."
            formula = "`cutoff_total_length = sum(lengths_m)`; материал/работы = `cutoff_total_length * unit_price`."
            example = f"{current} В ЮСВ все участки 250/400 мм суммируются в общую длину отсечки."
            check = "Подтвердить длины участков отсечной гидроизоляции по кладочному плану."
            why = "Каждый участок входит в суммарную длину материала и работ."
        elif "pallet_volume" in key:
            used = "Строки закупки газобетонных блоков."
            formula = "`ordered_volume = ceil(required_volume / pallet_volume_m3) * pallet_volume_m3`."
            example = f"{current} Объем поддона нужен для округления закупки блоков."
            check = "Подтвердить объем блока/поддона как данные поставщика."
            why = "Калькулятор закупает блоки упаковками, а не ровно проектным объемом."
        elif "adhesive_consumption" in key:
            used = "Строка клея для газобетона."
            formula = "`adhesive_bags = ceil(gas_block_volume_m3 * adhesive_consumption_bag_per_m3)`."
            example = f"{current} В ЮСВ расход клея применяется к объему кладки."
            check = "Подтвердить расход клея на 1 м3 кладки."
            why = "Это технологическая норма, влияющая на закупку мешков клея."
        elif "sand_concrete_bag_weight" in key:
            used = "Строка пескобетона/смеси для перемычек."
            formula = "Закупка мешков считается от расчетной потребности и веса одного мешка."
            example = f"{current} Вес мешка нужен для округления закупки."
            check = "Подтвердить вес стандартного мешка."
            why = "Это каталог/default материала."
        elif "lintel_lengths_m" in key:
            used = "Строки перемычек: изготовление/монтаж, бетон, арматура и опалубка перемычек."
            formula = "`lintel_total_length = sum(length_m * count)`; объем бетона = `lintel_total_length * section_width * section_height`."
            example = f"{current} В ЮСВ длина и количество каждой перемычки суммируются."
            check = "Подтвердить длину и количество перемычек по плану перемычек."
            why = "Это основа для объема бетона, работ и материалов перемычек."
        elif "gas_block_length_m" in key or "block_height_m" in key:
            used = "Блок кладки и армирования газобетонных стен."
            formula = "Размер блока используется для количества рядов/перевода геометрии кладки в расчетные длины и объемы."
            example = f"{current} В ЮСВ размер блока влияет на рядность и армирование кладки."
            check = "Подтвердить размер блока как материал/каталог."
            why = "Без размера блока нельзя корректно связать кладочный план с расчетом."
        elif "main_wall" in key and ("length" in key or "reinforcement" in key):
            used = "Строки армирования кладки несущих стен."
            formula = "`rebar_length = wall_length * reinforcement_rows * reinforcement_threads` с учетом выбранных стен 250/400 мм."
            example = f"{current} В ЮСВ длина стен, ряды и нитки формируют метраж арматуры кладки."
            check = "Подтвердить длину стен и схему армирования."
            why = "Эти параметры задают количество арматуры в кладке."
        elif "main_walls_crane_shifts" in key or "concrete_delivery_trips" in key or "parapet_crane_shifts" in key:
            used = "Строки техники/доставки для стен и перемычек."
            formula = "`line_total = quantity * unit_price`."
            example = f"{current} Количество смен/рейсов умножается на ставку."
            check = "Подтвердить количество смен или рейсов."
            why = "Это организационная часть сметы, обычно не берется из проекта напрямую."
        elif "lintel_rebar_items" in key:
            used = "Строки арматуры перемычек."
            formula = "Метаданные арматуры формируют название и price_code; закупка считается по весу/длине с округлением до хлыстов."
            example = f"{current} В ЮСВ код/класс/диаметр лучше брать из каталога, а не спрашивать вручную."
            check = "Подтвердить, что эти поля можно формировать из спецификации/каталога арматуры."
            why = "Это описание позиции, а не отдельное проектное количество."
        elif "lintel_section" in key:
            used = "Блок расчета бетонных перемычек."
            formula = "`lintel_concrete_volume = lintel_total_length * lintel_section_width_m * lintel_section_height_m`."
            example = f"{current} В ЮСВ сечение перемычки переводит длину перемычек в объем бетона."
            check = "Подтвердить ширину и высоту сечения перемычек."
            why = "Сечение напрямую влияет на объем бетона и опалубки."
        elif "lintel_concrete_min_order" in key:
            used = "Строка бетона перемычек."
            formula = "`concrete_order_volume = max(calculated_lintel_volume, lintel_concrete_min_order_volume_m3)`."
            example = f"{current} Минимальный заказ защищает от слишком малого закупочного объема."
            check = "Подтвердить минимальный заказ бетона."
            why = "Это закупочное правило, не проектный объем."
        elif any(token in key for token in ("parapet", "second_light", "vent_chimney")):
            used = "Дополнительные блоки кладки: парапет, второй свет, обкладка вентканалов."
            formula = "При включении блока объем кладки/длины сегментов переводятся в работы, блоки, клей и армирование."
            example = f"{current} В ЮСВ эти параметры включают или уточняют дополнительные участки кладки."
            check = "Подтвердить, входит ли этот блок в расчет и какие объемы/длины брать."
            why = "Это case-specific часть стен, ее нельзя скрывать без подтвержденного правила."

    elif section == "floor_slab_1":
        if "total_concrete_volume_from_spec" in key:
            used = "Строки бетонирования плиты, бетона М300, доставки бетона; также контроль разделения плита/балки."
            formula = "`slab_concrete_volume = total_concrete_volume_from_spec_m3 - beams_total_concrete_volume`; `slab_formwork_area = slab_concrete_volume / slab_thickness`."
            example = f"{current} В ЮСВ из общего объема вычитаются балки, затем считается площадь опалубки плиты."
            check = "Подтвердить общий объем бетона по спецификации плиты перекрытия."
            why = "Это базовый объем раздела перекрытия 1-го этажа."
        elif "slab_concrete_volume_m3_raw" in key or "slab_concrete_volume_m3_display" in key:
            used = "Контрольные raw/display значения для бетонирования плиты."
            formula = "`slab_concrete_volume_raw = total_concrete_volume_from_spec_m3 - beams_total_concrete_volume`; display = округленное значение для Excel."
            example = f"{current} В ЮСВ стоимость считается от raw, а не от display."
            check = "Не заполнять вручную; подтвердить исходные объемы, из которых это считается."
            why = "Raw/display нужны для точного совпадения с Excel."
        elif "slab_edge_perimeter" in key or "edge_formwork_height" in key:
            used = "Строка опалубки отбортовки плиты, фанера и пиломатериал."
            formula = "`edge_formwork_area = slab_edge_perimeter_m * edge_formwork_height_m`; потом `edge_and_beam_formwork_area = edge_formwork_area + beams_formwork_area`."
            example = f"{current} В ЮСВ торец плиты добавляется к площади опалубки балок."
            check = "Подтвердить периметр торца и высоту отбортовки."
            why = "Это влияет на фанеру, пиломатериал и нулевую строку опалубки."
        elif "beams.items" in key:
            used = "Строки `Бетонирование балки`, опалубка балок, фанера и пиломатериал."
            formula = "Для каждой балки: `concrete_volume = length * width * height`; `formwork_area = length * (width + 2 * height)`."
            example = f"{current} В ЮСВ балки Б-1/Б-2/Б-3 суммируются в общий объем и площадь опалубки."
            check = "Подтвердить длину, ширину и высоту каждой балки по доп. таблице/проекту."
            why = "Геометрия балок отделяет объем балок от плиты и задает опалубку."
        elif "formwork_supplier_quote" in key or "slab_2_formwork_area" in key:
            used = "Строка `Комплект опалубки` и справочный расчет ставки аренды."
            formula = "`raw_average_rate = supplier_quote_total / (floor_slab_1_area + floor_slab_2_area)`; в текущем Excel может использоваться фиксированная ставка."
            example = f"{current} В ЮСВ это контекст для проверки ставки опалубки."
            check = "Подтвердить предложение поставщика и площадь второго перекрытия для справочной ставки."
            why = "Нужно для проверки ставки аренды опалубки, но не должно быть слепым default."
        elif "crane_shift_rate" in key or "concrete_pump_rate" in key:
            used = "Строки автокрана или бетононасоса."
            formula = "`line_total = shifts * unit_price`."
            example = f"{current} Цена должна приходить из прайса, а количество смен задается отдельно."
            check = "Проверить наличие ставки в price_registry."
            why = "Это цена, не параметр проекта."
        elif "plywood_sheet_working_area" in key or "reserve_plywood_sheets" in key or "overhang_sheet_equivalent" in key:
            used = "Строки фанеры и пиломатериала для опалубки."
            formula = "`plywood_sheets = ceil(edge_and_beam_area / sheet_area + non_multiple_area / sheet_area + reserve)`; доп. объем пиломатериала учитывает некратные места и свесы."
            example = f"{current} В ЮСВ эти параметры участвуют в закупке фанеры/пиломатериала."
            check = "Подтвердить резерв/рабочую площадь/свесы как правило расчета."
            why = "Они переводят площадь опалубки в закупочные листы и кубы пиломатериала."
        elif "rebar_items" in key:
            used = "Строки арматуры перекрытия 1-го этажа и нулевая строка каркаса."
            formula = "`base_length = source_weight / kg_per_meter`; `length_with_waste = base_length * waste`; `rods = ceil(length_with_waste / rod_length)`; `order_length = rods * rod_length`; материал = `order_length * price_per_m`."
            example = f"{current} В ЮСВ веса по спецификации по диаметрам переводятся в закупочные метры."
            check = "Подтвердить веса арматуры по спецификации; метаданные брать из каталога."
            why = "Арматура закупается хлыстами, поэтому нужен вес по диаметрам."
        elif "rebar_weight_for_delivery" in key or "max_rebar_delivery_weight" in key:
            used = "Строка `Доставка арматуры, металла`."
            formula = "`trucks = ceil((floor_slab_1_rebar_weight + floor_slab_2_rebar_weight) / max_weight_per_truck)`."
            example = f"{current} В ЮСВ доставка металла должна считаться по общему весу, чтобы не задвоить ее."
            check = "Подтвердить весовой контекст и грузоподъемность машины."
            why = "Это влияет на количество машин доставки металла."
        elif "total_eps_volume_from_spec" in key:
            used = "Строки утепления торцов/низа плиты и ЭППС."
            formula = "`bottom_slab_eps_volume = total_eps_volume_from_spec - edge_and_beam_eps_volume`; площадь низа = `volume / thickness`."
            example = f"{current} В ЮСВ общий объем ЭППС делится на торцы/балки и низ плиты."
            check = "Подтвердить общий объем ЭППС по спецификации."
            why = "Это базовый объем утеплителя для раздела."
        elif "overheads" in key:
            used = "Строки `Логистика, и снабжение` и `Расходные материалы, амортизация инструмента`."
            formula = "`addon_total = raw_base_before_overheads * percent`."
            example = f"{current} В ЮСВ проценты считаются от raw-базы до overheads."
            check = "Подтвердить процент как системное правило."
            why = "Это не проектная геометрия, а сметное правило."
        elif "manual_lines" in key:
            used = "Ручные строки доставки/крана/бетононасоса."
            formula = "`line_total = manual_quantity * unit_price`."
            example = f"{current} Количество задается вручную, ставка берется из прайса/входных данных."
            check = "Подтвердить количество машин/смен для текущего объекта."
            why = "Это организация работ, не всегда выводится из геометрии."

    elif section == "floor_slab_2":
        if key in {"slab_length_m", "slab_width_m", "slab_area_m2", "slab_edge_perimeter_m"}:
            used = "Строки опалубки, утепления торца и контроля геометрии плиты 2-го этажа."
            formula = "`slab_area_m2 = slab_length_m * slab_width_m`; `slab_edge_perimeter_m = 2 * (slab_length_m + slab_width_m)`."
            example = f"{current} ЮСВ: 9 * 9.1 = 81.9 м2; периметр 36.2 м."
            check = "Подтвердить габариты, площадь и периметр плиты."
            why = "Эти параметры задают площадь опалубки и длину утепления торца."
        elif "main_formwork_area" in key:
            used = "Строка `Комплект опалубки`."
            formula = "`material_total = main_formwork_area_m2 * formwork_rental_used_rate_per_m2`."
            example = f"{current} ЮСВ: 81.9 * 850 = 69 615."
            check = "Подтвердить основную площадь опалубки."
            why = "Это базовое количество аренды опалубки."
        elif "formwork_rental_supplier_quote" in key:
            used = "Справочный контроль ставки аренды опалубки."
            formula = "`raw_supplier_rate = formwork_rental_supplier_quote_total / main_formwork_area_m2`; расчетная строка может использовать confirmed rate."
            example = f"{current} ЮСВ: 68 860 / 81.9 = 840.78; используется ставка 850."
            check = "Подтвердить сумму предложения поставщика."
            why = "Это проверка ставки, а не универсальный default."
        elif "delivery_trips" in key or "crane_shifts" in key or "concrete_pump_shifts" in key:
            used = "Строки доставки опалубки, автокрана или бетононасоса."
            formula = "`line_total = quantity * unit_price`."
            example = f"{current} В ЮСВ количество рейсов/смен умножается на ставку."
            check = "Подтвердить количество рейсов/смен."
            why = "Это организационный параметр."
        elif "edge_formwork_height" in key:
            used = "Строка торцевой опалубки, фанера и пиломатериал."
            formula = "`edge_formwork_area = slab_edge_perimeter_m * edge_formwork_height_m`."
            example = f"{current} ЮСВ: 36.2 * 0.2 = 7.24 м2."
            check = "Подтвердить высоту торцевой опалубки."
            why = "Нужна для фанеры и пиломатериала."
        elif "plywood" in key:
            used = "Строка фанеры и расчет пиломатериала."
            formula = "`plywood_sheets = ceil(edge_area / sheet_area + non_multiple_area / sheet_area + reserve)`."
            example = f"{current} ЮСВ: 7.24/2.3 + 16.38/2.3 + 5 = 15.27; заказ 16 листов."
            check = "Подтвердить рабочую площадь листа и резерв."
            why = "Это правило закупки фанеры."
        elif "rebar_items" in key:
            used = "Строки арматуры плиты 2-го этажа и нулевая строка каркаса."
            formula = "`order_length = ceil((source_weight / kg_per_meter * waste_coeff) / rod_length) * rod_length`; материал = `order_length * price_per_m`."
            example = f"{current} В ЮСВ веса арматуры переводятся в закупочные метры по диаметрам."
            check = "Подтвердить веса арматуры по спецификации; метаданные брать из каталога."
            why = "Арматура закупается хлыстами, поэтому нужен вес и диаметр."
        elif "concrete_mixer_volume" in key:
            used = "Строка `Доставка бетона до объекта`."
            formula = "`trips = ceil(concrete_order_volume_m3 / concrete_mixer_volume_m3)`."
            example = f"{current} ЮСВ: 17.5 / 9 = 1.94; заказ 2 рейса."
            check = "Подтвердить объем миксера."
            why = "Определяет количество рейсов бетона."
        elif "edge_insulation_height" in key:
            used = "Строки утепления торца и ЭППС 100 мм."
            formula = "`edge_insulation_area = slab_edge_perimeter_m * edge_insulation_height_m`; `eps_volume = area * eps_thickness`."
            example = f"{current} ЮСВ: 36.2 * 0.18 = 6.516 м2."
            check = "Подтвердить высоту утепления торца."
            why = "В названии раздела 200 мм, но для Excel сейчас используется 0.18 м; это надо подтвердить."
        elif "foam_min_cans" in key or "logistics_rate" in key or "consumables_rate" in key:
            used = "Строки клей-пены, логистики и расходников."
            formula = "Клей: `max(min_cans, ceil(edge_area / coverage))`; логистика/расходники: `raw_base * percent`."
            example = f"{current} В ЮСВ эти параметры являются системными правилами расчета."
            check = "Подтвердить как default/catalog правило."
            why = "Это не проектная геометрия, а правило закупки/накладных."

    elif section == "flat_roof":
        if "roof_area_level" in key or key == "roof_area_total_m2":
            used = "Строки пароизоляции, утепления ЭППС, геотекстиля, ПВХ мембраны и работ по кровле."
            formula = "`roof_area_total_m2 = roof_area_level_1_m2 + roof_area_level_2_m2`; работы/материалы считаются от общей площади с коэффициентами."
            example = f"{current} ЮСВ: 177.52 + 71.40 = 248.92 м2."
            check = "Подтвердить площади кровли по уровням."
            why = "Общая площадь кровли является базой большинства строк раздела."
        elif "parapet_length" in key or "vent_wall_abutment" in key or "parapet_and_abutment_total" in key:
            used = "Строки примыканий ПВХ мембраны, геотекстиля парапетов, алюминиевых реек."
            formula = "`parapet_and_abutment_total_length_m = parapet_length_level_1 + parapet_length_level_2 + vent_wall_abutments`."
            example = f"{current} ЮСВ: суммарная длина примыканий 138.62 м."
            check = "Подтвердить длины парапетов и примыканий по плану/узлам кровли."
            why = "Длина примыканий задает работы и материалы по вертикальным участкам кровли."
        elif "roll_area" in key or "roll_width" in key or "roll_length" in key:
            used = "Строки пароизоляции, геотекстиля или ПВХ мембраны."
            formula = "`rolls = ceil(required_area / roll_area)`; для ПВХ `roll_area = roll_width * roll_length`."
            example = f"{current} В ЮСВ рулонные материалы закупаются целыми рулонами."
            check = "Подтвердить параметры рулона как каталог материала."
            why = "Это закупочное округление, не проектная площадь."
        elif "supplier_required_volume" in key:
            used = "Строки ЭППС 50 мм и уклонных плит кровли."
            formula = "`packs = ceil(supplier_required_volume_m3 / pack_volume_m3)`; `ordered_volume = packs * pack_volume_m3`; сумма = `ordered_volume * price`."
            example = f"{current} В ЮСВ объемы уклонных плит берутся из раскладки поставщика/Технониколь."
            check = "Подтвердить объем по раскладке поставщика для текущей кровли."
            why = "Калькулятор пока не строит раскладку уклонных плит сам."
        elif "vent_shaft_abutment_count" in key:
            used = "Строка `Монтаж примыкания к вентшахтам`."
            formula = "`work_total = vent_shaft_abutment_count * rate_per_item`."
            example = f"{current} ЮСВ: 3 * 5 000 = 15 000."
            check = "Подтвердить количество вентшахт/примыканий."
            why = "Это отдельные штучные работы на кровле."
        elif "rail_piece_length" in key:
            used = "Строки прижимной и краевой алюминиевой рейки."
            formula = "`pieces = ceil(parapet_and_abutment_total_length_m / rail_piece_length_m)`; `ordered_length = pieces * rail_piece_length_m`."
            example = f"{current} ЮСВ: 138.62 / 3 = 46.2; заказ 47 шт = 141 м."
            check = "Подтвердить длину одной рейки как каталог."
            why = "Нужно для закупочного округления рейки."
        elif "pvc_membrane_expected_material_total" in key:
            used = "Строка `Полимерная мембрана ПВХ Logicroof V-RP`."
            formula = "В идеале: `rolls * price_per_roll`; в ЮСВ есть raw/display расхождение, поэтому сумма фиксируется как контрольная."
            example = f"{current} ЮСВ: 11 рулонов, Excel сумма 565 738."
            check = "Подтвердить источник суммы/цены ПВХ мембраны."
            why = "Нужно разобрать расхождение между отображаемой ценой и суммой Excel."
        elif "aerators_count" in key:
            used = "Строка `Аэратор кровельный PVC`."
            formula = "Материал = `count * unit_price`; работа = `count * installation_rate`."
            example = f"{current} ЮСВ: 3 аэратора."
            check = "Подтвердить количество кровельных аэраторов."
            why = "Это штучная позиция кровли."
        elif "installation_rate" in key or "drilling_rate" in key:
            used = "Строки установки аэраторов/воронок или пробивки отверстий."
            formula = "`work_total = count * installation_rate`."
            example = f"{current} Это ставка работы, ее надо хранить в price_registry."
            check = "Проверить ставку в прайсе."
            why = "Это цена/ставка, не проектный параметр."
        elif "drains_count" in key or "holes_count" in key or "internal_drain_height" in key:
            used = "Строки кровельных воронок, отверстий и внутреннего водостока."
            formula = "Воронки/отверстия: `count * unit_price/rate`; внутренний водосток: `internal_roof_drains_count * internal_drain_height_per_drain_m`."
            example = f"{current} В ЮСВ штучные количества и высота стояка задают работы и материалы."
            check = "Подтвердить количество воронок/отверстий и высоту внутреннего водостока."
            why = "Эти параметры задают штучные и погонные метры кровельного водоотвода."
        elif "roof_crane_lifting_shifts" in key:
            used = "Строка `Подъем материалов автокраном`."
            formula = "`line_total = roof_crane_lifting_shifts * crane_shift_price`."
            example = f"{current} В ЮСВ это количество смен крана для кровельных материалов."
            check = "Подтвердить количество смен автокрана."
            why = "Это организационная строка техники."
        elif "consumables" in key or "logistics" in key or "technical_supervision" in key or "procurement_storage" in key:
            used = "Строки расходников, логистики, технического надзора или заготовительно-складских расходов."
            formula = "Сумма берется как подтвержденная ставка/процент/ручная строка; после аудита должна уйти в price_registry/defaults или отдельное правило."
            example = f"{current} В ЮСВ это не геометрия кровли, а сметное правило/ставка."
            check = "Подтвердить источник: прайс, процент от базы или ручная строка."
            why = "Такие суммы нельзя оставлять непонятным ручным вводом."

    elif section == "schiedel_vent_channels":
        if "vent_channel_1_height" in key or "vent_channel_2_height" in key or "vent_channel_2_count" in key:
            used = "Строка `Кладка вентканалов Schiedel`."
            formula = "`schiedel_masonry_total_length_m = vent_channel_1_height_m + vent_channel_2_height_m * vent_channel_2_count`."
            example = f"{current} ЮСВ: 6.2 + 4.81 * 2 = 15.82 м."
            check = "Подтвердить высоты и количество вентканалов по разрезу."
            why = "Из этих параметров получается общая длина кладки."
        elif "schiedel_masonry_total_length" in key:
            used = "Строка `Кладка вентканалов Schiedel`."
            formula = "`work_total = schiedel_masonry_total_length_m * schiedel_masonry_work_rate_per_m`; значение может вычисляться из высот каналов."
            example = f"{current} ЮСВ: 15.82 * 5 000 = 79 100."
            check = "Не заполнять вручную, если подтверждены высоты каналов."
            why = "Это производное количество работ."
        elif "schiedel_delivery_trips" in key:
            used = "Строка `Доставка вентканалов`."
            formula = "Материалы/техника = `trips * truck_price`; работа = `trips * delivery_work_price`."
            example = f"{current} ЮСВ: 1 рейс * 15 000 + 1 * 2 500 = 17 500."
            check = "Подтвердить количество доставок Schiedel."
            why = "Это логистика, обычно задается сметчиком."
        elif "consumables_rate" in key:
            used = "Строка `Расходные материалы, амортизация инструмента`."
            formula = "`consumables_total = direct_cost_base_before_consumables * consumables_rate`."
            example = f"{current} ЮСВ: 114 456 * 0.03 = 3 433.68 -> 3 434."
            check = "Подтвердить процент расходников как системное правило."
            why = "Это процент от базы, не проектный параметр."

    if status == "PRICE_DATABASE" and used == "Связь со строками сметы требует уточнения.":
        used = "Строка сметы, где применяется соответствующая цена/ставка."
        formula = "`line_total = quantity * unit_price`; unit_price должен приходить из price_registry или project override."
        example = current
        check = "Проверить наличие price_code и актуальной цены в прайсе."
        why = "Цены и ставки не должны быть ручными параметрами таблицы проверки проекта."
    elif status == "DEFAULT_VALUE" and used == "Связь со строками сметы требует уточнения.":
        used = "Системное правило/каталог материала в соответствующем расчётном блоке."
        formula = "Используется как коэффициент, размер упаковки или технологическая норма в формуле закупочного количества."
        example = current
        check = "Подтвердить источник default: каталог, норматив или правило Елены."
        why = "Default можно скрыть от Елены только после подтверждения источника и области применения."
    elif status == "AUTO_CALCULATED" and used == "Связь со строками сметы требует уточнения.":
        used = "Производное поле внутри расчетного блока калькулятора."
        formula = clean(row.get("recommended_formula")) or "Вычисляется из других входных параметров."
        example = current
        check = "Проверить, какие исходные параметры нужны для автоматического расчета."
        why = "Производные поля не должны заполняться вручную."
    elif status == "AUTO_PROJECT" and used == "Связь со строками сметы требует уточнения.":
        used = "Строки сметы соответствующего раздела, где количество берется из проекта."
        formula = "Проектное значение используется как количество или участвует в геометрической формуле калькулятора."
        example = current
        check = "Подтвердить значение по проекту/спецификации."
        why = "Это проектный параметр, его нельзя заменить default без источника."

    return {
        "used_in_estimate_lines": used,
        "calculator_formula": formula,
        "usv_formula_example": example,
        "what_elena_should_check": check,
        "why_needed": why,
    }


def explanation_cells(row: dict[str, Any]) -> list[str]:
    ctx = context_pack(row)
    return [ctx[key] for key in EXPLANATION_KEYS]


def display_label(row: dict[str, Any]) -> str:
    key = clean(row.get("calculator_input_key"))
    label = clean(row.get("label"))
    if key in LABEL_OVERRIDES:
        return LABEL_OVERRIDES[key]

    match = re.fullmatch(r"non_insulated_edge_lengths_m\[(\d+)\]", key)
    if match:
        return f"Участок без утепления {int(match.group(1)) + 1}: длина"

    match = re.fullmatch(r"cutoff_waterproofing_wall_400_lengths_m\[(\d+)\]", key)
    if match:
        return f"Отсечная гидроизоляция стены 400 мм, участок {int(match.group(1)) + 1}: длина"

    match = re.fullmatch(r"cutoff_waterproofing_wall_250_lengths_m\[(\d+)\]", key)
    if match:
        return f"Отсечная гидроизоляция стены 250 мм, участок {int(match.group(1)) + 1}: длина"

    replacements = {
        "non insulated edge lengths m": "длина",
        "cutoff waterproofing wall 400 lengths m": "длина",
        "cutoff waterproofing wall 250 lengths m": "длина",
        "pit area m2": "Площадь котлована",
        "manual refinement depth m": "Глубина ручной доработки котлована",
        "trench volume m3": "Объем траншей",
        "sand truck step m3": "Шаг заказа песка машиной",
        "communications length m": "Длина коммуникаций",
        "geotextile laying area m2": "Площадь укладки геотекстиля",
        "slab formwork perimeter m": "Периметр бортов/опалубки плиты",
        "slab edge height m": "Высота борта плиты",
        "thermal insert length m": "Длина термовставок",
        "thermal insert piece length m": "Термовставка: длина элемента",
        "thermal insert piece width m": "Термовставка: ширина элемента",
        "thermal insert piece height m": "Термовставка: высота элемента",
        "main wall external length m": "Длина наружных несущих стен",
        "main wall reinforcement rows": "Количество рядов армирования кладки",
        "main wall 400 reinforcement threads": "Количество ниток армирования стены 400 мм",
        "main wall 250 reinforcement threads": "Количество ниток армирования стены 250 мм",
        "lintel section width m": "Перемычки: ширина сечения",
        "lintel section height m": "Перемычки: высота сечения",
        "parapet masonry volume m3": "Объем кладки парапета",
        "second light masonry volume m3": "Объем кладки зоны второго света",
        "total concrete volume from spec m3": "Объем бетона по спецификации",
        "total eps volume from spec m3": "Объем ЭППС по спецификации",
        "slab length m": "Длина плиты",
        "slab width m": "Ширина плиты",
        "slab area m2": "Площадь плиты",
        "parapet length level 1 m": "Длина парапета уровня 1",
        "parapet length level 2 m": "Длина парапета уровня 2",
        "vent wall abutment level 1 m": "Длина примыканий к вентшахтам/стенам уровня 1",
        "vent wall abutment level 2 m": "Длина примыканий к вентшахтам/стенам уровня 2",
        "eps50 supplier required volume m3": "Объем ЭППС 50 мм по раскладке поставщика",
        "slope plate a supplier required volume m3": "Объем уклонных плит A по раскладке поставщика",
        "slope plate b supplier required volume m3": "Объем уклонных плит B по раскладке поставщика",
        "slope plate j supplier required volume m3": "Объем уклонных плит J по раскладке поставщика",
        "slope plate k supplier required volume m3": "Объем уклонных плит K по раскладке поставщика",
        "vent channel 1 height m": "Высота вентканала 1",
        "vent channel 2 height m": "Высота вентканала 2",
        "vent channel 2 count": "Количество вентканалов типа 2",
        "schiedel masonry total length m": "Общая длина кладки вентканалов Schiedel",
    }
    for old, new in replacements.items():
        label = label.replace(old, new)
    return label


def bool_text(value: Any) -> str:
    if isinstance(value, bool):
        return "да" if value else "нет"
    text = clean(value).lower()
    if text in {"true", "1", "yes", "да"}:
        return "да"
    if text in {"false", "0", "no", "нет"}:
        return "нет"
    return clean(value)


def load_audit_rows() -> list[dict[str, Any]]:
    wb = load_workbook(AUDIT_XLSX, read_only=True, data_only=True)
    ws = wb["all_parameters_audit"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(v) for v in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:]]


def infer_project_source(row: dict[str, Any]) -> str:
    section = clean(row.get("section_code"))
    key = clean(row.get("calculator_input_key")).lower()
    label = clean(row.get("label")).lower()
    text = f"{key} {label}"
    if section in {"earthworks", "foundation_slab", "waterproofing"}:
        return "КР-1 / план / разрез / спецификация"
    if section in {"load_bearing_walls_lintels", "floor_slab_1", "floor_slab_2", "flat_roof", "schiedel_vent_channels"}:
        if "roof" in text or "кров" in text:
            return "КР-2 / план кровли / спецификация / узлы"
        if "vent" in text or "schiedel" in text or "вент" in text:
            return "КР-2 / разрез по вентканалам / спецификация"
        return "КР-2 / план / спецификация / разрез"
    return "неизвестно"


def infer_project_place(row: dict[str, Any]) -> str:
    key = clean(row.get("calculator_input_key")).lower()
    label = clean(row.get("label")).lower()
    text = f"{key} {label}"
    if any(token in text for token in ("rebar", "арматур", "weight", "вес", "concrete", "бетон", "eps", "утепл")):
        return "спецификация"
    if any(token in text for token in ("area", "площад", "perimeter", "периметр", "length", "длина")):
        return "план / схема"
    if any(token in text for token in ("height", "высота", "width", "ширина", "thickness", "толщина")):
        return "разрез / узел"
    if any(token in text for token in ("count", "количество")):
        return "спецификация / план"
    return "проект / уточнить источник"


def possible_price_code(row: dict[str, Any]) -> str:
    key = clean(row.get("calculator_input_key"))
    source = clean(row.get("recommended_registry_source"))
    if "price_code:" in source:
        return source.split("price_code:", 1)[1].strip()
    # The audit does not always know exact price_code. Keep blank for Elena/team decision.
    if any(token in key.lower() for token in ("concrete_pump", "crane", "rate", "price", "delivery", "technical_supervision")):
        return ""
    return ""


def row_reason(row: dict[str, Any]) -> str:
    return clean(row.get("reason_for_hiding")) or clean(row.get("comment")) or clean(row.get("source_evidence"))


def is_sporny(row: dict[str, Any]) -> bool:
    status = clean(row.get("recommended_source_status"))
    safe = clean(row.get("safe_for_new_projects"))
    risk = clean(row.get("risk_level"))
    if status in {"MANUAL_REQUIRED", "REQUIRES_VALIDATION"}:
        return True
    return safe == "requires_validation" or risk in {"medium", "high"}


def append_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = 0
        for cell in column_cells:
            max_len = max(max_len, min(len(clean(cell.value)), 70))
        ws.column_dimensions[letter].width = max(12, min(max_len + 2, 55))


def build_pack() -> None:
    global TEMPLATE_INPUTS
    TEMPLATE_INPUTS = load_template_inputs()
    audit_rows = load_audit_rows()
    wb = Workbook()
    wb.remove(wb.active)

    discussion = [
        row for row in audit_rows
        if clean(row.get("recommended_source_status")) in DISCUSSION_STATUSES and is_sporny(row)
    ]
    append_sheet(
        wb,
        "К_обсуждению_с_Еленой",
        [
            "section_name",
            "label",
            "calculator_input_key",
            *EXPLANATION_HEADERS,
            "current_status",
            "current_input_type",
            "recommended_source_status",
            "reason",
            "elena_decision",
            "source_of_truth",
            "can_override",
            "visible_to_elena",
            "risk_level",
            "comment",
        ],
        [
            [
                row.get("section_name"),
                display_label(row),
                row.get("calculator_input_key"),
                *explanation_cells(row),
                row.get("current_status"),
                row.get("current_input_type"),
                row.get("recommended_source_status"),
                row_reason(row),
                "",
                row.get("source_of_truth"),
                "да",
                bool_text(row.get("should_be_visible_to_elena")),
                row.get("risk_level"),
                "",
            ]
            for row in discussion
        ],
    )

    append_sheet(
        wb,
        "Ручные_решения_сметчика",
        [
            "section_name",
            "label",
            "calculator_input_key",
            *EXPLANATION_HEADERS,
            "current_status",
            "reason",
            "elena_decision",
            "source_of_truth",
            "risk_level",
            "comment",
        ],
        [
            [
                row.get("section_name"),
                display_label(row),
                row.get("calculator_input_key"),
                *explanation_cells(row),
                row.get("current_status"),
                row.get("comment"),
                "",
                row.get("source_of_truth"),
                row.get("risk_level"),
                "",
            ]
            for row in audit_rows
            if clean(row.get("recommended_source_status")) == "MANUAL_REQUIRED"
        ],
    )

    append_sheet(
        wb,
        "Подтвердить_defaults",
        [
            "section_name",
            "label",
            "calculator_input_key",
            *EXPLANATION_HEADERS,
            "proposed_default_value",
            "unit",
            "source",
            "applies_to",
            "can_override",
            "risk_level",
            "elena_approved",
            "comment",
        ],
        [
            [
                row.get("section_name"),
                display_label(row),
                row.get("calculator_input_key"),
                *explanation_cells(row),
                row.get("recommended_default_value"),
                row.get("unit"),
                row.get("source_of_truth"),
                "уточнить: все проекты / только текущий тип работ / только ЮСВ",
                "да",
                row.get("risk_level"),
                "",
                row.get("comment"),
            ]
            for row in audit_rows
            if clean(row.get("recommended_source_status")) == "DEFAULT_VALUE"
        ],
    )

    append_sheet(
        wb,
        "Данные_из_проекта",
        [
            "section_name",
            "label",
            "calculator_input_key",
            *EXPLANATION_HEADERS,
            "где искать в проекте",
            "КР-1 / КР-2 / АР / спецификация / план / разрез / неизвестно",
            "comment",
        ],
        [
            [
                row.get("section_name"),
                display_label(row),
                row.get("calculator_input_key"),
                *explanation_cells(row),
                infer_project_place(row),
                infer_project_source(row),
                row.get("comment"),
            ]
            for row in audit_rows
            if clean(row.get("recommended_source_status")) == "AUTO_PROJECT"
        ],
    )

    append_sheet(
        wb,
        "Авторасчет",
        [
            "section_name",
            "label",
            "calculator_input_key",
            *EXPLANATION_HEADERS,
            "formula_or_dependency",
            "missing_dependencies",
            "safe_for_new_projects",
            "comment",
        ],
        [
            [
                row.get("section_name"),
                display_label(row),
                row.get("calculator_input_key"),
                *explanation_cells(row),
                row.get("recommended_formula"),
                "уточнить при внедрении derived_parameters.py",
                row.get("safe_for_new_projects"),
                row.get("comment"),
            ]
            for row in audit_rows
            if clean(row.get("recommended_source_status")) == "AUTO_CALCULATED"
        ],
    )

    append_sheet(
        wb,
        "Прайс_и_ставки",
        [
            "section_name",
            "label",
            "calculator_input_key",
            *EXPLANATION_HEADERS,
            "price_code",
            "registry_source",
            "needs_price_registry_row",
            "comment",
        ],
        [
            [
                row.get("section_name"),
                display_label(row),
                row.get("calculator_input_key"),
                *explanation_cells(row),
                possible_price_code(row),
                row.get("recommended_registry_source") or "price_registry / project_price_overrides",
                "да" if not possible_price_code(row) else "проверить",
                row.get("comment"),
            ]
            for row in audit_rows
            if clean(row.get("recommended_source_status")) == "PRICE_DATABASE"
        ],
    )

    raw_rows = []
    for row in audit_rows:
        raw_row = dict(row)
        raw_row["label"] = display_label(row)
        raw_row.update(context_pack(row))
        raw_rows.append(raw_row)
    raw_headers = list(raw_rows[0].keys()) if raw_rows else []
    append_sheet(
        wb,
        "Все_параметры_raw",
        raw_headers,
        [[row.get(header) for header in raw_headers] for row in raw_rows],
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(PACK_XLSX)
    AGENDA_MD.write_text(build_agenda(audit_rows), encoding="utf-8")


def build_agenda(audit_rows: list[dict[str, Any]]) -> str:
    counts: dict[str, int] = {}
    for row in audit_rows:
        status = clean(row.get("recommended_source_status"))
        counts[status] = counts.get(status, 0) + 1
    manual = counts.get("MANUAL_REQUIRED", 0)
    total = len(audit_rows)
    can_move = total - manual
    return "\n".join([
        "# Agenda: review missing/manual параметров с Еленой",
        "",
        "## Цель созвона",
        "Подтвердить, какие параметры действительно должен заполнять сметчик, а какие должны уходить в проектное извлечение, авторасчет, defaults/material catalog или price_registry.",
        "",
        "## Главное",
        f"- Сейчас в аудите {total} missing/manual параметров.",
        f"- Не нужно вручную заполнять все {total} строк.",
        f"- Предварительно реально ручными оставлены {manual} параметра.",
        f"- Остальные {can_move} параметров нужно подтвердить как системные, проектные, расчетные или прайсовые.",
        "",
        "## Что обсуждаем по листам",
        "- `К_обсуждению_с_Еленой`: главный лист созвона, туда вносим решение.",
        "- `Ручные_решения_сметчика`: что, вероятно, остается ручным вводом.",
        "- `Подтвердить_defaults`: коэффициенты, упаковки и стандартные настройки, по которым нужен источник истины.",
        "- `Данные_из_проекта`: параметры, которые должны искаться в КР/АР/спецификациях.",
        "- `Авторасчет`: параметры, которые лучше вычислять из других входных данных.",
        "- `Прайс_и_ставки`: цены и ставки, которые должны быть в price_registry или project overrides.",
        "- `Все_параметры_raw`: полный технический список аудита.",
        "",
        "## Как читать строки",
        "В основных листах добавлены поясняющие колонки:",
        "- `Где используется в смете`: какая строка или блок сметы зависит от параметра.",
        "- `Формула калькулятора`: как параметр входит в расчет.",
        "- `Пример формулы ЮСВ`: как это выглядит на текущем кейсе ЮСВ.",
        "- `Что проверить Елене`: конкретное действие на созвоне.",
        "- `Зачем нужен параметр`: почему строка вообще есть в таблице.",
        "",
        "## Что НЕ делаем на созвоне",
        "- Не внедряем изменения в `section_schema.py`.",
        "- Не меняем калькуляторы и формулы.",
        "- Не меняем `expected.json`.",
        "- Не считаем demo_with_template_fallback production-расчетом.",
        "",
        "## Ожидаемый результат",
        "После созвона появится согласованное решение: какие параметры показывать Елене в reviewed_parameters.xlsx, а какие убрать в системные слои после отдельной задачи внедрения.",
        "",
    ])


if __name__ == "__main__":
    build_pack()
    print(f"Created: {PACK_XLSX}")
    print(f"Created: {AGENDA_MD}")
