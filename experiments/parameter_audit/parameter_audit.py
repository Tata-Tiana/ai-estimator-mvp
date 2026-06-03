"""Audit missing/manual parameters from the PDF parser pipeline.

This experiment does not change calculator inputs or schemas. It classifies
the current missing/manual review-sheet rows so the team can decide what should
stay visible to Elena and what should move to catalogs, pricing, defaults, or
derived-parameter logic.
"""

from __future__ import annotations

import csv
import json
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]

REVIEWED_PARAMETERS_PATH = REPO_ROOT / "experiments/pdf_parser_pipeline/output/mvp_usv_demo/reviewed_parameters.xlsx"
ELENA_MISSING_REPORT_PATH = REPO_ROOT / "experiments/pdf_parser_pipeline/output/mvp_usv_demo/elena_missing_parameters_by_section.md"
PRICE_REGISTRY_PATH = REPO_ROOT / "output/price_registry_filled_v3.xlsx"
LIVE_PRICING_REPORT_PATH = REPO_ROOT / "experiments/pricing/output/live_pricing_sections_report.md"
REQUIRED_PRICE_CODES_PATH = REPO_ROOT / "output/required_price_codes_v2.csv"


SOURCE_STATUSES = (
    "AUTO_PROJECT",
    "AUTO_CALCULATED",
    "DEFAULT_VALUE",
    "PRICE_DATABASE",
    "MANUAL_REQUIRED",
    "REQUIRES_VALIDATION",
)


@dataclass
class AuditRow:
    project: str
    section_code: str
    section_name: str
    parameter_code: str
    calculator_input_key: str
    label: str
    current_status: str
    current_input_type: str
    required: bool
    unit: str
    recommended_source_status: str
    should_be_visible_to_elena: bool
    recommended_default_value: str
    recommended_formula: str
    recommended_registry_source: str
    comment: str
    safe_for_new_projects: str
    source_of_truth: str
    source_evidence: str
    risk_level: str
    can_be_hidden_from_elena: bool
    reason_for_hiding: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "project": self.project,
            "section_code": self.section_code,
            "section_name": self.section_name,
            "parameter_code": self.parameter_code,
            "calculator_input_key": self.calculator_input_key,
            "label": self.label,
            "current_status": self.current_status,
            "current_input_type": self.current_input_type,
            "required": self.required,
            "unit": self.unit,
            "recommended_source_status": self.recommended_source_status,
            "should_be_visible_to_elena": self.should_be_visible_to_elena,
            "recommended_default_value": self.recommended_default_value,
            "recommended_formula": self.recommended_formula,
            "recommended_registry_source": self.recommended_registry_source,
            "comment": self.comment,
            "safe_for_new_projects": self.safe_for_new_projects,
            "source_of_truth": self.source_of_truth,
            "source_evidence": self.source_evidence,
            "risk_level": self.risk_level,
            "can_be_hidden_from_elena": self.can_be_hidden_from_elena,
            "reason_for_hiding": self.reason_for_hiding,
        }


def to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in {"true", "1", "yes", "да"}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(Decimal(str(value)).normalize())
    return str(value).strip()


def load_parameters(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["parameters"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean_text(v) for v in rows[0]]
    result: list[dict[str, Any]] = []
    for row in rows[1:]:
        item = dict(zip(headers, row))
        result.append(item)
    return result


def load_price_codes() -> set[str]:
    codes: set[str] = set()
    if REQUIRED_PRICE_CODES_PATH.exists():
        with REQUIRED_PRICE_CODES_PATH.open("r", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                code = clean_text(row.get("price_code"))
                if code:
                    codes.add(code)

    if PRICE_REGISTRY_PATH.exists():
        wb = load_workbook(PRICE_REGISTRY_PATH, read_only=True, data_only=True)
        for sheet_name in ("price_registry", "rows_to_add"):
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [clean_text(v) for v in rows[0]]
            try:
                idx = headers.index("price_code")
            except ValueError:
                continue
            for row in rows[1:]:
                code = clean_text(row[idx] if idx < len(row) else "")
                if code:
                    codes.add(code)
    return codes


def is_missing_or_manual(row: dict[str, Any]) -> bool:
    status = clean_text(row.get("elena_status")).lower()
    required = to_bool(row.get("required"))
    use_for_calculation = to_bool(row.get("use_for_calculation"))
    final_value = clean_text(row.get("final_value"))
    extracted_value = clean_text(row.get("extracted_value"))
    missing_reason = clean_text(row.get("missing_reason"))
    if status in {"missing", "manual_required"}:
        return True
    return required and use_for_calculation and not final_value and not extracted_value and bool(missing_reason)


def leaf_key(key: str) -> str:
    return re.sub(r"\[\d+\]", "", key.split(".")[-1]).lower()


def normalized_key(row: dict[str, Any]) -> str:
    key = clean_text(row.get("calculator_input_key"))
    parameter = clean_text(row.get("parameter_code"))
    label = clean_text(row.get("label"))
    return f"{key} {parameter} {label}".lower().replace("ё", "е")


def looks_like_price(text: str, input_type: str) -> bool:
    if input_type == "price":
        return True
    price_tokens = (
        "unit_price",
        "price",
        "rate_per",
        "work_rate",
        "installation_rate",
        "drilling_rate",
        "shift_rate",
        "truck_price",
        "unit_price_per",
        "material_unit_price",
        "work_unit_price",
        "стоимость",
        "ставка",
        "цена",
    )
    if any(token in text for token in price_tokens):
        percent_or_coeff = any(token in text for token in ("percent", "процент", "coeff", "коэффициент", "logistics_rate", "consumables_rate"))
        return not percent_or_coeff
    return False


def looks_like_default(text: str, input_type: str) -> bool:
    if input_type == "calculation_constant":
        return True
    tokens = (
        "coeff",
        "коэффициент",
        "waste",
        "overlap",
        "pack_volume",
        "roll_area",
        "roll_width",
        "roll_length",
        "sheet_width",
        "sheet_height",
        "sheet_working_area",
        "rod_length",
        "kg_per_meter",
        "coverage",
        "capacity",
        "round_step",
        "step",
        "min_order",
        "min_cans",
        "bag_weight",
        "bucket_weight",
        "consumption",
        "thickness",
        "layers",
        "percent",
        "процент",
        "площадь одного рулона",
        "объем одной упаковки",
        "длина одного хлыста",
        "вес 1 погонного метра",
        "per_membrane_roll",
    )
    return any(token in text for token in tokens)


def looks_like_calculated(text: str) -> bool:
    if any(token in text for token in ("from_spec", "specification", "source_weight", "supplier_required")):
        return False
    tokens = (
        "raw",
        "display",
        "total",
        "subtotal",
        "ordered",
        "order_",
        "packs",
        "rolls",
        "rods",
        "with_waste",
        "base_",
        "control",
        "supplier_rate",
        "average_rate",
        "equivalent",
        "density",
        "расчетное значение",
        "отображаемое значение",
        "общий вес металла",
    )
    return any(token in text for token in tokens)


def looks_like_supplier_or_manual_decision(text: str) -> bool:
    tokens = (
        "supplier_quote",
        "supplier required",
        "supplier_required",
        "manual",
        "override",
        "enabled",
        "strategy",
        "method",
        "case_specific",
        "shifts",
        "trips",
        "поставщик",
        "ручн",
        "смен",
        "рейс",
        "метод",
        "учитывать",
    )
    return any(token in text for token in tokens)


def looks_like_project(text: str, input_type: str) -> bool:
    if input_type == "parsed":
        return True
    project_tokens = (
        "area",
        "volume",
        "length",
        "perimeter",
        "height",
        "width",
        "weight",
        "count",
        "roof",
        "wall",
        "slab",
        "beam",
        "lintel",
        "rebar",
        "concrete",
        "sand",
        "geotextile",
        "membrane",
        "vent_channel",
        "площад",
        "объем",
        "длина",
        "периметр",
        "высота",
        "ширина",
        "вес",
        "количество",
        "арматура",
        "бетон",
        "плита",
        "балка",
        "перемычка",
        "парапет",
        "вентканал",
    )
    return any(token in text for token in project_tokens)


def looks_like_manual(text: str) -> bool:
    tokens = (
        "override",
        "manual",
        "enabled",
        "strategy",
        "method",
        "case_specific",
        "supplier_required",
        "supplier_quote",
        "shifts",
        "trips",
        "поставщик",
        "ручн",
        "смен",
        "рейс",
        "метод",
        "учитывать",
    )
    return any(token in text for token in tokens)


def formula_for_calculated(key: str, text: str) -> str:
    if "display" in text:
        return "display value = rounded raw value"
    if "total" in text or "subtotal" in text:
        return "total = sum of source line totals / calculated blocks"
    if "packs" in text or "rolls" in text or "rods" in text or "ordered" in text:
        return "ordered quantity = ceil(required quantity / package size) * package size"
    if "density" in text:
        return "control metric = weight / volume"
    if "area" in text and "volume" in text:
        return "area = volume / thickness"
    return "Derived from other calculator parameters; formula should be encoded in derived_parameters.py"


def classify(row: dict[str, Any], price_codes: set[str]) -> AuditRow:
    section = clean_text(row.get("section_code"))
    input_type = clean_text(row.get("input_type"))
    status = clean_text(row.get("elena_status"))
    key = clean_text(row.get("calculator_input_key"))
    label = clean_text(row.get("label"))
    text = normalized_key(row)
    default_value = clean_text(row.get("final_value")) or clean_text(row.get("extracted_value"))
    required = to_bool(row.get("required"))

    recommended = "MANUAL_REQUIRED"
    visible = True
    default = ""
    formula = ""
    registry_source = ""
    comment = ""
    safe = "requires_validation"
    source_of_truth = "reviewed_parameters.xlsx / estimator decision"
    evidence = "Current parameter is missing/manual in reviewed_parameters.xlsx"
    risk = "medium"
    can_hide = False
    hide_reason = ""

    if looks_like_price(text, input_type):
        recommended = "PRICE_DATABASE"
        visible = False
        registry_source = "price_registry / project_price_overrides by price_code"
        comment = "Цена/ставка должна приходить из price_registry, а не из таблицы ручной проверки."
        safe = "requires_validation"
        source_of_truth = "output/price_registry_filled_v3.xlsx"
        evidence = "Matches price/rate naming pattern; live-pricing layer already supports price_registry_with_fallback."
        risk = "medium"
        can_hide = True
        hide_reason = "Прайсовая позиция, не проектный параметр."
    elif any(token in text for token in ("technical_supervision_work_total", "технического надзора")):
        recommended = "PRICE_DATABASE"
        visible = False
        registry_source = "price_registry / project_price_overrides by price_code"
        comment = "Фиксированная услуга/ставка должна задаваться прайсом или override проекта."
        safe = "requires_validation"
        source_of_truth = "price_registry / project override"
        evidence = "Technical supervision is a fixed/manual service amount, not a project PDF quantity."
        risk = "medium"
        can_hide = True
        hide_reason = "Фиксированная сметная ставка, не проектный параметр."
    elif looks_like_default(text, input_type):
        recommended = "DEFAULT_VALUE"
        visible = False
        default = default_value
        comment = "Технологическая/каталожная константа. Не считать универсальной без подтверждения источника."
        safe = "requires_validation"
        source_of_truth = "future defaults_registry.py / material_catalog.py"
        evidence = "Matches coefficient/package/catalog naming pattern in calculator input."
        risk = "medium"
        can_hide = True
        hide_reason = "Системная настройка или каталог материала, не вопрос к проектировщику."
    elif looks_like_supplier_or_manual_decision(text):
        recommended = "MANUAL_REQUIRED"
        visible = True
        comment = "Похоже на решение сметчика, поставщика или case-specific параметр; оставить видимым до подтвержденного правила."
        safe = "no"
        source_of_truth = "Elena/manual project decision"
        evidence = "Matches supplier/manual/override/shifts/trips/case-specific naming pattern."
        risk = "high"
    elif looks_like_calculated(text):
        recommended = "AUTO_CALCULATED"
        visible = False
        formula = formula_for_calculated(key, text)
        comment = "Поле должно вычисляться из других параметров, raw/display или totals не нужно заполнять вручную."
        safe = "yes"
        source_of_truth = "derived_parameters.py / calculator formulas"
        evidence = "Matches raw/display/total/order/control derived naming pattern."
        risk = "low"
        can_hide = True
        hide_reason = "Вычисляемое поле."
    elif re.search(r"\.(code|name|steel_class|diameter_mm)$|rebar_items\[\d+\]\.(code|name|steel_class|diameter_mm)", key):
        recommended = "AUTO_CALCULATED"
        visible = False
        formula = "Build rebar item metadata from steel class and diameter catalog/specification mapping."
        comment = "Метаданные арматуры не должны быть ручными строками для Елены."
        safe = "requires_validation"
        source_of_truth = "future rebar catalog / specification mapping"
        evidence = "Special audit rule: do not show rebar_items[].code/name/diameter/steel_class as manual input."
        risk = "medium"
        can_hide = True
        hide_reason = "Каталожное описание позиции арматуры."
    elif looks_like_manual(text):
        recommended = "MANUAL_REQUIRED"
        visible = True
        comment = "Похоже на решение сметчика или case-specific параметр; оставить видимым до подтвержденного правила."
        safe = "no"
        source_of_truth = "Elena/manual project decision"
        evidence = "Matches manual/override/shifts/trips/supplier/case-specific naming pattern."
        risk = "high"
    elif looks_like_project(text, input_type):
        recommended = "AUTO_PROJECT"
        visible = True
        comment = "Проектный параметр: желательно извлекать из PDF/specification или задавать Елене, если parser не нашел."
        safe = "no"
        source_of_truth = "Project PDF/specification/review card"
        evidence = "Matches area/volume/length/weight/count/project quantity naming pattern."
        risk = "high"
    else:
        recommended = "REQUIRES_VALIDATION"
        visible = True
        comment = "Недостаточно уверенности для автоматического скрытия. Нужна ручная классификация."
        safe = "requires_validation"
        source_of_truth = "TBD"
        evidence = "No strong audit rule matched."
        risk = "high"

    if recommended == "PRICE_DATABASE":
        possible_code = re.sub(r"[^a-z0-9]+", "_", key.lower()).strip("_")
        if possible_code in price_codes:
            registry_source = f"price_code: {possible_code}"
            evidence += " Matching required/registry price_code exists."
        else:
            evidence += " Exact price_code is not inferred automatically; mapping is required."

    return AuditRow(
        project=clean_text(row.get("project")),
        section_code=section,
        section_name=clean_text(row.get("section_name")),
        parameter_code=clean_text(row.get("parameter_code")),
        calculator_input_key=key,
        label=label,
        current_status=status,
        current_input_type=input_type,
        required=required,
        unit=clean_text(row.get("unit")),
        recommended_source_status=recommended,
        should_be_visible_to_elena=visible,
        recommended_default_value=default,
        recommended_formula=formula,
        recommended_registry_source=registry_source,
        comment=comment,
        safe_for_new_projects=safe,
        source_of_truth=source_of_truth,
        source_evidence=evidence,
        risk_level=risk,
        can_be_hidden_from_elena=can_hide,
        reason_for_hiding=hide_reason,
    )


def autosize(ws) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = 0
        for cell in column_cells:
            value = clean_text(cell.value)
            max_len = max(max_len, min(len(value), 80))
        ws.column_dimensions[letter].width = max(12, min(max_len + 2, 60))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_table(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["empty"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    autosize(ws)


def write_excel(path: Path, audit_rows: list[AuditRow]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    rows = [row.as_dict() for row in audit_rows]
    write_table(wb, "all_parameters_audit", rows)

    by_section: list[dict[str, Any]] = []
    by_section_counter: dict[str, Counter[str]] = defaultdict(Counter)
    section_names: dict[str, str] = {}
    for row in audit_rows:
        by_section_counter[row.section_code][row.recommended_source_status] += 1
        section_names[row.section_code] = row.section_name
    for section_code in sorted(by_section_counter):
        counter = by_section_counter[section_code]
        item = {
            "section_code": section_code,
            "section_name": section_names.get(section_code, ""),
            "total": sum(counter.values()),
        }
        for status in SOURCE_STATUSES:
            item[status] = counter.get(status, 0)
        by_section.append(item)
    write_table(wb, "by_section_summary", by_section)

    sheet_map = {
        "DEFAULT_VALUE": "to_default_values",
        "PRICE_DATABASE": "to_price_database",
        "AUTO_CALCULATED": "to_auto_calculated",
        "AUTO_PROJECT": "to_auto_project",
        "MANUAL_REQUIRED": "still_manual_required",
    }
    for status, sheet in sheet_map.items():
        write_table(wb, sheet, [row.as_dict() for row in audit_rows if row.recommended_source_status == status])

    schema_changes = []
    for row in audit_rows:
        new_input_type = {
            "PRICE_DATABASE": "price",
            "DEFAULT_VALUE": "default",
            "AUTO_CALCULATED": "calculated",
            "AUTO_PROJECT": "parsed",
            "MANUAL_REQUIRED": "manual",
            "REQUIRES_VALIDATION": row.current_input_type,
        }[row.recommended_source_status]
        schema_changes.append({
            "section_code": row.section_code,
            "calculator_input_key": row.calculator_input_key,
            "old_input_type": row.current_input_type,
            "new_input_type": new_input_type,
            "visible_to_elena": row.should_be_visible_to_elena,
            "default_formula_or_source": row.recommended_default_value or row.recommended_formula or row.recommended_registry_source or row.source_of_truth,
            "safe_for_new_projects": row.safe_for_new_projects,
            "risk_level": row.risk_level,
            "comment": row.comment,
        })
    write_table(wb, "schema_changes", schema_changes)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    if not rows:
        return "_Нет строк._\n"
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(clean_text(v).replace("\n", " ") for v in row) + " |")
    return "\n".join(lines) + "\n"


def write_report(path: Path, audit_rows: list[AuditRow]) -> None:
    total = len(audit_rows)
    old_status = Counter(row.current_status for row in audit_rows)
    by_status = Counter(row.recommended_source_status for row in audit_rows)
    by_section: dict[str, Counter[str]] = defaultdict(Counter)
    section_names: dict[str, str] = {}
    for row in audit_rows:
        by_section[row.section_code][row.recommended_source_status] += 1
        section_names[row.section_code] = row.section_name

    remove_rows = [row for row in audit_rows if row.can_be_hidden_from_elena]
    manual_rows = [row for row in audit_rows if row.recommended_source_status == "MANUAL_REQUIRED"]
    reduction = (len(remove_rows) / total * 100) if total else 0

    lines = [
        "# Parameter Audit Report",
        "",
        "Аудит классифицирует текущие missing/manual параметры. Он не меняет калькуляторы, expected.json или section_schema.py.",
        "",
        "## Было",
        f"- Всего missing/manual: {total}",
        f"- manual_required: {old_status.get('manual_required', 0)}",
        f"- missing: {old_status.get('missing', 0)}",
        "",
        md_table(
            ["Раздел", "Всего", "missing", "manual_required"],
            [
                [
                    section_names.get(section, section),
                    sum(counter.values()),
                    sum(1 for row in audit_rows if row.section_code == section and row.current_status == "missing"),
                    sum(1 for row in audit_rows if row.section_code == section and row.current_status == "manual_required"),
                ]
                for section, counter in sorted(by_section.items())
            ],
        ),
        "## Рекомендованная классификация",
        md_table(
            ["source_status", "count", "explanation"],
            [
                ["AUTO_PROJECT", by_status.get("AUTO_PROJECT", 0), "Проектные площади, объемы, длины, веса, количества и спецификации."],
                ["AUTO_CALCULATED", by_status.get("AUTO_CALCULATED", 0), "Raw/display, totals, закупочные количества и производные значения."],
                ["DEFAULT_VALUE", by_status.get("DEFAULT_VALUE", 0), "Коэффициенты, упаковки, размеры стандартных материалов, технологические настройки."],
                ["PRICE_DATABASE", by_status.get("PRICE_DATABASE", 0), "Цены материалов, работ, техники и доставок из price_registry."],
                ["MANUAL_REQUIRED", by_status.get("MANUAL_REQUIRED", 0), "Реальные решения сметчика по объекту или пока неподтвержденные ручные параметры."],
                ["REQUIRES_VALIDATION", by_status.get("REQUIRES_VALIDATION", 0), "Нужна отдельная проверка, правило не уверенное."],
            ],
        ),
        "## Что уйдет из ручного ввода",
        md_table(
            ["section", "parameter", "old_status", "new_status", "reason"],
            [
                [row.section_code, row.label, row.current_status, row.recommended_source_status, row.reason_for_hiding]
                for row in remove_rows[:120]
            ],
        ),
    ]
    if len(remove_rows) > 120:
        lines.append(f"_Показаны первые 120 строк из {len(remove_rows)}. Полный список в Excel._\n")

    lines.extend([
        "## Что останется ручным",
        md_table(
            ["section", "parameter", "why_manual"],
            [[row.section_code, row.label, row.comment] for row in manual_rows[:120]],
        ),
    ])
    if len(manual_rows) > 120:
        lines.append(f"_Показаны первые 120 строк из {len(manual_rows)}. Полный список в Excel._\n")

    lines.extend([
        "## Рекомендации по изменению section_schema.py",
        md_table(
            ["calculator_input_key", "old_input_type", "new_input_type", "visible_to_elena", "default/formula/source"],
            [
                [
                    row.calculator_input_key,
                    row.current_input_type,
                    {
                        "PRICE_DATABASE": "price",
                        "DEFAULT_VALUE": "default",
                        "AUTO_CALCULATED": "calculated",
                        "AUTO_PROJECT": "parsed",
                        "MANUAL_REQUIRED": "manual",
                        "REQUIRES_VALIDATION": row.current_input_type,
                    }[row.recommended_source_status],
                    row.should_be_visible_to_elena,
                    row.recommended_default_value or row.recommended_formula or row.recommended_registry_source or row.source_of_truth,
                ]
                for row in audit_rows[:120]
            ],
        ),
    ])
    if len(audit_rows) > 120:
        lines.append(f"_Показаны первые 120 строк из {len(audit_rows)}. Полный список в Excel._\n")

    lines.extend([
        "## Ожидаемый эффект",
        f"- Было ручных/незаполненных: {total}",
        f"- Можно убрать из ручного ввода после подтверждения registry/catalog/defaults/derived logic: {len(remove_rows)}",
        f"- Останется реально ручных: {len(manual_rows)}",
        f"- Потенциальное сокращение: {reduction:.1f}%",
        "",
        "## Важные оговорки",
        "- DEFAULT_VALUE не означает автоматическую подстановку старой сметы как универсального норматива.",
        "- Для каждого default нужен источник истины, применимость, возможность override и оценка риска.",
        "- Проектные и case-specific параметры нельзя скрывать без надежного PDF/source или подтвержденной формулы.",
        "- demo_with_template_fallback не является production-расчетом.",
    ])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def write_result_json(path: Path, audit_rows: list[AuditRow], output_xlsx: Path, output_report: Path) -> None:
    by_status = Counter(row.recommended_source_status for row in audit_rows)
    by_section: dict[str, Counter[str]] = defaultdict(Counter)
    section_names: dict[str, str] = {}
    for row in audit_rows:
        by_section[row.section_code][row.recommended_source_status] += 1
        section_names[row.section_code] = row.section_name

    result = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "reviewed_parameters_path": str(REVIEWED_PARAMETERS_PATH.relative_to(REPO_ROOT)),
        "elena_missing_report_path": str(ELENA_MISSING_REPORT_PATH.relative_to(REPO_ROOT)),
        "parameters_audited": len(audit_rows),
        "recommended_classification": dict(by_status),
        "sections": [
            {
                "section_code": section,
                "section_name": section_names.get(section, ""),
                "total": sum(counter.values()),
                **{status: counter.get(status, 0) for status in SOURCE_STATUSES},
            }
            for section, counter in sorted(by_section.items())
        ],
        "outputs": {
            "parameter_audit_result_xlsx": str(output_xlsx.relative_to(REPO_ROOT)),
            "parameter_audit_report_md": str(output_report.relative_to(REPO_ROOT)),
        },
        "warnings": [
            "Audit only: calculators, expected.json, and section_schema.py were not changed.",
            "DEFAULT_VALUE recommendations require source validation before hiding from Elena.",
        ],
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def run_audit(case_dir: Path) -> dict[str, Path]:
    parameters = load_parameters(REVIEWED_PARAMETERS_PATH)
    price_codes = load_price_codes()
    audit_rows = [classify(row, price_codes) for row in parameters if is_missing_or_manual(row)]

    case_dir.mkdir(parents=True, exist_ok=True)
    output_dir = REPO_ROOT / "experiments/parameter_audit/output" / case_dir.name
    output_dir.mkdir(parents=True, exist_ok=True)

    case_xlsx = case_dir / "parameter_audit_result.xlsx"
    case_report = case_dir / "parameter_audit_report.md"
    case_json = case_dir / "result.json"
    output_xlsx = output_dir / "parameter_audit_result.xlsx"
    output_report = output_dir / "parameter_audit_report.md"
    output_json = output_dir / "result.json"

    write_excel(case_xlsx, audit_rows)
    write_report(case_report, audit_rows)
    write_result_json(case_json, audit_rows, case_xlsx, case_report)

    shutil.copy2(case_xlsx, output_xlsx)
    shutil.copy2(case_report, output_report)
    shutil.copy2(case_json, output_json)

    return {
        "case_xlsx": case_xlsx,
        "case_report": case_report,
        "case_json": case_json,
        "output_xlsx": output_xlsx,
        "output_report": output_report,
        "output_json": output_json,
    }
