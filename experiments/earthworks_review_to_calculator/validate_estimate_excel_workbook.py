#!/usr/bin/env python3
"""
validate_estimate_excel_workbook.py

Layout-aware validator for AI-estimator Excel workbooks.
No hardcoded row numbers — section_row and data_start_row come from CLI or layout manifest.

Correct layout usage:
    python validate_estimate_excel_workbook.py \
        --workbook "...earthworks_formula_review.xlsx" \
        --formula-ready "...formula_ready_result.json" \
        --calc-result "...calculation_result/result.json" \
        --expected-sheet-title "21.06.2026" \
        --section-code earthworks \
        --section-number 2 \
        --section-title "Земляные работы" \
        --section-row 11 \
        --data-start-row 12

Artificial layout test (should FAIL — proves no hardcoded rows):
    python validate_estimate_excel_workbook.py ... \
        --section-row 20 --data-start-row 21
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


# ── Constants ─────────────────────────────────────────────────────────────────

EXCEL_ERRORS = frozenset({
    "#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A",
    "#NUM!", "#NULL!", "#ERROR!",
})

# Strings that must NOT appear in any cell value — JSON metadata leak detection.
# "source" / "confidence" / "formula_ready" are internal JSON field names; they
# have no place in a finished estimate cell.
FORBIDDEN_STRINGS = frozenset({
    "pipe_items",
    "price_code",
    "calc_price_key",
    "helper_id",
    "source",
    "confidence",
    "formula_ready",
    "Ключ цены",
    "Метод расчета",
    "raw JSON",
})

# Must match WHITE_ZONE_MIRROR in export_formula_ready_to_excel.py
WHITE_ZONE_MIRROR: dict[str, str] = {
    "D": "J", "E": "K", "F": "L",
    "G": "M", "H": "N", "I": "O",
}

# Helper zone P:AA (P=col16 … AA=col27) — includes mini-table zone W:AA
HELPER_COL_START = 16   # P
HELPER_COL_END = 27     # AA
HELPER_COLS = [get_column_letter(i) for i in range(HELPER_COL_START, HELPER_COL_END + 1)]
MINI_TABLE_COL_START = 23  # W

# ── Earthworks row registry ───────────────────────────────────────────────────
# Each logical key maps to matching rules applied against the item_key part of row_id
# (i.e. the portion after "section_code.").
# Matching order: exact first, then contains substring.
# If no match → FAIL (required row missing).
# If >1 match → FAIL (ambiguous — add a more specific alias to exact list).
REQUIRED_EARTHWORKS_ROWS: dict[str, dict[str, list[str]]] = {
    "axis_marking": {
        "exact": ["axis_marking"],
        "contains": [],
    },
    "excavator": {
        "exact": ["machine_excavation", "excavator", "excavator_jcb"],
        "contains": ["excavator"],
    },
    "manual_excavation": {
        "exact": ["manual_excavation"],
        "contains": [],
    },
    "geotextile_laying": {
        "exact": ["geotextile_laying"],
        "contains": [],
    },
    "geotextile_material": {
        "exact": ["geotextile_material"],
        "contains": [],
    },
    "sand_filling": {
        "exact": ["sand_filling"],
        "contains": [],
    },
    "sand_material": {
        "exact": ["sand_material"],
        "contains": [],
    },
    "sand_manual_moving": {
        "exact": ["sand_manual_moving"],
        "contains": [],
    },
    "communications_work": {
        "exact": ["communications_work"],
        "contains": [],
    },
    "communications_material": {
        "exact": ["communications_material"],
        "contains": [],
    },
    "consumables": {
        "exact": ["consumables"],
        "contains": [],
    },
}


# ── Report collector ──────────────────────────────────────────────────────────

class ValidationReport:
    def __init__(self) -> None:
        self.checks: list[dict[str, str]] = []
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.infos: list[dict[str, Any]] = []

    def ok(self, name: str, detail: str = "") -> None:
        self.checks.append({"name": name, "status": "OK", "detail": detail})

    def fail(self, name: str, detail: str = "") -> None:
        self.checks.append({"name": name, "status": "FAIL", "detail": detail})
        self.errors.append(f"{name}: {detail}")

    def warn(self, name: str, detail: str = "") -> None:
        self.checks.append({"name": name, "status": "WARN", "detail": detail})
        self.warnings.append(f"{name}: {detail}")

    def info(self, key: str, value: Any) -> None:
        self.infos.append({"key": key, "value": value})

    def is_pass(self) -> bool:
        return len(self.errors) == 0

    def render_markdown(self) -> str:
        verdict = "PASS" if self.is_pass() else "FAIL"
        lines = [
            "# Excel Workbook Validation Report",
            "",
            f"## Verdict: {verdict}",
            "",
        ]
        if self.errors:
            lines += ["## Errors", ""]
            for e in self.errors:
                lines.append(f"- ERROR: {e}")
            lines.append("")
        if self.warnings:
            lines += ["## Warnings", ""]
            for w in self.warnings:
                lines.append(f"- WARN: {w}")
            lines.append("")
        if self.infos:
            lines += ["## Info", ""]
            for i in self.infos:
                lines.append(f"- {i['key']}: {i['value']}")
            lines.append("")
        lines += ["## All checks", ""]
        lines.append("| Check | Status | Detail |")
        lines.append("|---|---|---|")
        for c in self.checks:
            lines.append(f"| {c['name']} | {c['status']} | {c.get('detail', '')} |")
        return "\n".join(lines)


# ── Cell helpers ──────────────────────────────────────────────────────────────

def _cv(ws: Any, col_letter: str, row: int) -> Any:
    return ws[f"{col_letter.upper()}{row}"].value


def _is_formula(val: Any) -> bool:
    return isinstance(val, str) and val.startswith("=")


def _is_numeric(val: Any) -> bool:
    return isinstance(val, (int, float)) and not isinstance(val, bool)


def _is_excel_error(val: Any) -> bool:
    return isinstance(val, str) and val.strip() in EXCEL_ERRORS


def _has_forbidden(val: Any) -> str | None:
    if not isinstance(val, str):
        return None
    for token in FORBIDDEN_STRINGS:
        if token in val:
            return token
    return None


def _scan_numeric_range(ws: Any, row: int, col_start: int, col_end: int) -> list[float]:
    """Return numeric cell values (in order) across col_start..col_end at given row."""
    return [
        float(v)
        for ci in range(col_start, col_end + 1)
        if _is_numeric(v := _cv(ws, get_column_letter(ci), row))
    ]


# ── Earthworks row matching ───────────────────────────────────────────────────

def _matches_rule(item_key: str, rule: dict[str, list[str]]) -> bool:
    if item_key in rule["exact"]:
        return True
    return any(pattern in item_key for pattern in rule["contains"])


def _build_logical_row_map(
    rows: list[dict[str, Any]],
    section_code: str,
    data_start_row: int,
    rpt: ValidationReport,
) -> dict[str, int]:
    """
    Build logical_key -> excel_row map for earthworks section.
    Uses REQUIRED_EARTHWORKS_ROWS rules against item_key (the part of row_id after section_code.).
    Adds FAIL to rpt for each missing or ambiguous logical key.
    """
    prefix = f"{section_code}."
    item_to_row: dict[str, tuple[int, str]] = {}
    for idx, row_data in enumerate(rows):
        row_id = str(row_data.get("row_id", ""))
        if not row_id.startswith(prefix):
            rpt.warn(
                "earthworks_row_prefix",
                f"row_id {row_id!r} does not start with {prefix!r} — skipped",
            )
            continue
        item_key = row_id[len(prefix):]
        item_to_row[item_key] = (data_start_row + idx, row_id)

    logical_row_map: dict[str, int] = {}
    for logical_key, rule in REQUIRED_EARTHWORKS_ROWS.items():
        matches = [
            (ik, er, rid)
            for ik, (er, rid) in item_to_row.items()
            if _matches_rule(ik, rule)
        ]
        if len(matches) == 0:
            rpt.fail(
                f"earthworks_row[{logical_key}]",
                f"required row not found; exact={rule['exact']}, contains={rule['contains']}",
            )
        elif len(matches) > 1:
            rpt.fail(
                f"earthworks_row[{logical_key}]",
                f"ambiguous match: {[m[0] for m in matches]}",
            )
        else:
            ik, excel_row, row_id = matches[0]
            logical_row_map[logical_key] = excel_row
            rpt.ok(
                f"earthworks_row[{logical_key}]",
                f"row_id={row_id!r} → excel_row={excel_row}",
            )

    found = len(logical_row_map)
    total = len(REQUIRED_EARTHWORKS_ROWS)
    if found == total:
        rpt.ok("earthworks_row_map", f"all {total} required rows resolved")
    return logical_row_map


# ── Earthworks semantic helper checks ────────────────────────────────────────

def _check_earthworks_semantics(
    ws: Any,
    logical_row_map: dict[str, int],
    rpt: ValidationReport,
) -> None:
    """
    Semantic checks for earthworks-specific helper cell patterns.
    All row lookups go through logical_row_map — zero hardcoded row numbers.
    """

    def _check_paa_empty(logical_key: str, label: str) -> None:
        """All P:AA must be empty (no value) for rows that have no helper data."""
        if logical_key not in logical_row_map:
            return
        r = logical_row_map[logical_key]
        non_empty = [
            f"{get_column_letter(ci)}{r}"
            for ci in range(HELPER_COL_START, HELPER_COL_END + 1)
            if _cv(ws, get_column_letter(ci), r) is not None
        ]
        if non_empty:
            rpt.fail(
                f"earthworks_helper[{logical_key}.P_AA_empty]",
                f"{label}: expected P:AA empty, non-empty: {non_empty}",
            )
        else:
            rpt.ok(
                f"earthworks_helper[{logical_key}.P_AA_empty]",
                f"{label}: P:AA all empty ✓",
            )

    def _check_p_filled_qaa_empty(logical_key: str, label: str) -> None:
        """P must have a numeric value; Q:AA must be empty."""
        if logical_key not in logical_row_map:
            return
        r = logical_row_map[logical_key]
        p_val = _cv(ws, "P", r)
        if not _is_numeric(p_val):
            rpt.fail(
                f"earthworks_helper[{logical_key}.P_filled]",
                f"{label}: P{r} expected numeric, got {p_val!r}",
            )
        else:
            rpt.ok(
                f"earthworks_helper[{logical_key}.P_filled]",
                f"{label}: P{r}={p_val}",
            )
        non_empty_qaa = [
            f"{get_column_letter(ci)}{r}"
            for ci in range(17, HELPER_COL_END + 1)  # Q(17):AA(27)
            if _cv(ws, get_column_letter(ci), r) is not None
        ]
        if non_empty_qaa:
            rpt.fail(
                f"earthworks_helper[{logical_key}.Q_AA_empty]",
                f"{label}: expected Q:AA empty, non-empty: {non_empty_qaa}",
            )
        else:
            rpt.ok(
                f"earthworks_helper[{logical_key}.Q_AA_empty]",
                f"{label}: Q:AA empty ✓",
            )

    # axis_marking: no helpers → P:AA empty
    _check_paa_empty("axis_marking", "axis_marking")

    # manual_excavation: if W:AA has numeric values → sum of routes = total, total ≈ S
    if "manual_excavation" in logical_row_map:
        r = logical_row_map["manual_excavation"]
        waa = _scan_numeric_range(ws, r, MINI_TABLE_COL_START, HELPER_COL_END)
        if len(waa) >= 2:
            routes_sum = sum(waa[:-1])
            route_total = waa[-1]
            if abs(routes_sum - route_total) > 0.01:
                rpt.fail(
                    "earthworks_helper[manual_excavation.routes_sum]",
                    f"sum of route values {routes_sum:.2f} != route total {route_total:.2f}",
                )
            else:
                rpt.ok(
                    "earthworks_helper[manual_excavation.routes_sum]",
                    f"routes sum {routes_sum:.2f} == total {route_total:.2f} ✓",
                )
            s_val = _cv(ws, "S", r)
            if _is_numeric(s_val):
                if abs(route_total - float(s_val)) > 0.01:
                    rpt.fail(
                        "earthworks_helper[manual_excavation.total_vs_S]",
                        f"route total {route_total:.2f} != S{r}={s_val}",
                    )
                else:
                    rpt.ok(
                        "earthworks_helper[manual_excavation.total_vs_S]",
                        f"route total {route_total:.2f} == S{r}={s_val} ✓",
                    )
        elif len(waa) == 1:
            rpt.warn(
                "earthworks_helper[manual_excavation.routes_sum]",
                f"only 1 numeric value in W:AA at row {r} — expected route list + total",
            )
        else:
            rpt.ok(
                "earthworks_helper[manual_excavation.routes_sum]",
                f"no route mini-table at row {r} (no trench routes)",
            )

    # sand_material: P filled, Q:AA empty
    _check_p_filled_qaa_empty("sand_material", "sand_material")

    # sand_manual_moving: no helpers → P:AA empty
    _check_paa_empty("sand_manual_moving", "sand_manual_moving")

    # communications_work: P and Q filled; if W:AA has numbers → sum = total, total ≈ P
    if "communications_work" in logical_row_map:
        r = logical_row_map["communications_work"]
        p_val = _cv(ws, "P", r)
        q_val = _cv(ws, "Q", r)
        if not _is_numeric(p_val):
            rpt.fail(
                "earthworks_helper[communications_work.P_filled]",
                f"P{r} expected numeric (total length), got {p_val!r}",
            )
        else:
            rpt.ok(
                "earthworks_helper[communications_work.P_filled]",
                f"P{r}={p_val} (total comms length)",
            )
        if q_val is None:
            rpt.fail(
                "earthworks_helper[communications_work.Q_filled]",
                f"Q{r} is empty — expected pipe count",
            )
        else:
            rpt.ok(
                "earthworks_helper[communications_work.Q_filled]",
                f"Q{r}={q_val} (pipe count)",
            )
        waa = _scan_numeric_range(ws, r, MINI_TABLE_COL_START, HELPER_COL_END)
        if len(waa) >= 2:
            pipes_sum = sum(waa[:-1])
            pipe_total = waa[-1]
            if abs(pipes_sum - pipe_total) > 0.01:
                rpt.fail(
                    "earthworks_helper[communications_work.pipes_sum]",
                    f"sum of pipe lengths {pipes_sum:.2f} != total {pipe_total:.2f}",
                )
            else:
                rpt.ok(
                    "earthworks_helper[communications_work.pipes_sum]",
                    f"pipes sum {pipes_sum:.2f} == total {pipe_total:.2f} ✓",
                )
            if _is_numeric(p_val) and abs(pipe_total - float(p_val)) > 0.01:
                rpt.fail(
                    "earthworks_helper[communications_work.total_vs_P]",
                    f"pipe total {pipe_total:.2f} != P{r}={p_val}",
                )
            elif _is_numeric(p_val):
                rpt.ok(
                    "earthworks_helper[communications_work.total_vs_P]",
                    f"pipe total {pipe_total:.2f} == P{r}={p_val} ✓",
                )
        else:
            rpt.ok(
                "earthworks_helper[communications_work.pipes_sum]",
                f"no pipe mini-table at row {r}",
            )

    # communications_material: P filled, Q:AA empty
    _check_p_filled_qaa_empty("communications_material", "communications_material")

    # consumables: no helpers → P:AA empty
    _check_paa_empty("consumables", "consumables")


# ── Layout manifest helpers ───────────────────────────────────────────────────

def _apply_layout_manifest(
    manifest: dict[str, Any],
    section_code: str,
    section_row_default: int,
    data_start_row_default: int,
) -> dict[str, Any]:
    """Return manifest entry for section_code, or {} if not found."""
    for s in manifest.get("sections", []):
        if s.get("section_code") == section_code:
            return {
                "section_row": int(s.get("section_row", section_row_default)),
                "data_start_row": int(s.get("data_start_row", data_start_row_default)),
                "section_number": s.get("section_number"),
                "section_title": s.get("section_title"),
                "data_end_row": s.get("data_end_row"),
                "total_row": s.get("total_row"),
            }
    return {}


# ── Main validator ────────────────────────────────────────────────────────────

def validate(
    wb_path: Path,
    formula_ready_path: Path,
    calc_result_path: Path | None,
    section_row: int,
    data_start_row: int,
    expected_sheet_title: str | None = None,
    section_code_arg: str | None = None,
    section_number_arg: str | None = None,
    section_title_arg: str | None = None,
    manifest_data: dict[str, Any] | None = None,
) -> ValidationReport:
    rpt = ValidationReport()

    # ── Load formula_ready ────────────────────────────────────────────────────
    data = json.loads(formula_ready_path.read_text(encoding="utf-8"))
    rows: list[dict[str, Any]] = data.get("rows", [])
    section_code: str = data.get("section_code", "")
    section_title: str = data.get("section_title", "")
    item_count = len(rows)
    data_end_row = data_start_row + item_count - 1
    total_row = data_end_row + 1

    rpt.info("section_code", section_code)
    rpt.info("section_title", section_title)
    rpt.info("section_row", section_row)
    rpt.info("data_start_row", data_start_row)
    rpt.info("data_end_row", data_end_row)
    rpt.info("total_row", total_row)
    rpt.info("item_count", item_count)

    # Expected totals from calc_result (info only — xlsx has no cached formula values)
    if calc_result_path and calc_result_path.exists():
        cr = json.loads(calc_result_path.read_text(encoding="utf-8"))
        it = cr.get("internal_totals", {})
        rpt.info("expected_material_total_from_json", it.get("internal_materials_total", "N/A"))
        rpt.info("expected_work_total_from_json", it.get("internal_works_total", "N/A"))
        rpt.info("expected_section_total_from_json", it.get("internal_section_total", "N/A"))
        rpt.info("excel_O_total_formula", f"=SUM(O{data_start_row}:O{data_end_row})")
        rpt.info(
            "totals_note",
            "numeric values not cached in xlsx; formula structure verified, not values",
        )

    # ── Layout manifest cross-check (if provided) ─────────────────────────────
    if manifest_data:
        m_sr = manifest_data.get("section_row")
        m_dsr = manifest_data.get("data_start_row")
        m_der = manifest_data.get("data_end_row")
        m_tr = manifest_data.get("total_row")
        m_sn = manifest_data.get("section_number")
        m_st = manifest_data.get("section_title")

        if m_sr is not None and int(m_sr) != section_row:
            rpt.warn("manifest_section_row", f"manifest={m_sr}, CLI={section_row}")
        if m_dsr is not None and int(m_dsr) != data_start_row:
            rpt.warn("manifest_data_start_row", f"manifest={m_dsr}, CLI={data_start_row}")
        if m_der is not None and int(m_der) != data_end_row:
            rpt.warn(
                "manifest_data_end_row",
                f"manifest={m_der}, computed={data_end_row} "
                f"(data_start_row + item_count - 1 = {data_start_row} + {item_count} - 1)",
            )
        if m_tr is not None and int(m_tr) != total_row:
            rpt.warn("manifest_total_row", f"manifest={m_tr}, computed={total_row}")
        if m_sn is not None and section_number_arg is not None:
            if str(m_sn).strip() != str(section_number_arg).strip():
                rpt.warn(
                    "manifest_section_number",
                    f"manifest={m_sn!r}, --section-number={section_number_arg!r}",
                )
        if m_st is not None and section_title_arg is not None:
            if str(m_st).strip() != str(section_title_arg).strip():
                rpt.warn(
                    "manifest_section_title",
                    f"manifest={m_st!r}, --section-title={section_title_arg!r}",
                )

    # ── 1. File accessible ────────────────────────────────────────────────────
    if not wb_path.exists():
        rpt.fail("file_accessible", f"not found: {wb_path}")
        return rpt
    rpt.ok("file_accessible", wb_path.name)

    try:
        wb = openpyxl.load_workbook(str(wb_path), data_only=False)
    except Exception as exc:
        rpt.fail("workbook_openable", str(exc))
        return rpt
    rpt.ok("workbook_openable")

    # ── 2. Sheet title ────────────────────────────────────────────────────────
    sheet_names = wb.sheetnames
    rpt.info("sheets", str(sheet_names))
    if not sheet_names:
        rpt.fail("sheet_count", "workbook has no sheets")
        return rpt

    ws = wb.active
    actual_title = ws.title

    if expected_sheet_title is not None:
        if actual_title != expected_sheet_title:
            rpt.fail(
                "sheet_title",
                f"expected {expected_sheet_title!r}, got {actual_title!r}",
            )
        else:
            rpt.ok("sheet_title", actual_title)
    else:
        rpt.ok("sheet_title", f"(not checked) active sheet: {actual_title!r}")

    if len(sheet_names) > 1:
        rpt.warn("sheet_count", f"expected 1 sheet, found {len(sheet_names)}: {sheet_names}")
    else:
        rpt.ok("sheet_count", f"1 sheet: {sheet_names[0]!r}")

    # ── 3. section_code check ─────────────────────────────────────────────────
    if section_code_arg is not None:
        if section_code != section_code_arg:
            rpt.fail(
                "section_code",
                f"formula_ready.section_code={section_code!r} != --section-code={section_code_arg!r}",
            )
        else:
            rpt.ok("section_code", section_code)
    else:
        rpt.ok("section_code", f"(not checked) formula_ready.section_code={section_code!r}")

    # ── 4. Comments ───────────────────────────────────────────────────────────
    comment_count = sum(
        1 for row_cells in ws.iter_rows() for cell in row_cells if cell.comment is not None
    )
    if comment_count > 0:
        rpt.fail("no_comments", f"{comment_count} comment(s) found — not allowed")
    else:
        rpt.ok("no_comments")

    # ── 5. Excel errors and forbidden strings ─────────────────────────────────
    error_cells: list[str] = []
    forbidden_cells: list[str] = []
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            val = cell.value
            if _is_excel_error(val):
                error_cells.append(f"{cell.coordinate}={val}")
            found = _has_forbidden(val)
            if found:
                forbidden_cells.append(f"{cell.coordinate} contains {found!r}")

    if error_cells:
        rpt.fail("no_excel_errors", "; ".join(error_cells[:5]))
    else:
        rpt.ok("no_excel_errors")

    if forbidden_cells:
        rpt.fail("no_forbidden_strings", "; ".join(forbidden_cells[:5]))
    else:
        rpt.ok("no_forbidden_strings")

    # ── 6. Section row — exact checks ────────────────────────────────────────
    a_val = _cv(ws, "A", section_row)
    b_val = _cv(ws, "B", section_row)

    if section_number_arg is not None:
        # Exact match: compare both as strings to handle int vs str from xlsx
        if str(a_val).strip() == str(section_number_arg).strip():
            rpt.ok("section_row_A", f"A{section_row}={a_val!r} == {section_number_arg!r}")
        else:
            rpt.fail(
                "section_row_A",
                f"A{section_row}={a_val!r} != expected {section_number_arg!r}",
            )
    else:
        if a_val is None or str(a_val).strip() == "":
            rpt.fail("section_row_A", f"A{section_row} is empty — expected section number")
        else:
            rpt.ok("section_row_A", f"A{section_row}={a_val!r}")

    effective_section_title = section_title_arg if section_title_arg is not None else section_title
    if b_val is None or str(b_val).strip() == "":
        rpt.fail("section_row_B", f"B{section_row} is empty — expected {effective_section_title!r}")
    elif str(b_val).strip() != effective_section_title.strip():
        rpt.fail(
            "section_row_B",
            f"B{section_row}={b_val!r} != expected {effective_section_title!r}",
        )
    else:
        rpt.ok("section_row_B", f"B{section_row}={b_val!r}")

    # ── 7. Item count (B cells in data range) ────────────────────────────────
    actual_items = sum(
        1 for r in range(data_start_row, data_end_row + 1)
        if _cv(ws, "B", r) not in (None, "")
    )
    if actual_items != item_count:
        rpt.fail(
            "item_count",
            f"B column has {actual_items} non-empty rows in [{data_start_row}:{data_end_row}], "
            f"expected {item_count}",
        )
    else:
        rpt.ok("item_count", f"{item_count} rows in [{data_start_row}:{data_end_row}]")

    # ── 8. Row names match formula_ready ─────────────────────────────────────
    name_fails: list[str] = []
    for idx, row_data in enumerate(rows):
        excel_row = data_start_row + idx
        expected = str(row_data.get("estimate_line", "")).strip()
        actual_val = _cv(ws, "B", excel_row)
        actual = str(actual_val).strip() if actual_val is not None else ""
        if actual != expected:
            name_fails.append(
                f"B{excel_row}: got {actual!r} != {expected!r} ({row_data.get('row_id')})"
            )
    if name_fails:
        for msg in name_fails:
            rpt.fail("row_name", msg)
    else:
        rpt.ok("row_names", f"all {item_count} row names match formula_ready")

    # ── 9. White zone D:I mirrors calc zone J:O ───────────────────────────────
    white_fails: list[str] = []
    for idx in range(item_count):
        r = data_start_row + idx
        for white_col, calc_col in WHITE_ZONE_MIRROR.items():
            expected = f"={calc_col}{r}"
            actual = _cv(ws, white_col, r)
            if actual != expected:
                white_fails.append(f"{white_col}{r}: expected {expected!r}, got {actual!r}")
    if white_fails:
        for msg in white_fails:
            rpt.fail("white_mirror", msg)
    else:
        rpt.ok(
            "white_zone_mirrors",
            f"all {item_count * len(WHITE_ZONE_MIRROR)} white zone cells correct (D:I=J:O)",
        )

    # ── 10. Calc zone L/N/O formulas per data row (J*K, J*M, L+N) ────────────
    calc_fails: list[str] = []
    for idx in range(item_count):
        r = data_start_row + idx
        for col, expected in [
            ("L", f"=J{r}*K{r}"),
            ("N", f"=J{r}*M{r}"),
            ("O", f"=L{r}+N{r}"),
        ]:
            actual = _cv(ws, col, r)
            if actual != expected:
                calc_fails.append(f"{col}{r}: expected {expected!r}, got {actual!r}")
    if calc_fails:
        for msg in calc_fails:
            rpt.fail("calc_formula", msg)
    else:
        rpt.ok("calc_zone_formulas", f"all {item_count * 3} L/N/O formulas correct")

    # ── 11. Quantity cell J: formula if formula_model, else numeric value ──────
    qty_fails: list[str] = []
    qty_warns: list[str] = []
    for idx, row_data in enumerate(rows):
        r = data_start_row + idx
        row_id = str(row_data.get("row_id", ""))
        cz_qty = row_data.get("calc_zone", {}).get("quantity", {})
        has_fm = cz_qty.get("formula_model") is not None
        is_exportable = cz_qty.get("excel_formula_exportable") is not False
        actual = _cv(ws, "J", r)

        if has_fm and is_exportable:
            if not _is_formula(actual):
                qty_fails.append(
                    f"J{r} ({row_id}): expected formula (has formula_model), got {actual!r}"
                )
        else:
            if _is_formula(actual):
                qty_warns.append(
                    f"J{r} ({row_id}): has formula {actual!r} but no formula_model"
                )
            elif actual is None:
                qty_fails.append(f"J{r} ({row_id}): quantity cell is empty")
            elif not _is_numeric(actual):
                qty_warns.append(f"J{r} ({row_id}): non-numeric value {actual!r}")

    for msg in qty_fails:
        rpt.fail("qty_formula", msg)
    for msg in qty_warns:
        rpt.warn("qty_formula", msg)
    if not qty_fails:
        rpt.ok("qty_formulas", f"all {item_count} J quantity cells correct")

    # ── 12. Helper cells (P:AA): formula or numeric value as expected ──────────
    helper_fails: list[str] = []
    helper_warns: list[str] = []
    for idx, row_data in enumerate(rows):
        r = data_start_row + idx
        row_id = str(row_data.get("row_id", ""))
        for hcell in row_data.get("helper_zone", {}).get("cells", []):
            target_col = str(hcell.get("target_col_role", ""))
            helper_id = str(hcell.get("helper_id", ""))
            if target_col not in HELPER_COLS:
                helper_warns.append(
                    f"{row_id}.{helper_id}: target_col {target_col!r} outside P:AA — skipped"
                )
                continue

            has_fm = hcell.get("formula_model") is not None
            is_exportable = hcell.get("excel_formula_exportable") is not False
            actual = _cv(ws, target_col, r)
            addr = f"{target_col}{r}"

            if actual is None:
                helper_fails.append(f"{addr} ({row_id}.{helper_id}): cell is empty")
                continue
            if _is_excel_error(actual):
                helper_fails.append(f"{addr} ({row_id}.{helper_id}): Excel error {actual!r}")
                continue

            if has_fm and is_exportable:
                if not _is_formula(actual):
                    helper_fails.append(
                        f"{addr} ({row_id}.{helper_id}): expected formula (has formula_model), "
                        f"got value {actual!r}"
                    )
            else:
                if _is_formula(actual):
                    helper_warns.append(
                        f"{addr} ({row_id}.{helper_id}): formula {actual!r} but "
                        f"excel_formula_exportable=false"
                    )
                elif not _is_numeric(actual):
                    helper_warns.append(
                        f"{addr} ({row_id}.{helper_id}): non-numeric value {actual!r}"
                    )

    for msg in helper_fails:
        rpt.fail("helper_cell", msg)
    for msg in helper_warns:
        rpt.warn("helper_cell", msg)
    total_helper_cells = sum(
        len(r.get("helper_zone", {}).get("cells", [])) for r in rows
    )
    if not helper_fails:
        rpt.ok(
            "helper_cells",
            f"{total_helper_cells} helper cells correct across {item_count} rows",
        )

    # ── 13. Total row formulas ─────────────────────────────────────────────────
    expected_total: dict[str, str] = {
        "F": f"=L{total_row}",
        "H": f"=N{total_row}",
        "I": f"=O{total_row}",
        "L": f"=SUM(L{data_start_row}:L{data_end_row})",
        "N": f"=SUM(N{data_start_row}:N{data_end_row})",
        "O": f"=SUM(O{data_start_row}:O{data_end_row})",
    }
    total_fails: list[str] = []
    for col, expected in expected_total.items():
        actual = _cv(ws, col, total_row)
        if actual != expected:
            total_fails.append(f"{col}{total_row}: expected {expected!r}, got {actual!r}")
    if total_fails:
        for msg in total_fails:
            rpt.fail("total_row_formula", msg)
    else:
        rpt.ok("total_row_formulas", f"all 6 total row formulas correct at row {total_row}")

    # ── 14. Earthworks-specific checks ────────────────────────────────────────
    run_earthworks = (
        section_code_arg == "earthworks"
        or (section_code_arg is None and section_code == "earthworks")
    )
    if run_earthworks:
        logical_row_map = _build_logical_row_map(rows, section_code, data_start_row, rpt)
        _check_earthworks_semantics(ws, logical_row_map, rpt)

    return rpt


# ── CLI ───────────────────────────────────────────────────────────────────────

def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Validate AI-estimator Excel workbook (layout-aware, no hardcoded rows)"
    )
    p.add_argument("--workbook", required=True, help="Path to .xlsx file")
    p.add_argument("--formula-ready", required=True, help="Path to formula_ready_result.json")
    p.add_argument(
        "--calc-result", default=None,
        help="Path to calculation_result/result.json (optional, info only)",
    )
    p.add_argument(
        "--expected-sheet-title", default=None,
        help="Expected Excel sheet title (e.g. '21.06.2026')",
    )
    p.add_argument(
        "--section-code", default=None,
        help="Section code to validate against formula_ready (e.g. 'earthworks')",
    )
    p.add_argument(
        "--section-number", default=None,
        help="Expected value in A{section_row} (e.g. '2')",
    )
    p.add_argument(
        "--section-title", default=None,
        help="Expected value in B{section_row} (e.g. 'Земляные работы')",
    )
    p.add_argument(
        "--section-row", type=int, default=11,
        help="Row number of section header (default: 11)",
    )
    p.add_argument(
        "--data-start-row", type=int, default=12,
        help="First data row number (default: 12)",
    )
    p.add_argument(
        "--layout-manifest", default=None,
        help="Path to layout manifest JSON — overrides row args for the section",
    )
    p.add_argument("--out", default=None, help="Save validation report to this .md path")
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)

    formula_ready_path = Path(args.formula_ready).resolve()
    wb_path = Path(args.workbook).resolve()
    calc_result_path = Path(args.calc_result).resolve() if args.calc_result else None

    if not formula_ready_path.exists():
        print(f"error: formula_ready not found: {formula_ready_path}", file=sys.stderr)
        return 1

    section_row = args.section_row
    data_start_row = args.data_start_row
    manifest_data: dict[str, Any] | None = None

    if args.layout_manifest:
        try:
            manifest = json.loads(Path(args.layout_manifest).read_text(encoding="utf-8"))
            data = json.loads(formula_ready_path.read_text(encoding="utf-8"))
            section_code = data.get("section_code", "")
            entry = _apply_layout_manifest(manifest, section_code, section_row, data_start_row)
            if entry:
                section_row = entry["section_row"]
                data_start_row = entry["data_start_row"]
                manifest_data = entry
                print(
                    f"layout manifest: section_row={section_row}, "
                    f"data_start_row={data_start_row}",
                    file=sys.stderr,
                )
            else:
                print(
                    f"warning: section_code {section_code!r} not found in manifest",
                    file=sys.stderr,
                )
        except Exception as exc:
            print(f"warning: could not apply layout manifest: {exc}", file=sys.stderr)

    rpt = validate(
        wb_path=wb_path,
        formula_ready_path=formula_ready_path,
        calc_result_path=calc_result_path,
        section_row=section_row,
        data_start_row=data_start_row,
        expected_sheet_title=args.expected_sheet_title,
        section_code_arg=args.section_code,
        section_number_arg=args.section_number,
        section_title_arg=args.section_title,
        manifest_data=manifest_data,
    )
    report_md = rpt.render_markdown()
    print(report_md)

    if args.out:
        out_path = Path(args.out)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(report_md, encoding="utf-8")
        print(f"report saved: {out_path}", file=sys.stderr)

    verdict = "PASS" if rpt.is_pass() else "FAIL"
    print(f"validation: {verdict}", file=sys.stderr)
    return 0 if rpt.is_pass() else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
