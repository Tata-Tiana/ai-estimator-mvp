from __future__ import annotations

import argparse
import io
import json
import os
import shlex
import subprocess
import sys
from pathlib import Path

from dotenv import load_dotenv

REPO_ROOT = Path(__file__).resolve().parents[2]
BASE_DIR = Path(__file__).resolve().parent
STAGE1_DIR = REPO_ROOT / "experiments" / "earthworks_parser_google_stage1"
STAGE1_ENV = STAGE1_DIR / ".env"

EXPORT_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOWNLOAD_FILENAME = "downloaded_review_workbook.xlsx"

RUN_FULL_FLOW = BASE_DIR / "run_full_review_flow.py"


def _load_metadata(stage1_job_dir: Path) -> dict:
    metadata_path = stage1_job_dir / "google" / "google_sheet_metadata.json"
    if not metadata_path.exists():
        raise FileNotFoundError(f"google_sheet_metadata.json not found: {metadata_path}")
    data = json.loads(metadata_path.read_text(encoding="utf-8"))
    if not data.get("spreadsheet_id"):
        raise ValueError(f"spreadsheet_id is empty in {metadata_path}")
    if data.get("status") != "published":
        raise ValueError(
            f"Google Sheet metadata status is not published (got: {data.get('status')!r})"
        )
    return data


def _build_credentials():
    load_dotenv(STAGE1_ENV)

    token_path_raw = os.getenv("GOOGLE_TOKEN_PATH", "")
    if not token_path_raw:
        raise RuntimeError("GOOGLE_TOKEN_PATH not set in .env")
    token_path = Path(token_path_raw)
    if not token_path.is_absolute():
        token_path = REPO_ROOT / token_path
    if not token_path.exists():
        raise FileNotFoundError(f"token.json not found: {token_path}")

    try:
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request
    except ImportError as exc:
        raise RuntimeError(f"Google auth libraries not installed: {exc}") from exc

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive.file",
    ]
    creds = Credentials.from_authorized_user_file(str(token_path), scopes)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
    return creds


def _download_spreadsheet(spreadsheet_id: str, out_path: Path) -> None:
    try:
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaIoBaseDownload
    except ImportError as exc:
        raise RuntimeError(f"Google API client not installed: {exc}") from exc

    creds = _build_credentials()
    drive = build("drive", "v3", credentials=creds)
    request = drive.files().export_media(fileId=spreadsheet_id, mimeType=EXPORT_MIME)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with io.FileIO(str(out_path), mode="wb") as fh:
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()


def _restore_hidden_columns(downloaded_path: Path, local_path: Path) -> None:
    """Google Sheets strips hidden columns on export. Restore technical cols K:N
    (calc_price_key, price_registry_code, fallback_key, selected_price_source) from
    the original local review_workbook.xlsx.  User-visible columns A:J (including any
    price overrides Elena entered) are left untouched."""
    from openpyxl import load_workbook

    SHEET = "02_Цены себестоимости"
    HIDDEN_COLS = range(11, 15)   # K=11, L=12, M=13, N=14

    local_wb = load_workbook(local_path, data_only=True)
    dl_wb = load_workbook(downloaded_path)

    local_ws = local_wb[SHEET]
    dl_ws = dl_wb[SHEET]

    for row_idx in range(1, local_ws.max_row + 1):
        for col_idx in HIDDEN_COLS:
            dl_ws.cell(row_idx, col_idx).value = local_ws.cell(row_idx, col_idx).value

    for col_letter in ("K", "L", "M", "N"):
        dl_ws.column_dimensions[col_letter].hidden = True

    dl_wb.save(str(downloaded_path))


def _run_full_flow(
    review_workbook: Path,
    out_dir: Path,
    section_number: int,
    estimate_date: str | None,
    section_row: int,
    data_start_row: int,
) -> int:
    python = sys.executable
    cmd = [
        python,
        str(RUN_FULL_FLOW),
        "--review-workbook", str(review_workbook),
        "--out-dir", str(out_dir),
        "--section-number", str(section_number),
        "--section-row", str(section_row),
        "--data-start-row", str(data_start_row),
    ]
    if estimate_date:
        cmd.extend(["--estimate-date", estimate_date])

    parts = [str(p) for p in cmd]
    parts[0] = "python"
    print(f"$ {shlex.join(parts)}")
    result = subprocess.run(cmd, cwd=str(REPO_ROOT))
    return result.returncode


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Download reviewed Google Sheet and build earthworks estimate"
    )
    parser.add_argument(
        "--stage1-job-dir",
        required=True,
        help="Path to Stage1 job dir (contains google/google_sheet_metadata.json)",
    )
    parser.add_argument("--out-dir", required=True, help="Directory for outputs")
    parser.add_argument("--section-number", type=int, default=2)
    parser.add_argument("--estimate-date", default=None, help="DD.MM.YYYY for Excel tab")
    parser.add_argument("--section-row", type=int, default=11)
    parser.add_argument("--data-start-row", type=int, default=12)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    stage1_job_dir = Path(args.stage1_job_dir).resolve()
    out_dir = Path(args.out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"[download] reading metadata from {stage1_job_dir / 'google' / 'google_sheet_metadata.json'}")
    metadata = _load_metadata(stage1_job_dir)
    spreadsheet_id = metadata["spreadsheet_id"]
    print(f"[download] spreadsheet_id = {spreadsheet_id}")
    print(f"[download] url = {metadata.get('url', '—')}")

    xlsx_path = out_dir / DOWNLOAD_FILENAME
    print(f"[download] exporting from Google Drive → {xlsx_path.name}")
    _download_spreadsheet(spreadsheet_id, xlsx_path)
    print(f"[download] ok — {xlsx_path.stat().st_size} bytes")

    local_review_workbook = stage1_job_dir / "google" / "review_workbook.xlsx"
    if local_review_workbook.exists():
        print(f"[patch] restoring hidden technical columns from local review_workbook.xlsx")
        _restore_hidden_columns(xlsx_path, local_review_workbook)
        print(f"[patch] ok")
    else:
        print(f"[patch] WARNING: local review_workbook.xlsx not found — skipping hidden-column restore", file=sys.stderr)

    print()
    rc = _run_full_flow(
        review_workbook=xlsx_path,
        out_dir=out_dir,
        section_number=args.section_number,
        estimate_date=args.estimate_date,
        section_row=args.section_row,
        data_start_row=args.data_start_row,
    )

    if rc != 0:
        print(f"\n[build_from_google_sheet] FAIL (full_flow returncode={rc})", file=sys.stderr)
        return rc
    print("\n[build_from_google_sheet] PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
