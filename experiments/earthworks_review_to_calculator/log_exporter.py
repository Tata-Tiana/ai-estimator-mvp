"""
Standalone Excel log exporter for the Telegram bot.
Does NOT import telegram_bot. Takes all paths as parameters.
"""
from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

# ── limits ──────────────────────────────────────────────────────────────────
MAX_EVENTS_ROWS = 10_000
MAX_EXTRA_JSON_CHARS = 1_000
MAX_ERROR_CHARS = 2_000
MAX_PREVIEW_CHARS = 2_000

# ── stale timeouts (mirror telegram_bot constants, no import) ────────────────
DONE_REQUESTED_TIMEOUT_MINUTES = 15
QUEUED_TIMEOUT_MINUTES = 15
PROCESSING_TIMEOUT_MINUTES = 60

# ── redaction patterns ───────────────────────────────────────────────────────
_REDACT_PATTERNS: list[tuple[re.Pattern, str]] = [
    (re.compile(r'(TELEGRAM_BOT_TOKEN\s*[=:]\s*)\S+', re.I), r'\g<1>[REDACTED]'),
    (re.compile(r'(OPENAI_API_KEY\s*[=:]\s*)\S+', re.I), r'\g<1>[REDACTED]'),
    (re.compile(r'(ANTHROPIC_API_KEY\s*[=:]\s*)\S+', re.I), r'\g<1>[REDACTED]'),
    (re.compile(r'(HTTP_PROXY\s*[=:]\s*)\S+', re.I), r'\g<1>[REDACTED]'),
    (re.compile(r'(HTTPS_PROXY\s*[=:]\s*)\S+', re.I), r'\g<1>[REDACTED]'),
    (re.compile(r'(ALL_PROXY\s*[=:]\s*)\S+', re.I), r'\g<1>[REDACTED]'),
    (re.compile(r'(Bearer\s+)\S+', re.I), r'\g<1>[REDACTED]'),
    (re.compile(r'-----BEGIN PRIVATE KEY-----.*?-----END PRIVATE KEY-----', re.S), '[PRIVATE KEY REDACTED]'),
    (re.compile(r'"private_key"\s*:\s*"[^"]*"'), '"private_key": "[REDACTED]"'),
    # http/https proxy with credentials
    (re.compile(r'(https?://)[^:@/\s]+:[^@\s]+@', re.I), r'\g<1>[REDACTED]@'),
    # socks5 / socks4 proxy with credentials
    (re.compile(r'(socks[45]h?://)[^:@/\s]+:[^@\s]+@', re.I), r'\g<1>[REDACTED]@'),
]


def _redact_string(s: str) -> str:
    for pattern, repl in _REDACT_PATTERNS:
        s = pattern.sub(repl, s)
    return s


def _sanitize(value: Any) -> Any:
    if isinstance(value, str):
        return _redact_string(value)
    if isinstance(value, dict):
        return {k: _sanitize(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_sanitize(item) for item in value]
    return value


def _pdf_count(data: dict) -> int | str:
    if isinstance(data.get("pdfs"), list):
        return len(data["pdfs"])
    return data.get("pdf_count", "")


def _safe_str(value: Any, max_chars: int | None = None) -> str:
    if value is None:
        return ""
    s = str(value)
    if max_chars and len(s) > max_chars:
        s = s[:max_chars] + "…"
    return s


# ── period helpers ───────────────────────────────────────────────────────────

def _date_set_for_period(period: str) -> set[str] | None:
    """Return a set of YYYY-MM-DD strings to include, or None for 'all'."""
    today = datetime.now(timezone.utc).date()
    if period == "today":
        return {today.isoformat()}
    if period == "yesterday":
        return {(today - timedelta(days=1)).isoformat()}
    if period == "7d":
        return {(today - timedelta(days=i)).isoformat() for i in range(7)}
    return None  # all


def _date_from_jsonl_filename(fname: str) -> str | None:
    """Extract YYYY-MM-DD from filename like '2026-07-01.jsonl'."""
    stem = Path(fname).stem
    if re.match(r'^\d{4}-\d{2}-\d{2}$', stem):
        return stem
    return None


def _date_from_cmdlog_filename(fname: str) -> str | None:
    """Extract date from first 8 chars of filename like '20260701_123456_cmd.json'."""
    stem = Path(fname).stem
    if len(stem) >= 8 and stem[:8].isdigit():
        raw = stem[:8]
        try:
            d = datetime.strptime(raw, "%Y%m%d").date()
            return d.isoformat()
        except ValueError:
            return None
    return None


# ── stale logic (mirrors telegram_bot._stale_reason exactly) ─────────────────

def _parse_dt(value: Any) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    try:
        dt = datetime.fromisoformat(str(value))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except (ValueError, TypeError):
        return None


def _stale_reason(session: dict) -> str | None:
    status = session.get("status", "")
    now = datetime.now(timezone.utc)

    if status == "done_requested":
        start = _parse_dt(session.get("done_requested_at") or session.get("status_changed_at"))
        if start and (now - start).total_seconds() / 60 > DONE_REQUESTED_TIMEOUT_MINUTES:
            return f"done_requested>{DONE_REQUESTED_TIMEOUT_MINUTES}min"

    elif status == "queued":
        start = _parse_dt(session.get("queued_at") or session.get("status_changed_at"))
        if start and (now - start).total_seconds() / 60 > QUEUED_TIMEOUT_MINUTES:
            return f"queued>{QUEUED_TIMEOUT_MINUTES}min"

    elif status == "processing":
        start = _parse_dt(session.get("processing_started_at") or session.get("status_changed_at"))
        if start and (now - start).total_seconds() / 60 > PROCESSING_TIMEOUT_MINUTES:
            return f"processing>{PROCESSING_TIMEOUT_MINUTES}min"

    return None


# ── archive filename parser ──────────────────────────────────────────────────

def _parse_archive_filename(fname: str) -> dict:
    """
    Expected format: {chat_id}_{...session_id...}_{reason...}_{YYYYMMDD}_{HHMMSS}.json
    Archive timestamp = last two parts (YYYYMMDD_HHMMSS).
    If parts[1] is 8 digits and parts[2] is 6 digits → session_id = parts[1]+'_'+parts[2],
    reason = '_'.join(parts[3:-2]).
    Otherwise → session_id='unknown', reason='_'.join(parts[1:-2]).
    """
    stem = Path(fname).stem
    parts = stem.split("_")
    result: dict = {
        "chat_id": "",
        "session_id": "unknown",
        "archive_reason_from_filename": "",
        "archive_timestamp": "",
    }
    if len(parts) < 3:
        result["archive_reason_from_filename"] = stem
        return result

    result["chat_id"] = parts[0]
    archive_ts = ""
    if len(parts) >= 2:
        last_date = parts[-2]
        last_time = parts[-1]
        if len(last_date) == 8 and last_date.isdigit() and len(last_time) == 6 and last_time.isdigit():
            archive_ts = (
                f"{last_date[:4]}-{last_date[4:6]}-{last_date[6:8]} "
                f"{last_time[:2]}:{last_time[2:4]}:{last_time[4:6]}"
            )
    result["archive_timestamp"] = archive_ts

    if (
        len(parts) >= 5
        and len(parts[1]) == 8 and parts[1].isdigit()
        and len(parts[2]) == 6 and parts[2].isdigit()
    ):
        result["session_id"] = parts[1] + "_" + parts[2]
        result["archive_reason_from_filename"] = "_".join(parts[3:-2])
    else:
        result["archive_reason_from_filename"] = "_".join(parts[1:-2])

    return result


# ── data readers ─────────────────────────────────────────────────────────────

def _read_events(events_dir: Path, period: str) -> tuple[list[dict], int]:
    """
    Returns (exported_rows, total_found).
    Collects all events for the period first, then keeps the LAST MAX_EVENTS_ROWS.
    """
    date_set = _date_set_for_period(period)
    all_rows: list[dict] = []

    jsonl_files = sorted(events_dir.glob("*.jsonl"))
    if date_set is not None:
        jsonl_files = [f for f in jsonl_files if _date_from_jsonl_filename(f.name) in date_set]

    for path in jsonl_files:
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                event = json.loads(line)
            except json.JSONDecodeError:
                continue
            all_rows.append(event)

    total_found = len(all_rows)
    exported = all_rows[-MAX_EVENTS_ROWS:] if total_found > MAX_EVENTS_ROWS else all_rows
    return exported, total_found


def _read_command_logs(logs_dir: Path, period: str) -> list[dict]:
    date_set = _date_set_for_period(period)
    rows: list[dict] = []

    for chat_dir in sorted(logs_dir.iterdir()):
        if not chat_dir.is_dir():
            continue
        chat_id = chat_dir.name
        for log_file in sorted(chat_dir.glob("*.json")):
            file_date = _date_from_cmdlog_filename(log_file.name)
            if date_set is not None:
                if file_date is None:
                    continue  # skip unparseable for today/yesterday/7d
                if file_date not in date_set:
                    continue
            try:
                data = json.loads(log_file.read_text(encoding="utf-8", errors="replace"))
            except (OSError, json.JSONDecodeError):
                continue
            if not isinstance(data, dict):
                continue
            data = _sanitize(data)
            data.setdefault("_chat_id", chat_id)
            data.setdefault("_filename", log_file.name)
            rows.append(data)

    return rows


def _read_active_sessions(sessions_dir: Path) -> list[dict]:
    rows: list[dict] = []
    for f in sorted(sessions_dir.glob("*.json")):
        try:
            data = json.loads(f.read_text(encoding="utf-8", errors="replace"))
        except (OSError, json.JSONDecodeError):
            continue
        if not isinstance(data, dict):
            continue
        rows.append(_sanitize(data))
    return rows


def _read_archived_sessions(backups_dir: Path) -> list[dict]:
    rows: list[dict] = []
    for f in sorted(backups_dir.glob("*.json")):
        parsed = _parse_archive_filename(f.name)
        try:
            data = json.loads(f.read_text(encoding="utf-8", errors="replace"))
        except (OSError, json.JSONDecodeError):
            data = {}
        if not isinstance(data, dict):
            data = {}
        data = _sanitize(data)
        row = {
            "archive_file": f.name,
            "archive_timestamp": parsed["archive_timestamp"],
            "archive_reason_from_filename": parsed["archive_reason_from_filename"],
            "chat_id": data.get("chat_id", parsed["chat_id"]),
            "session_id": data.get("session_id", parsed["session_id"]),
            "status": data.get("status", ""),
            "created_at": data.get("created_at", ""),
            "updated_at": data.get("updated_at", ""),
            "status_changed_at": data.get("status_changed_at", ""),
            "done_requested_at": data.get("done_requested_at", ""),
            "queued_at": data.get("queued_at", ""),
            "processing_started_at": data.get("processing_started_at", ""),
            "completed_at": data.get("completed_at", ""),
            "failed_at": data.get("failed_at", ""),
            "cancelled_at": data.get("cancelled_at", ""),
            "stale_detected_at": data.get("stale_detected_at", ""),
            "stale_reason": data.get("stale_reason", ""),
            "job_id": data.get("job_id", ""),
            "processing_run_id": data.get("processing_run_id", ""),
            "pdf_count": _pdf_count(data),
            "project_name": data.get("project_name", ""),
            "spreadsheet_url": data.get("spreadsheet_url", ""),
            "last_error_type": data.get("last_error_type", ""),
            "last_error": _safe_str(data.get("last_error", ""), MAX_ERROR_CHARS),
            "reset_by_admin_chat_id": data.get("reset_by_admin_chat_id", ""),
            "reset_at": data.get("reset_at", ""),
            "reset_reason": data.get("reset_reason", ""),
        }
        rows.append(row)
    return rows


# ── sheet builders ────────────────────────────────────────────────────────────

_NO_DATA_ROW = "no data"


def _append_no_data_if_empty(ws) -> None:
    if ws.max_row <= 1:
        ws.append([_NO_DATA_ROW])


def _build_events_sheet(ws, events: list[dict]) -> None:
    columns = [
        "timestamp", "level", "event", "chat_id", "session_id", "job_id",
        "processing_run_id", "safe_message", "pdf_count", "filename",
        "project_name", "error_type", "error", "stale_reason", "command",
        "admin_chat_id", "target_chat_id", "backup_path", "local_path", "extra_json",
    ]
    _write_header(ws, columns)
    known_keys = set(columns) - {"extra_json"}
    for event in events:
        ev = _sanitize(event)
        extra: dict = {k: v for k, v in ev.items() if k not in known_keys}
        extra_str = _safe_str(json.dumps(extra, ensure_ascii=False) if extra else "", MAX_EXTRA_JSON_CHARS)
        row = [
            _safe_str(ev.get("timestamp")),
            _safe_str(ev.get("level", "INFO")),
            _safe_str(ev.get("event")),
            _safe_str(ev.get("chat_id")),
            _safe_str(ev.get("session_id")),
            _safe_str(ev.get("job_id")),
            _safe_str(ev.get("processing_run_id")),
            _safe_str(ev.get("safe_message"), MAX_PREVIEW_CHARS),
            _safe_str(ev.get("pdf_count")),
            _safe_str(ev.get("filename")),
            _safe_str(ev.get("project_name")),
            _safe_str(ev.get("error_type")),
            _safe_str(ev.get("error"), MAX_ERROR_CHARS),
            _safe_str(ev.get("stale_reason")),
            _safe_str(ev.get("command")),
            _safe_str(ev.get("admin_chat_id")),
            _safe_str(ev.get("target_chat_id")),
            _safe_str(ev.get("backup_path")),
            _safe_str(ev.get("local_path")),
            extra_str,
        ]
        ws.append(row)
    _append_no_data_if_empty(ws)


def _build_errors_sheet(ws, events: list[dict]) -> None:
    error_events = []
    for ev in events:
        level = str(ev.get("level", "")).upper()
        event_name = str(ev.get("event", "")).lower()
        if level in ("ERROR", "WARNING") or any(
            kw in event_name for kw in ("failed", "error", "stale", "access_denied")
        ):
            error_events.append(ev)

    columns = [
        "timestamp", "level", "event", "chat_id", "session_id", "job_id",
        "error_type", "error", "command", "backup_path", "safe_message", "extra_json",
    ]
    _write_header(ws, columns)
    known_keys = set(columns) - {"extra_json"}
    for ev in error_events:
        ev = _sanitize(ev)
        extra: dict = {k: v for k, v in ev.items() if k not in known_keys}
        extra_str = _safe_str(json.dumps(extra, ensure_ascii=False) if extra else "", MAX_EXTRA_JSON_CHARS)
        row = [
            _safe_str(ev.get("timestamp")),
            _safe_str(ev.get("level", "INFO")),
            _safe_str(ev.get("event")),
            _safe_str(ev.get("chat_id")),
            _safe_str(ev.get("session_id")),
            _safe_str(ev.get("job_id")),
            _safe_str(ev.get("error_type")),
            _safe_str(ev.get("error"), MAX_ERROR_CHARS),
            _safe_str(ev.get("command")),
            _safe_str(ev.get("backup_path")),
            _safe_str(ev.get("safe_message"), MAX_PREVIEW_CHARS),
            extra_str,
        ]
        ws.append(row)
    _append_no_data_if_empty(ws)


def _build_sessions_sheet(ws, sessions: list[dict]) -> None:
    columns = [
        "chat_id", "session_id", "status", "created_at", "updated_at",
        "status_changed_at", "done_requested_at", "queued_at",
        "processing_started_at", "completed_at", "failed_at", "cancelled_at",
        "last_pdf_added_at", "job_id", "processing_run_id", "pdf_count",
        "project_name", "spreadsheet_url", "last_error_type", "last_error",
        "stale_candidate", "stale_reason",
    ]
    _write_header(ws, columns)
    for session in sessions:
        reason = _stale_reason(session)
        row = [
            _safe_str(session.get("chat_id")),
            _safe_str(session.get("session_id")),
            _safe_str(session.get("status")),
            _safe_str(session.get("created_at")),
            _safe_str(session.get("updated_at")),
            _safe_str(session.get("status_changed_at")),
            _safe_str(session.get("done_requested_at")),
            _safe_str(session.get("queued_at")),
            _safe_str(session.get("processing_started_at")),
            _safe_str(session.get("completed_at")),
            _safe_str(session.get("failed_at")),
            _safe_str(session.get("cancelled_at")),
            _safe_str(session.get("last_pdf_added_at")),
            _safe_str(session.get("job_id")),
            _safe_str(session.get("processing_run_id")),
            _safe_str(_pdf_count(session)),
            _safe_str(session.get("project_name")),
            _safe_str(session.get("spreadsheet_url")),
            _safe_str(session.get("last_error_type")),
            _safe_str(session.get("last_error"), MAX_ERROR_CHARS),
            "yes" if reason else "no",
            reason or "",
        ]
        ws.append(row)
    _append_no_data_if_empty(ws)


def _build_command_logs_sheet(ws, logs: list[dict]) -> None:
    all_keys: list[str] = ["_chat_id", "_filename"]
    seen: set[str] = set(all_keys)
    for entry in logs:
        for k in entry.keys():
            if k not in seen:
                all_keys.append(k)
                seen.add(k)

    _write_header(ws, all_keys)
    for entry in logs:
        row = [_safe_str(entry.get(k), MAX_PREVIEW_CHARS) for k in all_keys]
        ws.append(row)
    _append_no_data_if_empty(ws)


def _build_archived_sessions_sheet(ws, archived: list[dict]) -> None:
    columns = [
        "archive_file", "archive_timestamp", "archive_reason_from_filename",
        "chat_id", "session_id", "status", "created_at", "updated_at",
        "status_changed_at", "done_requested_at", "queued_at",
        "processing_started_at", "completed_at", "failed_at", "cancelled_at",
        "stale_detected_at", "stale_reason",
        "job_id", "processing_run_id", "pdf_count", "project_name",
        "spreadsheet_url", "last_error_type", "last_error",
        "reset_by_admin_chat_id", "reset_at", "reset_reason",
    ]
    _write_header(ws, columns)
    for row_data in archived:
        row = [_safe_str(row_data.get(k), MAX_PREVIEW_CHARS) for k in columns]
        ws.append(row)
    _append_no_data_if_empty(ws)


def _build_summary_sheet(
    ws,
    period: str,
    events: list[dict],
    total_events_found: int,
    sessions: list[dict],
    archived: list[dict],
    command_logs: list[dict],
    generated_at: str,
) -> None:
    ws.append(["Параметр", "Значение"])
    ws["A1"].font = _bold_font()
    ws["B1"].font = _bold_font()

    events_exported = len(events)
    events_truncated = "yes" if total_events_found > MAX_EVENTS_ROWS else "no"

    error_count = sum(
        1 for ev in events
        if str(ev.get("level", "")).upper() in ("ERROR", "WARNING")
        or any(kw in str(ev.get("event", "")).lower() for kw in ("failed", "error", "stale", "access_denied"))
    )
    stale_count = sum(1 for s in sessions if _stale_reason(s))

    unique_chat_ids = len({str(ev.get("chat_id", "")) for ev in events if ev.get("chat_id")})
    unique_session_ids = len({str(ev.get("session_id", "")) for ev in events if ev.get("session_id")})
    unique_job_ids = len({str(ev.get("job_id", "")) for ev in events if ev.get("job_id")})

    event_type_counts = Counter(str(ev.get("event", "unknown")) for ev in events)
    top_event_types = event_type_counts.most_common(20)

    base_rows = [
        ("Период выгрузки", period),
        ("Сформирован", generated_at),
        ("", ""),
        ("events_total_found", str(total_events_found)),
        ("events_exported", str(events_exported)),
        ("events_truncated", events_truncated),
        ("unique_chat_ids (events)", str(unique_chat_ids)),
        ("unique_session_ids (events)", str(unique_session_ids)),
        ("unique_job_ids (events)", str(unique_job_ids)),
        ("Ошибок/предупреждений в events", str(error_count)),
        ("", ""),
        ("Активных сессий", str(len(sessions))),
        ("Сессий-кандидатов на stale", str(stale_count)),
        ("Архивных сессий", str(len(archived))),
        ("Фильтр архивных сессий", "all archived sessions included"),
        ("Лог-файлов команд", str(len(command_logs))),
        ("", ""),
        ("MAX_EVENTS_ROWS", str(MAX_EVENTS_ROWS)),
        ("DONE_REQUESTED_TIMEOUT_MINUTES", str(DONE_REQUESTED_TIMEOUT_MINUTES)),
        ("QUEUED_TIMEOUT_MINUTES", str(QUEUED_TIMEOUT_MINUTES)),
        ("PROCESSING_TIMEOUT_MINUTES", str(PROCESSING_TIMEOUT_MINUTES)),
        ("", ""),
        ("Топ типов событий (event)", "count"),
    ]

    for label, value in base_rows:
        ws.append([label, value])

    for event_name, count in top_event_types:
        ws.append([f"  {event_name}", str(count)])

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 40


# ── formatting helpers ────────────────────────────────────────────────────────

def _bold_font():
    from openpyxl.styles import Font
    return Font(bold=True)


def _write_header(ws, columns: list[str]) -> None:
    from openpyxl.styles import Font
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _auto_column_widths(ws, max_width: int = 80) -> None:
    for col_cells in ws.columns:
        width = 0
        for cell in col_cells:
            if cell.value:
                width = max(width, min(len(str(cell.value)), max_width))
        col_letter = col_cells[0].column_letter
        ws.column_dimensions[col_letter].width = max(width + 2, 10)


# ── main export function ──────────────────────────────────────────────────────

def export_logs_to_xlsx(
    period: str,
    events_dir: Path,
    sessions_dir: Path,
    backups_dir: Path,
    logs_dir: Path,
    output_dir: Path,
) -> Path:
    """
    Build a diagnostics Excel workbook and return the path to the created file.

    period: 'today' | 'yesterday' | '7d' | 'all'
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError("openpyxl is required: pip install openpyxl") from exc

    output_dir.mkdir(parents=True, exist_ok=True)

    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    filename = "bot_logs_" + datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S") + ".xlsx"
    output_path = output_dir / filename

    events, total_events_found = _read_events(events_dir, period)
    sessions = _read_active_sessions(sessions_dir)
    archived = _read_archived_sessions(backups_dir)
    command_logs = _read_command_logs(logs_dir, period)

    wb = openpyxl.Workbook()

    # Summary — first sheet (default)
    ws_summary = wb.active
    ws_summary.title = "Summary"
    _build_summary_sheet(
        ws_summary, period, events, total_events_found,
        sessions, archived, command_logs, generated_at,
    )

    ws_events = wb.create_sheet("Events")
    _build_events_sheet(ws_events, events)
    _auto_column_widths(ws_events)

    ws_errors = wb.create_sheet("Errors")
    _build_errors_sheet(ws_errors, events)
    _auto_column_widths(ws_errors)

    ws_sessions = wb.create_sheet("Sessions")
    _build_sessions_sheet(ws_sessions, sessions)
    _auto_column_widths(ws_sessions)

    ws_cmdlogs = wb.create_sheet("CommandLogs")
    _build_command_logs_sheet(ws_cmdlogs, command_logs)
    _auto_column_widths(ws_cmdlogs)

    ws_archived = wb.create_sheet("ArchivedSessions")
    _build_archived_sessions_sheet(ws_archived, archived)
    _auto_column_widths(ws_archived)

    wb.save(output_path)
    return output_path
