from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from constants import (
    NORMALIZED_JSON_FILENAME,
    REQUIRED_CALC_PRICE_KEYS,
    REQUIRED_PARAMETERS,
    SECTION_CODE,
    SECTION_NAME_RU,
    SHEET_01_NAME,
    SHEET_02_NAME,
    SHEET_03_NAME,
    STAGE1_JOBS_DIR,
)
from normalization import cell_text, display_number, is_blank, parse_boolean, parse_number, safe_float_sum


VALID_SELECTED_PRICE_SOURCES = {"price_registry", "fallback"}
VALID_EFFECTIVE_PRICE_SOURCES = {"price_registry", "fallback", "manual_override"}
REQUIRED_PARAMETER_POSITIVE = {
    "pit_area_m2",
    "pit_excavation_depth_m",
    "sand_base_volume_m3",
    "geotextile_area_m2",
    "geotextile_laying_area_m2",
}
REQUIRED_PARAMETER_NON_NEGATIVE = {
    "trench_volume_m3",
    "communications_length_m",
}


def resolve_workbook_path(workbook_path: str | Path) -> Path:
    path = Path(workbook_path)
    if path.exists():
        return path

    if path.name == "review_workbook.xlsx" and STAGE1_JOBS_DIR.exists():
        candidates = sorted(
            STAGE1_JOBS_DIR.glob("*/google/review_workbook.xlsx"),
            key=lambda item: item.stat().st_mtime,
            reverse=True,
        )
        if candidates:
            return candidates[0]

    raise FileNotFoundError(f"Workbook not found: {path}")


def find_header_row(ws, required_headers: list[str]) -> int:
    required = {cell_text(header).lower() for header in required_headers}
    for row_idx in range(1, ws.max_row + 1):
        row_values = {cell_text(ws.cell(row_idx, col).value).lower() for col in range(1, ws.max_column + 1)}
        if required.issubset(row_values):
            return row_idx
    raise ValueError(f"Cannot find header row for {required_headers} in sheet {ws.title}")


def header_map(ws, header_row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header = cell_text(ws.cell(header_row, col).value)
        if header:
            mapping[header] = col
    return mapping


def sheet_row_as_text(ws, row_idx: int, col_map: dict[str, int], header: str) -> str:
    return cell_text(ws.cell(row_idx, col_map[header]).value) if header in col_map else ""


def build_parameter_row(ws, row_idx: int, col_map: dict[str, int]) -> dict[str, Any] | None:
    technical_key = sheet_row_as_text(ws, row_idx, col_map, "technical_key")
    if not technical_key:
        return None
    if technical_key not in REQUIRED_PARAMETERS:
        return None

    original_value = sheet_row_as_text(ws, row_idx, col_map, "Найдено в проекте")
    override_value = sheet_row_as_text(ws, row_idx, col_map, "Исправить / ввести значение")
    selected_value = override_value if not is_blank(override_value) else original_value
    selected_column = "Исправить / ввести значение" if not is_blank(override_value) else "Найдено в проекте"
    parsed_value = parse_number(selected_value)
    value = parsed_value if parsed_value is not None else selected_value

    return {
        "value": value,
        "unit": sheet_row_as_text(ws, row_idx, col_map, "Ед."),
        "source_sheet": SHEET_01_NAME,
        "source_column": selected_column,
        "override_used": not is_blank(override_value),
        "original_value": original_value,
        "override_value": override_value,
    }


def read_parameters_sheet(wb) -> tuple[dict[str, dict[str, Any]], list[str]]:
    ws = wb[SHEET_01_NAME]
    header_row = find_header_row(ws, ["technical_key"])
    col_map = header_map(ws, header_row)

    parameters: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        row = build_parameter_row(ws, row_idx, col_map)
        if row is None:
            continue
        key = sheet_row_as_text(ws, row_idx, col_map, "technical_key")
        parameters[key] = row

    return parameters, warnings


def build_price_row(ws, row_idx: int, col_map: dict[str, int]) -> dict[str, Any] | None:
    estimate_line = sheet_row_as_text(ws, row_idx, col_map, "Строка сметы")
    if not estimate_line:
        return None

    price_registry_value = parse_number(sheet_row_as_text(ws, row_idx, col_map, "Цена из прайса"))
    fallback_value = parse_number(sheet_row_as_text(ws, row_idx, col_map, "Цена fallback"))
    price_for_calculation = parse_number(sheet_row_as_text(ws, row_idx, col_map, "Цена для расчета"))
    override_value_text = sheet_row_as_text(ws, row_idx, col_map, "Исправить цену")
    override_value = parse_number(override_value_text) if override_value_text else None
    selected_price = override_value if override_value is not None else price_for_calculation
    calc_price_key = sheet_row_as_text(ws, row_idx, col_map, "calc_price_key")
    price_registry_code = sheet_row_as_text(ws, row_idx, col_map, "price_registry_code")
    fallback_key = sheet_row_as_text(ws, row_idx, col_map, "fallback_key")
    selected_price_source = sheet_row_as_text(ws, row_idx, col_map, "selected_price_source").lower()
    price_source = sheet_row_as_text(ws, row_idx, col_map, "Источник цены")
    if selected_price_source not in VALID_SELECTED_PRICE_SOURCES:
        if price_registry_code:
            selected_price_source = "price_registry"
        elif "price_registry" in price_source.lower():
            selected_price_source = "price_registry"
        else:
            selected_price_source = "fallback"
    effective_price_source = "manual_override" if override_value is not None else selected_price_source

    return {
        "estimate_line": estimate_line,
        "price_role": sheet_row_as_text(ws, row_idx, col_map, "Что это за цена"),
        "unit": sheet_row_as_text(ws, row_idx, col_map, "Ед."),
        "price_registry_value": price_registry_value,
        "fallback_value": fallback_value,
        "price_for_calculation": price_for_calculation,
        "override_value": override_value,
        "selected_price": selected_price,
        "override_used": override_value is not None,
        "calc_price_key": calc_price_key,
        "price_registry_code": price_registry_code,
        "fallback_key": fallback_key,
        "selected_price_source": selected_price_source,
        "effective_price_source": effective_price_source,
        "price_source": price_source,
        "needs_attention": sheet_row_as_text(ws, row_idx, col_map, "Нужно внимание"),
        "comment": sheet_row_as_text(ws, row_idx, col_map, "Комментарий"),
    }


def read_prices_sheet(wb) -> list[dict[str, Any]]:
    ws = wb[SHEET_02_NAME]
    header_row = find_header_row(ws, ["Строка сметы", "Цена для расчета"])
    col_map = header_map(ws, header_row)

    prices: list[dict[str, Any]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        row = build_price_row(ws, row_idx, col_map)
        if row is None:
            continue
        prices.append(row)
    return prices


def build_trench_item(ws, row_idx: int, col_map: dict[str, int]) -> dict[str, Any] | None:
    item_type = sheet_row_as_text(ws, row_idx, col_map, "Тип")
    if item_type != "Траншея":
        return None

    return {
        "name": sheet_row_as_text(ws, row_idx, col_map, "Наименование"),
        "length_m": parse_number(sheet_row_as_text(ws, row_idx, col_map, "Длина, м")),
        "depth_m": parse_number(sheet_row_as_text(ws, row_idx, col_map, "Глубина, м")),
        "width_m": parse_number(sheet_row_as_text(ws, row_idx, col_map, "Ширина, м")),
        "volume_m3": parse_number(sheet_row_as_text(ws, row_idx, col_map, "Объем, м3")),
        "source": sheet_row_as_text(ws, row_idx, col_map, "Источник"),
        "fragment": sheet_row_as_text(ws, row_idx, col_map, "Фрагмент проекта"),
    }


def build_communication_item(ws, row_idx: int, col_map: dict[str, int]) -> dict[str, Any] | None:
    item_type = sheet_row_as_text(ws, row_idx, col_map, "Тип")
    if item_type != "Коммуникация":
        return None

    included_raw = sheet_row_as_text(ws, row_idx, col_map, "Включено")
    name = sheet_row_as_text(ws, row_idx, col_map, "Наименование")
    diameter_mm = int(parse_number(sheet_row_as_text(ws, row_idx, col_map, "Диаметр, мм")) or 0)
    if diameter_mm == 0 and name:
        # Extract diameter from pipe name: Ф110, ф110, Ø110, D110, 110мм
        m = re.search(r"[ФфØD][\s]?(\d+)|(\d+)\s*мм", name, re.IGNORECASE)
        if m:
            diameter_mm = int(m.group(1) or m.group(2))
    return {
        "name": name,
        "diameter_mm": diameter_mm,
        "pipe_length_m": parse_number(sheet_row_as_text(ws, row_idx, col_map, "Длина одной, м")),
        "quantity": parse_number(sheet_row_as_text(ws, row_idx, col_map, "Количество")),
        "total_length_m": parse_number(sheet_row_as_text(ws, row_idx, col_map, "Итоговая длина, м")),
        "included": parse_boolean(included_raw),
        "source": sheet_row_as_text(ws, row_idx, col_map, "Источник"),
        "fragment": sheet_row_as_text(ws, row_idx, col_map, "Фрагмент проекта"),
    }


def read_details_sheet(wb) -> dict[str, Any]:
    ws = wb[SHEET_03_NAME]
    header_row = find_header_row(
        ws,
        ["Тип", "Наименование", "Источник", "Фрагмент проекта"],
    )
    col_map = header_map(ws, header_row)

    trench_routes: list[dict[str, Any]] = []
    communications_pipe_items: list[dict[str, Any]] = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        trench_item = build_trench_item(ws, row_idx, col_map)
        if trench_item is not None:
            trench_routes.append(trench_item)
            continue

        comm_item = build_communication_item(ws, row_idx, col_map)
        if comm_item is not None:
            communications_pipe_items.append(comm_item)

    communications_total_length_m = safe_float_sum(
        [item.get("total_length_m") for item in communications_pipe_items]
    )

    return {
        "trench_routes": trench_routes,
        "communications_pipe_items": communications_pipe_items,
        "summary": {
            "trench_routes_count": len(trench_routes),
            "communications_pipe_items_count": len(communications_pipe_items),
            "communications_total_length_m": communications_total_length_m,
        },
    }


def find_price_row(
    prices: list[dict[str, Any]],
    estimate_line: str,
    price_role: str | None = None,
) -> dict[str, Any] | None:
    for row in prices:
        if row.get("estimate_line") != estimate_line:
            continue
        if price_role is not None and row.get("price_role") != price_role:
            continue
        return row
    return None


def validate_normalized_review(data: dict[str, Any]) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []

    parameters = data.get("parameters", {})
    prices = data.get("prices", [])
    details = data.get("details", {})
    summary = details.get("summary", {})

    for key in REQUIRED_PARAMETERS:
        parameter = parameters.get(key)
        if parameter is None:
            errors.append(f"missing required parameter: {key}")
            continue
        value = parse_number(parameter.get("value"))
        if value is None:
            errors.append(f"required parameter must be numeric: {key}")
            continue
        if key in REQUIRED_PARAMETER_POSITIVE and value <= 0:
            errors.append(f"required parameter must be > 0: {key}")
        if key in REQUIRED_PARAMETER_NON_NEGATIVE and value < 0:
            errors.append(f"required parameter must be >= 0: {key}")

    calc_price_keys = [cell_text(row.get("calc_price_key")) for row in prices]
    if any(not key for key in calc_price_keys):
        errors.append("all price rows must have calc_price_key")
    if len(set(calc_price_keys)) != len(prices):
        errors.append("calc_price_key values must be unique across price rows")
    price_by_key = {cell_text(row.get("calc_price_key")): row for row in prices if cell_text(row.get("calc_price_key"))}
    missing_calc_keys = [key for key in REQUIRED_CALC_PRICE_KEYS if key not in price_by_key]
    if missing_calc_keys:
        errors.append(f"missing required calc_price_keys: {', '.join(missing_calc_keys)}")

    for row in prices:
        calc_price_key = cell_text(row.get("calc_price_key"))
        selected_source = cell_text(row.get("selected_price_source")).lower()
        effective_source = cell_text(row.get("effective_price_source")).lower()
        price_registry_code = cell_text(row.get("price_registry_code"))
        fallback_key = cell_text(row.get("fallback_key"))
        override_used = bool(row.get("override_used"))
        selected_price = parse_number(row.get("selected_price"))
        if selected_source not in VALID_SELECTED_PRICE_SOURCES:
            errors.append(f"invalid selected_price_source for {calc_price_key}: {selected_source}")
        if effective_source not in VALID_EFFECTIVE_PRICE_SOURCES:
            errors.append(f"invalid effective_price_source for {calc_price_key}: {effective_source}")
        if selected_price is None:
            errors.append(f"selected_price must be numeric for {calc_price_key}")
        elif selected_price < 0:
            errors.append(f"selected_price must be >= 0 for {calc_price_key}")
        if override_used and effective_source != "manual_override":
            errors.append(f"{calc_price_key}: override_used rows must have effective_price_source manual_override")
        if not override_used and effective_source != selected_source:
            errors.append(
                f"{calc_price_key}: effective_price_source must equal selected_price_source when no override is used"
            )
        if selected_source == "fallback":
            if not fallback_key.startswith("fallback."):
                errors.append(f"{calc_price_key}: fallback rows must have fallback_key starting with fallback.")
            if price_registry_code:
                errors.append(f"{calc_price_key}: fallback rows must not have price_registry_code")
        if selected_source == "price_registry" and not price_registry_code:
            warnings.append(f"price_registry_code is empty for price_registry row: {calc_price_key}")

    trench_routes = details.get("trench_routes", [])
    if not isinstance(trench_routes, list):
        errors.append("trench_routes must be a list")
        trench_routes = []
    for index, item in enumerate(trench_routes):
        if not isinstance(item, dict):
            errors.append(f"trench_routes[{index}] must be a mapping")
            continue
        if not cell_text(item.get("name")):
            errors.append(f"trench_routes[{index}].name is required")
        for field in ("length_m", "depth_m", "width_m", "volume_m3"):
            value = parse_number(item.get(field))
            if value is None:
                errors.append(f"trench_routes[{index}].{field} must be numeric")
            elif value < 0:
                errors.append(f"trench_routes[{index}].{field} must be >= 0")

    communications = details.get("communications_pipe_items", [])
    if not isinstance(communications, list):
        errors.append("communications_pipe_items must be a list")
        communications = []
    for index, item in enumerate(communications):
        if not isinstance(item, dict):
            errors.append(f"communications_pipe_items[{index}] must be a mapping")
            continue
        if not cell_text(item.get("name")):
            errors.append(f"communications_pipe_items[{index}].name is required")
        total_length = parse_number(item.get("total_length_m"))
        if total_length is None:
            errors.append(f"communications_pipe_items[{index}].total_length_m must be numeric")
        elif total_length < 0:
            errors.append(f"communications_pipe_items[{index}].total_length_m must be >= 0")
        diameter = item.get("diameter_mm")
        if diameter not in {None, ""}:
            diameter_value = parse_number(diameter)
            if diameter_value is None:
                errors.append(f"communications_pipe_items[{index}].diameter_mm must be numeric")
            elif diameter_value <= 0:
                errors.append(f"communications_pipe_items[{index}].diameter_mm must be > 0")
        quantity = item.get("quantity")
        if quantity not in {None, ""}:
            quantity_value = parse_number(quantity)
            if quantity_value is None:
                errors.append(f"communications_pipe_items[{index}].quantity must be numeric")
            elif quantity_value < 0:
                errors.append(f"communications_pipe_items[{index}].quantity must be >= 0")
        included = item.get("included")
        included_bool = included if isinstance(included, bool) else parse_boolean(str(included))
        if included_bool is None:
            errors.append(f"communications_pipe_items[{index}].included must be boolean-like")
        item["included"] = bool(included_bool)

    total_length = safe_float_sum(
        [row.get("total_length_m") for row in communications if row.get("included") is True]
    )
    summary_total = parse_number(summary.get("communications_total_length_m"))
    if summary_total is None:
        errors.append("communications_total_length_m must be numeric")
    elif abs(summary_total - total_length) > 0.001:
        errors.append(
            f"communications_total_length_m must equal sum of pipe rows ({total_length}), got {summary_total}"
        )

    return errors, warnings


def build_review_data(workbook_path: str | Path) -> dict[str, Any]:
    resolved_path = resolve_workbook_path(workbook_path)
    wb = load_workbook(resolved_path, data_only=True)

    parameters, parameter_warnings = read_parameters_sheet(wb)
    prices = read_prices_sheet(wb)
    details = read_details_sheet(wb)

    data = {
        "meta": {
            "source_type": "local_review_workbook",
            "workbook_path": str(resolved_path),
            "section_code": SECTION_CODE,
            "section_name": SECTION_NAME_RU,
            "created_at": datetime.now().isoformat(timespec="seconds"),
        },
        "parameters": parameters,
        "prices": prices,
        "details": details,
        "validation": {
            "errors": [],
            "warnings": parameter_warnings,
        },
    }

    errors, warnings = validate_normalized_review(data)
    data["validation"]["errors"] = errors
    data["validation"]["warnings"].extend(warnings)
    return data


def render_review_report(data: dict[str, Any]) -> str:
    meta = data["meta"]
    parameters = data.get("parameters", {})
    prices = data.get("prices", [])
    details = data.get("details", {})
    validation = data.get("validation", {})
    summary = details.get("summary", {})

    lines = [
        "# Review reader report",
        "",
        "## Source",
        f"- workbook: {meta.get('workbook_path', '')}",
        f"- section: {meta.get('section_code', '')}",
        "",
        "## Parameters",
    ]

    for key in REQUIRED_PARAMETERS:
        parameter = parameters.get(key, {})
        value = parameter.get("value")
        unit = parameter.get("unit", "")
        lines.append(f"- {key}: {display_number(value)} {unit}".rstrip())

    lines.extend(
        [
            "",
            "## Prices",
            f"- rows read: {len(prices)}",
            f"- overrides used: {sum(1 for row in prices if row.get('override_used'))}",
            f"- rows needing attention: {sum(1 for row in prices if cell_text(row.get('needs_attention')).lower() not in {'', 'нет'})}",
            "",
            "## Price keys",
            f"- rows with calc_price_key: {sum(1 for row in prices if cell_text(row.get('calc_price_key')))}",
            f"- selected_price_source=fallback: {sum(1 for row in prices if cell_text(row.get('selected_price_source')).lower() == 'fallback')}",
            f"- selected_price_source=price_registry: {sum(1 for row in prices if cell_text(row.get('selected_price_source')).lower() == 'price_registry')}",
            f"- manual overrides: {sum(1 for row in prices if row.get('override_used'))}",
            f"- missing calc_price_key: {sum(1 for row in prices if not cell_text(row.get('calc_price_key')))}",
            "",
            "### Price key samples",
            f"- axis_marking_work_unit_price: {cell_text(next((row for row in prices if cell_text(row.get('calc_price_key')) == 'axis_marking_work_unit_price'), {}).get('selected_price_source') if next((row for row in prices if cell_text(row.get('calc_price_key')) == 'axis_marking_work_unit_price'), None) else '')} / {cell_text(next((row for row in prices if cell_text(row.get('calc_price_key')) == 'axis_marking_work_unit_price'), {}).get('fallback_key') if next((row for row in prices if cell_text(row.get('calc_price_key')) == 'axis_marking_work_unit_price'), None) else '')}",
            f"- geotextile_material_unit_price: {cell_text(next((row for row in prices if cell_text(row.get('calc_price_key')) == 'geotextile_material_unit_price'), {}).get('selected_price_source') if next((row for row in prices if cell_text(row.get('calc_price_key')) == 'geotextile_material_unit_price'), None) else '')} / {cell_text(next((row for row in prices if cell_text(row.get('calc_price_key')) == 'geotextile_material_unit_price'), {}).get('price_registry_code') if next((row for row in prices if cell_text(row.get('calc_price_key')) == 'geotextile_material_unit_price'), None) else '')}",
            f"- consumables_amount: {cell_text(next((row for row in prices if cell_text(row.get('calc_price_key')) == 'consumables_amount'), {}).get('selected_price_source') if next((row for row in prices if cell_text(row.get('calc_price_key')) == 'consumables_amount'), None) else '')} / {cell_text(next((row for row in prices if cell_text(row.get('calc_price_key')) == 'consumables_amount'), {}).get('fallback_key') if next((row for row in prices if cell_text(row.get('calc_price_key')) == 'consumables_amount'), None) else '')}",
            "## Details",
            f"- trench_routes: {len(details.get('trench_routes', []))}",
            f"- communications_pipe_items: {len(details.get('communications_pipe_items', []))}",
            f"- communications_total_length_m: {display_number(summary.get('communications_total_length_m'))}",
            "",
            "## Validation",
            f"- errors: {len(validation.get('errors', []))}",
            f"- warnings: {len(validation.get('warnings', []))}",
        ]
    )

    if validation.get("errors"):
        lines.extend(["", "### Errors"])
        lines.extend(f"- {error}" for error in validation["errors"])
    if validation.get("warnings"):
        lines.extend(["", "### Warnings"])
        lines.extend(f"- {warning}" for warning in validation["warnings"])

    return "\n".join(lines) + "\n"


def dump_review_json(data: dict[str, Any], output_path: Path) -> None:
    output_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
