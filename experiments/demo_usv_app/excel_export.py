from __future__ import annotations

import html
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


APP_DIR = Path(__file__).resolve().parent
LOGO_PATH = APP_DIR / "assets" / "brick_house_logo.png"


WHITE_HEADERS = [
    "№",
    "Наименование работ",
    "Ед. изм.",
    "Кол-во",
    "Материалы/механизмы цена",
    "Материалы/механизмы сумма",
    "Работы цена",
    "Работы сумма",
    "Итого",
]
GREY_HEADERS = [
    "Кол-во",
    "Материалы/механизмы цена",
    "Материалы/механизмы сумма",
    "Работы цена",
    "Работы сумма",
    "Итого",
]


def normalize_section_title(section_title: str) -> str:
    return section_title.replace("(250мм, 300мм)", "(250мм,300мм)")


def excel_address(address: str) -> str:
    text = re.sub(r"<br\s*/?>", " ", address, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    text = html.unescape(text)
    return " ".join(text.split())


def qty(line: dict[str, Any]) -> Any:
    return line.get("display_quantity", line.get("quantity"))


def make_workbook(result: dict[str, Any], section_title: str, address: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Фундаментная плита"

    grey_fill = PatternFill("solid", fgColor="B7B7B7")
    header_fill = PatternFill("solid", fgColor="FFFFFF")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="EFEFEF")
    vertical_fill = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="303030")
    blue = Side(style="thick", color="1D4ED8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    split_border = Border(left=blue, right=thin, top=thin, bottom=thin)
    section_label = normalize_section_title(section_title)

    for row in range(1, 8):
        for col in range(10, 16):
            sheet.cell(row, col).fill = grey_fill

    sheet.merge_cells("A1:I4")
    if LOGO_PATH.exists():
        logo = OpenpyxlImage(LOGO_PATH)
        logo.width = 300
        logo.height = 50
        sheet.add_image(logo, "C2")
    else:
        sheet["A1"] = "Brick House"
        sheet["A1"].font = Font(size=18, bold=True)
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells("A5:I5")
    sheet["A5"] = "СМЕТНЫЙ РАСЧЕТ НА СТРОИТЕЛЬСТВО ДОМА"
    sheet["A5"].font = Font(size=12, bold=True)
    sheet["A5"].alignment = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells("A6:I6")
    sheet["A6"] = excel_address(address)
    sheet["A6"].font = Font(size=10, bold=False)
    sheet["A6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    sheet.row_dimensions[6].height = 24

    group_header_row = 8
    subheader_row = 9

    header_merges = [
        (group_header_row, 1, subheader_row, 1, "№"),
        (group_header_row, 2, subheader_row, 2, "Наименование работ"),
        (group_header_row, 3, subheader_row, 3, "Ед. изм."),
        (group_header_row, 4, subheader_row, 4, "Кол-во"),
        (
            group_header_row,
            5,
            group_header_row,
            6,
            "Стоимость материалов,\nмашин и механизмов, руб.",
        ),
        (group_header_row, 7, group_header_row, 8, "Стоимость работ, руб."),
        (group_header_row, 9, subheader_row, 9, "Итого, руб."),
        (group_header_row, 10, subheader_row, 10, "Кол-во"),
        (
            group_header_row,
            11,
            group_header_row,
            12,
            "Стоимость материалов,\nмашин и механизмов, руб.",
        ),
        (group_header_row, 13, group_header_row, 14, "Стоимость работ, руб."),
        (group_header_row, 15, subheader_row, 15, "Итого, руб."),
    ]
    for start_row, start_col, end_row, end_col, value in header_merges:
        sheet.merge_cells(
            start_row=start_row,
            start_column=start_col,
            end_row=end_row,
            end_column=end_col,
        )
        sheet.cell(start_row, start_col).value = value

    subheaders = {
        5: "За ед",
        6: "Итого",
        7: "За ед",
        8: "Итого",
        11: "За ед",
        12: "Итого",
        13: "За ед",
        14: "Итого",
    }
    for col, value in subheaders.items():
        sheet.cell(subheader_row, col).value = value

    for row in [group_header_row, subheader_row]:
        for col in range(1, 16):
            cell = sheet.cell(row, col)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = split_border if col == 10 else border
            cell.fill = grey_fill if col >= 10 else header_fill
    sheet.row_dimensions[group_header_row].height = 42
    sheet.row_dimensions[subheader_row].height = 22

    section_row = subheader_row + 1
    for col in range(1, 16):
        cell = sheet.cell(section_row, col)
        cell.fill = grey_fill if col >= 10 else section_fill
        cell.border = split_border if col == 10 else border
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.cell(section_row, 1).value = 3
    sheet.cell(section_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.cell(section_row, 2).value = section_label
    sheet.cell(section_row, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    estimate_lines = result.get("estimate_lines", [])
    first_estimate_row = section_row + 1
    row_index = first_estimate_row
    for line in estimate_lines:
        values = [
            None,
            line.get("name"),
            line.get("unit"),
            qty(line),
            line.get("material_unit_price"),
            line.get("material_total"),
            line.get("work_unit_price"),
            line.get("work_total"),
            line.get("line_total"),
            qty(line),
            line.get("material_unit_price"),
            line.get("material_total"),
            line.get("work_unit_price"),
            line.get("work_total"),
            line.get("line_total"),
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col)
            cell.value = value
            cell.border = split_border if col == 10 else border
            cell.alignment = Alignment(
                horizontal="center" if col in {1, 3, 4, 10} else None,
                vertical="top",
                wrap_text=col == 2,
            )
            if col >= 10:
                cell.fill = grey_fill
        row_index += 1

    last_estimate_row = row_index - 1
    if estimate_lines and first_estimate_row <= last_estimate_row:
        sheet.merge_cells(
            start_row=first_estimate_row,
            start_column=1,
            end_row=last_estimate_row,
            end_column=1,
        )
        section_cell = sheet.cell(first_estimate_row, 1)
        section_cell.value = section_label
        section_cell.font = Font(bold=True)
        section_cell.fill = vertical_fill
        section_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            text_rotation=90,
            wrap_text=True,
        )
        section_cell.border = border
        for merged_row in range(first_estimate_row + 1, last_estimate_row + 1):
            sheet.cell(merged_row, 1).border = border
            sheet.cell(merged_row, 1).fill = vertical_fill

    totals = result.get("internal_totals", {})
    total_values = [
        "",
        "Итого по разделу:",
        "",
        "",
        "",
        totals.get("internal_materials_total"),
        "",
        totals.get("internal_works_total"),
        totals.get("internal_section_total"),
        "",
        "",
        totals.get("internal_materials_total"),
        "",
        totals.get("internal_works_total"),
        totals.get("internal_section_total"),
    ]
    for col, value in enumerate(total_values, start=1):
        cell = sheet.cell(row_index, col)
        cell.value = value
        cell.font = Font(bold=True)
        cell.fill = grey_fill if col >= 10 else total_fill
        cell.border = split_border if col == 10 else border

    widths = {
        1: 6,
        2: 56,
        3: 10,
        4: 12,
        5: 16,
        6: 16,
        7: 14,
        8: 14,
        9: 14,
        10: 12,
        11: 16,
        12: 16,
        13: 14,
        14: 14,
        15: 14,
    }
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width
    frozen_columns = 4
    sheet.freeze_panes = f"{get_column_letter(frozen_columns + 1)}{first_estimate_row}"
    return workbook


def build_excel_bytes(result: dict[str, Any], section_title: str, address: str) -> bytes:
    workbook = make_workbook(result, section_title, address)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def save_excel(path: Path, result: dict[str, Any], section_title: str, address: str) -> bytes:
    payload = build_excel_bytes(result, section_title, address)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)
    return payload
