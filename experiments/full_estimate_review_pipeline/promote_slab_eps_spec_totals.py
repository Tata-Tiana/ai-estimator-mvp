"""Promote explicitly verified raw EPS totals into new extraction/review copies."""
from __future__ import annotations

import argparse
from copy import copy, deepcopy
import json
import math
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

sys.path.insert(0, str(Path(__file__).resolve().parent / "review_to_calculator"))
from core.workbook_reader import JSON_COLUMN_LETTER  # noqa: E402


FIELD = "eps_material_spec_volume_m3"


def promote(extraction_path, workbook_path, out_extraction, out_workbook, totals):
    outputs = [Path(out_extraction), Path(out_workbook)]
    if outputs[0].resolve() == outputs[1].resolve() or any(p.exists() for p in outputs):
        raise ValueError("Output paths must be distinct new files; existing files are never overwritten")
    extraction = json.loads(Path(extraction_path).read_text(encoding="utf-8"))
    section = extraction["sections"]["floor_slabs"]
    found = section["found"]
    evidence_by_zone = {}
    for zone_id, table_id, volume in totals:
        if zone_id in evidence_by_zone or not math.isfinite(volume) or volume < 0:
            raise ValueError("Each zone needs one finite, nonnegative verified volume")
        raw = [row for row in section["raw_table_rows"] if row.get("table_id") == table_id]
        zones = [row for row in found if row.get("target_code") == "floor_slab_zones"
                 and row.get("value", {}).get("zone_id") == zone_id]
        if len(raw) != 1 or len(zones) != 1:
            raise ValueError(f"Expected exactly one raw table and zone for {zone_id}/{table_id}")
        evidence = raw[0]
        if evidence.get("normalized_unit") != "m3":
            raise ValueError(f"Raw total {table_id} must be explicitly in m3")
        previous = zones[0]["value"].get(FIELD)
        if previous is not None and previous != volume:
            raise ValueError(f"Conflicting existing spec volume for {zone_id}")
        zones[0]["value"][FIELD] = volume
        source = f"{evidence['source_pdf']}, стр. {evidence['page_number']}: {evidence['raw_text']}"
        zones[0]["notes"] = (zones[0].get("notes") or "") + f"\n{FIELD}={volume:g}; {source}"
        evidence["mapped_target_codes"] = list(dict.fromkeys(
            evidence.get("mapped_target_codes", []) + ["floor_slab_zones"]
        ))
        evidence["notes"] = (evidence.get("notes") or "") + f"\nPromoted to {zone_id}.{FIELD}"
        evidence_by_zone[zone_id] = (volume, evidence)

    wb = load_workbook(workbook_path)
    ws = wb["01_Проверка проекта"]
    headers = next({cell.value: cell.column for cell in row if cell.value is not None}
                   for row in ws.iter_rows()
                   if any(cell.value == "technical_key" for cell in row))
    json_col = column_index_from_string(JSON_COLUMN_LETTER)
    templates = {}
    for row in ws.iter_rows():
        if row[headers["section_code"] - 1].value != "floor_slabs":
            continue
        if row[headers["technical_key"] - 1].value != "floor_slab_zones":
            continue
        cell = row[json_col - 1]
        item = json.loads(cell.value)
        zone_id = item.get("zone_id")
        if zone_id not in evidence_by_zone:
            continue
        if not item.get("_sheet_item_id"):
            raise ValueError("Migration needs a field-per-row workbook with _sheet_item_id")
        volume, _ = evidence_by_zone[zone_id]
        if item.get(FIELD) is not None and item[FIELD] != volume:
            raise ValueError(f"Conflicting review spec volume for {zone_id}")
        item[FIELD] = volume
        templates.setdefault(zone_id, (row[0].row, deepcopy(item)))
        cell.value = json.dumps(item, ensure_ascii=False)
    if set(templates) != set(evidence_by_zone):
        raise ValueError("Every requested zone must exist in the review workbook")

    # Append new review fields rather than shifting existing rows and their manual overrides.
    for zone_id, (volume, evidence) in evidence_by_zone.items():
        source_row, item = templates[zone_id]
        row_idx = ws.max_row + 1
        for source_cell in ws[source_row]:
            target = ws.cell(row_idx, source_cell.column, source_cell.value)
            target._style = copy(source_cell._style)
            target.alignment = copy(source_cell.alignment)
            target.number_format = source_cell.number_format
        ws.row_dimensions[row_idx].height = ws.row_dimensions[source_row].height
        item["_sheet_field_key"] = FIELD
        values = {
            "Что проверяем": f"{item.get('display_name', zone_id)} — ЭППС 100 мм, готовый объем материала",
            "Найдено в проекте": volume, "Ед.": "м3", "Статус": "Найдено",
            "Уверенность": evidence.get("confidence"),
            "Что нужно сделать": "Готовый итог материала заменяет сумму частей; работы считаются отдельно.",
            "Источник": f"{evidence['source_pdf']}, стр. {evidence['page_number']}",
            "Фрагмент проекта": evidence["raw_text"],
            "Исправить / ввести значение": None, "Комментарий Елены": None,
        }
        for name, value in values.items():
            ws.cell(row_idx, headers[name]).value = value
        for col in range(15, json_col):
            ws.cell(row_idx, col).value = None
        ws.cell(row_idx, json_col).value = json.dumps(item, ensure_ascii=False)

    for path in outputs:
        path.parent.mkdir(parents=True, exist_ok=True)
    outputs[0].write_text(json.dumps(extraction, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    wb.save(outputs[1])


def verified_total(value):
    try:
        zone_id, table_id, volume = value.split(":")
        return zone_id, table_id, float(volume)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Use zone_id:raw_table_id:verified_volume_m3") from exc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--extraction", required=True)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--out-extraction", required=True)
    parser.add_argument("--out-workbook", required=True)
    parser.add_argument("--total", type=verified_total, action="append", required=True,
                        help="Explicitly verified zone_id:raw_table_id:volume_m3; no guessing from raw text")
    args = parser.parse_args()
    promote(args.extraction, args.workbook, args.out_extraction, args.out_workbook, args.total)
    print(args.out_extraction)
    print(args.out_workbook)


if __name__ == "__main__":
    main()
