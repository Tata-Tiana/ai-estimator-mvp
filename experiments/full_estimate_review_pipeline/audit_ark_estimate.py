"""Explicit ARK row mapping for a read-only estimate audit against Elena's workbook."""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from export_calculator_results_to_estimate_workbook import (
    _cost_parts, _load_floor_slabs_blocks, _load_lines, _num,
)


# Explicit semantic matches, not fuzzy name matching. Multiple engine rows may map to one
# Elena row (rebar pooling, separate U/monolithic lintels, multiple parapet zones).
MAPPINGS = {
    "earthworks": dict(zip([
        "axis_marking", "excavator_jcb", "manual_excavation", "geotextile_laying",
        "geotextile_material", "sand_filling", "sand_material", "sand_manual_moving",
        "communications_work", "communications_material", "consumables", "technical_supervision",
        "procurement_warehouse_costs", "overhead_general_business_costs", "estimated_profit",
    ], range(26, 41))),
    "foundation_slab": {
        "planter_membrane_installation": 76, "planter_standard_material": 77,
        "planterband_material": 78, "eps50_laying_under_slab": 79,
        "thermal_insert_items_installation": 80, "thermal_insert_material_item_0": 81,
        "eps50_penoplex_geo_material": 82, "formwork_installation": 83,
        "formwork_plywood": 84, "formwork_timber": 85, "rebar_frame_assembly": 86,
        "rebar_a500_d10_m_pooled": 88, "foundation_slab_rebar_a500_d12_m_8": 87,
        "foundation_slab_rebar_a240_d6_m_9": 89, "foundation_slab_concreting_work": 90,
        "concrete_b22_5_m300_material": 91, "concrete_delivery": 92, "concrete_pump_32m": 93,
        "formwork_dismantling": 94, "logistics_and_supply": 96, "consumables_tool_amortization": 97,
        "technical_supervision": 98, "procurement_warehouse_costs_excel_structure": 99,
        "overhead_general_business_costs_excel_structure": 100, "estimated_profit_excel_structure": 101,
    },
    "waterproofing": dict(zip([
        "waterproofing_bitumen_mastic_work", "bitumen_primer_aquamast_18l",
        "bitumen_mastic_aquamast_18kg", "eps100_wall_insulation_work",
        "eps100_wall_penoplex_geo_material", "eps_glue_foam", "waterproofing_logistics_and_supply",
        "waterproofing_consumables_tool_amortization", "technical_supervision",
        "procurement_warehouse_costs", "overhead_general_business_costs", "estimated_profit",
    ], range(104, 116))),
    "load_bearing_walls_lintels_p6": {
        "scaffolding_setup_dismantling": 137, "scaffolding_timber_material": 138,
        "main_walls_cutoff_waterproofing_under_first_row_blocks": 140, "main_walls_masonry_work": 141,
        "main_walls_block_outer_d400": 142, "main_walls_block_inner_d500": 143,
        "main_walls_u_block_lintel_cutting": 144, "main_walls_sand_concrete_m300_first_row": 145,
        "main_walls_block_adhesive": 146, "main_walls_chasing_for_reinforcement": 147,
        "main_walls_masonry_chasing_rebar_а500с_d10": 148, "main_walls_other_rebar_а500с_d10": 148,
        "main_walls_blocks_delivery": 149, "main_walls_blocks_unloading_manipulator": 150,
        "main_walls_blocks_crane_moving": 151, "main_walls_lintel_formwork_installation": 152,
        "main_walls_lintel_formwork_plywood_material": 153, "main_walls_lintel_formwork_timber_material": 154,
        "main_walls_lintel_rebar_frame_assembly": 155, "main_walls_lintels_rebar_а500с_d16": 156,
        "main_walls_lintels_rebar_а500с_d12": 157, "main_walls_lintels_rebar_а240_d6": 158,
        "main_walls_u_block_lintel_concreting_work": 159, "main_walls_monolithic_lintel_concreting_work": 159,
        "main_walls_lintel_concrete_b22_5_m300_material": 160, "main_walls_lintel_concrete_delivery": 161,
        "main_walls_manual_concrete_lifting": 162, "main_walls_lintel_formwork_dismantling": 163,
        "main_walls_lintel_edge_insulation_work": 164, "main_walls_lintel_edge_insulation_eps_material": 165,
        "main_walls_lintel_edge_insulation_glue_foam": 166, "superstructure_masonry_work": 168,
        "superstructure_block_superstructure_d400": 169, "superstructure_block_adhesive": 170,
        "superstructure_chasing_for_reinforcement": 171, "superstructure_masonry_chasing_rebar_а500с_d10": 172,
        "upper_parapet_vent_blocks_delivery": 173, "upper_parapet_vent_blocks_unloading_manipulator": 174,
        "upper_parapet_vent_blocks_crane_moving": 175, "parapet_masonry_work": 177,
        "vent_chimney_cladding_gas_block_cladding_work": 178, "parapet_block_parapet_pm1_d400": 179,
        "vent_chimney_cladding_block_vent_d400": 180, "parapet_block_adhesive": 181,
        "parapet_chasing_for_reinforcement": 182, "parapet_parapet_chasing_rebar_а500с_d10": 183,
        "parapet_blocks_crane_moving": 184, "walls_consumables_tool_amortization": 185,
        "construction_waste_removal": 186, "walls_technical_supervision": 187,
        "procurement_warehouse_costs": 188, "overhead_general_business_costs": 189, "estimated_profit": 190,
    },
    "schiedel_vent_channels": {
        "schiedel_masonry_work": 279, "schiedel_vent_channel_4x": 280, "schiedel_vent_channel_3x": 281,
        "schiedel_delivery": 282, "schiedel_consumables_tool_depreciation": 283,
        "technical_supervision_zero": 284, "procurement_storage_zero": 285, "overhead_zero": 286,
        "profit_zero": 287,
    },
    "flat_roof": {
        "roof_base_preparation_control": 290, "vapor_barrier_installation": 291,
        "vapor_barrier_film_technonikol_120mk": 292, "eps_roof_insulation_installation": 293,
        "eps100_technonikol_carbon_eco": 294, "eps50_technonikol_carbon_eco": 295,
        "eps_slope_2_1_plate_a": 296, "eps_slope_2_1_plate_b": 297, "eps_slope_4_2_plate_j": 298,
        "eps_slope_4_2_plate_k": 299, "fiberglass_mat_technonikol_100gr": 300,
        "geotextile_prof_150_parapet": 301, "pvc_membrane_flat_installation": 302,
        "pvc_membrane_abutment_installation": 303, "roof_abutment_strip_installation_control": 304,
        "aluminum_pressure_rail_3m": 305, "aluminum_edge_rail_3m": 306,
        "pvc_membrane_logicroof_vrp_1_5mm_gray": 307, "roof_pvc_aerator_75x375": 308,
        "internal_roof_drain_with_heating": 309, "parapet_roof_drain_installation": 310,
        "internal_drain_pvc_110mm": 311, "roof_crane_lifting": 312, "roof_logistics_and_supply": 313,
        "roof_consumables_tool_depreciation": 314, "roof_waste_removal": 315,
        "technical_supervision": 316, "procurement_storage": 317, "overhead_zero": 318, "profit_zero": 319,
    },
}

SLAB_MAP = {
    "slab_formwork_installation_control": (212, 248), "formwork_set_rental_material": (213, 249),
    "formwork_delivery_return_manipulator": (214, 250), "formwork_rebar_crane_supply": (215, 251),
    "formwork_consumables": (216, 252), "edge_beam_formwork_installation_control": (217, 253),
    "plywood_fk_18mm_for_edges_and_non_multiple_places": (218, 254), "formwork_timber_gost": (219, 255),
    "floor_slab_rebar_frame_assembly_control": (222, 256), "rebar_a500c_d20_z1": (223, 257),
    "rebar_a500c_d16_z1": (224, 258), "rebar_a500c_d12_z1": (225, None),
    "rebar_a500c_d10_z1": (226, 259), "rebar_a240_d6_z1": (227, 260),
    "rebar_metal_delivery": (228, None), "beam_concreting_work_tall": (229, None),
    "beam_concreting_work": (230, 261), "floor_slab_concreting_work": (231, 262),
    "concrete_b22_5_m300_material": (232, 263), "concrete_delivery": (233, 264),
    "concrete_pump_32m": (234, 265), "formwork_dismantling_zero_internal": (235, 266),
    "bottom_slab_insulation_work": (236, 267), "edge_beam_insulation_work": (237, 268),
    "eps_penoplex_osnova_100mm": (238, 269), "eps_glue_foam": (239, 270),
    "logistics_and_supply": (240, 271), "consumables_tool_depreciation": (241, 272),
    "technical_supervision": (242, 273), "procurement_warehouse_costs": (243, 274),
    "overhead_general_business_costs": (244, 275), "estimated_profit": (245, 276),
}

RANGES = {
    "earthworks": (26, 40, 41), "foundation_slab": (76, 101, 102),
    "waterproofing": (104, 115, 116), "load_bearing_walls_lintels_p6": (137, 190, 191),
    "pm1": (212, 245, 246), "pm2": (248, 276, 277),
    "schiedel_vent_channels": (279, 287, 288), "flat_roof": (290, 319, 320),
}

NOTES = {
    28: "Ручная глубина траншей 0,4 м сохранена; небольшая разница базы, ставка 1400 вместо 1300.",
    38: "ЗСР у нас ноль, у Елены 5000; отчет 01 отдельно оговаривает эту строку.",
    77: "Отчет 02: перенос остатка PLANTER у Елены, универсальная закупка у нас 9 рулонов.",
    78: "Отчет 02: подтверждено универсальное правило 4 ленты на рулон.",
    85: "Отчет 02: у нас доска 50 мм и округление до 0,1 м3; у Елены расчет 40 мм без округления.",
    95: "Строки нет у нас; отчет 02 оставил как исключение АРК, не универсальную норму.",
    104: "Разный охват: у Елены ростверк + плита, у нас только плита; отчет 03.",
    105: "Разный охват гидроизоляции, нельзя свести к разнице цены.",
    106: "Разный охват гидроизоляции, нельзя свести к разнице цены.",
    107: "У Елены ростверк + плита; у нас только плита.",
    108: "Разный охват утепления.", 109: "Производная от разного охвата утепления.",
    140: "Отчет 04: Елена подтвердила приоритет готовой площади 43,7 м2 из проекта.",
    142: "Одинаковая D400 паллета 2,15 м3; остатки общие с надстройкой и парапетами, отчет 08.",
    144: "Отчет 04: потеряна перемычка 1,7 м у Елены; у нас 8,5/0,6 округлено до 15 резов.",
    145: "Производная от подтвержденной проектной площади первого ряда.",
    147: "Отчет 04: проект разделяет несущие стены/перегородки и исключает окна.",
    148: "Наши два ряда арматуры сгруппированы; разница проектной базы и старого ручного расчета.",
    152: "Отчет 04: наша площадь 3,93; у Елены не включена ПБ2, сравнение не требует уменьшать нашу.",
    159: "У нас U и монолитные перемычки раздельно (разные ставки); у Елены одна строка и пропуск 1,7 м.",
    163: "Строка демонтажа добавлена, равна монтажу 3,93; внутренние суммы нулевые.",
    169: "Остатки общие по проекту; закупка этапа, не объем его кладки. Отчет 08 отменяет старое решение 04.",
    179: "Последний этап D400 использует остатки: 11 паллет; всего 37+7+11=55, не отдельный запас каждому этапу.",
    180: "В локальном review исправлена подтвержденная опечатка D400 -> D500; 1,8 м3 и цена не изменились.",
    185: "Расходники у нас 3%, у Елены 2% по формуле O185; разница методики, не только цены.",
    218: "Google-review сохраняет резерв 33 листа; он отсутствовал в локальном review отчета 05.",
    220: "ЗД-1 нет в нашей смете, ранее отмечено в отчете 05.",
    221: "Лист 16 мм нет в нашей смете, ранее отмечено в отчете 05.",
    229: "Готовые проектные объемы наших балок 3,83; у Елены ручной расчет 3,915.",
    231: "Чистый бетон плиты 57,8 из спецификации; Елена в J231 поставила 54,8. Не вычитать балки повторно.",
    235: "Исправлено: демонтаж = обе строки монтажа, 321,111111 + 72,28 = 393,391111.",
    238: "Проектный итог 7,0 м3 вместо сборки частей, 27 упаковок по 0,2773; работы не менялись.",
    241: "Расходники у нас 3%, у Елены 2%; коэффициентная разница.",
    254: "Google-review сохраняет резерв 15 листов; он отсутствовал в локальном review отчета 05.",
    259: "НЕ полное совпадение марки: у нас А500С D10, у Елены надпись D12 при одинаковой длине 2749,5 м.",
    263: "20 м3 у нас против 19,5 у Елены: округление заказа до целого м3, балки добавлены один раз.",
    266: "Исправлено: демонтаж = 97,222222 + 15,2 = 112,422222.",
    269: "Проектный итог 3,8 м3, 15 упаковок по 0,2773 = 4,1595; у Елены упаковка 0,2776 = 4,164.",
    272: "Расходники у нас 3%, у Елены 2%; коэффициентная разница.",
    282: "Количество совпадает; у нас 15000 материал/машина + 2500 работа, у Елены 10000.",
    283: "Разные базы/коэффициенты: у нас 3% прямого раздела, у Елены формула 4% SUM(O279:O282).",
    294: "Готовые проектные кубы 62,2; у Елены закупка 64,05984. В проекте/Елене CARBON PROF, у нас CARBON ECO: соответствие продукта и цены требует проверки.",
    295: "Количество из заполненного КП сохранено. В проекте/Елене CARBON PROF 50 мм, у нас CARBON ECO; соответствие продукта и цены требует проверки.",
    300: "Новое правило: 364 м2 кровли <=430, один рулон 400; у Елены тоже 400.",
    301: "Новое правило: примыкания 168,2 * 1,1 -> 200 м2; у Елены тоже 200.",
    303: "Наши проектные примыкания 168,2 против ручного обмера Елены 156,9; не только ставки.",
    307: "У Елены шт, у нас рул: один рулон того же формата 2,1x20 м, физическое количество совпадает.",
    316: "Ручной технадзор в Google-review 10000, у Елены 5000.",
}


def audit(results_dir: Path, original_results: Path, elena_path: Path, out_dir: Path):
    wb_elena = load_workbook(elena_path, data_only=True)
    elena = wb_elena.active
    blocks = [(code, _load_lines(results_dir, code), MAPPINGS[code]) for code in MAPPINGS]
    slabs = _load_floor_slabs_blocks(results_dir)
    if len(slabs) != 2:
        raise ValueError("This explicit ARK audit expects PM1 and PM2, in that order")
    blocks += [(f"pm{i+1}", lines, {code: rows[i] for code, rows in SLAB_MAP.items()})
               for i, (_, lines) in enumerate(slabs)]
    details, summary = [], []
    for section, lines, mapping in blocks:
        start, end, total_row = RANGES[section]
        grouped = {}
        for line in lines:
            row = mapping.get(line["code"])
            grouped.setdefault(row if row is not None else line["code"], []).append(line)
        for row in range(start, end + 1):
            if elena.cell(row, 3).value is not None:
                grouped.setdefault(row, [])
        for key, ours in grouped.items():
            row = key if isinstance(key, int) else None
            ours_qty = sum(_num(x.get("quantity_raw", x.get("quantity"))) for x in ours)
            costs = [_cost_parts(x) for x in ours]
            ours_total = sum(x[-1] for x in costs)
            elena_qty = _num(elena.cell(row, 10).value) if row else None
            elena_total = _num(elena.cell(row, 15).value) if row else None
            ours_units = sorted({x["unit"] for x in ours})
            unit = elena.cell(row, 3).value if row else None
            price_pairs = sorted({(x[1], x[3]) for x in costs})
            ep = (_num(elena.cell(row, 11).value), _num(elena.cell(row, 13).value)) if row else None
            marks = ""
            if not ours:
                status = "Нет строки у нас"
            elif row is None:
                status = "Дополнительная строка у нас"
            elif row in (104, 105, 106, 107, 108, 109):
                status = "Разный охват"
            elif row == 259:
                status = "Разная марка арматуры"
            elif row in (294, 295):
                status = "Разная номенклатура ЭППС"
            elif abs(ours_qty - elena_qty) > 0.00001:
                status = "Разница количества"
            elif ours_total == 0 and elena_total == 0:
                status = "Нулевая стоимость, количество совпало"
            elif abs(ours_total - elena_total) <= 1:
                status = "Совпало"
            else:
                status = "Разница цены/базы расходов"
            if unit is not None and ours_units and unit not in ours_units:
                marks = "Разное обозначение единицы; сопоставлено вручную."
            details.append({
                "section": section, "elena_row": row, "codes": [x["code"] for x in ours],
                "ours_names": [x["name"] for x in ours],
                "elena_name": elena.cell(row, 2).value if row else None,
                "ours_unit": "; ".join(ours_units), "elena_unit": unit,
                "ours_quantity": ours_qty if ours else None, "elena_quantity": elena_qty,
                "ours_prices": price_pairs, "elena_prices": ep,
                "ours_total": ours_total if ours else None, "elena_total": elena_total,
                "delta": ours_total - (elena_total or 0), "status": status,
                "note": " ".join(x for x in [NOTES.get(row, ""), marks] if x),
            })
        ours_sum = sum(_cost_parts(x)[-1] for x in lines)
        elena_sum = _num(elena.cell(total_row, 15).value)
        if abs(sum(x["delta"] for x in details if x["section"] == section) - (ours_sum - elena_sum)) > 1:
            raise ValueError(f"Row reconciliation failed for {section}")
        summary.append({"section": section, "ours_total": ours_sum, "elena_total": elena_sum,
                        "delta": ours_sum - elena_sum, "delta_percent": (ours_sum / elena_sum - 1) * 100,
                        "elena_total_row": total_row})
    outside = [{"name": elena.cell(header, 2).value, "total_row": row,
                "total": _num(elena.cell(row, 15).value)}
               for header, row in [(11, 24), (42, 74), (117, 135), (192, 210), (321, 337)]]
    original_lines = [_load_lines(original_results, code) for code in MAPPINGS]
    original_lines += [lines for _, lines in _load_floor_slabs_blocks(original_results)]
    original_total = sum(_cost_parts(line)[-1] for lines in original_lines for line in lines)
    payload = {"summary": summary, "rows": details, "outside_coverage": outside,
               "as_received_total": original_total,
               "verified_input_change_delta": sum(x["ours_total"] for x in summary) - original_total,
               "status_counts": dict(Counter(x["status"] for x in details)),
               "ours_total": sum(x["ours_total"] for x in summary),
               "elena_comparison_total": sum(x["elena_total"] for x in summary),
               "elena_full_internal_total": sum(x["elena_total"] for x in summary) + sum(x["total"] for x in outside)}
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "comparison.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    output = Workbook()
    sheet = output.active
    sheet.title = "Сводка"
    sheet.append(["Раздел", "У нас, руб.", "У Елены, руб.", "Разница, руб.", "Разница, %", "Итоговая строка Елены"])
    for item in summary:
        sheet.append(list(item.values()))
    sheet.append(["Итого по сравниваемым разделам", payload["ours_total"], payload["elena_comparison_total"],
                  payload["ours_total"] - payload["elena_comparison_total"]])
    sheet = output.create_sheet("Построчно")
    sheet.append(["Раздел", "Строка Елены", "Коды у нас", "Наименование у нас", "Наименование Елены",
                  "Ед. у нас", "Ед. Елены", "Количество у нас", "Количество Елены", "Ставки у нас (мат.; раб.)",
                  "Ставки Елены (мат.; раб.)", "Стоимость у нас", "Стоимость Елены", "Разница", "Статус", "Комментарий"])
    for item in details:
        sheet.append([item["section"], item["elena_row"], "; ".join(item["codes"]),
                      "; ".join(item["ours_names"]), item["elena_name"], item["ours_unit"], item["elena_unit"],
                      item["ours_quantity"], item["elena_quantity"], str(item["ours_prices"]), str(item["elena_prices"]),
                      item["ours_total"], item["elena_total"], item["delta"], item["status"], item["note"]])
    sheet = output.create_sheet("Вне охвата")
    sheet.append(["Раздел Елены", "Итоговая строка", "Внутренняя стоимость, руб."])
    for item in outside:
        sheet.append(list(item.values()))
    sheet = output.create_sheet("Изменения входов")
    sheet.append(["Изменение", "Основание", "Влияние"])
    sheet.append(["ПМ1 eps_material_spec_volume_m3 = 7,0", "КР2 стр.19, raw pm1_eps_total; отчет 11",
                  "ЭППС 8,319 -> 7,4871 м3; работы не менялись"])
    sheet.append(["ПМ2 eps_material_spec_volume_m3 = 3,8", "КР2 стр.21, raw pm2_eps_total; отчет 11",
                  "ЭППС 3,6049 -> 4,1595 м3; работы не менялись"])
    sheet.append(["Обкладка вентшахт D400 -> D500", "Подтверждение Елены; отчет 04 находка 7",
                  "Название/марка; 1,8 м3, цена и стоимость без изменений"])
    sheet.append(["Источник Google", "1Y7TH2_q4Bw-Ai-Ld43caM6Yq2DLcv6bjdjELGK4velE", "Удаленная таблица не изменялась"])
    sheet.append(["SHA256 исходного XLSX", hashlib.sha256((out_dir / "review_workbook_google.xlsx").read_bytes()).hexdigest(),
                  "Сохранен отдельный неизмененный снимок"])
    sheet.append(["Смета строго из Google-review", original_total, "До переноса подтвержденных проектных кубов и марки блока"])
    sheet.append(["Смета после проверенных правок", payload["ours_total"], "Все семь разделов пересчитаны заново"])
    sheet.append(["Ставки в построчной сверке", "Представление экспортера", "При отсутствии явной ставки используется raw-стоимость / raw-количество; если raw-стоимости нет, округленная стоимость / количество."])
    for sheet in output:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="327D72")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(70, max(18, max(len(str(x.value or "")) for x in column) // 2))
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 48
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.0000"
    output.save(out_dir / "ark_vs_elena_all_sections.xlsx")
    print(json.dumps({k: v for k, v in payload.items() if k != "rows"}, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-dir", type=Path, required=True)
    parser.add_argument("--original-results", type=Path, required=True)
    parser.add_argument("--elena", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    audit(args.results_dir, args.original_results, args.elena, args.out_dir)


if __name__ == "__main__":
    main()
