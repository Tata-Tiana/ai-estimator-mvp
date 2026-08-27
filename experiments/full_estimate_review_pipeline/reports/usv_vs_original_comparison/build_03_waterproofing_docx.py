"""One-off builder for 03_waterproofing_sources_for_elena.docx - same style/format as
build_02_foundation_slab_docx.py (see that file's docstring + feedback_sources_for_elena_docx_workflow
memory for the exact rules: Table Grid, bold header only, Excel-style qty formatting via
format_qty()). Numbers sourced from output/usv_build/section_results_live/waterproofing_result.json
(built 2026-08-26, fresh Sheet download, after the "Технический надзор" line was added to
waterproofing_calculator.py)."""

from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_PATH = Path(__file__).resolve().parent / "03_waterproofing_sources_for_elena.docx"


def format_qty(value: float, unit: str) -> str:
    """Excel-style fixed 2-decimal, comma-separator display, NO thousands separator:
    222.3 -> '222,30 мп'. Uses Decimal(str(value)) + ROUND_HALF_UP, not raw float f"{value:.2f}" -
    plain float formatting silently rounds DOWN on values whose binary float representation sits
    just under the true decimal (39.555 -> "39.55" instead of the correct "39.56") - caught by the
    walls-section J-column verification (2026-08-26)."""
    quantized = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return str(quantized).replace(".", ",") + f" {unit}"


# (Строка, число (полная точность = J в смете), единица, Откуда)
_RAW_ROWS = [
    ("Гидроизоляция фундаментной плиты битумной мастикой в 2 слоя", 23.58, "м2",
     "Из проекта напрямую - площадь опалубки бортов фундаментной плиты, используется как "
     "заменитель площади гидроизоляции (в PDF нет отдельной готовой строки под гидроизоляцию, по "
     "методике так и берётся)."),
    ("Праймер битумный AquaMast, 18 л", 1, "шт",
     "Посчитано. Площадь (23.58 м²) x расход праймера (0.3 л/м²) = 7.074 л -> округление вверх до "
     "кратности канистры (18 л): 7.074 / 18 = 0.39 -> 1 шт."),
    ("Мастика гидроизоляционная битумная для фундаментов AquaMast, 18 кг", 3, "шт",
     "Посчитано. Площадь (23.58 м²) x расход мастики (1 кг/м²/слой) x количество слоёв (2) = "
     "47.16 кг -> округление вверх до кратности ведра (18 кг): 47.16 / 18 = 2.62 -> 3 шт."),
    ("Утепление торца/борта фундаментной плиты ЭППС 100 мм", 17.5, "м2",
     "Из проекта напрямую - готовая строка спецификации PDF (\"Площадь утепления торцов\")."),
    ("Пеноплэкс ГЕО 100 мм", 1.9432, "м3",
     "Посчитано. Площадь (17.5 м²) x толщина (0.1м) x запас (1.05) = 1.8375 м3 -> округление вверх "
     "до кратности упаковки (0.2776 м3): 1.8375 / 0.2776 = 6.62 -> 7 упаковок x 0.2776 = 1.9432 м3."),
    ("Клей-пена для ЭППС", 2, "баллон",
     "Посчитано. Утепляемая площадь (17.5 м², ЭППС 50мм у этого проекта нет) / покрытие одного "
     "баллона (10 м²) = 1.75 -> округление вверх = 2 баллона (минимум 1 баллон по нормативу)."),
    ("Логистика и снабжение", 1, "(-)",
     "Посчитано (сама сумма - не эта цифра). 2% (ручной норматив) от суммы всех строк раздела выше: "
     "46 547 ₽ x 2% = 930.94 ₽."),
    ("Расходные материалы, амортизация инструмента", 1, "(комплект)",
     "Посчитано (сама сумма - не эта цифра). 3% (ручной норматив) от суммы всех строк раздела выше: "
     "46 547 ₽ x 3% = 1 396.41 ₽."),
    ("Технический надзор", 1, "(-)",
     "Ручное поле, по умолчанию 0 ₽ (строка добавлена в калькулятор 2026-08-26 - раньше "
     "отсутствовала вообще, у Елены в реальной смете эта строка тоже всегда пустая/0 во всех "
     "проверенных проектах)."),
    ("Заготовительно-складские расходы", 1, "(-)",
     "Строка есть для структуры сметы, у нас всегда 0 - считаем только внутреннюю себестоимость, "
     "без клиентской наценки."),
    ("Накладные и общехозяйственные расходы", 1, "(-)", "Та же логика - 0, без наценки."),
    ("Сметная прибыль", 1, "(-)", "Та же логика - 0, без наценки."),
]

ROWS = [(name, format_qty(value, unit), source) for name, value, unit, source in _RAW_ROWS]

doc = Document()

doc.add_heading("Гидроизоляция, утепление бортов плит — откуда взято каждое число (для Елены)", level=1)
doc.add_paragraph(
    "Раздел «Гидроизоляция, утепление бортов плит», столбец «Кол-во». По каждой строке: число и "
    "откуда оно взялось — из проекта напрямую, или посчитано (и по какой формуле)."
)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
header = table.rows[0].cells
for cell, text in zip(header, ["Строка", "Кол-во", "Откуда"]):
    run = cell.paragraphs[0].add_run(text)
    run.bold = True

for name, qty, source in ROWS:
    row = table.add_row().cells
    row[0].text = name
    row[1].text = qty
    row[2].text = source

doc.add_heading("Что из проекта, а что вручную/по умолчанию", level=2)

summary_items = [
    ("Из проекта (спецификация PDF): ",
     "площадь гидроизоляции/опалубки бортов (23.58 м²), площадь утепления торцов ЭППС 100мм "
     "(17.5 м²)."),
    ("Ручные поля (по умолчанию, можно менять под проект): ",
     "расход праймера (0.3 л/м²), расход мастики (1 кг/м²/слой, 2 слоя), запас на потери ЭППС "
     "(1.05), покрытие баллона клей-пены (10 м²), нормативы логистики (2%) и расходных материалов "
     "(3%), технический надзор (0 ₽)."),
    ("Вручную поправлено конкретно по этому проекту: ",
     "не потребовалось - все поля этого раздела на типовых значениях."),
]

for lead, rest in summary_items:
    p = doc.add_paragraph()
    r1 = p.add_run(lead)
    r1.bold = True
    p.add_run(rest)

for style_name in ["Normal"]:
    style = doc.styles[style_name]
    style.font.size = Pt(11)

doc.save(OUT_PATH)
print(f"saved: {OUT_PATH}")
for name, qty, _ in ROWS:
    print(f"{name!r:80} -> {qty}")
