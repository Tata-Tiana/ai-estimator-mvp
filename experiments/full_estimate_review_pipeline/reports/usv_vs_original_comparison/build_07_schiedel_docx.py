"""One-off builder for 07_schiedel_vent_channels_sources_for_elena.docx - same style/format as
build_06_flat_roof_docx.py (see feedback_sources_for_elena_docx_workflow memory: Table Grid,
bold header only, Excel-style qty formatting via format_qty()). No zone subheaders here - the
section has no physical zones/floors, it's a single flat list matching the real smeta's own
layout. Numbers sourced from output/usv_build/section_results_live/schiedel_vent_channels_result.json
(rebuilt 2026-08-27, same live Sheet as before - only the calculator changed: added the
missing "Технический надзор" zero line, see report finding №1)."""

from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_PATH = Path(__file__).resolve().parent / "07_schiedel_vent_channels_sources_for_elena.docx"


def format_qty(value: float, unit: str) -> str:
    """Excel-style fixed 2-decimal, comma-separator display, NO thousands separator:
    222.3 -> '222,30 мп'. Uses Decimal(str(value)) + ROUND_HALF_UP, not raw float f"{value:.2f}" -
    plain float formatting silently rounds DOWN on values whose binary float representation sits
    just under the true decimal."""
    quantized = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return str(quantized).replace(".", ",") + f" {unit}"


# (Строка, число (полная точность = J в смете), единица, Откуда)
_RAW_ROWS: list[tuple[str, float, str, str]] = [
    ("Кладка вентканалов Schiedel", 15.82, "мп",
     "Из проекта напрямую - готовая длина кладки вентканалов из спецификации PDF "
     "(schiedel_masonry_total_length_m)."),
    ("Вентиляционный канал 2х,36/25 см Schiedel", 24, "шт",
     "Из проекта напрямую - сумма строк спецификации вентканалов (schiedel_channel_items) с "
     "типом «2х»."),
    ("Вентиляционный канал 3х,52/25 см Schiedel", 8, "шт",
     "Из проекта напрямую - сумма строк спецификации вентканалов (schiedel_channel_items) с "
     "типом «3х»."),
    ("Доставка вентканалов", 1, "маш",
     "Ручное поле - количество рейсов доставки, нет источника в PDF/формуле."),
    ("Расходные материалы, амортизация инструмента", 1, "комплект",
     "Посчитано (сама сумма - не эта цифра). Норматив 3% от суммы прямых затрат раздела до этой "
     "строки (114 753,60₽ x 0,03 = 3 442,61₽)."),
    ("Технический надзор", 1, "-",
     "Строка есть для структуры сметы, у нас всегда 0 - технический надзор по этому разделу "
     "нигде не считается отдельной суммой (у Елены тоже 0₽ в реальной смете)."),
    ("Заготовительно-складские расходы", 1, "-",
     "Строка есть для структуры сметы, у нас всегда 0 - считаем только внутреннюю себестоимость, "
     "без клиентской наценки."),
    ("Накладные и общехозяйственные расходы", 1, "-", "Та же логика - 0, без наценки."),
    ("Сметная прибыль", 1, "-", "Та же логика - 0, без наценки."),
]

ROWS = [(name, format_qty(value, unit), source) for name, value, unit, source in _RAW_ROWS]

doc = Document()

doc.add_heading("Вентиляционные каналы Schiedel — откуда взято каждое число (для Елены)", level=1)
doc.add_paragraph(
    "Раздел «Вентиляционные каналы Schiedel», столбец «Кол-во». По каждой строке: число и откуда "
    "оно взялось - из проекта напрямую, вручную, или посчитано (и по какой формуле). В разделе нет "
    "физических зон/этажей - одна плоская таблица, как и в самой смете."
)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
header = table.rows[0].cells
for cell, text in zip(header, ["Строка", "Кол-во", "Откуда"]):
    run = cell.paragraphs[0].add_run(text)
    run.bold = True

for name, qty, source in ROWS:
    row = table.add_row().cells
    row[0].text = name
    row[1].text = qty
    row[2].text = source

doc.add_heading("Что из проекта, а что вручную/по умолчанию", level=2)

summary_items = [
    ("Из проекта (спецификация PDF): ",
     "длина кладки вентканалов, количество и типы вентканалов по спецификации."),
    ("Ручные поля (по умолчанию, можно менять под проект): ",
     "количество рейсов доставки."),
    ("Вручную поправлено конкретно по этому проекту: ",
     "ничего - все значения раздела совпали с реальной сметой без ручных правок на листе 01."),
]

for lead, rest in summary_items:
    p = doc.add_paragraph()
    r1 = p.add_run(lead)
    r1.bold = True
    p.add_run(rest)

for style_name in ["Normal"]:
    style = doc.styles[style_name]
    style.font.size = Pt(11)

doc.save(OUT_PATH)
print(f"saved: {OUT_PATH}")
print(f"rows: {len(ROWS)}")
