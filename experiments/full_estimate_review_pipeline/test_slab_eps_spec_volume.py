from __future__ import annotations

from decimal import Decimal
import json
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE / "review_to_calculator"))
sys.path.insert(0, str(BASE.parent / "floor_slab_1_calculator"))
from sections.floor_slabs.build_input import _resolve_insulation  # noqa: E402
from floor_slab_calculator import calculate_insulation_context  # noqa: E402
from populate_review_workbook_from_extraction import floor_slab_visible_field_keys  # noqa: E402
from promote_slab_eps_spec_totals import promote  # noqa: E402


class SpecEpsVolumeTests(unittest.TestCase):
    def rows(self):
        return [
            {"role": "slab_bottom", "area_m2": 23.3, "thickness_mm": 100},
            {"role": "slab_edge", "length_m": 41.8, "height_m": 0.18, "thickness_mm": 100},
        ]

    def zone(self, **extra):
        return {"zone_id": "slab", "display_name": "Slab", "beams_eps_material_area_m2": 1.62, **extra}

    def test_spec_total_replaces_parts_including_beams(self):
        resolved = _resolve_insulation(self.zone(eps_material_spec_volume_m3=3.8), self.rows())
        self.assertEqual(resolved["total_eps_volume_from_spec_m3"], 3.8)
        self.assertEqual(resolved["eps_material_volume_source"], "spec_total")
        self.assertEqual(resolved["eps_material_derived_volume_m3"], 3.2444)

    def test_work_and_foam_quantities_stay_independent(self):
        before = _resolve_insulation(self.zone(), self.rows())
        after = _resolve_insulation(self.zone(eps_material_spec_volume_m3=3.8), self.rows())
        for key in ["bottom_slab_eps_work_area_m2", "slab_outer_edge_eps_work_length_m",
                    "slab_edge_eps_material_area_m2", "beams_eps_material_area_m2"]:
            self.assertEqual(before[key], after[key])

    def test_absent_or_null_total_retains_fallback(self):
        for zone in [self.zone(), self.zone(eps_material_spec_volume_m3=None)]:
            with self.subTest(zone=zone):
                r = _resolve_insulation(zone, self.rows())
                self.assertEqual(r["total_eps_volume_from_spec_m3"], 3.2444)
                self.assertEqual(r["eps_material_volume_source"], "derived_parts")

    def test_zero_total_is_not_treated_as_missing(self):
        r = _resolve_insulation(self.zone(eps_material_spec_volume_m3=0), self.rows())
        self.assertEqual(r["total_eps_volume_from_spec_m3"], 0)
        self.assertEqual(r["eps_material_volume_source"], "spec_total")

    def test_invalid_total_rejected(self):
        for value in [-1, float("nan"), float("inf")]:
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "finite and >= 0"):
                _resolve_insulation(self.zone(eps_material_spec_volume_m3=value), self.rows())
        for value in ["not a number", True]:
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "must be a number"):
                _resolve_insulation(self.zone(eps_material_spec_volume_m3=value), self.rows())

    def test_ready_volume_does_not_invent_missing_work(self):
        with self.assertRaisesRegex(ValueError, "work quantities cannot be honestly split"):
            _resolve_insulation(self.zone(eps_material_spec_volume_m3=3.8), [])

    def test_spec_volume_with_combined_row_not_added_twice(self):
        rows = self.rows() + [{"role": "combined_bottom_and_edge", "volume_m3": 3.8}]
        r = _resolve_insulation(self.zone(eps_material_spec_volume_m3=3.8), rows)
        self.assertEqual(r["total_eps_volume_from_spec_m3"], 3.8)

    def test_pack_rounding_and_waste_once(self):
        r = _resolve_insulation(self.zone(eps_material_spec_volume_m3=3.8), self.rows())
        r.update(eps_pack_volume_m3=0.2773, eps_waste_coeff=1.05, eps_thickness_m=0.1,
                 foam_coverage_m2_per_can=10)
        context, _ = calculate_insulation_context(r, None, Decimal("0.18"))
        self.assertEqual(context["required_eps_volume_m3_raw"], 3.99)
        self.assertEqual(context["eps_packs_ordered"], 15)
        self.assertEqual(context["order_eps_volume_m3_raw"], 4.1595)

    def test_spec_volume_visible_in_review(self):
        self.assertIn("eps_material_spec_volume_m3", floor_slab_visible_field_keys("floor_slab_zones", []))


class SpecEpsPromotionTests(unittest.TestCase):
    def prepare(self, directory):
        root = Path(directory)
        source = root / "source.json"
        workbook = root / "source.xlsx"
        item = {"zone_id": "slab", "display_name": "Slab", "manual_concrete_pump_shifts": None,
                "_sheet_item_id": "slab", "_sheet_field_key": "manual_concrete_pump_shifts"}
        source.write_text(json.dumps({"sections": {"floor_slabs": {
            "found": [{"target_code": "floor_slab_zones", "value": {"zone_id": "slab"}}],
            "raw_table_rows": [{"table_id": "eps", "normalized_unit": "m3", "source_pdf": "source.pdf",
                                "page_number": 1, "raw_text": "EPS 3.8 m3", "mapped_target_codes": []}],
        }}}))
        wb = Workbook()
        ws = wb.active
        ws.title = "01_Проверка проекта"
        ws.append(["Что проверяем", "Найдено в проекте", "Ед.", "Статус", "Уверенность",
                   "Что нужно сделать", "Источник", "Фрагмент проекта", "Исправить / ввести значение",
                   "Комментарий Елены", "section_code", "technical_key", "source_class", "target_code"])
        ws.append(["Pump", None, "shift", "Found", None, None, None, None, 2, "Keep comment",
                   "floor_slabs", "floor_slab_zones", "AUTO_PROJECT", "floor_slab_zones"])
        ws.cell(2, 19, json.dumps(item))
        wb.create_sheet("Prices")["A1"] = 9020
        wb.save(workbook)
        return source, workbook, root / "promoted.json", root / "promoted.xlsx"

    def test_promotion_preserves_manual_values_and_originals(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = self.prepare(directory)
            original = [p.read_bytes() for p in paths[:2]]
            promote(*paths, [("slab", "eps", 3.8)])
            self.assertEqual(original, [p.read_bytes() for p in paths[:2]])
            wb = load_workbook(paths[3])
            ws = wb["01_Проверка проекта"]
            self.assertEqual(ws["I2"].value, 2)
            self.assertEqual(ws["J2"].value, "Keep comment")
            self.assertEqual(wb["Prices"]["A1"].value, 9020)
            self.assertEqual(ws["B3"].value, 3.8)
            self.assertEqual(json.loads(ws["S3"].value)["_sheet_field_key"], "eps_material_spec_volume_m3")
            with self.assertRaisesRegex(ValueError, "never overwritten"):
                promote(*paths, [("slab", "eps", 3.8)])

    def test_missing_evidence_rejected_without_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = self.prepare(directory)
            with self.assertRaisesRegex(ValueError, "exactly one raw table"):
                promote(*paths, [("slab", "missing", 3.8)])
            self.assertFalse(paths[2].exists())
            self.assertFalse(paths[3].exists())


if __name__ == "__main__":
    unittest.main()
