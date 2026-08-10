from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_LOGO_PATH = REPO_ROOT / "experiments" / "earthworks_review_to_calculator" / "assets" / "brick_house_logo.png"

TITLE_TEXT = "СМЕТНЫЙ РАСЧЕТ НА СТРОИТЕЛЬСТВО ДОМА"
GROUP_HEADER_ROW = 9
SUBHEADER_ROW = 10
FIRST_SECTION_ROW = 11

SECTION_ORDER: list[tuple[str, str]] = [
    ("earthworks", "ЗЕМЛЯНЫЕ РАБОТЫ"),
    ("foundation_slab", "УСТРОЙСТВО ФУНДАМЕНТНОЙ ПЛИТЫ"),
    ("waterproofing", "ГИДРОИЗОЛЯЦИЯ, УТЕПЛЕНИЕ БОРТОВ ПЛИТ"),
    ("load_bearing_walls_lintels", "ВНЕШНИЕ И ВНУТРЕННИЕ НЕСУЩИЕ СТЕНЫ, ПЕРЕМЫЧКИ НАД ПРОЕМАМИ"),
    # floor_slab_1/floor_slab_2 removed here (P3, FLOOR_SLAB_UNIFICATION_PLAN.md) - they used to
    # be two fixed entries in this exact spot. Now a dynamic number of blocks (as many as there
    # are real pours - 2 for every project today, since no real extraction has zone_context yet)
    # gets spliced in right after load_bearing_walls_lintels - see FLOOR_SLABS_INSERT_AFTER and
    # _all_section_blocks() below. "Honestly reflects reality: however many plates exist, that
    # many blocks land in the smeta" - the user's own framing when choosing this design.
    ("flat_roof", "ПЛОСКАЯ КРОВЛЯ"),
    ("schiedel_vent_channels", "ВЕНТИЛЯЦИОННЫЕ КАНАЛЫ"),
]

FLOOR_SLABS_INSERT_AFTER = "load_bearing_walls_lintels"
FLOOR_SLABS_RESULT_FILENAME = "floor_slabs_result.json"

CALC_COLS = {
    "quantity": "J",
    "material_unit_price": "K",
    "material_total": "L",
    "work_unit_price": "M",
    "work_total": "N",
    "row_total": "O",
}
WHITE_ZONE_MIRROR = {
    "D": "J",
    "E": "K",
    "F": "L",
    "G": "M",
    "H": "N",
    "I": "O",
}

_GREY = PatternFill("solid", fgColor="FFB7B7B7")
_WHITE = PatternFill("solid", fgColor="FFFFFFFF")
_SECTION = PatternFill("solid", fgColor="FFEFEFEF")
_TOTAL = PatternFill("solid", fgColor="FFD9EAF7")
_THIN = Side(style="thin", color="303030")
_BLUE = Side(style="thick", color="1D4ED8")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_SPLIT = Border(left=_BLUE, right=_THIN, top=_THIN, bottom=_THIN)


def _num(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(_text(item) for item in value if item is not None)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _line_value(line: dict[str, Any], key: str) -> Any:
    if key in line and line.get(key) not in (None, ""):
        return line.get(key)
    internal_cost = line.get("internal_cost") or {}
    if isinstance(internal_cost, dict):
        return internal_cost.get(key)
    return None


def _quantity(line: dict[str, Any]) -> float:
    for key in ("quantity", "quantity_display", "quantity_raw"):
        if line.get(key) not in (None, ""):
            return _num(line.get(key))
    return 0.0


def _cost_parts(line: dict[str, Any]) -> tuple[float, float, float, float, float, float]:
    quantity = _quantity(line)
    material_total = _num(_line_value(line, "material_total"))
    work_total = _num(_line_value(line, "work_total"))
    row_total = _num(_line_value(line, "line_total")) or material_total + work_total

    material_unit_price = _num(_line_value(line, "material_unit_price"))
    work_unit_price = _num(_line_value(line, "work_unit_price"))

    # Some production calculators return percentage/fixed rows as totals, not as a
    # meaningful unit price. The Excel estimate should still keep live row formulas,
    # so represent such rows as quantity * derived unit price.
    if quantity and material_total and not material_unit_price:
        material_unit_price = material_total / quantity
    if quantity and work_total and not work_unit_price:
        work_unit_price = work_total / quantity

    return quantity, material_unit_price, material_total, work_unit_price, work_total, row_total


def _load_lines(results_dir: Path, section_code: str) -> list[dict[str, Any]]:
    path = results_dir / f"{section_code}_result.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    result = data.get("result", data)
    return list(result.get("estimate_lines") or result.get("lines") or [])


def _load_floor_slabs_blocks(results_dir: Path) -> list[tuple[str, list[dict[str, Any]]]]:
    """P3: reads floor_slabs_result.json (produced by build_floor_slabs_result_json.py, see that
    script's own docstring) - {"pours": [{"title", "estimate_lines"}, ...]}. Raises on a missing
    file, same as _load_lines() does for every other section's JSON - a smeta silently missing
    both floor-slab sections because the file was never generated is worse than a loud crash."""
    path = results_dir / FLOOR_SLABS_RESULT_FILENAME
    data = json.loads(path.read_text(encoding="utf-8"))
    return [(pour["title"], list(pour.get("estimate_lines") or [])) for pour in data.get("pours") or []]


def _all_section_blocks(results_dir: Path) -> list[tuple[str, list[dict[str, Any]]]]:
    """Full ordered list of (title, lines) blocks to render - SECTION_ORDER's 6 fixed sections
    plus the dynamic floor-slabs block spliced in right after FLOOR_SLABS_INSERT_AFTER, in the
    exact spot floor_slab_1/floor_slab_2 used to occupy as two fixed entries."""
    blocks: list[tuple[str, list[dict[str, Any]]]] = []
    for section_code, section_title in SECTION_ORDER:
        blocks.append((section_title, _load_lines(results_dir, section_code)))
        if section_code == FLOOR_SLABS_INSERT_AFTER:
            blocks.extend(_load_floor_slabs_blocks(results_dir))
    return blocks


def _extract_project_address(review_workbook: Path | None) -> str:
    if review_workbook is None or not review_workbook.exists():
        return "Адрес объекта: —"
    wb = load_workbook(review_workbook, data_only=True)
    if "01_Проверка проекта" not in wb.sheetnames:
        return "Адрес объекта: —"
    ws = wb["01_Проверка проекта"]
    for row in ws.iter_rows(values_only=True):
        label = _text(row[0] if row else "")
        if "адрес" not in label.lower():
            continue
        for value in row[1:4]:
            candidate = _text(value).strip()
            if candidate:
                return candidate
    return "Адрес объекта: —"


def _setup_top_header(ws: Any, logo_path: Path, address: str) -> None:
    for row in range(1, 9):
        for col in range(10, 23):
            ws.cell(row, col).fill = _GREY

    ws.merge_cells("A1:I4")
    if logo_path.exists():
        logo = OpenpyxlImage(str(logo_path))
        logo.width = 300
        logo.height = 50
        ws.add_image(logo, "C2")

    ws.merge_cells("A5:I5")
    ws["A5"].value = TITLE_TEXT
    ws["A5"].font = Font(size=12, bold=True)
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A6:I6")
    ws["A6"].value = address
    ws["A6"].font = Font(size=10)
    ws["A6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    ws.row_dimensions[6].height = 24

    header_merges = [
        (GROUP_HEADER_ROW, 1, SUBHEADER_ROW, 1, "№"),
        (GROUP_HEADER_ROW, 2, SUBHEADER_ROW, 2, "Наименование работ"),
        (GROUP_HEADER_ROW, 3, SUBHEADER_ROW, 3, "Ед. изм."),
        (GROUP_HEADER_ROW, 4, SUBHEADER_ROW, 4, "Кол-во"),
        (GROUP_HEADER_ROW, 5, GROUP_HEADER_ROW, 6, "Стоимость материалов,\nмашин и механизмов, руб."),
        (GROUP_HEADER_ROW, 7, GROUP_HEADER_ROW, 8, "Стоимость работ, руб."),
        (GROUP_HEADER_ROW, 9, SUBHEADER_ROW, 9, "Итого, руб."),
        (GROUP_HEADER_ROW, 10, SUBHEADER_ROW, 10, "Кол-во"),
        (GROUP_HEADER_ROW, 11, GROUP_HEADER_ROW, 12, "Стоимость материалов,\nмашин и механизмов, руб."),
        (GROUP_HEADER_ROW, 13, GROUP_HEADER_ROW, 14, "Стоимость работ, руб."),
        (GROUP_HEADER_ROW, 15, SUBHEADER_ROW, 15, "Итого, руб."),
    ]
    for sr, sc, er, ec, value in header_merges:
        ws.merge_cells(start_row=sr, start_column=sc, end_row=er, end_column=ec)
        ws.cell(sr, sc).value = value

    for col, value in {
        5: "За ед",
        6: "Итого",
        7: "За ед",
        8: "Итого",
        11: "За ед",
        12: "Итого",
        13: "За ед",
        14: "Итого",
    }.items():
        ws.cell(SUBHEADER_ROW, col).value = value

    for row in (GROUP_HEADER_ROW, SUBHEADER_ROW):
        for col in range(1, 16):
            cell = ws.cell(row, col)
            cell.fill = _GREY if col >= 10 else _WHITE
            cell.border = _SPLIT if col == 10 else _BORDER
            cell.font = Font(size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[GROUP_HEADER_ROW].height = 46
    ws.row_dimensions[SUBHEADER_ROW].height = 22


def _setup_dimensions(ws: Any) -> None:
    widths = {
        1: 6,
        2: 58,
        3: 10,
        4: 12,
        5: 16,
        6: 16,
        7: 14,
        8: 14,
        9: 14,
        10: 12,
        11: 16,
        12: 16,
        13: 14,
        14: 14,
        15: 14,
    }
    for col_num, width in widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width
    for col_num in range(16, 23):
        ws.column_dimensions[get_column_letter(col_num)].width = 10
    ws.freeze_panes = "E12"


def _write_section_header(ws: Any, row_num: int, section_number: int, section_title: str) -> None:
    for col in range(1, 23):
        cell = ws.cell(row_num, col)
        cell.fill = _GREY if col >= 10 else _SECTION
        cell.border = _SPLIT if col == 10 else _BORDER
        cell.font = Font(size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row_num, 1).value = section_number
    ws.cell(row_num, 2).value = section_title
    ws.cell(row_num, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row_num].height = 22


def _write_data_row(ws: Any, row_num: int, line: dict[str, Any]) -> None:
    quantity, material_unit, _material_total, work_unit, _work_total, _row_total = _cost_parts(line)

    ws.cell(row_num, 2).value = _text(line.get("name"))
    ws.cell(row_num, 3).value = _text(line.get("unit"))

    for white_col, grey_col in WHITE_ZONE_MIRROR.items():
        cell = ws[f"{white_col}{row_num}"]
        cell.value = f"={grey_col}{row_num}"
        cell.fill = _WHITE
        cell.border = _BORDER
        cell.font = Font(size=10)
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        cell.number_format = "0.00" if white_col == "D" else "0"

    values = {
        "J": quantity,
        "K": material_unit,
        "L": f"=J{row_num}*K{row_num}",
        "M": work_unit,
        "N": f"=J{row_num}*M{row_num}",
        "O": f"=L{row_num}+N{row_num}",
    }
    for col, value in values.items():
        cell = ws[f"{col}{row_num}"]
        cell.value = value
        cell.fill = _GREY
        cell.border = _SPLIT if col == "J" else _BORDER
        cell.font = Font(size=10)
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        cell.number_format = "0.00" if col == "J" else "0"

    for col in range(1, 23):
        cell = ws.cell(row_num, col)
        if col in (1, 16, 17, 18, 19, 20, 21, 22):
            cell.fill = _GREY if col >= 16 else _WHITE
            cell.border = _BORDER
        if col in (2, 3):
            cell.fill = _WHITE
            cell.border = _BORDER
            cell.font = Font(size=10)
            cell.alignment = Alignment(horizontal="left" if col == 2 else "center", vertical="top", wrap_text=True)

    ws.row_dimensions[row_num].height = 18


def _write_section_total(ws: Any, row_num: int, first_data_row: int, last_data_row: int) -> None:
    totals = {
        2: "Итого по разделу:",
        6: f"=L{row_num}",
        8: f"=N{row_num}",
        9: f"=O{row_num}",
        12: f"=SUM(L{first_data_row}:L{last_data_row})",
        14: f"=SUM(N{first_data_row}:N{last_data_row})",
        15: f"=SUM(O{first_data_row}:O{last_data_row})",
    }
    for col in range(1, 23):
        cell = ws.cell(row_num, col)
        cell.value = totals.get(col, "")
        cell.fill = _GREY if col >= 10 else _TOTAL
        cell.border = _SPLIT if col == 10 else _BORDER
        cell.font = Font(size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.number_format = "0"
    ws.row_dimensions[row_num].height = 22


def _write_grand_total(ws: Any, row_num: int, section_total_rows: list[int]) -> None:
    material_refs = ",".join(f"L{row}" for row in section_total_rows)
    work_refs = ",".join(f"N{row}" for row in section_total_rows)
    total_refs = ",".join(f"O{row}" for row in section_total_rows)
    totals = {
        2: "ИТОГО ПО СМЕТЕ:",
        6: f"=L{row_num}",
        8: f"=N{row_num}",
        9: f"=O{row_num}",
        12: f"=SUM({material_refs})",
        14: f"=SUM({work_refs})",
        15: f"=SUM({total_refs})",
    }
    for col in range(1, 23):
        cell = ws.cell(row_num, col)
        cell.value = totals.get(col, "")
        cell.fill = _GREY if col >= 10 else _TOTAL
        cell.border = _SPLIT if col == 10 else _BORDER
        cell.font = Font(size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.number_format = "0"
    ws.row_dimensions[row_num].height = 24


def build_workbook(results_dir: Path, review_workbook: Path | None, logo_path: Path) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = dt.date.today().strftime("%d.%m.%Y")

    _setup_top_header(ws, logo_path=logo_path, address=_extract_project_address(review_workbook))
    _setup_dimensions(ws)

    row_num = FIRST_SECTION_ROW
    section_total_rows: list[int] = []
    for section_number, (section_title, lines) in enumerate(_all_section_blocks(results_dir), start=2):
        _write_section_header(ws, row_num, section_number, section_title)
        row_num += 1
        first_data_row = row_num
        for line in lines:
            _write_data_row(ws, row_num, line)
            row_num += 1
        last_data_row = row_num - 1
        if first_data_row <= last_data_row:
            ws.merge_cells(start_row=first_data_row, start_column=1, end_row=last_data_row, end_column=1)
            section_cell = ws.cell(first_data_row, 1)
            section_cell.value = section_title
            section_cell.font = Font(size=10, bold=True)
            section_cell.fill = _WHITE
            section_cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True)
            section_cell.border = _BORDER
        _write_section_total(ws, row_num, first_data_row, last_data_row)
        section_total_rows.append(row_num)
        row_num += 1

    row_num += 1
    _write_grand_total(ws, row_num, section_total_rows)
    ws.auto_filter.ref = f"A{GROUP_HEADER_ROW}:O{row_num}"

    return wb


def main() -> int:
    parser = argparse.ArgumentParser(description="Export all calculator section results to an estimate-style workbook.")
    parser.add_argument("--results-dir", required=True, help="Directory with <section>_result.json files.")
    parser.add_argument("--review-workbook", default=None, help="Optional filled review workbook for project metadata.")
    parser.add_argument("--out", required=True, help="Output .xlsx path.")
    parser.add_argument("--logo-path", default=str(DEFAULT_LOGO_PATH), help="Path to Brick House logo.")
    args = parser.parse_args()

    results_dir = Path(args.results_dir)
    review_workbook = Path(args.review_workbook) if args.review_workbook else None
    out_path = Path(args.out)
    logo_path = Path(args.logo_path)

    wb = build_workbook(results_dir=results_dir, review_workbook=review_workbook, logo_path=logo_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(out_path.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
