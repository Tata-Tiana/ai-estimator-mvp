"""Regenerates output/manual_values_registry.xlsx from the current section contracts.

This registry is a universal manual-input template, not a project sample. It may preserve
already-filled values only when those values are approved as project-independent defaults.
Do not store quantities, fixed amounts, or comments copied from ARK/TRC/USV estimates here.

Matched by (section_code, key); drops rows for fields that are no longer MANUAL_REVIEW/
SUPPLIER_INPUT (e.g. once a field becomes AUTO_CALCULATED); adds newly-manual fields blank.

NOTE (2026-08-08): Elena now edits typical values in a Google Sheet, not in this local file
directly (see download_manual_values_registry.py). Running this script still works locally
(useful to see which new blank rows a contract change adds), but any new row it produces must
be copied into the Google Sheet by hand, or it will be silently dropped the next time
download_manual_values_registry.py overwrites this file.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from build_review_workbook_from_contracts import (
    default_contract_paths,
    load_yaml_contract,
    manual_values_rows_for_contract,
    section_code,
    section_name,
)

PIPELINE_DIR = Path(__file__).resolve().parent
DEFAULT_REGISTRY_PATH = PIPELINE_DIR.parents[1] / "output" / "manual_values_registry.xlsx"

HEADERS = ["Раздел", "Наименование", "Ед. изм.", "Типовое значение", "Комментарий", "section_code", "key"]


def load_existing_values(path: Path) -> dict[tuple[str, str], dict[str, Any]]:
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["manual_values"] if "manual_values" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    existing: dict[tuple[str, str], dict[str, Any]] = {}
    for raw in rows:
        row = dict(zip(headers, raw))
        sec_code = row.get("section_code")
        key = row.get("key")
        if not sec_code or not key:
            continue
        existing[(str(sec_code).strip(), str(key).strip())] = {
            "value": row.get("Типовое значение"),
            "comment": row.get("Комментарий") or "",
        }
    return existing


def build_registry(contract_paths: list[Path], existing: dict[tuple[str, str], dict[str, Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "manual_values"
    ws.append(HEADERS)

    kept = 0
    added = 0
    dropped_keys = set(existing.keys())

    for path in contract_paths:
        contract = load_yaml_contract(path)
        sec_code = section_code(contract)
        for param in manual_values_rows_for_contract(contract):
            key = param.get("key", "")
            entry = existing.get((sec_code, key))
            if entry is not None:
                dropped_keys.discard((sec_code, key))
                value, comment = entry["value"], entry["comment"]
                kept += 1
            else:
                value, comment = None, ""
                added += 1
            ws.append([
                section_name(contract),
                param.get("label_ru", key),
                param.get("unit", ""),
                value,
                comment,
                sec_code,
                key,
            ])

    for column, width in [("A", 26), ("B", 40), ("C", 10), ("D", 16), ("E", 60), ("F", 20), ("G", 30)]:
        ws.column_dimensions[column].width = width
    for column in ["F", "G"]:
        ws.column_dimensions[column].hidden = True
    ws.freeze_panes = "A2"

    print(f"kept {kept} existing values, added {added} new blank rows, dropped {len(dropped_keys)} stale rows")
    if dropped_keys:
        for sec_code, key in sorted(dropped_keys):
            print(f"  dropped: {sec_code}.{key}")
    return wb


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default=str(DEFAULT_REGISTRY_PATH))
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_path = Path(args.output)
    existing = load_existing_values(output_path)
    wb = build_registry(default_contract_paths(), existing)
    wb.save(output_path)
    print(f"written: {output_path}")


if __name__ == "__main__":
    main()
