"""Builds the review workbook the same way build_review_workbook_from_contracts.py
does, but fills sheet 01 with real values from a chat-extraction extraction_output.json
instead of leaving it blank. This is the missing "step 2" from ADAPTER_BUILD_PLAN.md's
"Полный путь PDF -> смета" section - the first real test of it against a real JSON.

Sheet 01 is filled from the extraction JSON, sheet 02 (prices) is filled from a
price_registry_filled_v*.xlsx (see build_rebar_lookup for the one templated-code
expansion that needs project data; see reports/step_27_sheet02_price_fill_plan.md for
the rest). Sheet 03 is filled from raw_table_rows in the extraction JSON. Sheets 00/04-06
are unchanged, built the same way as the empty-template script.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "experiments" / "box_calculator"))

from metal_delivery_allocator import MetalSection, allocate_metal_deliveries  # noqa: E402

from build_review_workbook_from_contracts import (  # noqa: E402
    COMPACT_ROW_HEIGHT,
    FILL_FOUND,
    FILL_HEADER,
    FILL_INPUT,
    FILL_MISSING,
    FILL_REVIEW,
    FILL_WHITE,
    FONT_NAME,
    ITEM_BLOCK_HEADERS,
    PROJECT_HEADERS,
    apply_table_style,
    append_block,
    autofit_row_heights,
    append_section_band,
    build_constructor_sheet,
    build_contracts_summary_sheet,
    build_instruction_sheet,
    build_manual_values_sheet,
    build_prices_sheet,
    build_raw_contracts_sheet,
    correction_headers,
    default_contract_paths,
    diagnostic_repeated_row_params,
    hide_diagnostic_sheets,
    load_manual_values_registry,
    load_price_registry,
    load_yaml_contract,
    merge_row_full_width,
    production_repeated_row_params,
    rebar_group_keys_for_contract,
    rebar_item_weight_kg,
    restyle_block_sheet,
    restyle_section_bands,
    scalar_review_rows_for_contract,
    section_code,
    section_name,
    set_widths,
    style_block_title_row,
    style_header_row,
)

ROOT = Path(__file__).resolve().parents[2]
PIPELINE_DIR = Path(__file__).resolve().parent
DEFAULT_PRICE_REGISTRY = ROOT / "output" / "price_registry_filled_v4.xlsx"
DEFAULT_MANUAL_VALUES_REGISTRY = ROOT / "output" / "manual_values_registry.xlsx"


def project_status_fill(status: str):
    """Visual status colors for sheet 01 rows."""
    if status.startswith("Найдено"):
        return FILL_FOUND
    if status.startswith("Проверьте"):
        return FILL_REVIEW
    if status.startswith("Не найдено"):
        return FILL_MISSING
    if status.startswith("Не требуется"):
        return FILL_WHITE
    return FILL_REVIEW


# Narrow, purpose-built support for the "sum(included <group>.<field>)" auto_calculated formula
# shape only — not a general expression evaluator. Most auto_calculated entries in
# section_contract.yaml are internal computation notes never meant to reach sheet 01 (see plan
# section 34); only entries a review_parameters row explicitly opts into via cross_check_key get
# rendered, and only this one formula shape is understood. If a future cross-check needs a
# different formula shape, extend this function narrowly rather than building a generic parser.
_SUM_INCLUDED_RE = re.compile(r"^sum\(included (\w+)\.(\w+)\)$")


def compute_cross_check_value(formula: str, found_groups: dict[str, list[Any]]) -> float | None:
    match = _SUM_INCLUDED_RE.match(formula.strip())
    if not match:
        return None
    group_key, field = match.groups()
    rows = found_groups.get(group_key) or []
    total = 0.0
    found_any = False
    for item in rows:
        value = item.get("value") or {}
        include = value.get("include_in_communications", value.get("include", True))
        if not include:
            continue
        explicit_total = value.get(field)
        if explicit_total is not None:
            total += float(explicit_total)
            found_any = True
            continue
        # Mirror earthworks_calculator.py's own pipe_items fallback: pipe_length_m * quantity when
        # no explicit total_length_m is given on the row (e.g. straight pipe segments). Rows with
        # no length component at all (elbows, tees, plugs) contribute 0, not an error.
        pipe_length = value.get("pipe_length_m")
        quantity = value.get("quantity")
        if pipe_length is not None and quantity is not None:
            total += float(pipe_length) * float(quantity)
            found_any = True
    return total if found_any else None


def group_items(found_groups: dict[str, list[Any]], group_key: str) -> list[dict[str, Any]]:
    return [item for item in found_groups.get(group_key, []) if isinstance(item, dict)]


def sum_group_numeric_field(found_groups: dict[str, list[Any]], group_key: str, field_key: str) -> float | None:
    total = 0.0
    found_any = False
    for item in group_items(found_groups, group_key):
        value = item.get("value") or {}
        field_value = value.get(field_key)
        if field_value is None:
            continue
        total += float(field_value)
        found_any = True
    return round(total, 3) if found_any else None


def sum_communications_pipe_items(found_groups: dict[str, list[Any]]) -> float | None:
    total = 0.0
    found_any = False
    for item in group_items(found_groups, "communications_pipe_items"):
        value = item.get("value") or {}
        if not value.get("include_in_communications", True):
            continue
        explicit_total = value.get("total_length_m")
        if explicit_total is not None:
            total += float(explicit_total)
            found_any = True
            continue
        pipe_length = value.get("pipe_length_m")
        quantity = value.get("quantity")
        if pipe_length is not None and quantity is not None:
            total += float(pipe_length) * float(quantity)
            found_any = True
    return round(total, 3) if found_any else None


def min_group_confidence(found_groups: dict[str, list[Any]], group_key: str) -> str:
    values: list[float] = []
    for item in group_items(found_groups, group_key):
        try:
            values.append(float(item.get("confidence")))
        except (TypeError, ValueError):
            continue
    return f"{min(values):.2f}" if values else ""


def earthworks_alternative_scalar(
    target_code: str,
    found_groups: dict[str, list[Any]],
) -> tuple[Any, str, str, str, str] | None:
    """Return a sheet-01 scalar replacement when earthworks data is present as item rows."""
    if target_code == "pit_excavation_depth_m" and group_items(found_groups, "pit_items"):
        total = sum_group_numeric_field(found_groups, "pit_items", "volume_m3")
        return (
            "",
            "Не требуется (есть готовые объёмы выемки ниже)",
            "",
            f"Глубина не используется: объём механизированной выемки берётся из pit_items. Сумма: {total:g} м3.",
            min_group_confidence(found_groups, "pit_items"),
        )
    if target_code == "sand_base_volume_m3" and group_items(found_groups, "sand_items"):
        total = sum_group_numeric_field(found_groups, "sand_items", "volume_m3")
        return (
            "",
            "Не требуется (есть готовые объёмы песка ниже)",
            "",
            "Песок берётся из sand_items по категориям, не из одной строки 'под основание'. "
            f"Сумма до коэффициента: {total:g} м3.",
            min_group_confidence(found_groups, "sand_items"),
        )
    if target_code == "communications_length_m" and group_items(found_groups, "communications_pipe_items"):
        total = sum_communications_pipe_items(found_groups)
        return (
            total,
            "Найдено (автосумма строк труб ниже)",
            "",
            "Сумма включённых позиций communications_pipe_items: "
            f"{total:g} м. Позиции без линейной длины (углы, тройники, заглушки) не добавляют метры.",
            min_group_confidence(found_groups, "communications_pipe_items"),
        )
    if target_code == "trench_volume_m3" and group_items(found_groups, "trench_routes"):
        total = sum_group_numeric_field(found_groups, "trench_routes", "volume_m3")
        return (
            total,
            "Найдено (автосумма маршрутов ниже)",
            "",
            f"Сумма volume_m3 из trench_routes: {total:g} м3.",
            min_group_confidence(found_groups, "trench_routes"),
        )
    return None


def candidate_sum_scalar(item: dict[str, Any] | None, target_code: str) -> tuple[float, str, str] | None:
    if item is None or item.get("value") is not None:
        return None
    candidates = [candidate for candidate in item.get("candidates") or [] if isinstance(candidate, dict)]
    matching_values: list[float] = []
    fragments: list[str] = []
    confidences: list[float] = []
    for candidate in candidates:
        if candidate.get("target_code") != target_code:
            continue
        value = candidate.get("value")
        if value is None:
            continue
        try:
            matching_values.append(float(value))
        except (TypeError, ValueError):
            continue
        raw_text = candidate.get("raw_text")
        if raw_text:
            fragments.append(str(raw_text))
        try:
            confidences.append(float(candidate.get("confidence")))
        except (TypeError, ValueError):
            pass
    if len(matching_values) < 2:
        return None
    confidence = f"{min(confidences):.2f}" if confidences else display_confidence(item)
    return round(sum(matching_values), 3), "; ".join(fragments), confidence


def candidate_component_fragments(item: dict[str, Any], limit: int = 4) -> str:
    fragments: list[str] = []
    for candidate in item.get("candidates") or []:
        if not isinstance(candidate, dict):
            continue
        value = candidate.get("value")
        unit = candidate.get("unit") or candidate.get("normalized_unit") or ""
        raw_text = candidate.get("raw_text") or candidate.get("notes") or candidate.get("item_name")
        parts = []
        if value is not None:
            parts.append(f"{value}{' ' + str(unit) if unit else ''}")
        if raw_text:
            parts.append(str(raw_text))
        if parts:
            fragments.append(" — ".join(parts))
    if len(fragments) > limit:
        fragments = fragments[:limit] + [f"... ещё {len(fragments) - limit} кандидатов"]
    return "; ".join(fragments)


def unresolved_candidate_fragment(item: dict[str, Any]) -> str:
    fragments = candidate_component_fragments(item)
    note = item.get("notes") or "Есть несколько candidates, но для этого target_code нет универсального правила автосуммы."
    if not fragments:
        return str(note)
    return f"{note} Автосумма не применена; проверьте компоненты: {fragments}"


def item_fragment(item: dict[str, Any] | None) -> str:
    if item is None:
        return ""
    raw_text = str(item.get("raw_text") or "")
    notes = str(item.get("notes") or "")
    if raw_text and notes and notes not in raw_text:
        return f"{raw_text}\nNotes: {notes}"
    return raw_text or notes


AUTO_SUM_CANDIDATE_TARGETS = {
    "foundation_slab": {
        "membrane_area_m2": "компонентов мембраны",
    },
    "flat_roof": {
        "roof_internal_drains_count": "внутренних кровельных воронок",
        "roof_parapet_drains_count": "парапетных воронок",
    },
    "load_bearing_walls_lintels": {
        # Elena's rule 2026-07-25: external + internal load-bearing wall areas are the same
        # continuous wall body (differ only by block thickness) - sum them when the PDF gives no
        # ready total. Already in the extraction prompt (calculator_targets_compact.json), but the
        # model doesn't always apply it - this is the workbook-side safety net so it doesn't just
        # sit as an unresolved needs_review when it doesn't. See
        # cutoff_waterproofing_external_internal_sum memory.
        "cutoff_waterproofing_load_bearing_walls_area": "наружных и внутренних несущих стен",
    },
    "floor_slab_1": {
        # Real ТРЦ case, 2026-08-05: PDF gives the slab's edge perimeter and under-slab formwork
        # area as two zone components (main slab + kitchen/dining) with no combined total for either,
        # same shape as the existing slab_zones concrete-volume case. Safe to auto-sum: both are
        # simple linear/area components of one physical slab, not alternative readings of the same
        # thing.
        "floor_slab_1_slab_edge_perimeter": "периметра торца плиты по зонам",
        "floor_slab_1_under_slab_formwork_area": "площади опалубки под плитой по зонам",
        # Real ТРЦ case, 2026-08-05: PDF gives EPS-100 volume as three distinct, non-overlapping
        # components of the SAME total (edge insulation + under-slab insulation + kitchen zone
        # insulation), not alternative readings of one measurement — safe to sum into the single
        # total_eps_volume_from_spec_m3 target.
        "floor_slab_1_eps100_volume": "объёма ЭППС-100 по компонентам утепления",
    },
}


def group_sources(found_groups: dict[str, list[Any]], group_key: str) -> str:
    sources: list[str] = []
    for item in group_items(found_groups, group_key):
        parts = []
        if item.get("source_pdf"):
            parts.append(str(item["source_pdf"]))
        if item.get("page_number") is not None:
            parts.append(f"стр. {item['page_number']}")
        source = ", ".join(parts)
        if source and source not in sources:
            sources.append(source)
    return "; ".join(sources)


def group_fragments(found_groups: dict[str, list[Any]], group_key: str, limit: int = 5) -> str:
    fragments: list[str] = []
    for item in group_items(found_groups, group_key):
        raw_text = item.get("raw_text") or item.get("notes")
        if raw_text:
            fragments.append(str(raw_text))
    if len(fragments) > limit:
        fragments = fragments[:limit] + [f"... ещё {len(fragments) - limit} строк"]
    return "; ".join(fragments)


THERMAL_INSERT_SCALAR_TARGETS = {
    "thermal_insert_50_length",
    "thermal_insert_100_length",
    "thermal_insert_combined_length_m",
    "thermal_insert_50_material_spec_qty",
    "thermal_insert_100_material_spec_qty",
}


def foundation_alternative_scalar(
    target_code: str,
    found: dict[str, Any] | None,
    found_groups: dict[str, list[Any]],
    found_by_target: dict[str, Any],
) -> tuple[Any, str, str, str, str] | None:
    if target_code == "membrane_area_m2":
        label = AUTO_SUM_CANDIDATE_TARGETS.get("foundation_slab", {}).get(target_code)
        if not label:
            return None
        candidate_sum = candidate_sum_scalar(found, target_code)
        if candidate_sum is None:
            return None
        total, fragments, confidence = candidate_sum
        return (
            total,
            "Проверьте (автосумма компонентов)",
            found.get("source_pdf") or "",
            f"Автосумма {label}: {fragments}. Общего итога в PDF нет.",
            confidence,
        )
    if target_code == "concrete_project_volume" and group_items(found_groups, "slab_zones"):
        total = sum_group_numeric_field(found_groups, "slab_zones", "concrete_volume_m3")
        return (
            total,
            "Найдено (через slab_zones)",
            group_sources(found_groups, "slab_zones"),
            "Единый итог бетона в PDF не указан; для расчёта используется сумма строк slab_zones. "
            f"Строки: {group_fragments(found_groups, 'slab_zones')}",
            min_group_confidence(found_groups, "slab_zones"),
        )
    if target_code in THERMAL_INSERT_SCALAR_TARGETS and group_items(found_groups, "thermal_insert_items"):
        length_total = sum_group_numeric_field(found_groups, "thermal_insert_items", "length_m")
        material_total = sum_group_numeric_field(found_groups, "thermal_insert_items", "material_spec_qty_m3")
        totals = []
        if length_total is not None:
            totals.append(f"длина {length_total:g} мп")
        if material_total is not None:
            totals.append(f"материал {material_total:g} м3")
        totals_text = f" ({', '.join(totals)})" if totals else ""
        return (
            "",
            "Не требуется (есть thermal_insert_items)",
            group_sources(found_groups, "thermal_insert_items"),
            "Скалярные поля 50/100 мм не заполняются, потому что PDF дал термовставки "
            f"произвольными типоразмерами; рабочий источник — строки thermal_insert_items{totals_text}. "
            f"Строки: {group_fragments(found_groups, 'thermal_insert_items')}",
            min_group_confidence(found_groups, "thermal_insert_items"),
        )
    combined_length_found = found_by_target.get("thermal_insert_combined_length_m")
    if target_code in {"thermal_insert_50_length", "thermal_insert_100_length"} and combined_length_found is not None:
        combined_value = combined_length_found.get("value")
        value_text = f"{combined_value:g} мп" if isinstance(combined_value, (int, float)) else ""
        return (
            "",
            "Не требуется (есть общая длина термовставок)",
            combined_length_found.get("source_pdf") or "",
            "Раздельная длина монтажа для этой толщины не заполняется: PDF даёт одну общую длину узла"
            f"{(' ' + value_text) if value_text else ''} на оба слоя термовставки — копировать её в оба "
            "раздельных поля нельзя, это задвоило бы работу по монтажу. Рабочий источник — "
            "thermal_insert_combined_length_m.",
            display_confidence(combined_length_found),
        )
    return None


def load_bearing_walls_lintels_alternative_scalar(
    target_code: str,
    found: dict[str, Any] | None,
    found_groups: dict[str, list[Any]],
) -> tuple[Any, str, str, str, str] | None:
    if target_code == "floors_count" and group_items(found_groups, "wall_block_items"):
        return (
            "",
            "Не требуется (этажность выводится из wall_block_items)",
            group_sources(found_groups, "wall_block_items"),
            "Текстовая этажность PDF не используется как главный источник: при наличии wall_block_items "
            "калькулятор определяет второй уровень по строкам кладки.",
            min_group_confidence(found_groups, "wall_block_items"),
        )
    if target_code in {
        "parapet_masonry_volume",
        "parapet_gas_block_d500_250_volume",
        "main_wall_gas_block_400_spec_volume",
        "main_wall_gas_block_250_spec_volume",
        "floor_2_masonry_volume",
    } and group_items(found_groups, "wall_block_items"):
        return (
            "",
            "Не требуется (объём есть в wall_block_items)",
            group_sources(found_groups, "wall_block_items"),
            "Фиксированный scalar объёма кладки не требуется, если объёмы пришли строками "
            "wall_block_items (по ролям main_walls/floor_2/parapet). Калькулятор берёт "
            "production repeated rows, а не старые одиночные поля.",
            min_group_confidence(found_groups, "wall_block_items"),
        )
    if target_code == "cutoff_waterproofing_load_bearing_walls_area":
        label = AUTO_SUM_CANDIDATE_TARGETS.get("load_bearing_walls_lintels", {}).get(target_code)
        candidate_sum = candidate_sum_scalar(found, target_code) if label else None
        if candidate_sum is None:
            return None
        total, fragments, confidence = candidate_sum
        return (
            total,
            "Проверьте (автосумма компонентов)",
            found.get("source_pdf") or "" if found else "",
            f"Автосумма {label}: {fragments}. Общего итога в PDF нет — наружные и внутренние "
            "несущие стены суммируются как одна стена (Elena, 2026-07-25); перегородки в сумму "
            "не входят.",
            confidence,
        )
    return None


def floor_slab_1_alternative_scalar(
    target_code: str,
    found: dict[str, Any] | None,
    found_groups: dict[str, list[Any]],
    found_by_target: dict[str, Any],
) -> tuple[Any, str, str, str, str] | None:
    if target_code == "floor_slab_1_concrete_volume" and group_items(found_groups, "slab_zones"):
        total = sum_group_numeric_field(found_groups, "slab_zones", "concrete_volume_m3")
        if total is None:
            return None
        return (
            total,
            "Найдено (через slab_zones)",
            group_sources(found_groups, "slab_zones"),
            "Единый итог бетона плиты в PDF не указан; для расчёта используется сумма строк "
            f"slab_zones. Строки: {group_fragments(found_groups, 'slab_zones')}",
            min_group_confidence(found_groups, "slab_zones"),
        )
    if target_code == "floor_slab_1_beams_concrete_volume" and group_items(found_groups, "beam_items"):
        total = sum_group_numeric_field(found_groups, "beam_items", "concrete_volume_m3")
        if total is None:
            return None
        return (
            total,
            "Найдено (автосумма beam_items)",
            group_sources(found_groups, "beam_items"),
            "Готовый итог бетона балок в PDF не указан; для расчёта используется сумма "
            f"beam_items.concrete_volume_m3: {total:g} м3. Это не применяется к утеплению балок.",
            min_group_confidence(found_groups, "beam_items"),
        )
    if target_code in {
        "floor_slab_1_slab_edge_perimeter",
        "floor_slab_1_under_slab_formwork_area",
        "floor_slab_1_eps100_volume",
    }:
        label = AUTO_SUM_CANDIDATE_TARGETS.get("floor_slab_1", {}).get(target_code)
        candidate_sum = candidate_sum_scalar(found, target_code) if label else None
        if candidate_sum is None:
            return None
        total, fragments, confidence = candidate_sum
        return (
            total,
            "Проверьте (автосумма компонентов)",
            found.get("source_pdf") or "" if found else "",
            f"Автосумма {label}: {fragments}. Общего итога в PDF нет — плита дана по зонам "
            "(основная + кухня/гостиная), компоненты суммируются как один физический торец/площадь/объём.",
            confidence,
        )
    if target_code == "floor_slab_1_beams_formwork_area" and group_items(found_groups, "beam_items"):
        total = sum_group_numeric_field(found_groups, "beam_items", "formwork_area_m2")
        if total is None:
            return None
        return (
            total,
            "Найдено (автосумма beam_items)",
            group_sources(found_groups, "beam_items"),
            "Готовый итог площади опалубки балок в PDF не указан или ненадёжен; для расчёта "
            f"используется сумма beam_items.formwork_area_m2: {total:g} м2 (контракт уже "
            "документирует этот fallback как поведение калькулятора).",
            min_group_confidence(found_groups, "beam_items"),
        )
    if target_code == "floor_slab_1_edge_formwork_height":
        thickness_found = found_by_target.get("floor_slab_1_slab_thickness")
        if thickness_found is None or thickness_found.get("value") is None:
            return None
        thickness_value = thickness_found.get("value")
        return (
            thickness_value,
            "Найдено (=толщина плиты)",
            thickness_found.get("source_pdf") or "",
            "Отдельной строки высоты торцевой опалубки в PDF нет; для контрольного расчёта "
            f"принимается равной толщине плиты: {thickness_value} м (см. floor_slab_1_slab_thickness). "
            "На стоимость не влияет — используется только для контрольной сверки площади торца.",
            display_confidence(thickness_found),
        )
    return None


def floor_slab_2_alternative_scalar(
    target_code: str,
    found_groups: dict[str, list[Any]],
) -> tuple[Any, str, str, str, str] | None:
    if target_code == "floor_slab_2_beams_concrete_volume" and group_items(
        found_groups, "floor_slab_2_beam_items"
    ):
        total = sum_group_numeric_field(found_groups, "floor_slab_2_beam_items", "concrete_volume_m3")
        if total is None:
            return None
        return (
            total,
            "Найдено (автосумма floor_slab_2_beam_items)",
            group_sources(found_groups, "floor_slab_2_beam_items"),
            "Готовый итог бетона балок в PDF не указан; для расчёта используется сумма "
            f"floor_slab_2_beam_items.concrete_volume_m3: {total:g} м3. "
            "Это не применяется к утеплению балок.",
            min_group_confidence(found_groups, "floor_slab_2_beam_items"),
        )
    return None


ROOF_ZONES_FALLBACK_ONLY_TARGETS = {
    "roof_area_level_1",
    "roof_area_level_2",
    "roof_parapet_length_level_1",
    "roof_parapet_length_level_2",
    "roof_vent_wall_abutment_level_1",
    "roof_vent_wall_abutment_level_2",
}


def flat_roof_alternative_scalar(
    target_code: str,
    found: dict[str, Any] | None,
    found_groups: dict[str, list[Any]],
) -> tuple[Any, str, str, str, str] | None:
    if target_code in ROOF_ZONES_FALLBACK_ONLY_TARGETS and group_items(found_groups, "roof_zones"):
        return (
            "",
            "Не требуется (есть roof_zones)",
            group_sources(found_groups, "roof_zones"),
            "Fallback-поле уровня 1/2 не заполняется: площадь/парапет/примыкания уже пришли по зонам "
            f"в roof_zones ниже. Строки: {group_fragments(found_groups, 'roof_zones')}",
            min_group_confidence(found_groups, "roof_zones"),
        )
    label = AUTO_SUM_CANDIDATE_TARGETS.get("flat_roof", {}).get(target_code)
    if not label:
        return None
    candidate_sum = candidate_sum_scalar(found, target_code)
    if candidate_sum is None:
        return None
    total, fragments, confidence = candidate_sum
    return (
        total,
        "Проверьте (автосумма компонентов)",
        found.get("source_pdf") or "",
        f"Автосумма {label}: {fragments}. Общего итога в PDF нет.",
        confidence,
    )


def earthworks_group_total_row(
    sec_code: str,
    group_key: str,
    found_groups: dict[str, list[Any]],
    headers_len: int,
) -> list[Any] | None:
    if sec_code != "earthworks":
        return None
    if group_key == "pit_items":
        total = sum_group_numeric_field(found_groups, group_key, "volume_m3")
        unit = "м3"
        note = "Сумма готовых объёмов выемки грунта; используется вместо площадь*глубина."
    elif group_key == "sand_items":
        total = sum_group_numeric_field(found_groups, group_key, "volume_m3")
        unit = "м3"
        note = "Сумма готовых объёмов песка до коэффициента уплотнения."
    elif group_key == "trench_routes":
        total = sum_group_numeric_field(found_groups, group_key, "volume_m3")
        unit = "м3"
        note = "Сумма готовых объёмов траншей из строк маршрутов."
    elif group_key == "communications_pipe_items":
        total = sum_communications_pipe_items(found_groups)
        unit = "м"
        note = "Сумма включённых линейных труб; фитинги без длины в метры не входят."
    else:
        return None
    if total is None:
        return None

    row = [
        "Итого по строкам блока",
        total,
        unit,
        "Итого (авто по строкам блока)",
        min_group_confidence(found_groups, group_key),
        "",
        "",
        note,
        "",
        "",
        "",
        "",
        "",
        "",
    ]
    return row + [""] * max(0, headers_len - len(row))


def auto_calculated_by_key(contract: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {entry["key"]: entry for entry in (contract.get("auto_calculated") or []) if entry.get("key")}


# 2026-07-29: sections/fields that only exist for some projects (U-block lintels, monolithic
# lintels — per floor, independently) currently show the exact same "Не найдено" status whether
# the field is safely inapplicable or genuinely missing money-critical data. Real ARK case that
# motivated this: floor_1_lintel_monolithic_total_length_m was found (5.4, proving monolithic
# lintels exist on this floor) but floor_1_lintel_monolithic_concrete_volume_m3 was missing — the
# calculator used to silently price concrete material as 0 while still billing the concreting
# work by length (see plan section 41 / calculate_monolithic_lintel_block fix). Each inner list is
# a "presence pair": if extraction found a truthy value for ANY code in the pair, the others are
# confirmed required — this project definitely has this construction, so a still-missing sibling
# is a real gap, not a safely-blank optional field. Scoped to the 4 pairs with a proven money-loss
# bug (U-block + monolithic lintels, floor 1 and floor 2); other conditional sections (floor_2
# masonry, parapet, vent chimney cladding, beam_items presence) don't have a sibling-field risk
# today, so are not included here — extend this table if a similar bug is found for them.
SECTION_PRESENCE_PAIRS: dict[str, list[list[str]]] = {
    "load_bearing_walls_lintels": [
        ["lintel_total_length", "lintel_concrete_volume"],
        ["floor_2_lintel_total_length", "floor_2_lintel_concrete_volume"],
        ["floor_1_lintel_monolithic_total_length", "floor_1_lintel_monolithic_concrete_volume"],
        ["floor_2_lintel_monolithic_total_length", "floor_2_lintel_monolithic_concrete_volume"],
    ],
}

CONDITIONAL_ABSENT_TARGETS: dict[str, dict[str, str]] = {
    "load_bearing_walls_lintels": {
        "vent_chimney_gas_block_150_volume": "Не блокер, если в проекте нет обкладки вентканалов/дымохода газобетоном 150 мм. Это не Schiedel.",
        "floor_2_lintel_total_length": "Не блокер, если на 2-м этаже нет перемычек в U-блоках.",
        "floor_2_lintel_concrete_volume": "Не блокер, если на 2-м этаже нет перемычек в U-блоках.",
        "floor_2_lintel_monolithic_concrete_volume": "Не блокер, если на 2-м этаже нет монолитных перемычек.",
        "floor_2_lintel_monolithic_total_length": "Не блокер, если на 2-м этаже нет монолитных перемычек.",
        "floor_2_lintel_insulation_length": "Не блокер, если на 2-м этаже нет утепляемых монолитных перемычек.",
        "floor_2_lintel_formwork_horizontal_area": "Не блокер, если на 2-м этаже нет монолитных перемычек.",
        "floor_2_lintel_formwork_vertical_area": "Не блокер, если на 2-м этаже нет монолитных перемычек.",
        "floor_2_lintel_insulation_eps_volume": "Не блокер, если на 2-м этаже нет утепляемых монолитных перемычек.",
    },
}

DIAGNOSTIC_ONLY_MISSING_TARGETS: dict[str, dict[str, str]] = {
    "load_bearing_walls_lintels": {
        "floor_1_lintel_groove_length": "Диагностическое поле: перемычки в штробе пока фиксируются для будущей доработки и не участвуют в смете.",
        "floor_2_lintel_groove_length": "Диагностическое поле: перемычки в штробе пока фиксируются для будущей доработки и не участвуют в смете.",
    },
}


def compute_confirmed_required(found_by_target: dict[str, Any], sec_code: str) -> set[str]:
    """Returns target_codes that are confirmed required for this project because a sibling
    code in the same presence pair was found with a truthy value. See SECTION_PRESENCE_PAIRS."""
    confirmed: set[str] = set()
    for pair in SECTION_PRESENCE_PAIRS.get(sec_code, []):
        signal_fired = False
        for code in pair:
            found = found_by_target.get(code)
            if found is not None and found.get("value") not in (None, 0, 0.0):
                signal_fired = True
                break
        if signal_fired:
            confirmed.update(pair)
    return confirmed


def classified_missing_target(sec_code: str, target_code: str) -> tuple[str, str] | None:
    diagnostic_note = DIAGNOSTIC_ONLY_MISSING_TARGETS.get(sec_code, {}).get(target_code)
    if diagnostic_note:
        return "Не требуется (диагностическое поле)", diagnostic_note
    conditional_note = CONDITIONAL_ABSENT_TARGETS.get(sec_code, {}).get(target_code)
    if conditional_note:
        return "Не требуется / условно отсутствует", conditional_note
    return None


def index_extraction_section(extraction: dict[str, Any], sec_code: str) -> tuple[dict, dict, set]:
    section = (extraction.get("sections") or {}).get(sec_code) or {}
    found_by_target: dict[str, Any] = {}
    found_groups: dict[str, list[Any]] = {}
    for item in section.get("found") or []:
        group_code = item.get("group_code")
        if group_code:
            found_groups.setdefault(group_code, []).append(item)
        else:
            target_code = item.get("target_code")
            if target_code:
                found_by_target[target_code] = item
    missing = set(section.get("missing") or [])
    return found_by_target, found_groups, missing


def display_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def display_confidence(item: dict[str, Any] | None) -> str:
    if item is None:
        return ""
    confidence = item.get("confidence")
    if confidence is None:
        return ""
    try:
        return f"{float(confidence):.2f}"
    except (TypeError, ValueError):
        return str(confidence)


DETAIL_RAW_HEADERS = [
    "№",
    "Исходные колонки PDF",
    "Исходные ячейки PDF",
    "Сырая строка PDF",
    "Ед.",
    "Норм. ед.",
    "Target codes",
    "Needs review",
    "Комментарий parser",
    "Комментарий Елены",
    "section_code",
    "source_pdf",
    "page_number",
    "page_title",
    "table_id",
    "row_index",
    "mapped_target_codes_json",
]


def joined(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return " | ".join("" if item is None else str(item) for item in value)
    return str(value)


def table_title(row: dict[str, Any]) -> str:
    title = row.get("table_title")
    if title:
        return str(title)
    page_title = row.get("page_title") or "Лист без названия"
    table_id = row.get("table_id") or "unknown_table"
    return f"{page_title} — таблица {table_id}"


def table_group_key(row: dict[str, Any]) -> tuple[Any, ...]:
    return (
        row.get("section_code") or "",
        row.get("source_pdf") or "",
        row.get("page_number") or "",
        row.get("page_title") or "",
        row.get("table_id") or "",
        row.get("table_title") or "",
    )


def iter_raw_table_groups(extraction: dict[str, Any]) -> list[tuple[tuple[Any, ...], list[dict[str, Any]]]]:
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for section in (extraction.get("sections") or {}).values():
        for row in section.get("raw_table_rows") or []:
            if not isinstance(row, dict):
                continue
            groups.setdefault(table_group_key(row), []).append(row)
    return list(groups.items())


def build_details_sheet_from_extraction(wb: Workbook, extraction: dict[str, Any]) -> dict[str, int]:
    ws = wb.create_sheet("03_Детали объемов")
    groups = iter_raw_table_groups(extraction)
    max_col = len(DETAIL_RAW_HEADERS)

    for _group_key, rows in groups:
        first = rows[0]
        title = table_title(first)
        context = (
            f"Раздел: {first.get('section_code') or ''} | "
            f"PDF: {first.get('source_pdf') or ''} | "
            f"стр. {first.get('page_number') or ''} | "
            f"Лист: {first.get('page_title') or ''}"
        )

        if ws.max_row == 1 and not ws.cell(1, 1).value:
            ws.cell(1, 1).value = title
        else:
            ws.append([])
            ws.append([title])
        style_block_title_row(ws, ws.max_row, max_col)
        ws.append([context])
        for cell in ws[ws.max_row]:
            cell.fill = FILL_WHITE
            cell.font = Font(name=FONT_NAME, size=10, italic=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        ws.append(DETAIL_RAW_HEADERS)
        style_header_row(ws, ws.max_row, max_col)

        for row in rows:
            mapped = row.get("mapped_target_codes") or []
            ws.append([
                row.get("row_index"),
                joined(row.get("columns")),
                joined(row.get("cells")),
                row.get("raw_text") or "",
                row.get("unit") or "",
                row.get("normalized_unit") or "",
                joined(mapped),
                "да" if row.get("needs_review") else "нет",
                row.get("notes") or "",
                "",
                row.get("section_code") or "",
                row.get("source_pdf") or "",
                row.get("page_number") or "",
                row.get("page_title") or "",
                row.get("table_id") or "",
                row.get("row_index") or "",
                json.dumps(mapped, ensure_ascii=False),
            ])
            for cell in ws[ws.max_row]:
                cell.font = Font(name=FONT_NAME, size=11)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    if not groups:
        ws.append(["Сырые таблицы parser не найдены в extraction JSON."])
        style_block_title_row(ws, 1, 1)

    set_widths(ws, {
        "A": 8,
        "B": 44,
        "C": 58,
        "D": 72,
        "E": 12,
        "F": 12,
        "G": 32,
        "H": 14,
        "I": 52,
        "J": 30,
        "K": 20,
        "L": 34,
        "M": 10,
        "N": 34,
        "O": 28,
        "P": 12,
        "Q": 32,
    })
    for column in ["K", "L", "M", "N", "O", "P", "Q"]:
        ws.column_dimensions[column].hidden = True
    ws.freeze_panes = "A1"

    return {"detail_table_blocks": len(groups), "raw_detail_rows": sum(len(rows) for _, rows in groups)}


def build_project_sheet_from_extraction(
    wb: Workbook, contracts: list[dict[str, Any]], extraction: dict[str, Any]
) -> dict[str, int]:
    ws = wb.create_sheet("01_Проверка проекта")
    project_address = (extraction.get("project_name") or "").strip()
    project_title = f"Разбор проекта: {project_address}" if project_address else "Разбор проекта"
    ws.append([project_title, "", "", "", "", "", "", "", "", "", "", "", "", ""])
    ws.append([])
    ws.append([])
    ws.append(PROJECT_HEADERS)

    counts = {
        "found": 0,
        "needs_review": 0,
        "missing": 0,
        "missing_confirmed_required": 0,
        "conditional_absent": 0,
        "item_rows": 0,
    }

    rebar_weights_by_section = compute_rebar_weights_by_section(contracts, extraction)
    rebar_metal_delivery_allocation = compute_rebar_metal_delivery_allocation(rebar_weights_by_section)
    box_total_metal_weight_kg = round(sum(rebar_weights_by_section.values()), 1)

    for contract in contracts:
        sec_code = section_code(contract)
        found_by_target, found_groups, missing = index_extraction_section(extraction, sec_code)
        confirmed_required = compute_confirmed_required(found_by_target, sec_code)
        append_section_band(ws, [section_name(contract)], len(PROJECT_HEADERS))

        box_delivery_key = REBAR_METAL_DELIVERY_FIELD_BY_SECTION.get(sec_code)

        for param in scalar_review_rows_for_contract(contract):
            review_behavior = param.get("review_behavior") or {}
            target_code = param.get("target_code", "")
            found = found_by_target.get(target_code)
            row_fill = FILL_INPUT
            param_key = param.get("key")

            box_row: tuple[Any, str, str] | None = None
            if box_delivery_key is not None and param_key == box_delivery_key:
                # Box-calculator-allocated delivery trucks (2026-07-30) - computed from this
                # project's real rebar weight, not looked up in found_by_target/missing at all.
                # Crane-shift fields are NOT special-cased here anymore (Elena's ruling: no real
                # crane-shift-count formula exists, they render through the normal found/missing
                # path below like any other manual field). See
                # compute_rebar_metal_delivery_allocation/REBAR_METAL_DELIVERY_FIELD_BY_SECTION.
                box_row = (
                    rebar_metal_delivery_allocation.get(sec_code, 0),
                    "",
                    (
                        f"Вес арматуры раздела: {rebar_weights_by_section.get(sec_code, 0):g} кг. "
                        f"Автораспределение по накоплению {int(METAL_TRUCK_CAPACITY_KG // 1000)} т — "
                        "проверьте и поправьте при необходимости."
                    ),
                )
            elif param_key == BOX_TOTAL_METAL_WEIGHT_KEY:
                box_row = (
                    box_total_metal_weight_kg,
                    "",
                    "Сумма веса арматуры по всем 4 разделам с арматурой (включая стены/перемычки, "
                    "у которых нет отдельной строки доставки) - используется калькулятором фундаментной "
                    "плиты только для контрольного предупреждения.",
                )

            if box_row is not None:
                found_value, source, fragment = box_row
                status = "Найдено (авто, box-калькулятор)"
                row_fill = project_status_fill(status)
                counts["found"] += 1
                ws.append([
                    param.get("label_ru", param.get("key", "")),
                    found_value,
                    param.get("unit", ""),
                    status,
                    "",
                    review_behavior.get("action_ru", "Проверьте значение."),
                    source,
                    fragment,
                    "",
                    "",
                    sec_code,
                    param.get("key", ""),
                    param.get("source_class", ""),
                    target_code,
                ])
                for cell in ws[ws.max_row]:
                    cell.fill = row_fill
                ws.row_dimensions[ws.max_row].height = COMPACT_ROW_HEIGHT
                continue

            confidence = display_confidence(found)
            alternative_scalar = earthworks_alternative_scalar(target_code, found_groups) if sec_code == "earthworks" else None
            if alternative_scalar is None and sec_code == "foundation_slab":
                alternative_scalar = foundation_alternative_scalar(target_code, found, found_groups, found_by_target)
            if alternative_scalar is None and sec_code == "load_bearing_walls_lintels":
                alternative_scalar = load_bearing_walls_lintels_alternative_scalar(target_code, found, found_groups)
            if alternative_scalar is None and sec_code == "floor_slab_1":
                alternative_scalar = floor_slab_1_alternative_scalar(target_code, found, found_groups, found_by_target)
            if alternative_scalar is None and sec_code == "floor_slab_2":
                alternative_scalar = floor_slab_2_alternative_scalar(target_code, found_groups)
            if alternative_scalar is None and sec_code == "flat_roof":
                alternative_scalar = flat_roof_alternative_scalar(target_code, found, found_groups)
            is_unresolved_needs_review = found is not None and found.get("value") is None
            if alternative_scalar is not None and (found is None or target_code in missing or is_unresolved_needs_review):
                found_value, status, source, fragment, confidence = alternative_scalar
                row_fill = project_status_fill(status)
                counts["needs_review" if status.startswith("Проверьте") else "found"] += 1
            elif target_code in confirmed_required and (is_unresolved_needs_review or (target_code in missing and found is None)):
                # A sibling field in the same presence pair was found — this project definitely
                # has this construction, so a still-blank value here (whether never attempted, in
                # section.missing, or a needs_review entry that only has candidates — e.g. the real
                # ARK ПБ1/ПБ2 case where extraction correctly refused to sum 0.21+0.16 itself) is a
                # real, money-relevant gap, not a safely-skippable optional field. Escalate instead
                # of following the normal found/missing branches below. See SECTION_PRESENCE_PAIRS.
                found_value = None
                status = "Не найдено — ОБЯЗАТЕЛЬНО (раздел точно есть в проекте)"
                row_fill = FILL_MISSING
                source = found.get("source_pdf") if found is not None else ""
                fragment = item_fragment(found)
                counts["missing_confirmed_required"] += 1
            elif found is not None:
                found_value = display_value(found.get("value"))
                needs_review = bool(found.get("needs_review"))
                status = "Проверьте (needs_review)" if needs_review else "Найдено"
                row_fill = project_status_fill(status)
                source = found.get("source_pdf") or ""
                if found.get("candidates") and found.get("value") is None:
                    fragment = unresolved_candidate_fragment(found)
                else:
                    fragment = item_fragment(found)
                if (
                    sec_code == "load_bearing_walls_lintels"
                    and target_code == "floors_count"
                    and group_items(found_groups, "wall_block_items")
                ):
                    # Value stays visible (real cross-check value - a mismatch with wall_block_items'
                    # own floor_2 presence would be worth catching), but the calculator itself never
                    # reads this field once wall_block_items is present - see
                    # load_bearing_walls_lintels_alternative_scalar's floors_count branch, which only
                    # fires when this field is missing, not when it's found like here.
                    fragment = (fragment + "\n" if fragment else "") + (
                        "Справочно: калькулятор определяет этажность по наличию объёма 2-го этажа "
                        "в wall_block_items, а не по этому текстовому значению из PDF."
                    )
                counts["needs_review" if needs_review else "found"] += 1
            elif target_code and target_code in missing:
                found_value = None
                source = ""
                classified_missing = classified_missing_target(sec_code, target_code)
                if classified_missing is not None:
                    status, fragment = classified_missing
                    counts["conditional_absent"] += 1
                    row_fill = FILL_WHITE
                else:
                    status = "Не найдено"
                    fragment = ""
                    row_fill = project_status_fill(status)
                    counts["missing"] += 1
            else:
                found_value = None
                status = "Проверьте"
                row_fill = project_status_fill(status)
                source = ""
                fragment = ""

            ws.append([
                param.get("label_ru", param.get("key", "")),
                found_value,
                param.get("unit", ""),
                status,
                confidence,
                review_behavior.get("action_ru", "Проверьте значение."),
                source,
                fragment,
                "",
                "",
                sec_code,
                param.get("key", ""),
                param.get("source_class", ""),
                target_code,
            ])
            for cell in ws[ws.max_row]:
                cell.fill = row_fill
            ws.row_dimensions[ws.max_row].height = COMPACT_ROW_HEIGHT

        for param in production_repeated_row_params(contract) + diagnostic_repeated_row_params(contract):
            review_behavior = param.get("review_behavior") or {}
            action_ru = review_behavior.get("action_ru", "Проверьте позиции построчно.")
            group_key = param.get("key")
            title = f"{section_name(contract)} — {param.get('label_ru', group_key)} ({action_ru})"
            headers = ITEM_BLOCK_HEADERS + correction_headers(param) + ["row_data_json"]

            item_label_columns = param.get("item_label_columns") or []
            correction_columns = param.get("correction_columns") or []
            columns_by_key = {c["key"]: c for c in (param.get("columns") or [])}

            rows: list[list[Any]] = []
            for item in found_groups.get(group_key, []):
                value = item.get("value") or {}
                needs_review = bool(item.get("needs_review"))

                label_parts = [str(value[k]) for k in item_label_columns if value.get(k) not in (None, "")]
                label = " ".join(label_parts) or item.get("item_name") or group_key

                summary_parts = []
                for key in correction_columns:
                    v = value.get(key)
                    if v is None:
                        continue
                    unit = columns_by_key.get(key, {}).get("unit", "")
                    summary_parts.append(f"{v}{' ' + unit if unit else ''}")
                summary = ", ".join(summary_parts)

                status = "Проверьте (needs_review)" if needs_review else "Найдено"
                row_fill = project_status_fill(status)
                counts["needs_review" if needs_review else "found"] += 1
                counts["item_rows"] += 1

                row = [
                    label,
                    summary,
                    param.get("unit", ""),
                    status,
                    display_confidence(item),
                    "",  # action_ru already stated once in the block title above, not per row -
                    # repeating a ~100-char sentence on every item row was the main cause of
                    # tall wrapped rows (2026-07-29 design fix)
                    item.get("source_pdf") or "",
                    item_fragment(item),
                    "",
                    "",
                    sec_code,
                    group_key,
                    param.get("source_class", ""),
                    param.get("target_code", ""),
                ]
                row += ["", "", "", ""]  # correction columns - Elena fills these, not the extraction
                row.append(json.dumps(value, ensure_ascii=False))
                rows.append(row)

            append_block(ws, title, headers, rows)
            if rows:
                for row_idx in range(ws.max_row - len(rows) + 1, ws.max_row + 1):
                    status = str(ws.cell(row_idx, 4).value or "")
                    row_fill = project_status_fill(status)
                    for cell in ws[row_idx]:
                        cell.fill = row_fill
                    ws.row_dimensions[row_idx].height = COMPACT_ROW_HEIGHT
                total_row = earthworks_group_total_row(sec_code, group_key, found_groups, len(headers))
                if total_row:
                    ws.append(total_row)
                    for cell in ws[ws.max_row]:
                        cell.fill = FILL_HEADER
                        cell.font = Font(name=FONT_NAME, bold=True, size=10)
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    ws.row_dimensions[ws.max_row].height = COMPACT_ROW_HEIGHT

    # Итог по коробке (2026-07-30): one cross-section summary row after all 8 sections, so
    # Elena can see the real total driving the box-calculator's delivery-truck allocation above
    # (crane shifts are manual and not related to this total - see REBAR_METAL_DELIVERY_FIELD_BY_SECTION).
    # See compute_rebar_weights_by_section/compute_rebar_metal_delivery_allocation.
    section_names_by_code = {section_code(c): section_name(c) for c in contracts}
    breakdown = "; ".join(
        f"{section_names_by_code.get(code, code)}: {weight:g} кг"
        for code, weight in rebar_weights_by_section.items()
        if weight > 0
    )
    append_section_band(ws, ["Итог по коробке"], len(PROJECT_HEADERS))
    ws.append([
        "Общий вес арматуры по проекту (все разделы с арматурой), кг",
        box_total_metal_weight_kg,
        "кг",
        "Найдено (авто, box-калькулятор)",
        "",
        "Справочно — сумма веса, из которой считается автораспределение машин доставки арматуры/металла по разделам.",
        "",
        breakdown,
        "",
        "",
        "",
        "total_rebar_weight_kg",
        "AUTO_CALCULATED",
        "",
    ])
    for cell in ws[ws.max_row]:
        cell.fill = project_status_fill("Найдено (авто, box-калькулятор)")
    ws.row_dimensions[ws.max_row].height = COMPACT_ROW_HEIGHT

    ws.cell(1, 1).font = Font(name=FONT_NAME, bold=True, size=13)
    ws.cell(1, 1).fill = FILL_HEADER
    apply_table_style(ws, header_row=4)
    restyle_section_bands(ws)
    restyle_block_sheet(ws)
    # Merged last, using the sheet's true final ws.max_column (grows as item blocks with extra
    # correction columns get added above) - merging earlier at a narrower width risks the same
    # overlapping-merge corruption already fixed once for section bands/block titles.
    merge_row_full_width(ws, 1, ws.max_column)
    ws.freeze_panes = "A5"
    set_widths = {
        "A": 30, "B": 26, "C": 10, "D": 20, "E": 20, "F": 62, "G": 30, "H": 64,
        "I": 28, "J": 28, "K": 20, "L": 28, "M": 18, "N": 24,
        "O": 24, "P": 24, "Q": 24, "R": 24,
    }
    for col, width in set_widths.items():
        ws.column_dimensions[col].width = width
    for column in ["K", "L", "M", "N", "S"]:
        ws.column_dimensions[column].hidden = True

    # Column widths are final now - autofit reads them to grow any row whose wrapped text needs
    # more than COMPACT_ROW_HEIGHT's one line (see estimate_row_height docstring).
    autofit_row_heights(ws, 1, ws.max_row)

    return counts


def build_rebar_lookup(
    contracts: list[dict[str, Any]], extraction: dict[str, Any]
) -> dict[str, list[dict[str, Any]]]:
    """section_code -> real rebar rows (steel_class/diameter_mm) found by extraction, pooled
    across every rebar-shaped group in that section (see rebar_group_keys_for_contract). Used to
    expand sheet 02's rebar_<class>_d<diameter>_m templated price row into one row per
    diameter/class actually present in this project."""
    lookup: dict[str, list[dict[str, Any]]] = {}
    for contract in contracts:
        sec_code = section_code(contract)
        group_keys = rebar_group_keys_for_contract(contract)
        if not group_keys:
            continue
        _, found_groups, _ = index_extraction_section(extraction, sec_code)
        items = [
            item.get("value") or {}
            for group_key in group_keys
            for item in found_groups.get(group_key, [])
        ]
        if items:
            lookup[sec_code] = items
    return lookup


# Crane-shift counts are NOT automated (Elena's 2026-07-30 ruling: no real crane-shift-count
# formula exists anywhere in the codebase - see rebar_crane_manual_and_box_delivery_final memory).
# Only rebar/metal DELIVERY TRUCKS are automated here, which is what metal_delivery_allocator.py
# was actually designed and named for.
#
# section_code -> the one review_parameters/supplier_inputs key that holds "delivery trucks for
# rebar/metal" in that section. load_bearing_walls_lintels is deliberately absent because this
# section does not have its own dedicated billed metal-delivery line in the current estimate
# structure; its rebar weight is still counted in METAL_SECTION_ORDER's box-wide total below since
# it still needs to physically arrive on site.
REBAR_METAL_DELIVERY_FIELD_BY_SECTION = {
    "foundation_slab": "rebar_metal_delivery_trucks",
    "floor_slab_1": "rebar_metal_delivery_trucks",
    "floor_slab_2": "rebar_metal_delivery_trucks",
}
# All 4 rebar-bearing sections (experiments/box_calculator/section_registry.py convention) - used
# for the weight total and the box-wide truck allocation, even though load_bearing_walls_lintels
# has no field of its own to render an allocated count into.
METAL_SECTION_ORDER = ["foundation_slab", "load_bearing_walls_lintels", "floor_slab_1", "floor_slab_2"]
METAL_TRUCK_CAPACITY_KG = 10000.0
# foundation_slab's own field: the box-wide total weight, used only for its internal calculator
# warning (see foundation_slab_calculator.py's suggested_box_metal_delivery_trucks check) - same
# number as the "Итог по коробке" summary row below, not a separate calculation.
BOX_TOTAL_METAL_WEIGHT_KEY = "box_total_metal_weight_kg"


def compute_rebar_weights_by_section(
    contracts: list[dict[str, Any]], extraction: dict[str, Any]
) -> dict[str, float]:
    """section_code -> total rebar weight in kg for that section, length x rate summed across
    every rebar item found (see rebar_item_weight_kg - rate is the item's own kg_per_meter if
    given, else the fixed GOST catalog by diameter). Only the 4 rebar-bearing sections can appear;
    a section with items but zero computable weight (no length/diameter data at all) still gets an
    entry of 0.0, not omitted, so downstream code doesn't have to guess whether "missing" means
    "no rebar" or "rebar present but unweighable"."""
    lookup = build_rebar_lookup(contracts, extraction)
    weights: dict[str, float] = {}
    for sec_code in METAL_SECTION_ORDER:
        total = 0.0
        for item in lookup.get(sec_code, []):
            weight = rebar_item_weight_kg(item)
            if weight is not None:
                total += weight
        weights[sec_code] = total
    return weights


def compute_rebar_metal_delivery_allocation(weights_by_section: dict[str, float]) -> dict[str, int]:
    """section_code -> automatically allocated delivery trucks for rebar/metal, using the
    existing box_calculator threshold-by-10-tonnes logic (experiments/box_calculator/
    metal_delivery_allocator.py) - the first truck goes to the first section with any weight,
    later trucks go to whichever section's cumulative weight crosses the next 10-tonne boundary.
    unit_price is 0 here - this call only needs allocated_trucks, not a cost (pricing for these
    lines already comes from the normal price_keys mechanism on sheet 02)."""
    sections = [
        MetalSection(section_code=code, section_name=code, metal_weight_kg=weights_by_section.get(code, 0.0))
        for code in METAL_SECTION_ORDER
    ]
    result = allocate_metal_deliveries(sections=sections, capacity_kg=METAL_TRUCK_CAPACITY_KG, unit_price=0.0)
    return {item["section_code"]: item["allocated_trucks"] for item in result["sections"]}


def build_workbook_from_extraction(
    contract_paths: list[Path],
    extraction_path: Path,
    output_path: Path,
    price_registry_path: Path | None = DEFAULT_PRICE_REGISTRY,
    manual_values_registry_path: Path | None = DEFAULT_MANUAL_VALUES_REGISTRY,
) -> dict[str, Any]:
    contracts = [load_yaml_contract(path) for path in contract_paths]
    extraction = json.loads(extraction_path.read_text(encoding="utf-8"))

    price_registry = None
    if price_registry_path is not None and price_registry_path.exists():
        price_registry = load_price_registry(price_registry_path)
    manual_values_registry = None
    if manual_values_registry_path is not None and manual_values_registry_path.exists():
        manual_values_registry = load_manual_values_registry(manual_values_registry_path)
    rebar_lookup = build_rebar_lookup(contracts, extraction)

    wb = Workbook()
    build_constructor_sheet(wb, contracts)
    counts = build_project_sheet_from_extraction(wb, contracts, extraction)
    build_manual_values_sheet(wb, contracts, manual_values_registry=manual_values_registry)
    build_prices_sheet(wb, contracts, price_registry=price_registry, rebar_lookup=rebar_lookup)
    detail_counts = build_details_sheet_from_extraction(wb, extraction)
    build_instruction_sheet(wb)
    build_contracts_summary_sheet(wb, contracts)
    build_raw_contracts_sheet(wb, contracts)
    hide_diagnostic_sheets(wb)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return {"output_path": str(output_path), **counts, **detail_counts}


def next_versioned_path(path: Path) -> Path:
    """Appends/bumps a _vN suffix so repeated builds never overwrite the previous one (a real
    review session rebuilds this file many times while iterating) - review_workbook.xlsx ->
    ..._v1.xlsx, then _v2.xlsx, etc., based on the highest _vN already present in the target
    directory. Re-versions cleanly if the given path already ends in _vN."""
    match = re.match(r"^(.*)_v(\d+)$", path.stem)
    base_stem = match.group(1) if match else path.stem
    max_version = 0
    if path.parent.exists():
        for existing in path.parent.glob(f"{base_stem}_v*{path.suffix}"):
            existing_match = re.match(rf"^{re.escape(base_stem)}_v(\d+)$", existing.stem)
            if existing_match:
                max_version = max(max_version, int(existing_match.group(1)))
    return path.parent / f"{base_stem}_v{max_version + 1}{path.suffix}"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("extraction_json", help="Path to a real extraction_output.json")
    parser.add_argument(
        "--output",
        default=str(PIPELINE_DIR / "output" / "step_20_populated_from_real_extraction.xlsx"),
    )
    parser.add_argument(
        "--price-registry",
        default=str(DEFAULT_PRICE_REGISTRY),
        help="Path to a price_registry_filled_v*.xlsx; pass an empty string to skip price fill.",
    )
    parser.add_argument(
        "--manual-values-registry",
        default=str(DEFAULT_MANUAL_VALUES_REGISTRY),
        help="Path to manual_values_registry.xlsx; pass an empty string to skip sheet 01-1 fill.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    price_registry_path = Path(args.price_registry) if args.price_registry else None
    manual_values_registry_path = Path(args.manual_values_registry) if args.manual_values_registry else None
    output_path = next_versioned_path(Path(args.output))
    result = build_workbook_from_extraction(
        default_contract_paths(),
        Path(args.extraction_json),
        output_path,
        price_registry_path,
        manual_values_registry_path,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
