"""Downloads manual_values_registry.xlsx from its Google Sheets master copy.

Elena now edits typical manual-input values (crane shifts, delivery trips, pump
shifts, etc.) directly in a Google Sheet instead of a locally-hand-edited xlsx.
This script fetches the current sheet and writes it to the same local path every
other script in the pipeline already reads (output/manual_values_registry.xlsx),
so nothing downstream needs to change - only this file's origin does.

Run this before build_review_workbook_from_contracts.py /
populate_review_workbook_from_extraction.py whenever you need the latest typical
values. Local edits to output/manual_values_registry.xlsx are NOT preserved -
they get silently overwritten on the next download. Edit the Google Sheet, not
the local file.
"""
from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from openpyxl import load_workbook

PIPELINE_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT_PATH = PIPELINE_DIR.parents[1] / "output" / "manual_values_registry.xlsx"

# From https://docs.google.com/spreadsheets/d/16d4Nm_1EeQ1ssB4t9nsJVVpU4et63wFa/edit?usp=sharing...
DEFAULT_SHEET_ID = "16d4Nm_1EeQ1ssB4t9nsJVVpU4et63wFa"
EXPORT_URL_TEMPLATE = "https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"

EXPECTED_HEADERS = ["Раздел", "Наименование", "Ед. изм.", "Типовое значение", "Комментарий", "section_code", "key"]


class DownloadError(RuntimeError):
    pass


def download_bytes(sheet_id: str) -> bytes:
    # Uses the system `curl` (not urllib) - this venv's Python has no working local CA
    # bundle for SSL verification, while macOS's own curl uses the system keychain and
    # just works. Args are passed as a list (no shell=True), so sheet_id can't inject.
    url = EXPORT_URL_TEMPLATE.format(sheet_id=sheet_id)
    try:
        result = subprocess.run(
            ["curl", "-sL", "--fail", "-A", "Mozilla/5.0", url],
            capture_output=True,
            timeout=30,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise DownloadError(f"Не удалось запустить curl для скачивания таблицы: {exc}") from exc

    if result.returncode != 0:
        raise DownloadError(
            f"curl завершился с ошибкой (код {result.returncode}) при скачивании sheet_id={sheet_id}. "
            "Проверьте, что ссылка на таблицу открыта на просмотр всем, у кого есть ссылка, и что есть "
            f"интернет-соединение. stderr: {result.stderr.decode(errors='replace')[:500]}"
        )

    data = result.stdout
    if data[:64].lstrip().lower().startswith(b"<!doctype html") or data[:64].lstrip().lower().startswith(b"<html"):
        raise DownloadError(
            "Google вернул HTML вместо xlsx - обычно значит, что таблица недоступна по ссылке "
            "(доступ ограничен) или sheet_id неверный. Проверьте настройки доступа таблицы."
        )
    return data


def validate_workbook(path: Path) -> int:
    """Returns the number of real data rows. Raises DownloadError on structural mismatch."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DownloadError(f"Скачанный файл не открывается как xlsx: {exc}") from exc

    if "manual_values" not in wb.sheetnames:
        raise DownloadError(
            f"В скачанной таблице нет листа 'manual_values' (есть: {wb.sheetnames}). "
            "Проверьте, что ссылка ведёт на правильную таблицу."
        )
    ws = wb["manual_values"]
    header = [cell.value for cell in ws[1]]
    if header != EXPECTED_HEADERS:
        raise DownloadError(
            f"Заголовок листа не совпадает с ожидаемым.\n  Ожидалось: {EXPECTED_HEADERS}\n  Получено: {header}"
        )

    row_count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        section_code, key = row[5], row[6]
        if section_code and key:
            row_count += 1
    return row_count


def download_registry(sheet_id: str, output_path: Path) -> int:
    data = download_bytes(sheet_id)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # openpyxl rejects files by extension, not just content - the temp file must keep
    # the .xlsx suffix (e.g. ".manual_values_registry.download.xlsx", not "...xlsx.tmp").
    tmp_path = output_path.with_name(f".{output_path.stem}.download.xlsx")
    tmp_path.write_bytes(data)
    try:
        row_count = validate_workbook(tmp_path)
    except DownloadError:
        tmp_path.unlink(missing_ok=True)
        raise
    tmp_path.replace(output_path)
    return row_count


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sheet-id", default=DEFAULT_SHEET_ID, help="Google Sheets file ID (from the sharing URL).")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH, help="Where to write the xlsx.")
    args = parser.parse_args()

    row_count = download_registry(args.sheet_id, args.output)
    print(f"manual_values_registry скачан -> {args.output} ({row_count} строк с данными)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
