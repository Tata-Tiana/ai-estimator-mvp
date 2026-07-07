from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
REPO_ROOT = ROOT.parents[1]
DEFAULT_TARGETS_PATH = ROOT / "data" / "calculator_targets_compact.json"
DEFAULT_OUTPUT_PATH = ROOT / "outputs" / "review_workbook_preview.xlsx"
DEFAULT_ELENA_AUDIT_PATH = REPO_ROOT / "docs" / "elena_parameter_review_pack_audit.md"

SECTION_NAMES_RU = {
    "earthworks": "Земляные работы",
    "foundation_slab": "Фундаментная плита",
    "waterproofing": "Гидроизоляция",
    "load_bearing_walls_lintels": "Несущие стены и перемычки",
    "floor_slab_1": "Плита перекрытия 1",
    "floor_slab_2": "Плита перекрытия 2",
    "flat_roof": "Плоская кровля",
    "schiedel_vent_channels": "Вентканалы Schiedel",
}

AUDIT_SECTION_CODES = {
    "Земляные работы": "earthworks",
    "Фундаментная плита": "foundation_slab",
    "Гидроизоляция": "waterproofing",
    "Несущие стены и перемычки": "load_bearing_walls_lintels",
    "Плита перекрытия 1-го этажа": "floor_slab_1",
    "Плита перекрытия 2-го этажа": "floor_slab_2",
    "Плоская кровля": "flat_roof",
    "Вентиляционные каналы Schiedel": "schiedel_vent_channels",
}

AUTO_PROJECT_ALIASES = {
    ("earthworks", "pit_area_m2"): ["pit_area"],
    ("earthworks", "trench_volume_m3"): ["trench_volume_total", "trench_routes"],
    ("earthworks", "communications_length_m"): ["communications_pipe_items"],
    ("earthworks", "geotextile_laying_area_m2"): ["geotextile_laying_area"],
    ("foundation_slab", "thermal_insert_50_length_m"): ["thermal_insert_50_length"],
    ("foundation_slab", "thermal_insert_100_length_m"): ["thermal_insert_100_length"],
    ("foundation_slab", "thermal_insert_50_material_spec_qty"): ["thermal_insert_50_material_spec_qty"],
    ("foundation_slab", "thermal_insert_100_material_spec_qty"): ["thermal_insert_100_material_spec_qty"],
    ("load_bearing_walls_lintels", "main_wall_reinforcement_rows"): ["main_wall_reinforcement_rows"],
    ("load_bearing_walls_lintels", "main_wall_400_reinforcement_threads"): ["main_wall_400_reinforcement_threads"],
    ("load_bearing_walls_lintels", "lintel_total_length_m"): ["lintel_items"],
    ("load_bearing_walls_lintels", "lintel_concrete_spec_volume_m3"): ["lintel_concrete_volume"],
    ("load_bearing_walls_lintels", "vent_chimney_gas_block_spec_volume_m3"): ["gas_block_d500_150_volume", "vent_chimney_gas_block_spec_volume"],
    ("load_bearing_walls_lintels", "floor_2_masonry_volume_m3"): ["gas_block_d500_250_volume"],
    ("load_bearing_walls_lintels", "parapet_masonry_volume_m3"): ["parapet_masonry_volume"],
    ("load_bearing_walls_lintels", "cutoff_waterproofing_load_bearing_walls_area_m2"): ["waterproofing_cutoff_material_area"],
    ("load_bearing_walls_lintels", "cutoff_waterproofing_partitions_area_m2"): ["waterproofing_cutoff_material_area"],
    ("floor_slab_1", "insulation.total_eps_volume_from_spec_m3"): ["floor_slab_1_eps100_volume"],
    ("floor_slab_1", "main_formwork_area_m2"): ["floor_slab_1_main_formwork_area"],
    ("floor_slab_1", "beams_formwork_area_m2"): ["beam_items"],
    ("floor_slab_1", "beams_concrete_volume_m3"): ["floor_slab_1_beam_concrete_volume", "beam_items"],
    ("floor_slab_1", "slab_concrete_volume_m3"): ["floor_slab_1_concrete_volume"],
    ("floor_slab_2", "slab_edge_perimeter_m"): ["floor_slab_2_slab_edge_perimeter"],
    ("floor_slab_2", "main_formwork_area_m2"): ["floor_slab_2_main_formwork_area"],
    ("floor_slab_2", "edge_formwork_area_m2"): ["floor_slab_2_edge_formwork_area_outside_targets"],
    ("floor_slab_2", "slab_concrete_volume_m3"): ["floor_slab_2_concrete_volume"],
    ("flat_roof", "roof_area_level_1_m2"): ["roof_area_level_1"],
    ("flat_roof", "roof_area_level_2_m2"): ["roof_area_level_2"],
    ("flat_roof", "parapet_length_level_1_m"): ["roof_parapet_abutment_total_length"],
    ("flat_roof", "parapet_length_level_2_m"): ["roof_parapet_abutment_total_length"],
    ("flat_roof", "vent_wall_abutment_level_1_m"): ["roof_wall_abutment_total_length"],
    ("flat_roof", "vent_wall_abutment_level_2_m"): ["roof_wall_abutment_total_length"],
    ("flat_roof", "internal_drains_count"): ["roof_internal_drains_count"],
    ("flat_roof", "internal_roof_drains_count"): ["roof_internal_drains_count"],
    ("flat_roof", "parapet_roof_drains_count"): ["roof_parapet_drains_count"],
    ("schiedel_vent_channels", "vent_channel_2_count"): ["vent_channel_2_count", "schiedel_vent_channel_2x_count"],
    ("schiedel_vent_channels", "schiedel_vent_2_count"): ["schiedel_vent_channel_2x_count"],
    ("schiedel_vent_channels", "schiedel_vent_3_count"): ["schiedel_vent_channel_3x_count"],
    ("schiedel_vent_channels", "vent_chimney_gas_block_spec_volume_m3"): ["schiedel_gas_block_150_volume"],
}

SHEET_00 = "00_Конструктор сметы"
SHEET_01 = "01_Проверка проекта"
SHEET_02 = "02_Цены себестоимости"
SHEET_03 = "03_Детали объемов"
SHEET_04 = "04_Инструкция"
SHEET_05 = "05_Кандидаты parser"
SHEET_06 = "06_Сырые данные parser"

VISIBLE_REVIEW_HEADERS = [
    "Что проверяем",
    "Найдено в проекте",
    "Ед.",
    "Статус",
    "Что нужно сделать",
    "Источник",
    "Фрагмент проекта",
    "Исправить / ввести значение",
    "Комментарий Елены",
]

TECH_REVIEW_HEADERS = [
    "section_code",
    "target_code",
    "group_code",
    "item_name",
    "row_kind",
    "value_kind",
    "calculator_input_key",
    "extraction_status",
    "confidence",
    "source_json_path",
    "adapter_status",
]

PRICE_HEADERS = [
    "Строка сметы",
    "Что это за цена",
    "Ед.",
    "Цена из прайса",
    "Цена fallback",
    "Цена для расчета",
    "Исправить цену",
    "Источник цены",
    "Нужно внимание",
    "Комментарий",
    "calc_price_key",
    "price_registry_code",
    "fallback_key",
    "selected_price_source",
]

HELPER_TARGET_CODES = {
    "waterproofing_cutoff_material_area",
    "floor_slab_1_beam_concrete_volume",
    "floor_slab_2_edge_formwork_area_outside_targets",
    "roof_parapet_abutment_total_length",
    "roof_wall_abutment_total_length",
    "vent_chimney_cladding_length_total",
    "project_spec_roof_area",
}

DETAIL_GROUP_CODES = {
    "trench_routes",
    "communications_pipe_items",
    "foundation_rebar_items",
    "floor_slab_1_rebar_items",
    "floor_slab_2_rebar_items",
    "lintel_rebar_items",
    "beam_items",
    "lintel_items",
}

REBAR_GROUP_CODES = {
    "foundation_rebar_items",
    "floor_slab_1_rebar_items",
    "floor_slab_2_rebar_items",
    "lintel_rebar_items",
}

FONT_NAME = "Arial"
FILL_HEADER = PatternFill("solid", fgColor="D9D9D9")
FILL_BLOCK = PatternFill("solid", fgColor="E7E6E6")
FILL_REVIEW = PatternFill("solid", fgColor="FCE4D6")
FILL_MISSING = PatternFill("solid", fgColor="F4CCCC")
FILL_HELPER = PatternFill("solid", fgColor="FFF2CC")
FILL_DETAIL = PatternFill("solid", fgColor="DAE8FC")
FILL_OK = PatternFill("solid", fgColor="D9EAD3")
FILL_FOUND = FILL_OK
FILL_GRAY = PatternFill("solid", fgColor="E7E6E6")
BORDER_THIN = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)

DETAIL_HEADERS = [
    "Тип",
    "Наименование",
    "Длина, м",
    "Глубина, м",
    "Ширина, м",
    "Объем, м3",
    "Диаметр, мм",
    "Длина одной, м",
    "Количество",
    "Итоговая длина, м",
    "Включено",
    "Источник",
    "Фрагмент проекта",
    "Комментарий Елены",
]


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def dumps_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def display_number(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).replace(".", ",")
    return str(value).replace(".", ",")


def excel_value(value: Any) -> Any:
    """Keep numeric workbook cells numeric so Excel does not show text-number warnings."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        stripped = value.strip()
        normalized = stripped.replace(" ", "").replace(",", ".")
        if re.fullmatch(r"-?\d+(?:\.\d+)?", normalized):
            number = float(normalized)
            return int(number) if number.is_integer() else number
    return value


def source_label(item: dict[str, Any]) -> str:
    source_pdf = item.get("source_pdf") or ""
    page = item.get("page_number") or ""
    title = item.get("page_title") or item.get("table_title") or ""
    parts = []
    if title:
        parts.append(str(title))
    if source_pdf:
        parts.append(f"источник: {source_pdf}")
    if page:
        parts.append(f"стр. {page}")
    return " / ".join(parts)


def value_kind(value: Any) -> str:
    if isinstance(value, dict):
        return "object"
    if isinstance(value, list):
        return "list"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    if value is None:
        return "empty"
    return "string"


def status_for_item(item: dict[str, Any], *, helper: bool = False) -> str:
    if helper:
        return "HELPER_ONLY"
    if item.get("needs_review"):
        return "NEEDS_REVIEW"
    return "FOUND"


def action_for_item(item: dict[str, Any], *, row_kind: str, helper: bool = False) -> str:
    target_code = item.get("target_code") or ""
    if helper:
        return "Показано для проверки; не передавать в calculator input без явного adapter mapping."
    if row_kind == "detail_ref":
        return "Проверить детали на листе 03_Детали объемов."
    if target_code in {"thermal_insert_50_length", "thermal_insert_100_length"} and item.get("needs_review"):
        return "В PDF одна общая длина термовставок без разбивки 50/100 мм."
    if item.get("needs_review"):
        return "Проверьте значение; при необходимости заполните исправление или комментарий."
    return "Проверьте при необходимости; если всё верно, ничего не меняйте."


def human_status_for_item(item: dict[str, Any], *, helper: bool = False) -> str:
    if helper:
        return "🟧 Проверьте"
    if item.get("needs_review"):
        return "🟧 Проверьте"
    return "✅ Найдено"


def human_status_for_missing(target_code: str) -> str:
    if any(token in target_code for token in ("height", "length")):
        return "🟥 Требуется ручной ввод"
    return "🟥 Не найдено"


def extraction_status_for_human_status(status: str) -> str:
    if "Не найдено" in status or "ручной" in status:
        return "missing"
    if "Проверьте" in status:
        return "needs_review"
    return "found"


def compact_target_index(path: Path) -> tuple[dict[str, dict[str, dict[str, Any]]], set[str]]:
    data = load_json(path)
    by_section: dict[str, dict[str, dict[str, Any]]] = {}
    all_codes: set[str] = set()
    for section in data.get("sections", []):
        section_code = section.get("section_code")
        by_section.setdefault(section_code, {})
        for target in section.get("targets", []):
            code = target.get("code")
            if not code:
                continue
            by_section[section_code][code] = target
            all_codes.add(code)
    return by_section, all_codes


def parse_auto_project_targets(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    text = path.read_text(encoding="utf-8")
    match = re.search(r"## AUTO_PROJECT high risk\n(.*?)(?:\n## |\Z)", text, re.S)
    if not match:
        return []
    rows = []
    pattern = re.compile(r"^- (.*?) :: `([^`]+)` — (.*?)(?: \((.*)\))?$", re.M)
    for order, (section_name, key, label, note) in enumerate(pattern.findall(match.group(1)), start=1):
        rows.append(
            {
                "order": order,
                "section_name": section_name,
                "section_code": AUDIT_SECTION_CODES.get(section_name, section_name),
                "calculator_input_key": key,
                "label": label,
                "unit": infer_unit_from_key(key),
                "note": note or "",
            }
        )
    return rows


def infer_unit_from_key(key: str) -> str:
    tail = key.split(".")[-1]
    tail = re.sub(r"\[\d+\]", "", tail)
    if tail.endswith("_m2"):
        return "м2"
    if tail.endswith("_m3"):
        return "м3"
    if tail.endswith("_mm"):
        return "мм"
    if tail.endswith("_kg"):
        return "кг"
    if tail.endswith("_m") or tail in {"length_m", "height_m", "width_m", "depth_m"}:
        return "м"
    if tail.endswith("_count") or tail.endswith("_threads") or tail.endswith("_rows"):
        return "шт"
    if tail == "diameter_mm":
        return "мм"
    return ""


def code_stem(code: str) -> str:
    stem = code.replace(".", "_")
    stem = re.sub(r"\[\d+\]", "", stem)
    return re.sub(r"_(m2|m3|m|mm|kg|pcs|count|qty)$", "", stem)


def iter_section_items(extraction: dict[str, Any], section_code: str) -> list[dict[str, Any]]:
    section = extraction.get("sections", {}).get(section_code, {})
    return list(section.get("found") or [])


def item_value_for_auto_project(item: dict[str, Any], key: str) -> Any:
    value = item.get("value")
    code = item_code(item)
    if isinstance(value, dict):
        field = key.split(".")[-1]
        field = re.sub(r"\[\d+\]", "", field)
        if field in value:
            return value.get(field)
        if field == "weight_parts_kg":
            return value.get("weight_kg")
        if field in {"source_weight_parts_kg", "source_weight_kg"}:
            return value.get("weight_kg")
        if field == "spec_length_m":
            return value.get("length_m")
        if field == "length_m" and "length_m" in value:
            return value.get("length_m")
        if field == "diameter_mm" and "diameter_mm" in value:
            return value.get("diameter_mm")
        if field == "steel_class" and "steel_class" in value:
            return value.get("steel_class")
        if code == "floor_slab_1_eps100_volume":
            if key.endswith("total_eps_volume_from_spec_m3"):
                bottom = value.get("bottom_volume_m3") or 0
                edge = value.get("edge_volume_m3") or 0
                return round(float(bottom) + float(edge), 3)
            if key.endswith("bottom_slab_eps_work_area_m2"):
                return value.get("bottom_volume_m3")
            if key.endswith("slab_edge_eps_material_area_m2"):
                return value.get("edge_volume_m3")
        return value
    return value


def detail_group_for_key(section_code: str, key: str) -> str:
    if key.startswith("rebar_items["):
        return {
            "foundation_slab": "foundation_rebar_items",
            "floor_slab_1": "floor_slab_1_rebar_items",
            "floor_slab_2": "floor_slab_2_rebar_items",
        }.get(section_code, "rebar_items")
    if key.startswith("lintel_rebar_items["):
        return "lintel_rebar_items"
    if key.startswith("beams.items["):
        return "beam_items"
    return ""


def detail_index_for_key(key: str) -> int | None:
    match = re.search(r"\[(\d+)\]", key)
    if not match:
        return None
    return int(match.group(1))


def match_auto_project_target(target: dict[str, Any], extraction: dict[str, Any]) -> tuple[dict[str, Any] | None, Any, str]:
    section_code = target["section_code"]
    key = target["calculator_input_key"]
    items = iter_section_items(extraction, section_code)

    group = detail_group_for_key(section_code, key)
    if group:
        grouped = [item for item in items if item_code(item) == group]
        idx = detail_index_for_key(key)
        if idx is not None and idx < len(grouped):
            item = grouped[idx]
            return item, item_value_for_auto_project(item, key), f"detail:{group}[{idx}]"
        return None, "", f"detail:{group}"

    aliases = [key, code_stem(key), *AUTO_PROJECT_ALIASES.get((section_code, key), [])]
    aliases = [alias for alias in aliases if alias]
    for item in items:
        code = item_code(item)
        candidates = {code, code_stem(code)}
        if any(alias in candidates for alias in aliases):
            return item, item_value_for_auto_project(item, key), f"code:{code}"

    if section_code == "load_bearing_walls_lintels" and key in {
        "cutoff_waterproofing_load_bearing_walls_area_m2",
        "cutoff_waterproofing_partitions_area_m2",
    }:
        needle = "несущ" if "load_bearing" in key else "перегород"
        for item in iter_section_items(extraction, "waterproofing"):
            name = f"{item.get('item_name') or ''} {item.get('raw_text') or ''}".lower()
            if item_code(item) == "waterproofing_cutoff_material_area" and needle in name:
                return item, item_value_for_auto_project(item, key), "cross_section:waterproofing"

    return None, "", "not_found"


def target_label(
    section_code: str,
    target_code: str | None,
    item_name: str | None,
    target_index: dict[str, dict[str, dict[str, Any]]],
) -> str:
    if item_name:
        return item_name
    if target_code and target_code in target_index.get(section_code, {}):
        return target_index[section_code][target_code].get("label_ru") or target_code
    return target_code or ""


def expected_unit(section_code: str, target_code: str | None, item: dict[str, Any], target_index: dict[str, dict[str, dict[str, Any]]]) -> str:
    return (
        item.get("unit")
        or item.get("normalized_unit")
        or (target_index.get(section_code, {}).get(target_code, {}).get("unit") if target_code else "")
        or ""
    )


def item_code(item: dict[str, Any]) -> str:
    return item.get("group_code") or item.get("target_code") or ""


def is_rebar_item(item: dict[str, Any]) -> bool:
    code = item_code(item)
    value = item.get("value")
    return code in REBAR_GROUP_CODES or (
        code == "masonry_rebar_a500_d10_weight"
        and isinstance(value, dict)
        and value.get("length_m") not in (None, "")
    )


def is_detail_item(item: dict[str, Any]) -> bool:
    code = item_code(item)
    value = item.get("value")
    if code in {"trench_routes", "communications_pipe_items", "beam_items", "lintel_items"}:
        return True
    if is_rebar_item(item):
        return True
    if code == "floor_slab_1_eps100_volume" and isinstance(value, dict):
        return True
    return False


def is_helper_item(item: dict[str, Any], section_code: str, target_codes: set[str]) -> bool:
    code = item_code(item)
    if code in HELPER_TARGET_CODES:
        return True
    if code and code not in target_codes and not is_detail_item(item):
        return True
    return False


def status_fill(status: str) -> PatternFill:
    if status == "MISSING" or status == "MANUAL_NEEDED":
        return FILL_MISSING
    if status == "NEEDS_REVIEW":
        return FILL_REVIEW
    if status == "HELPER_ONLY":
        return FILL_HELPER
    return FILL_OK


def style_sheet(ws, header_row: int = 1, hide_after_col: int | None = None) -> None:
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = Font(name=FONT_NAME, bold=cell.font.bold, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER_THIN
    for cell in ws[header_row]:
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(header_row, col_idx).value or "")
        width = min(max(len(header) + 4, 14), 34)
        if any(token in header for token in ("Фрагмент", "Комментарий", "Что нужно", "raw_text", "cells", "columns")):
            width = 55
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    if hide_after_col:
        for col_idx in range(hide_after_col + 1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].hidden = True


def apply_table_theme(ws, max_row: int | None = None, max_col: int | None = None) -> None:
    max_row = max_row or ws.max_row
    max_col = max_col or ws.max_column
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.font = Font(name=FONT_NAME, bold=cell.font.bold, size=cell.font.sz or 10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER_THIN


def set_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def style_header_row(ws, row_idx: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row_idx, col)
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER_THIN


def style_sheet_basic(ws, header_row: int = 1, widths: dict[str, float] | None = None) -> None:
    style_header_row(ws, header_row, ws.max_column)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    apply_table_theme(ws)
    for idx in range(1, ws.max_column + 1):
        header = str(ws.cell(header_row, idx).value or "")
        width = min(max(len(header) + 4, 14), 42)
        if "Фрагмент" in header or "Комментарий" in header or "Что нужно" in header:
            width = 56
        ws.column_dimensions[get_column_letter(idx)].width = width
    if widths:
        set_widths(ws, widths)


def append_table(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([dumps_cell(row.get(header, "")) for header in headers])
    style_sheet_basic(
        ws,
        widths={
            "A": 28,
            "B": 20,
            "C": 18,
            "D": 32,
            "E": 56,
            "F": 26,
        },
    )


def style_block_title_row(ws, row_idx: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row_idx, col)
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.fill = FILL_GRAY
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER_THIN


def write_block_table(
    ws,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    start_row: int,
) -> int:
    ws.cell(start_row, 1).value = title
    style_block_title_row(ws, start_row, max(len(headers), 1))
    header_row = start_row + 1
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(header_row, col_idx).value = header
    style_header_row(ws, header_row, len(headers))
    row_idx = header_row + 1
    for data_row in rows:
        for col_idx, value in enumerate(data_row, start=1):
            ws.cell(row_idx, col_idx).value = dumps_cell(value)
        row_idx += 1
    for idx in range(start_row + 2, row_idx):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(idx, col)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER_THIN
    return row_idx + 1


def append_block_title(ws, title: str, width: int) -> None:
    ws.append([title] + [""] * (width - 1))
    row_idx = ws.max_row
    for col_idx in range(1, width + 1):
        cell = ws.cell(row_idx, col_idx)
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.fill = FILL_BLOCK
        cell.border = BORDER_THIN


def append_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    row_idx = ws.max_row
    for cell in ws[row_idx]:
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER_THIN


def append_review_rows(
    ws,
    extraction: dict[str, Any],
    target_index: dict[str, dict[str, dict[str, Any]]],
    compact_codes: set[str],
    report: dict[str, Any],
) -> None:
    ws.append(VISIBLE_REVIEW_HEADERS + TECH_REVIEW_HEADERS)
    status_col = VISIBLE_REVIEW_HEADERS.index("Статус") + 1

    seen_found_keys: set[tuple[str, str, str]] = set()
    for section_code, section in extraction.get("sections", {}).items():
        section_name = SECTION_NAMES_RU.get(section_code, section_code)
        for index, item in enumerate(section.get("found") or []):
            code = item_code(item)
            helper = is_helper_item(item, section_code, compact_codes)
            if helper:
                continue
            target_code = item.get("target_code")
            group_code = item.get("group_code")
            value = item.get("value")
            row_kind = "detail_ref" if is_detail_item(item) else "scalar"
            if row_kind == "detail_ref":
                found_value = "см. лист 03_Детали объемов"
            else:
                found_value = display_number(value) if not isinstance(value, (dict, list)) else dumps_cell(value)
            status = status_for_item(item)
            seen_found_keys.add((section_code, code, dumps_cell(value)))
            ws.append(
                [
                    section_name,
                    target_label(section_code, target_code or group_code, item.get("item_name"), target_index),
                    found_value,
                    expected_unit(section_code, target_code or group_code, item, target_index),
                    status,
                    action_for_item(item, row_kind=row_kind),
                    source_label(item),
                    item.get("raw_text") or "",
                    "",
                    "",
                    section_code,
                    target_code or "",
                    group_code or "",
                    item.get("item_name") or "",
                    row_kind,
                    value_kind(value),
                    "",
                    "needs_review" if item.get("needs_review") else "found",
                    item.get("confidence") or "",
                    f"sections.{section_code}.found[{index}]",
                    "preview_only",
                ]
            )
            ws.cell(ws.max_row, status_col).fill = status_fill(status)

        for index, missing in enumerate(section.get("missing") or []):
            if isinstance(missing, dict):
                target_code = missing.get("target_code") or missing.get("code") or ""
                note = missing.get("notes") or ""
            else:
                target_code = str(missing)
                note = ""
            label = target_label(section_code, target_code, None, target_index)
            status = "MANUAL_NEEDED" if any(token in target_code for token in ("height", "length")) else "MISSING"
            ws.append(
                [
                    section_name,
                    label,
                    "",
                    target_index.get(section_code, {}).get(target_code, {}).get("unit", ""),
                    status,
                    "Значение не найдено в extraction JSON; введите вручную или оставьте комментарий.",
                    "",
                    note,
                    "",
                    "",
                    section_code,
                    target_code,
                    "",
                    "",
                    "missing",
                    "empty",
                    "",
                    "missing",
                    "",
                    f"sections.{section_code}.missing[{index}]",
                    "manual_required",
                ]
            )
            ws.cell(ws.max_row, status_col).fill = status_fill(status)


def append_review_sheet_earthworks_like(
    ws,
    extraction: dict[str, Any],
    target_index: dict[str, dict[str, dict[str, Any]]],
    compact_codes: set[str],
    report: dict[str, Any],
    auto_project_targets: list[dict[str, Any]] | None = None,
) -> None:
    ws.append(["Разбор проекта:\nВсе разделы из GPT extraction", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
    ws.append(VISIBLE_REVIEW_HEADERS + TECH_REVIEW_HEADERS)

    status_col = VISIBLE_REVIEW_HEADERS.index("Статус") + 1
    found = review = missing_count = manual = 0

    def append_section_row(section_code: str, section_name: str) -> None:
        ws.append(
            [
                section_name,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                section_code,
                "",
                "",
                "",
                "section_header",
                "",
                "",
                "",
                "",
                f"sections.{section_code}",
                "section_header",
            ]
        )
        for cell in ws[ws.max_row]:
            cell.fill = FILL_GRAY
            cell.font = Font(name=FONT_NAME, bold=True, size=10)

    def append_row(row: list[Any], status: str) -> None:
        nonlocal found, review, missing_count, manual
        ws.append(row)
        if "Найдено" in status:
            found += 1
            fill = FILL_FOUND
        elif "Проверьте" in status:
            review += 1
            fill = FILL_REVIEW
        else:
            if "ручной" in status:
                manual += 1
            else:
                missing_count += 1
            fill = FILL_MISSING
        for cell in ws[ws.max_row]:
            cell.fill = fill

    if auto_project_targets:
        report["auto_project_targets_total"] = len(auto_project_targets)
        by_section: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for target in auto_project_targets:
            by_section[target["section_code"]].append(target)
        for section_code, targets in by_section.items():
            section_name = SECTION_NAMES_RU.get(section_code, targets[0]["section_name"])
            append_section_row(section_code, section_name)
            for target in targets:
                key = target["calculator_input_key"]
                item, matched_value, match_rule = match_auto_project_target(target, extraction)
                row_kind = "auto_project"
                raw_text = target.get("note") or ""
                source = ""
                confidence = ""
                group_code = ""
                item_name = ""
                if item:
                    group_code = item.get("group_code") or ""
                    item_name = item.get("item_name") or ""
                    source = source_label(item)
                    confidence = item.get("confidence") or ""
                    raw_text = item.get("raw_text") or raw_text
                value = matched_value
                if item and isinstance(value, (dict, list)):
                    row_kind = "detail_ref"
                    found_value = "см. лист 03_Детали объемов"
                elif item and value not in (None, ""):
                    found_value = excel_value(value)
                else:
                    found_value = ""

                split_by_level = key in {
                    "parapet_length_level_1_m",
                    "parapet_length_level_2_m",
                    "vent_wall_abutment_level_1_m",
                    "vent_wall_abutment_level_2_m",
                }
                if item and found_value != "":
                    status = "🟧 Проверьте" if split_by_level or item.get("needs_review") else "✅ Найдено"
                    action = (
                        "GPT нашел общий показатель; разнесите по уровням перед calculator input."
                        if split_by_level
                        else "Проверьте при необходимости; если всё верно, ничего не меняйте."
                    )
                else:
                    status = human_status_for_missing(key)
                    action = "AUTO_PROJECT из решения Елены не найден в GPT extraction; введите вручную или оставьте комментарий."
                    report["manual_logic_needed"].append(f"{section_code}:{key}: missing AUTO_PROJECT")

                unit = target.get("unit") or ""
                if not unit and item and not key.endswith(".steel_class"):
                    unit = expected_unit(section_code, item_code(item), item, target_index)
                append_row(
                    [
                        target["label"],
                        found_value,
                        unit,
                        status,
                        action,
                        source,
                        raw_text,
                        "",
                        "",
                        section_code,
                        key,
                        group_code,
                        item_name,
                        row_kind,
                        value_kind(value),
                        key,
                        extraction_status_for_human_status(status),
                        confidence,
                        match_rule,
                        "auto_project_from_elena_audit",
                    ],
                    status,
                )
    else:
        for section_code, section in extraction.get("sections", {}).items():
            section_name = SECTION_NAMES_RU.get(section_code, section_code)
            append_section_row(section_code, section_name)
            for index, item in enumerate(section.get("found") or []):
                helper = is_helper_item(item, section_code, compact_codes)
                if helper:
                    continue
                target_code = item.get("target_code")
                group_code = item.get("group_code")
                code = target_code or group_code
                value = item.get("value")
                row_kind = "detail_ref" if is_detail_item(item) else "scalar"
                found_value = "см. лист 03_Детали объемов" if row_kind == "detail_ref" else (
                    excel_value(value) if not isinstance(value, (dict, list)) else dumps_cell(value)
                )
                label = target_label(section_code, code, item.get("item_name"), target_index)
                status = human_status_for_item(item)
                append_row(
                    [
                        label,
                        found_value,
                        expected_unit(section_code, code, item, target_index),
                        status,
                        action_for_item(item, row_kind=row_kind),
                        source_label(item),
                        "См. лист 03_Детали объемов." if row_kind == "detail_ref" else item.get("raw_text") or "",
                        "",
                        "",
                        section_code,
                        target_code or "",
                        group_code or "",
                        item.get("item_name") or "",
                        row_kind,
                        value_kind(value),
                        "",
                        extraction_status_for_human_status(status),
                        item.get("confidence") or "",
                        f"sections.{section_code}.found[{index}]",
                        "preview_only",
                    ],
                    status,
                )

            for index, missing in enumerate(section.get("missing") or []):
                if isinstance(missing, dict):
                    target_code = missing.get("target_code") or missing.get("code") or ""
                    note = missing.get("notes") or ""
                else:
                    target_code = str(missing)
                    note = ""
                status = human_status_for_missing(target_code)
                label = target_label(section_code, target_code, None, target_index)
                append_row(
                    [
                        label,
                        "",
                        target_index.get(section_code, {}).get(target_code, {}).get("unit", ""),
                        status,
                        "Значение не найдено в extraction JSON; введите вручную в колонку “Исправить / ввести значение”.",
                        "",
                        note,
                        "",
                        "",
                        section_code,
                        target_code,
                        "",
                        "",
                        "missing",
                        "empty",
                        "",
                        "missing",
                        "",
                        f"sections.{section_code}.missing[{index}]",
                        "manual_required",
                    ],
                    status,
                )

    ws.cell(2, 1).value = f"Найдено уверенно: {found}"
    ws.cell(2, 2).value = f"Проверьте: {review}"
    ws.cell(2, 3).value = f"Не найдено: {missing_count}"
    ws.cell(2, 4).value = f"Ручной ввод: {manual}"
    ws.cell(1, 1).font = Font(name=FONT_NAME, bold=True, size=13)
    ws.cell(1, 1).fill = FILL_HEADER
    style_sheet_basic(
        ws,
        header_row=4,
        widths={
            "A": 42,
            "B": 22,
            "C": 10,
            "D": 18,
            "E": 68,
            "F": 50,
            "G": 64,
            "H": 28,
            "I": 28,
            "J": 22,
            "K": 22,
            "L": 22,
            "M": 26,
            "N": 20,
            "O": 16,
            "P": 22,
            "Q": 20,
            "R": 14,
            "S": 36,
            "T": 18,
        },
    )
    ws.freeze_panes = "A5"
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[4].height = 28
    for col_idx in range(len(VISIBLE_REVIEW_HEADERS) + 1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].hidden = True


def append_details_sheet_earthworks_like(ws, extraction: dict[str, Any], report: dict[str, Any]) -> None:
    ws.append(DETAIL_HEADERS)

    def block(title: str) -> None:
        ws.append([title] + [""] * (len(DETAIL_HEADERS) - 1))
        for cell in ws[ws.max_row]:
            cell.fill = FILL_GRAY
            cell.font = Font(name=FONT_NAME, bold=True, size=10)

    block("Земляные работы — траншеи")
    for item in iter_found_items(extraction):
        if item_code(item) != "trench_routes" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        route_name = value.get("route_name") or value.get("name") or ""
        ws.append([
            "Траншея",
            route_name,
            excel_value(value.get("length_m")),
            excel_value(value.get("depth_m")),
            excel_value(value.get("width_m")),
            excel_value(value.get("volume_m3")),
            "",
            "",
            "",
            "",
            "",
            source_label(item),
            item.get("raw_text") or f"Таблица траншей: {route_name}",
            "NEEDS_REVIEW" if item.get("needs_review") else "",
        ])
        report["detail_counts"]["trench_routes"] += 1

    block("Земляные работы — коммуникации")
    for item in iter_found_items(extraction):
        if item_code(item) != "communications_pipe_items" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        total_length = value.get("total_length_m")
        if total_length in (None, ""):
            piece = value.get("piece_length_m")
            qty = value.get("quantity_pcs") or value.get("quantity")
            linear = value.get("linear_length_m")
            if piece not in (None, "") and qty not in (None, ""):
                total_length = float(piece) * float(qty)
            elif linear not in (None, ""):
                total_length = linear
        ws.append([
            "Коммуникация",
            value.get("name") or "",
            "",
            "",
            "",
            "",
            excel_value(value.get("diameter_mm")),
            excel_value(value.get("piece_length_m")),
            excel_value(value.get("quantity_pcs") or value.get("quantity")),
            excel_value(total_length),
            "да" if total_length not in (None, "") else "нет",
            source_label(item),
            item.get("raw_text") or value.get("name") or "",
            "Фитинг без длины не участвует в сумме." if total_length in (None, "") else "",
        ])
        report["detail_counts"]["communications_pipe_items"] += 1

    block("Арматура")
    for section_code, item in iter_found_items(extraction, with_section=True):
        if not is_rebar_item(item) or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        length = value.get("length_m")
        mass_per_m = value.get("mass_per_m_kg")
        calculated = ""
        comment = ""
        if length not in (None, "") and mass_per_m not in (None, ""):
            calculated = round(float(length) * float(mass_per_m), 3)
            comment = f"Справочный вес: {display_number(calculated)} кг; не значение из PDF."
        elif value.get("weight_kg") in (None, ""):
            comment = "Нет mass_per_m_kg/weight_kg; нужна ручная логика перед calculator input."
        ws.append([
            "Арматура",
            f"{SECTION_NAMES_RU.get(section_code, section_code)} — {value.get('name') or item_code(item)}",
            excel_value(length),
            "",
            "",
            excel_value(value.get("weight_kg")) if value.get("weight_kg") not in (None, "") else "",
            excel_value(value.get("diameter_mm")),
            excel_value(mass_per_m),
            "",
            excel_value(calculated),
            "да",
            source_label(item),
            item.get("raw_text") or "",
            comment,
        ])
        report["detail_counts"]["rebar"] += 1

    block("Балки")
    for section_code, item in iter_found_items(extraction, with_section=True):
        if item_code(item) != "beam_items" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        ws.append([
            "Балка",
            f"{SECTION_NAMES_RU.get(section_code, section_code)} — {value.get('code') or value.get('name') or ''}",
            excel_value(value.get("length_m")),
            excel_value(value.get("height_m")),
            excel_value(value.get("width_m")),
            excel_value(value.get("concrete_m3")),
            "",
            "",
            "",
            excel_value(value.get("formwork_area_m2")),
            "да",
            source_label(item),
            item.get("raw_text") or "",
            "",
        ])
        report["detail_counts"]["beam_items"] += 1

    block("ЭППС плиты 1 этажа — детализация")
    for section_code, item in iter_found_items(extraction, with_section=True):
        if item_code(item) != "floor_slab_1_eps100_volume" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        bottom = value.get("bottom_volume_m3")
        edge = value.get("edge_volume_m3")
        total = float(bottom or 0) + float(edge or 0)
        for name, volume, comment in [
            ("ЭППС низ плиты", bottom, "child row"),
            ("ЭППС торец плиты", edge, "child row"),
            ("контрольная сумма bottom + edge", total, "справочно, подтвердить перед calculator input"),
        ]:
            ws.append([
                "ЭППС",
                f"{SECTION_NAMES_RU.get(section_code, section_code)} — {name}",
                "",
                "",
                "",
                excel_value(volume),
                "",
                "",
                "",
                "",
                "да",
                source_label(item),
                item.get("raw_text") or "",
                comment,
            ])
            report["detail_counts"]["floor_slab_1_eps_split"] += 1

    block("Прочие detail groups")
    handled_codes = {
        "trench_routes",
        "communications_pipe_items",
        "beam_items",
        "floor_slab_1_eps100_volume",
        *REBAR_GROUP_CODES,
        "masonry_rebar_a500_d10_weight",
    }
    for section_code, item in iter_found_items(extraction, with_section=True):
        code = item_code(item)
        value = item.get("value")
        if code in handled_codes:
            continue
        if code in DETAIL_GROUP_CODES or isinstance(value, (dict, list)):
            ws.append([
                "Detail",
                f"{SECTION_NAMES_RU.get(section_code, section_code)} — {code}",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "да",
                source_label(item),
                dumps_cell(value),
                "Нужен отдельный adapter mapping.",
            ])
            report["detail_counts"]["other_details"] += 1

    style_sheet_basic(
        ws,
        header_row=1,
        widths={
            "A": 24,
            "B": 52,
            "C": 12,
            "D": 12,
            "E": 12,
            "F": 14,
            "G": 14,
            "H": 14,
            "I": 12,
            "J": 16,
            "K": 12,
            "L": 44,
            "M": 78,
            "N": 32,
        },
    )
    ws.freeze_panes = "A2"
    for row_idx in range(2, ws.max_row + 1):
        row_type = str(ws.cell(row_idx, 1).value or "")
        if "—" in row_type or row_type in {"Арматура", "Балки"} and not ws.cell(row_idx, 2).value:
            for cell in ws[row_idx]:
                cell.fill = FILL_GRAY
                cell.font = Font(name=FONT_NAME, bold=True, size=10)


def append_candidates_sheet_earthworks_like(
    ws,
    extraction: dict[str, Any],
    compact_codes: set[str],
    report: dict[str, Any],
) -> None:
    ws.append(["Технический лист для разработчика. Здесь показаны helper rows и диагностические строки GPT extraction. Полные raw table rows находятся на листе 06."])
    ws.cell(1, 1).font = Font(name=FONT_NAME, bold=True, size=12)
    headers = [
        "Параметр",
        "technical_key",
        "Значение кратко",
        "Ед.",
        "Статус кандидата",
        "Уверенность",
        "Правило поиска",
        "Лист проекта",
        "Тип листа",
        "PDF",
        "Стр.",
        "Фрагмент короткий",
        "Комментарий parser-а",
        "evidence_id",
        "candidate_id",
    ]
    ws.append(headers)
    for section_code, item in iter_found_items(extraction, with_section=True):
        helper = is_helper_item(item, section_code, compact_codes)
        candidates = item.get("candidates")
        if not helper and not candidates:
            continue
        code = item_code(item)
        ws.append([
            f"{SECTION_NAMES_RU.get(section_code, section_code)} — {item.get('item_name') or code}",
            code,
            dumps_cell(item.get("value")),
            item.get("unit") or item.get("normalized_unit") or "",
            "helper_only" if helper else "candidate",
            item.get("confidence") or "",
            "not_auto_calculator_input",
            item.get("page_title") or item.get("table_context") or "",
            item.get("table_context") or "",
            item.get("source_pdf") or "",
            item.get("page_number") or "",
            item.get("raw_text") or "",
            item.get("notes") or "Visible for Elena; blocked until explicit adapter mapping.",
            "",
            "",
        ])
        report["helper_rows"].append(f"{section_code}:{code}")
    style_sheet_basic(
        ws,
        header_row=2,
        widths={
            "A": 36,
            "B": 34,
            "C": 32,
            "D": 10,
            "E": 18,
            "F": 14,
            "G": 28,
            "H": 28,
            "I": 28,
            "J": 22,
            "K": 10,
            "L": 64,
            "M": 48,
            "N": 20,
            "O": 20,
        },
    )


def append_raw_sheet_earthworks_like(ws, extraction: dict[str, Any], report: dict[str, Any]) -> None:
    summary_headers = ["Показатель", "Значение", "Комментарий"]
    sections = extraction.get("sections", {})
    raw_total = sum(len(section.get("raw_table_rows") or []) for section in sections.values())
    found_total = sum(len(section.get("found") or []) for section in sections.values())
    missing_total = sum(len(section.get("missing") or []) for section in sections.values())
    summary_rows = [
        ["project_name", extraction.get("project_name") or "", "имя проекта из GPT JSON"],
        ["source_type", "GPT/Claude extraction_output.json", "upstream для preview workbook"],
        ["sections_count", len(sections), "сколько разделов в JSON"],
        ["raw_table_rows_count", raw_total, "сколько исходных строк таблиц выгружено"],
        ["found_count", found_total, "сколько found items"],
        ["missing_count", missing_total, "сколько missing targets"],
    ]
    row_idx = write_block_table(ws, "Блок 0: Summary запуска", summary_headers, summary_rows, 1)
    headers = [
        "section_code",
        "table_id",
        "table_title",
        "page_number",
        "page_title",
        "row_index",
        "columns",
        "cells",
        "unit",
        "normalized_unit",
        "mapped_target_codes",
        "raw_text",
        "source_pdf",
        "notes",
    ]
    rows: list[list[Any]] = []
    for section_code, section in sections.items():
        for row in section.get("raw_table_rows") or []:
            rows.append([
                section_code,
                row.get("table_id") or "",
                row.get("table_title") or "",
                row.get("page_number") or "",
                row.get("page_title") or "",
                row.get("row_index") or "",
                dumps_cell(row.get("columns")),
                dumps_cell(row.get("cells")),
                row.get("unit") or "",
                row.get("normalized_unit") or "",
                dumps_cell(row.get("mapped_target_codes")),
                row.get("raw_text") or "",
                row.get("source_pdf") or "",
                row.get("notes") or "",
            ])
            report["raw_rows"] += 1
    row_idx = write_block_table(ws, "Блок 1: raw_table_rows GPT", headers, rows, row_idx)

    found_rows: list[list[Any]] = []
    for section_code, section in sections.items():
        for index, item in enumerate(section.get("found") or []):
            found_rows.append([
                section_code,
                index,
                item.get("target_code") or "",
                item.get("group_code") or "",
                item.get("item_name") or "",
                dumps_cell(item.get("value")),
                item.get("unit") or "",
                item.get("normalized_unit") or "",
                item.get("confidence") or "",
                "да" if item.get("needs_review") else "нет",
                source_label(item),
                item.get("raw_text") or "",
                item.get("notes") or "",
            ])
    write_block_table(
        ws,
        "Блок 2: found items GPT",
        ["section_code", "index", "target_code", "group_code", "item_name", "value", "unit", "normalized_unit", "confidence", "needs_review", "source", "raw_text", "notes"],
        found_rows,
        row_idx,
    )
    apply_table_theme(ws)
    set_widths(
        ws,
        {
            "A": 30,
            "B": 32,
            "C": 34,
            "D": 18,
            "E": 24,
            "F": 18,
            "G": 26,
            "H": 36,
            "I": 18,
            "J": 18,
            "K": 28,
            "L": 72,
            "M": 32,
            "N": 34,
        },
    )
    for idx in range(1, ws.max_row + 1):
        if str(ws.cell(idx, 1).value or "").startswith("Блок "):
            style_block_title_row(ws, idx, ws.max_column)
    ws.freeze_panes = "A3"


def append_details_sheet(ws, extraction: dict[str, Any], report: dict[str, Any]) -> None:
    max_width = 14

    append_block_title(ws, "Земляные работы — траншеи", max_width)
    headers = ["route_name", "length_m", "depth_m", "width_m", "volume_m3", "status", "source", "raw_text", "comment_for_elena"]
    append_header(ws, headers)
    for item in iter_found_items(extraction):
        if item_code(item) != "trench_routes" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        status = status_for_item(item)
        ws.append(
            [
                value.get("route_name") or value.get("name") or "",
                value.get("length_m"),
                value.get("depth_m"),
                value.get("width_m"),
                value.get("volume_m3"),
                status,
                source_label(item),
                item.get("raw_text") or "",
                "Проверить объем/размеры; preview ничего не исправляет." if item.get("needs_review") else "",
            ]
        )
        report["detail_counts"]["trench_routes"] += 1

    ws.append([])
    append_block_title(ws, "Земляные работы — коммуникации", max_width)
    headers = ["name", "diameter_mm", "piece_length_m", "quantity_pcs", "linear_length_m", "total_length_m", "unit", "status", "source", "raw_text"]
    append_header(ws, headers)
    for item in iter_found_items(extraction):
        if item_code(item) != "communications_pipe_items" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        total_length = value.get("total_length_m")
        if total_length in (None, ""):
            piece = value.get("piece_length_m")
            qty = value.get("quantity_pcs") or value.get("quantity")
            linear = value.get("linear_length_m")
            if piece not in (None, "") and qty not in (None, ""):
                total_length = float(piece) * float(qty)
            elif linear not in (None, ""):
                total_length = linear
        status = status_for_item(item)
        ws.append(
            [
                value.get("name") or "",
                value.get("diameter_mm"),
                value.get("piece_length_m"),
                value.get("quantity_pcs") or value.get("quantity"),
                value.get("linear_length_m"),
                total_length,
                item.get("unit") or item.get("normalized_unit") or "",
                status,
                source_label(item),
                item.get("raw_text") or "",
            ]
        )
        report["detail_counts"]["communications_pipe_items"] += 1

    ws.append([])
    append_block_title(ws, "Арматура", max_width)
    headers = [
        "section",
        "group_code",
        "item_name",
        "steel_class",
        "diameter_mm",
        "length_m",
        "mass_per_m_kg",
        "calculated_weight_kg",
        "weight_kg_from_pdf",
        "status",
        "source",
        "raw_text",
        "comment_for_elena",
    ]
    append_header(ws, headers)
    for section_code, item in iter_found_items(extraction, with_section=True):
        if not is_rebar_item(item) or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        length = value.get("length_m")
        mass_per_m = value.get("mass_per_m_kg")
        calculated = ""
        status = status_for_item(item)
        comment = ""
        if length not in (None, "") and mass_per_m not in (None, ""):
            calculated = round(float(length) * float(mass_per_m), 3)
            comment = "Справочный расчет preview; это не вес из PDF."
        elif value.get("weight_kg") in (None, ""):
            status = "MANUAL_NEEDED" if not item.get("needs_review") else "NEEDS_REVIEW"
            comment = "Нет mass_per_m_kg или weight_kg_from_pdf; нужна ручная логика/каталог."
        ws.append(
            [
                SECTION_NAMES_RU.get(section_code, section_code),
                item_code(item),
                value.get("name") or item.get("item_name") or "",
                value.get("steel_class") or "",
                value.get("diameter_mm") or "",
                length,
                mass_per_m,
                calculated,
                value.get("weight_kg") or "",
                status,
                source_label(item),
                item.get("raw_text") or "",
                comment,
            ]
        )
        report["detail_counts"]["rebar"] += 1

    ws.append([])
    append_block_title(ws, "Балки", max_width)
    headers = ["section", "code", "length_m", "width_m", "height_m", "concrete_m3", "formwork_area_m2", "status", "source", "raw_text", "comment_for_elena"]
    append_header(ws, headers)
    for section_code, item in iter_found_items(extraction, with_section=True):
        if item_code(item) != "beam_items" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        ws.append(
            [
                SECTION_NAMES_RU.get(section_code, section_code),
                value.get("code") or value.get("name") or "",
                value.get("length_m"),
                value.get("width_m"),
                value.get("height_m"),
                value.get("concrete_m3"),
                value.get("formwork_area_m2"),
                status_for_item(item),
                source_label(item),
                item.get("raw_text") or "",
                "",
            ]
        )
        report["detail_counts"]["beam_items"] += 1

    ws.append([])
    append_block_title(ws, "Floor slab 1 EPS split", max_width)
    headers = ["section", "target_code", "child_name", "volume_m3", "status", "source", "raw_text", "comment_for_elena"]
    append_header(ws, headers)
    for section_code, item in iter_found_items(extraction, with_section=True):
        if item_code(item) != "floor_slab_1_eps100_volume" or not isinstance(item.get("value"), dict):
            continue
        value = item["value"]
        bottom = value.get("bottom_volume_m3")
        edge = value.get("edge_volume_m3")
        total = None
        if bottom not in (None, "") and edge not in (None, ""):
            total = float(bottom) + float(edge)
        rows = [
            ("ЭППС низ плиты", bottom, "child row; не утвержденный calculator input"),
            ("ЭППС торец плиты", edge, "child row; не утвержденный calculator input"),
            ("контрольная сумма bottom + edge", total, "справочно, подтвердить перед calculator input"),
        ]
        for child_name, volume, comment in rows:
            ws.append(
                [
                    SECTION_NAMES_RU.get(section_code, section_code),
                    item_code(item),
                    child_name,
                    volume,
                    "NEEDS_REVIEW",
                    source_label(item),
                    item.get("raw_text") or "",
                    comment,
                ]
            )
            report["detail_counts"]["floor_slab_1_eps_split"] += 1

    ws.append([])
    append_block_title(ws, "Прочие detail groups", max_width)
    headers = ["section", "target_or_group_code", "item_name", "value_json", "unit", "status", "source", "raw_text", "comment_for_elena"]
    append_header(ws, headers)
    handled_codes = {
        "trench_routes",
        "communications_pipe_items",
        "beam_items",
        "floor_slab_1_eps100_volume",
        *REBAR_GROUP_CODES,
        "masonry_rebar_a500_d10_weight",
    }
    for section_code, item in iter_found_items(extraction, with_section=True):
        code = item_code(item)
        value = item.get("value")
        if code in handled_codes:
            continue
        if code in DETAIL_GROUP_CODES or isinstance(value, (dict, list)):
            ws.append(
                [
                    SECTION_NAMES_RU.get(section_code, section_code),
                    code,
                    item.get("item_name") or "",
                    dumps_cell(value),
                    item.get("unit") or item.get("normalized_unit") or "",
                    status_for_item(item),
                    source_label(item),
                    item.get("raw_text") or "",
                    "Preview сохранил как detail row; нужен отдельный adapter mapping.",
                ]
            )
            report["detail_counts"]["other_details"] += 1


def iter_found_items(extraction: dict[str, Any], with_section: bool = False):
    for section_code, section in extraction.get("sections", {}).items():
        for item in section.get("found") or []:
            if with_section:
                yield section_code, item
            else:
                yield item


def append_helper_rows(
    ws,
    extraction: dict[str, Any],
    compact_codes: set[str],
    report: dict[str, Any],
) -> None:
    headers = ["section", "target_code", "item_name", "value", "unit", "status", "reason", "source", "raw_text", "comment_for_elena"]
    ws.append(headers)
    status_col = headers.index("status") + 1
    for section_code, item in iter_found_items(extraction, with_section=True):
        helper = is_helper_item(item, section_code, compact_codes)
        candidates = item.get("candidates")
        if not helper and not candidates:
            continue
        code = item_code(item)
        reason = "helper/outside row; not sent to calculator input"
        if candidates:
            reason = f"{reason}; candidates present"
        if code == "project_spec_roof_area":
            reason = "control row from roof material/spec area; do not override roof levels automatically"
        status = "HELPER_ONLY"
        ws.append(
            [
                SECTION_NAMES_RU.get(section_code, section_code),
                code,
                item.get("item_name") or "",
                dumps_cell(item.get("value")),
                item.get("unit") or item.get("normalized_unit") or "",
                status,
                reason,
                source_label(item),
                item.get("raw_text") or "",
                "Visible for Elena; blocked until explicit adapter mapping.",
            ]
        )
        ws.cell(ws.max_row, status_col).fill = status_fill(status)
        report["helper_rows"].append(f"{section_code}:{code}")


def append_raw_rows(ws, extraction: dict[str, Any], report: dict[str, Any]) -> None:
    headers = [
        "section_code",
        "table_id",
        "table_title",
        "page_number",
        "page_title",
        "row_index",
        "columns",
        "cells",
        "unit",
        "normalized_unit",
        "mapped_target_codes",
        "raw_text",
        "source_pdf",
        "notes",
    ]
    ws.append(headers)
    for section_code, section in extraction.get("sections", {}).items():
        for row in section.get("raw_table_rows") or []:
            ws.append(
                [
                    section_code,
                    row.get("table_id") or "",
                    row.get("table_title") or "",
                    row.get("page_number") or "",
                    row.get("page_title") or "",
                    row.get("row_index") or "",
                    dumps_cell(row.get("columns")),
                    dumps_cell(row.get("cells")),
                    row.get("unit") or "",
                    row.get("normalized_unit") or "",
                    dumps_cell(row.get("mapped_target_codes")),
                    row.get("raw_text") or "",
                    row.get("source_pdf") or "",
                    row.get("notes") or "",
                ]
            )
            report["raw_rows"] += 1


def append_constructor(ws, extraction: dict[str, Any]) -> None:
    constructor_rows = [
        ["Земляные работы", "earthworks", "да", "GPT preview", "проверка проектных данных, цен и деталей", ""],
        ["Фундаментная плита", "foundation_slab", "да", "GPT preview", "проверка проектных данных и арматуры", ""],
        ["Гидроизоляция", "waterproofing", "да", "GPT preview", "проверка проектных данных", ""],
        ["Стены и перемычки", "load_bearing_walls_lintels", "да", "GPT preview", "проверка проектных данных и арматуры", ""],
        ["Перекрытие 1 этажа", "floor_slab_1", "да", "GPT preview", "проверка проектных данных, арматуры и балок", ""],
        ["Перекрытие 2 этажа", "floor_slab_2", "да", "GPT preview", "проверка проектных данных и арматуры", ""],
        ["Кровля", "flat_roof", "да", "GPT preview", "проверка проектных данных и helper rows", ""],
        ["Schiedel / вентканалы", "schiedel_vent_channels", "да", "GPT preview", "проверка проектных данных", ""],
    ]
    ws.append(["Раздел сметы", "Код раздела", "Включить в смету", "Статус готовности", "Что сейчас проверяется", "Комментарий"])
    for row in constructor_rows:
        ws.append(row)


def append_prices_placeholder(ws) -> None:
    ws.append(PRICE_HEADERS)
    ws.append(
        [
            "TODO",
            "Цены пока не подключены в preview из GPT JSON",
            "",
            "",
            "",
            "",
            "",
            "preview",
            "да",
            "Лист сохранен в структуре workbook, как в earthworks flow.",
            "",
            "",
            "",
            "",
        ]
    )


def append_instruction(ws) -> None:
    ws.append(["Раздел", "Инструкция"])
    lines = [
        "Начните с листа 00_Конструктор сметы.",
        "На листе 01_Проверка проекта посмотрите цветные строки.",
        "Зеленые строки GPT нашел уверенно. Если все верно, ничего делать не нужно.",
        "Оранжевые строки требуют проверки. Если значение проекта неверное, внесите исправление в колонку “Исправить / ввести значение”.",
        "Красные строки GPT не нашел. Внесите значение вручную в колонку “Исправить / ввести значение”.",
        "Где именно проверять оранжевую строку, написано в колонке “Что нужно сделать”.",
        "Детали траншей, коммуникаций, арматуры и балок находятся на листе 03_Детали объемов.",
        "Helper rows на листе 05_Кандидаты parser не идут в расчет автоматически.",
        "Цены пока не подключены; лист 02_Цены себестоимости оставлен в структуре workbook.",
        "Preview ничего не публикует в Google и не запускает calculators.",
    ]
    for idx, line in enumerate(lines, start=1):
        ws.append([idx, line])


def build_workbook(
    extraction: dict[str, Any],
    target_index: dict[str, dict[str, dict[str, Any]]],
    compact_codes: set[str],
    auto_project_targets: list[dict[str, Any]] | None = None,
) -> tuple[Workbook, dict[str, Any]]:
    report: dict[str, Any] = {
        "sheet_rows": {},
        "detail_counts": Counter(),
        "helper_rows": [],
        "raw_rows": 0,
        "unknown_target_codes": Counter(),
        "manual_logic_needed": [],
    }

    wb = Workbook()
    wb.remove(wb.active)
    ws00 = wb.create_sheet(SHEET_00)
    ws01 = wb.create_sheet(SHEET_01)
    ws02 = wb.create_sheet(SHEET_02)
    ws03 = wb.create_sheet(SHEET_03)
    ws04 = wb.create_sheet(SHEET_04)
    ws05 = wb.create_sheet(SHEET_05)
    ws06 = wb.create_sheet(SHEET_06)

    append_constructor(ws00, extraction)
    append_review_sheet_earthworks_like(ws01, extraction, target_index, compact_codes, report, auto_project_targets)
    append_prices_placeholder(ws02)
    append_details_sheet_earthworks_like(ws03, extraction, report)
    append_instruction(ws04)
    append_candidates_sheet_earthworks_like(ws05, extraction, compact_codes, report)
    append_raw_sheet_earthworks_like(ws06, extraction, report)

    for section_code, item in iter_found_items(extraction, with_section=True):
        code = item_code(item)
        if code and code not in compact_codes and not is_detail_item(item):
            report["unknown_target_codes"][f"{section_code}:{code}"] += 1
        if is_rebar_item(item) and isinstance(item.get("value"), dict):
            value = item["value"]
            if value.get("mass_per_m_kg") in (None, "") and value.get("weight_kg") in (None, ""):
                report["manual_logic_needed"].append(f"{section_code}:{code}:{value.get('name') or ''}")
        if item_code(item) in {"roof_parapet_abutment_total_length", "roof_wall_abutment_total_length"}:
            report["manual_logic_needed"].append(f"{section_code}:{item_code(item)}: split by roof level")
    for section_code, section in extraction.get("sections", {}).items():
        for missing in section.get("missing") or []:
            target_code = missing.get("target_code") if isinstance(missing, dict) else str(missing)
            report["manual_logic_needed"].append(f"{section_code}:{target_code}: missing")

    style_sheet(ws00)
    style_sheet(ws02, hide_after_col=10)
    style_sheet(ws04)

    for ws in wb.worksheets:
        report["sheet_rows"][ws.title] = max(ws.max_row - 1, 0)
    return wb, report


def render_report(output_path: Path, report: dict[str, Any]) -> str:
    lines = [
        "Review workbook preview created.",
        f"output: {output_path}",
        "",
    ]
    if report.get("auto_project_targets_total"):
        lines.extend([f"AUTO_PROJECT targets from Elena audit: {report['auto_project_targets_total']}", ""])
    lines.append("Sheet rows:")
    for sheet, count in report["sheet_rows"].items():
        lines.append(f"- {sheet}: {count}")
    lines.extend(["", "Detail rows:"])
    for key, count in sorted(report["detail_counts"].items()):
        lines.append(f"- {key}: {count}")
    lines.extend(["", "Helper rows:"])
    if report["helper_rows"]:
        for key in sorted(set(report["helper_rows"])):
            lines.append(f"- {key}")
    else:
        lines.append("- none")
    lines.extend(["", "Unknown target_code/helper candidates:"])
    if report["unknown_target_codes"]:
        for key, count in sorted(report["unknown_target_codes"].items()):
            lines.append(f"- {key}: {count}")
    else:
        lines.append("- none")
    lines.extend(["", "Manual logic needed before calculator input:"])
    for item in report["manual_logic_needed"][:40]:
        lines.append(f"- {item}")
    if len(report["manual_logic_needed"]) > 40:
        lines.append(f"- ... and {len(report['manual_logic_needed']) - 40} more")
    return "\n".join(lines) + "\n"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build local review_workbook_preview.xlsx from GPT extraction_output.json.")
    parser.add_argument("extraction_json", type=Path, help="Path to GPT/Claude extraction_output.json")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH, help="Output xlsx path")
    parser.add_argument("--targets", type=Path, default=DEFAULT_TARGETS_PATH, help="calculator_targets_compact.json path")
    parser.add_argument("--elena-audit", type=Path, default=DEFAULT_ELENA_AUDIT_PATH, help="Markdown audit with the 87 AUTO_PROJECT rows")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    extraction = load_json(args.extraction_json)
    target_index, compact_codes = compact_target_index(args.targets)
    auto_project_targets = parse_auto_project_targets(args.elena_audit)
    workbook, report = build_workbook(extraction, target_index, compact_codes, auto_project_targets)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(render_report(args.output, report))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
