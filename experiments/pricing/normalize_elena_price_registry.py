from __future__ import annotations

import argparse
import re
import shutil
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = ROOT / "output" / "Цены на материалы.xlsx"
DEFAULT_ARCHIVE_SOURCE = ROOT / "experiments" / "pricing" / "input" / "price_registry_elena_2026-07-28.xlsx"
DEFAULT_PREVIOUS_REGISTRY = ROOT / "output" / "price_registry_filled_v3.xlsx"
DEFAULT_OUTPUT = ROOT / "output" / "price_registry_filled_v4.xlsx"
DEFAULT_REPORT = ROOT / "experiments" / "pricing" / "output" / "price_registry_v4_mapping_report.md"

REGISTRY_HEADERS = [
    "Раздел",
    "Наименование",
    "Ед. изм.",
    "Минимальная комплектация",
    "Цена",
    "Дата обновления",
    "Комментарий",
    "price_code",
]


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).replace("\xa0", " ").strip()


def normalize_text(value: Any) -> str:
    text = clean_text(value).lower()
    text = text.replace("ё", "е")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_unit(value: Any) -> str:
    unit = normalize_text(value)
    aliases = {
        "1п.м": "мп",
        "1п. м": "мп",
        "п.м": "мп",
        "п. м": "мп",
        "м/п": "мп",
        "м.п": "мп",
        "м": "мп",
        "м2": "м2",
        "м²": "м2",
        "м3": "м3",
        "м³": "м3",
        "шт.": "шт",
        "смен": "смена",
    }
    return aliases.get(unit, unit)


def row_key(name: Any, unit: Any) -> tuple[str, str]:
    return normalize_text(name), normalize_unit(unit)


MANUAL_WORK_PRICE_CODES: dict[tuple[str, str], str] = {
    row_key(
        'Экскаватор-погрузчик JCB (раскопка "корыта" под устройство дороги) /Контроль за механизированной разработкой грунта',
        "смена",
    ): "excavator_jcb_shift",
    row_key(
        "Разработка грунта вручную (в т.ч. Копка траншей под коммуникации)",
        "м3",
    ): "manual_excavation_m3",
    row_key(
        "Отсыпка дна котлована песком с трамбованием механизировано с применением виброплиты (коэф.упл. 1,3)",
        "м3",
    ): "sand_filling_work_m3",
    row_key(
        "Закладка технологических входов коммуникаций до границы дома. Канализация, водоснабжение (ориентировочно)",
        "мп",
    ): "communications_installation_m",
    # Fixed 2026-08-06: was pointed at thermal_insert_installation_work_m, the LEGACY
    # registry_code for a disabled estimate line (thermal_insert_mode=standard_50_100 is
    # production and never uses it) - so this real 100 руб./мп price was never actually
    # reaching any live calculation. Retargeted to the combined-mode code TRC's real data
    # needs (thermal_insert_combined_length_m given, not split 50mm/100mm lengths).
    row_key("Устройство и монтаж термовкладыша", "мп"): "thermal_insert_combined_installation_work_m",
    # Added 2026-08-06: these are materials, normally on a separate sheet this script doesn't
    # read (extract_work_price_rows() only processes wb.sheetnames[0]) - moved onto the first
    # sheet ("Прайс по видам работ") in the live source file (output/Цены на материалы.xlsx,
    # rows 43-44) specifically so the existing single-sheet extraction picks them up
    # automatically instead of needing a DERIVED_PRICE_ROWS patch. Same physical product
    # already used for foundation_slab wall insulation, same price.
    row_key("Пеноплэкс ГЕО 50 мм (для термовставок)", "м3"): "thermal_insert_50_material_m3",
    row_key("Пеноплэкс ГЕО 100 мм (для термовставок)", "м3"): "thermal_insert_100_material_m3",
    # Added 2026-08-06: schiedel_masonry_gas_block_items[] (D400/D500 masonry AROUND the
    # Schiedel vent-shaft modules). Elena's raw price for this exact block was only on
    # "Блокпаротермшидель" sheet (rows 4-5: "Блок D 400"=6100, "Блок D 500"=5400 руб/м3, this
    # script only reads sheet 0) - moved onto sheet 0 (rows 45-46) same as the termovstavka
    # materials above. Note these are generic block prices, NOT the same block/price as
    # wall_block_items' main-wall D400/D500 (9500/9800) - that's a per-project supplier quote
    # entered directly on each row, not a shared registry price at all, so there is no
    # conflict. User confirmed 2026-08-06: use this value as-is even though it may be a
    # different block size than the vent-shaft's actual 150x250x650mm spec - it's what
    # Elena's price list gives for this line, don't block on the possible size mismatch.
    row_key("Блок D400 (для кладки шахты вентканалов)", "м3"): "schiedel_masonry_gas_block_d400_m3",
    row_key("Блок D500 (для кладки шахты вентканалов)", "м3"): "schiedel_masonry_gas_block_d500_m3",
    row_key(
        "Бетонирование фундаментной плиты в опалубке бетоном марки В22,5 (М300)",
        "м3",
    ): "concrete_placing_work_m3",
    row_key(
        "Гидроизоляция торца фундаментной плиты битумной мастикой в 2 слоя",
        "м2",
    ): "bitumen_waterproofing_work_m2",
    row_key("Утепление торца плиты ЭППС 100мм", "м2"): "eps_wall_insulation_work_m2",
    row_key(
        "Устройство лесов, подмостей для кладки, демонтаж лесов после завершения работ",
        "компл",
    ): "scaffolding_setup_dismantling_work_set",
    row_key(
        "Бетонирование перемычек в U блоке бетоном марки В22,5 (М300)",
        "мп",
    ): "lintel_concreting_work_m",
    row_key(
        "Бетонирование монолитных перемычек бетоном марки В22,5 (М300)",
        "мп",
    ): "lintel_monolithic_concreting_work_m",
    row_key(
        "Бетонирование балки бетоном марки В22,5 (М300) (высотой до 250мм)",
        "мп",
    ): "beam_concrete_placing_work_m",
    row_key("Вывоз мусора с объекта", "маш"): "waste_removal_loading_work_truck",
    row_key("Кладка парапета из газобетонных блоков", "м3"): "gas_block_masonry_work_m3",
    row_key(
        "Обкладка дымохода и вентканалов толщ. 150мм из из газобетонных блоков",
        "м2",
    ): "gas_block_cladding_work_m2",
    row_key("Пароизоляция основания плёнкой ПВХ", "м2"): "roof_vapor_barrier_installation_work_m2",
    row_key(
        "Утепление кровельного покрытия ЭППС (1 слой -100мм, 2 слой -100мм (50мм), 3 слой - разуклонка)",
        "м2",
    ): "roof_eps_insulation_installation_work_m2",
    row_key("Укладка ПВХ Мембраны", "м2"): "roof_pvc_membrane_installation_work_m2",
    row_key("Монтаж примыкания кровли из ПВХ мембраны", "мп"): "roof_pvc_membrane_abutment_work_m",
    row_key("Монтаж примыкания к вентшахтам", "шт"): "roof_vent_shaft_abutment_installation_item",
    # Fixed 2026-08-06: this used to map to roof_parapet_drain_item (the drain's MATERIAL
    # code) - wrong, "Установка..." (installation) is clearly work, not material. That bug
    # meant the material price silently got the installation rate (5000) instead of the real
    # material price (6001.23, see DERIVED_PRICE_ROWS below) for years. Corrected target.
    row_key("Установка воронки парапетной", "шт"): "roof_parapet_drain_installation_item",
    # New 2026-08-06: same class of installation-work row, previously never mapped at all -
    # roof_internal_drain_with_heating_item (material) happened to carry the same 5000 value
    # from an unrelated v3-carryforward match, masking the fact that this work price was
    # never actually sourced from anywhere.
    row_key(
        "Установка воронки кровельной (с обжимным мет. фланцем с обогревом 110х450мм) (без пробивки отверстий)",
        "шт",
    ): "roof_internal_drain_with_heating_installation_item",
    row_key("Подъем материалов автокраном / разгрузка материала в ручную", "смена"): "roof_crane_lifting_shift",
    row_key("Кладка вентканалов Schiedel", "мп"): "schiedel_masonry_work_m",
    # Removed 2026-08-06: this row ('Доставка вентканалов/ разгрузка в ручную на объекте' = 2500)
    # used to map here, but it's the combined delivery+unloading number, not delivery/material
    # alone - mapping it to schiedel_delivery_truck silently overwrote the real material price
    # (15000, carried forward from v3) with the work-only rate, double-counting against the new
    # separate schiedel_delivery_unloading_work code below. schiedel_delivery_truck is now a
    # DERIVED_PRICE_ROWS entry instead (real material figure, confirmed from real ЮСВ smeta).
}

IGNORED_WORK_PRICE_ROWS: set[tuple[str, str]] = {
    row_key(
        "Бетонирование балки бетоном марки В22,5 (М300) (высотой более 250мм)",
        "м3",
    ),
}

OBSOLETE_PREVIOUS_PRICE_CODES = {
    "beam_concrete_placing_work_m3",
    "waste_removal_truck",
    # Removed 2026-08-07: this was a stale placeholder carried forward untouched since v2/v3
    # (its own comment said "Перенесено из rows_to_add... требуется проверка Елены" - never
    # actually reviewed). Value was 9759.08 руб./м2, ~207x too high - real rate confirmed 47
    # руб./м2 from 6 rows across 2 real projects (TRC + ARK, see DERIVED_PRICE_ROWS below).
    # Caused two absurd ~1.2М/730К руб. "Расходные материалы для установки опалубки" lines in
    # a real TRC estimate (floor_slab_1/floor_slab_2, both use this same registry_code).
    "formwork_consumables_m2",
    # Removed 2026-08-08: same class of stale never-reviewed placeholder (1000 руб./м3, no
    # comment trail beyond "требуется проверка Елены"). Real price lives on a ZONE/SUPPLIER
    # MATRIX (Елена price list, sheet "Бетон + песок", rows 15-24) - most Moscow-region zones
    # are 1300 руб./м3 (20м3 truck), zone 1 (supplier Евгений) is 1400. TRC's own real smeta
    # uses 1400 - see DERIVED_PRICE_ROWS below. This mechanism can only hold one flat price per
    # code; a real per-project zone lookup is not built. If a future project is confirmed to be
    # in a different price zone, this DERIVED_PRICE_ROWS value needs a manual one-off override,
    # not a silent switch.
    "sand_m3",
    # Removed 2026-08-08: not wrong, just ugly - "Добавлено из калькулятора MVP" (never a real
    # price-list row at all, born from some very old board-price/volume division in the MVP
    # calculator, never cleaned up). Value itself was already cross-checked correct (matches
    # ARK's real smeta almost exactly, 21500 vs 21499.997319) - just rounding off the leftover
    # decimal tail here for a clean number.
    "timber_m3",
    # Removed 2026-08-08: stale "rows_to_add/needs_review" placeholders (6400/7500/38000 руб.),
    # never sourced from anything. Real prices are zone/project-dependent (concrete price zones
    # 6150-7690, delivery 5250-7200/рейс, pump size scales with pour height - see
    # concrete_pump_32m_shift's own comment below). No zone mechanism exists yet, so replaced
    # with a straight average across the 2-3 real reference projects we have (TRC/ARK/ЮСВ) as a
    # marked-temporary value, same approach as eps50_wall_insulation_work_unit_price earlier.
    "concrete_b22_5_m3",
    "concrete_delivery_trip",
    "concrete_pump_32m_shift",
}

PRICE_CODE_NAME_OVERRIDES = {
    "edge_insulation_work_m": "Устройство утепления по наружной стороне торцов плиты, балок, перемычек",
}

DERIVED_PRICE_ROWS = [
    {
        "section": "Работы",
        "name": "Устройство и монтаж термовставок 50 мм",
        "unit": "мп",
        "min_quantity": 1,
        "price": 100,
        "price_code": "thermal_insert_50_installation_work_m",
        "comment": "Добавлено 2026-08-07: новый прайс Елены содержит одну универсальную строку 'Устройство и монтаж термовкладыша' = 100 руб./мп. В реальных эталонных сметах АРК/ТРЦ/ЮСВ работа термовставок также считается по длине с этой ставкой; для production-раздельного режима 50/100 нужен отдельный price_code, поэтому цена размножена из той же универсальной строки прайса.",
    },
    {
        "section": "Работы",
        "name": "Устройство и монтаж термовставок 100 мм",
        "unit": "мп",
        "min_quantity": 1,
        "price": 100,
        "price_code": "thermal_insert_100_installation_work_m",
        "comment": "Добавлено 2026-08-07: та же универсальная работа 'Устройство и монтаж термовкладыша' = 100 руб./мп из нового прайса Елены. Толщина слоя меняет материал, но не ставку монтажной работы по м.п.",
    },
    {
        "section": "Работы",
        "name": "Устройство и монтаж термовставок (произвольный типоразмер)",
        "unit": "мп",
        "min_quantity": 1,
        "price": 100,
        "price_code": "thermal_insert_items_installation_work_m",
        "comment": "Добавлено 2026-08-07: для режима thermal_insert_mode='items' калькулятор считает одну работу по суммарной длине всех термовставок произвольных типоразмеров. Ставка берется из той же универсальной строки нового прайса Елены 'Устройство и монтаж термовкладыша' = 100 руб./мп.",
    },
    {
        "section": "Работы",
        "name": "Материал термовставок (произвольный типоразмер)",
        "unit": "м3",
        "min_quantity": 1,
        "price": 8900,
        "price_code": "thermal_insert_item_material_m3",
        "comment": "Добавлено 2026-08-07: общий fallback для thermal_insert_items произвольного типоразмера. Новый прайс Елены дает Пеноплэкс ГЕО 100 мм (для термовставок) = 8900 руб./м3 и Пеноплэкс ГЕО 50 мм = 8800 руб./м3; для произвольных типоразмеров без отдельного толщинного price_code используем базовую цену 100 мм как более частый материал термовставок. Если в будущем появятся разные цены по eps_size, заменить на построчный price_code.",
    },
    {
        "section": "Работы",
        "name": "Штробление блоков под дополнительное усиление, армирование арматурой диаметром 10 мм",
        "unit": "мп",
        "min_quantity": 1,
        "price": 250,
        "price_code": "block_chasing_reinforcement_work_m",
        "comment": "Добавлено 2026-08-07 по сверке АРК/ТРЦ/ЮСВ: в эталонных сметах это платная работа в правой части себестоимости. Временная production-ставка 250 руб./мп до отдельного подтверждения Елены.",
    },
    {
        "section": "Работы",
        "name": "Устройство утепления по наружной стороне торцов плиты, балок, перемычек",
        "unit": "мп",
        "min_quantity": 1,
        "price": 450,
        "price_code": "lintel_edge_insulation_work_m",
        "comment": "Добавлено по ответу Елены 2026-07-30: ставка такая же, как утепление балок и торца плиты.",
    },
    {
        "section": "Работы",
        "name": "Установка аэратора кровельного PVC, А75х375",
        "unit": "шт",
        "min_quantity": 1,
        "price": 2500,
        "price_code": "roof_pvc_aerator_75x375_installation_item",
        "comment": "Добавлено по ответу Елены 2026-07-30: 587 руб. — материал аэратора, работа отдельно 2500 руб./шт.",
    },
    {
        "section": "Работы",
        "name": "Монтаж примыкания кровли из ПВХ мембраны",
        "unit": "мп",
        "min_quantity": 1,
        "price": 700,
        "price_code": "roof_pvc_membrane_abutment_work_m",
        "comment": "Добавлено 2026-08-02 по сверке эталонных смет: общее линейное примыкание кровли к парапетам/стенам/ВК считается в м.п.; отдельное примыкание к вентшахтам в шт остается optional legacy.",
    },
    {
        "section": "Работы",
        "name": "Вывоз мусора с объекта, контейнер/машина",
        "unit": "маш",
        "min_quantity": 1,
        "price": 10000,
        "price_code": "waste_removal_container_truck",
        "comment": "Перенесено из v3 и уточнено 2026-07-30: контейнер/машина идет в материальной колонке сметы.",
    },
    {
        "section": "Работы",
        "name": "Монтаж опалубки из доски для отбортовки плиты",
        "unit": "м2",
        "min_quantity": 1,
        "price": 0,
        "price_code": "timber_formwork_installation_work_m2",
        "comment": "ИСПРАВЛЕНО 2026-08-06: изначально взяла 580 руб./м2 из ТРЦ (лист 'НС 29.06.26', строка 42), но это была ЛЕВАЯ/белая (клиентская) колонка сметы. У каждой строки в реальных сметах два параллельных блока колонок - левый (белый, с накруткой для клиента) и правый (серый, настоящая внутренняя себестоимость Елены; в ТРЦ правый блок физически закрашен заливкой, в АРК/ЮСВ заливки нет, но структура та же и подтверждается арифметикой - левое/правое = Коэф-т Рентабельности листа, обычно 1.32). Правая/серая колонка для этой строки = 0 в ТРЦ. Пользователь подтвердила 2026-08-06: 0 - это реальный ответ, эта работа у Елены отдельно не тарифицируется внутри (входит в другую работу/бригадо-день), клиенту просто показывается отдельной строкой для формы.",
    },
    {
        "section": "Земляные работы",
        "name": "Перемещение песка вручную",
        "unit": "м3",
        "min_quantity": 1,
        "price": 0,
        "price_code": "sand_manual_moving_m3",
        "comment": "ИСПРАВЛЕНО 2026-08-06: изначально взяла 200 руб./м3 из ТРЦ (лист 'НС 29.06.26', строка 32), но это была ЛЕВАЯ/белая (клиентская) колонка. Правая/серая (реальная внутренняя себестоимость) колонка для этой же строки = 0. См. подробное объяснение серая/белая структура в комментарии к timber_formwork_installation_work_m2 выше - тот же класс ошибки. Пользователь подтвердила 2026-08-06: 0 верно, работа не тарифицируется отдельно внутри.",
    },
    {
        "section": "Устройство фундаментной плиты",
        "name": "Демонтаж опалубки после завершения бетонирования",
        "unit": "м2",
        "min_quantity": 1,
        "price": 0,
        "price_code": "formwork_dismantling_work_m2",
        "comment": "ИСПРАВЛЕНО 2026-08-06: изначально взяла 120 руб./м2 из АРК/ЮСВ (левая/белая колонка на нескольких строках 'Демонтаж опалубки после завершения бетонирования'). Проверила правую/серую колонку в АРК на тех же строках - везде 0. Тот же класс ошибки, что и timber_formwork_installation_work_m2/sand_manual_moving_m3 (см. их комментарии). Пользователь подтвердила 2026-08-06: 0 верно.",
    },
    {
        "section": "Schiedel",
        "name": "Доставка вентканалов, разгрузка вручную",
        "unit": "маш",
        "min_quantity": 1,
        "price": 2500,
        "price_code": "schiedel_delivery_unloading_work",
        "comment": "Добавлено 2026-08-06: этой строки нет в прайсе Елены отдельно (там доставка+разгрузка одной строкой = 2500) - ставка взята из реальной сметы ЮСВ, где разгрузка вручную идет отдельной строкой от доставки и стабильно равна 2500 руб./маш независимо от цены самой доставки (проверено в 2 разных листах ЮСВ: 2500/2500 и 15000/2500). По решению пользователя 2026-08-06 материал и работа считаются раздельно как более универсальный подход, а не совмещенной строкой как в АРК.",
    },
    {
        "section": "Schiedel",
        "name": "Доставка вентканалов (машина/транспорт, без разгрузки)",
        "unit": "маш",
        "min_quantity": 1,
        "price": 15000,
        "price_code": "schiedel_delivery_truck",
        "comment": "ИСПРАВЛЕНО 2026-08-06: было ошибочно затёрто на 2500 через устаревшее правило MANUAL_WORK_PRICE_CODES, которое мапило комбинированную строку 'Доставка+разгрузка' (2500) на этот материальный код - двойной счёт с schiedel_delivery_unloading_work. Реальное значение 15000 подтверждено на листе ЮСВ 'АЛ 06.04 КР1,КР2 (ЕЧ) (ЮВ)', строка 235 'Доставка вентканалов': материал=15000, работа=2500 - оба параллельных блока колонок (левый/белый и правый/серый) сходятся на этой цифре, без разночтений.",
    },
    {
        "section": "Устройство фундаментной плиты",
        "name": "Расходные материалы для установки опалубки (смазка; звездочки ПВХ, трубки)",
        "unit": "м2",
        "min_quantity": 1,
        "price": 47,
        "price_code": "formwork_consumables_m2",
        "comment": "ИСПРАВЛЕНО 2026-08-07: старое значение 9759.08 руб./м2 было непроверенным плейсхолдером (см. formwork_consumables_m2 в OBSOLETE_PREVIOUS_PRICE_CODES) - в 207 раз больше реальной цены, раздувало 'Расходные материалы для установки опалубки' на floor_slab_1/floor_slab_2 до ~1.2М/730К руб. на реальном проекте ТРЦ. Реальная ставка 47 руб./м2 подтверждена на 6 строках в 2 независимых проектах: ТРЦ (5556/118.218=47.0, 1404/29.875=47.0, 4307/91.6292=47.0, 1105/23.5=47.0) и АРК (15098/321.24=47.0, 4573/97.3=47.0) - везде площадь монтажа опалубки берётся из соседней строки того же раздела сметы.",
    },
    {
        "section": "Земляные работы",
        "name": "Песок строительный",
        "unit": "м3",
        "min_quantity": 1,
        "price": 1400,
        "price_code": "sand_m3",
        "comment": "ИСПРАВЛЕНО 2026-08-08: старое значение 1000 руб./м3 было непроверенным плейсхолдером (то же 'требуется проверка Елены' что и у formwork_consumables_m2). Реальная цена - в прайсе Елены на листе 'Бетон + песок' (строки 15-24), это ЗОНОВАЯ матрица по Москве и поставщикам: большинство зон (Евгений зоны 2-4, Алексей, Олег) = 1300 руб./м3 (20м3 машина), зона 1 (Евгений) = 1400 руб./м3. Взяла 1400 - это ровно та цена, что в реальной смете ТРЦ (строка 'Песок строительный' = 1400 руб./м3). Зонозависимость этот механизм не поддерживает - если у нового проекта другая зона/поставщик, значение нужно поправить здесь вручную, автоматически не подтянется.",
    },
    {
        "section": "Земляные работы",
        "name": "Экскаватор-погрузчик JCB, аренда/материал (без работы)",
        "unit": "смена",
        "min_quantity": 1,
        "price": 26000,
        "price_code": "excavator_jcb_shift_material",
        "comment": "Добавлено 2026-08-07: у 'Экскаватор-погрузчик JCB...' на первом листе прайса Елены только ОДНА цена (3500 руб./смена) - это цена РАБОТЫ (оператор), уже верно замаплена на excavator_jcb_shift. Материальной/арендной части (стоимость самого экскаватора) на первом листе нет вообще. Реальная ставка 26000 руб./смена подтверждена в 8+ независимых снимках реальных смет за несколько месяцев: ТРЦ (2026-06-29), АРК (2026-04-27), и ЮСВ на 8 разных листах/датах с 2026-02-19 по 2026-04-06 - везде одна и та же пара 26000 материал / 3500 работа для строки 'Механизированная разработка грунта, Экскаватор JCB'. До этой правки контракт (excavator_material_unit_price) ошибочно указывал на тот же registry_code, что и работа (excavator_jcb_shift), поэтому материал всегда дублировал цену работы (3500 вместо 26000).",
    },
    {
        "section": "Гидроизоляция",
        "name": "Утепление стен ЭППС 50мм, работа (ВРЕМЕННО 0 - ждет ответа Елены)",
        "unit": "м2",
        "min_quantity": 1,
        "price": 0,
        "price_code": "eps_wall_insulation_work_50_m2",
        "comment": "ВРЕМЕННОЕ РЕШЕНИЕ 2026-08-06, требует уточнения у Елены: в смете ТРЦ утепление стен 50мм+100мм считается ОДНОЙ работой (680 руб./м2 на оба слоя сразу, лист 'НС 29.06.26' строка 73), а не двумя отдельными монтажными работами как в наших калькуляторах. АРК и ЮСВ вообще не имеют слоя 50мм на стенах (только 100мм), сверить не с чем. Поставлено 0, чтобы не задвоить деньги поверх уже существующей (возможно тоже неточной - см. eps_wall_insulation_work_m2) цены на монтаж 100мм-слоя. НЕ ИСПОЛЬЗОВАТЬ КАК ОКОНЧАТЕЛЬНОЕ ЗНАЧЕНИЕ - заменить, когда Елена ответит на вопрос от 2026-08-06 (один проход маляра на оба слоя или два отдельных).",
    },
    {
        "section": "Кровля",
        "name": "Воронка кровельная внутренняя с обогревом (материал)",
        "unit": "шт",
        "min_quantity": 1,
        "price": 5000,
        "price_code": "roof_internal_drain_with_heating_item",
        "comment": "Добавлено 2026-08-06: материалы не читаются автоматическим скриптом (только первый лист 'Прайс по видам работ', лист с материалами 'Утеплителипленкимастика (Олег)' сюда не входит) - цена 5000 руб./шт взята оттуда напрямую, строка 32 ('Воронка кровельная (с обжимным мет. фланцем с обогревом 110х450мм)'). До этой правки код случайно получал ту же цену через устаревшее (неверное) сопоставление с работой монтажа - см. исправление в MANUAL_WORK_PRICE_CODES выше. Цена совпала случайно (обе 5000), поэтому баг был незаметен.",
    },
    {
        "section": "Кровля",
        "name": "Воронка парапетная с листвоуловителем (материал)",
        "unit": "шт",
        "min_quantity": 1,
        "price": 6001.23,
        "price_code": "roof_parapet_drain_item",
        "comment": "Добавлено 2026-08-06: та же причина, что и у внутренней воронки - материалы не читаются автоматическим скриптом, цена взята напрямую из листа 'Утеплителипленкимастика (Олег)', строка 31 ('Воронка парапетная с листвоуловителем и отводом VC-PVC, для ПВХ мембран, 100х100х650') = 6001.23 руб./шт. До этой правки код ошибочно получал цену РАБОТЫ монтажа (5000 руб.) через устаревшее сопоставление в MANUAL_WORK_PRICE_CODES - реальная цена материала выше, это была настоящая ошибка в деньгах, не совпадение.",
    },
    {
        "section": "Кровля",
        "name": "ПВХ мембрана Logicroof V-GR 1,5 мм (эксплуатируемая кровля)",
        "unit": "рул",
        "min_quantity": 1,
        "price": 1146.91,
        "price_code": "roof_pvc_membrane_logicroof_vgr_1_5mm_gray_roll",
        "comment": "Добавлено 2026-08-06: такая же цена, как у обычной V-RP мембраны (1146.91 руб./рул, прайс Елены лист 'Утеплителипленкимастика (Олег)' строка 29 - в прайсе V-GR и V-RP указаны с одинаковой ценой). Пользователь вписывал эту цену вручную в гугл-таблицу 2026-08-06 после того, как для ТРЦ впервые понадобилась V-GR (эксплуатируемая зона кровли) - перенесено в реестр, чтобы не терялось при пересборке.",
    },
    {
        "section": "Schiedel",
        "name": "Вентиляционный канал одноходовой VENT 20/25, 0.33 пм",
        "unit": "шт",
        "min_quantity": 1,
        "price": 268.4,
        "price_code": "schiedel_vent_channel_1x_item",
        "comment": "Добавлено 2026-08-06: лист 'Блокпаротермшидель (Алексей)' прайса Елены, строка 27. h=330мм, габариты 200х250мм, вес 13кг, 63 шт на поддоне.",
    },
    {
        "section": "Schiedel",
        "name": "Вентиляционный канал четырехходовой VENT 36/50, 0.33 пм",
        "unit": "шт",
        "min_quantity": 1,
        "price": 902.8,
        "price_code": "schiedel_vent_channel_4x_item",
        "comment": "Добавлено 2026-08-06: лист 'Блокпаротермшидель (Алексей)' прайса Елены, строка 30. h=330мм, габариты 500х360мм, вес 38кг, 16 шт на поддоне.",
    },
    {
        "section": "Schiedel",
        "name": "Вентканал CVENT 26х26 (канал + вентблок-спутник)",
        "unit": "шт",
        "min_quantity": 1,
        "price": 2663.2,
        "price_code": "schiedel_vent_channel_cvent_item",
        "comment": "Добавлено 2026-08-06: было два кандидата на листе 'Блокпаротермшидель (Алексей)' прайса Елены - строка 33 'Вентблок CVENT 26*26 со спутником, 0.327' = 961.2 руб./шт и строка 34 'Вентиляционный канал Schiedel CVENT 500Х360Х327 26Х26' = 1702 руб./шт. Пользователь подтвердила 2026-08-06: это не два кандидата на один и тот же продукт, а два РАЗНЫХ компонента одного полного CVENT-узла (канал-модуль + его блок-спутник той же глубины 327мм) - берутся оба, цена = сумма 961.2 + 1702 = 2663.2 руб./шт. Калькулятор (CHANNEL_TYPE_SPECS['cvent'] в schiedel_vent_channels_calculator/calculator.py) хранит только ОДНУ цену на product_type=cvent, как и для 1x/2x/3x/4x - поэтому сумма зашита здесь одним числом, а не как два отдельных price_code; calculator.py не менялся. CVENT остаётся редкой строкой (Elena 2026-07-30: 'бывает достаточно редко'), ни один реальный проект её пока не использовал.",
    },
    {
        "section": "Пиломатериал",
        "name": "Фанера 1,52х1,52*18 (для опалубки монолитных перемычек)",
        "unit": "шт",
        "min_quantity": 1,
        "price": 1400,
        "price_code": "lintel_formwork_plywood_sheet",
        "comment": "Добавлено 2026-08-06: тот же лист фанеры, что и plywood_1520x1520_18mm_sheet (прайс Елены, лист 'Пиломатериалы (Олег)', строка 23) - отдельный price_code нужен калькулятору перемычек load_bearing_walls_lintels_calculator.py.",
    },
    {
        "section": "Земляные работы",
        "name": "Пиломатериал обрезной для устройства опалубки ГОСТ",
        "unit": "м3",
        "min_quantity": 1,
        "price": 21500,
        "price_code": "timber_m3",
        "comment": "ИСПРАВЛЕНО 2026-08-08: старое значение 21499.997319 не было ошибкой по сути (подтверждено по смете АРК - там та же позиция 21500 руб./м3), но само число было 'Добавлено из калькулятора MVP' - какой-то старый расчёт цена_доски/объём_доски, никогда не бывшее настоящей строкой прайса, с некрасивым хвостом из 6 знаков после запятой. Округлила до чистых 21500 - число то же самое по сути, просто без мусора.",
    },
    {
        "section": "Устройство фундаментной плиты",
        "name": "Бетон марки В22,5 (М300) (ВРЕМЕННО - среднее по проектам, до внедрения зон)",
        "unit": "м3",
        "min_quantity": 1,
        "price": 6275,
        "price_code": "concrete_b22_5_m3",
        "comment": "ВРЕМЕННО 2026-08-08: цена бетона зонозависима (актуальный прайс Елены даёт зоны 6150-7690 руб./м3), готового зонного механизма пока нет. Взято простое среднее по 2 реальным проектам: ТРЦ=6600, АРК=5950 -> (6600+5950)/2=6275. Заменить, когда появится зонная логика или прямое указание, в какой зоне проект.",
    },
    {
        "section": "Устройство фундаментной плиты",
        "name": "Доставка бетона до объекта (ВРЕМЕННО - среднее по проектам)",
        "unit": "рейс",
        "min_quantity": 1,
        "price": 6225,
        "price_code": "concrete_delivery_trip",
        "comment": "ВРЕМЕННО 2026-08-08: та же зонозависимость, что у concrete_b22_5_m3. Актуальный прайс Елены даёт цену доставки ЗА М3 (не за рейс, единицы не совпадают с этим price_code), поэтому усреднила напрямую реальные рейсовые цены из 2 проектов: ТРЦ=7200, АРК=5250 -> (7200+5250)/2=6225. Заменить, когда появится зонная логика.",
    },
    {
        "section": "Устройство фундаментной плиты",
        "name": "Работа бетононасоса (нижний уровень: фундамент/1-й этаж) (ВРЕМЕННО - среднее по проектам)",
        "unit": "смена",
        "min_quantity": 1,
        "price": 37500,
        "price_code": "concrete_pump_32m_shift",
        "comment": "ВРЕМЕННО 2026-08-08: нашли реальную систему - размер/цена насоса зависит от высоты заливки, не от проекта. В ТРЦ и АРК фундамент+1й этаж используют насос МЕНЬШЕГО размера, чем 2й этаж+ (см. concrete_pump_32m_shift_floor_slab_2). ЮСВ не различает уровни (один насос везде) - закономерность реальная, но не универсальная. Нижний уровень: ТРЦ='28м'=36500, АРК='36м'(фундамент)=38000, ЮСВ='32м+гаситель'=38000 -> среднее (36500+38000+38000)/3=37500. Раньше этот же код делился между foundation_slab/floor_slab_1/floor_slab_2 - теперь у каждого раздела свой code, см. floor_slab_1/floor_slab_2 контракты.",
    },
    {
        "section": "Ж/Б монолитная плита перекрытия 1-го этажа",
        "name": "Работа бетононасоса (нижний уровень: фундамент/1-й этаж) (ВРЕМЕННО - среднее по проектам)",
        "unit": "смена",
        "min_quantity": 1,
        "price": 37500,
        "price_code": "concrete_pump_32m_shift_floor_slab_1",
        "comment": "ВРЕМЕННО 2026-08-08: плита 1-го этажа заливается с той же высоты, что и фундамент - в ТРЦ и АРК используется насос того же (меньшего) размера, что для фундамента (см. concrete_pump_32m_shift). Та же группа проектных значений: ТРЦ='28м'=36500, АРК='36м'=38000, ЮСВ='32м+гаситель'=38000 -> среднее 37500. Новый registry_code, добавлен вместе с fixed 2026-08-08 разделением concrete_pump_32m_shift по разделам (было общее на все 3 раздела: foundation_slab/floor_slab_1/floor_slab_2).",
    },
    {
        "section": "Ж/Б монолитная плита перекрытия 2-го этажа",
        "name": "Работа бетононасоса (верхний уровень: 2-й этаж+) (ВРЕМЕННО - среднее по проектам)",
        "unit": "смена",
        "min_quantity": 1,
        "price": 39833,
        "price_code": "concrete_pump_32m_shift_floor_slab_2",
        "comment": "ВРЕМЕННО 2026-08-08: верхний уровень (2-й этаж и выше) - реально используется насос БОЛЬШЕГО размера, чем на фундаменте/1-м этаже (см. concrete_pump_32m_shift). ТРЦ='36м'=39500, АРК='40м'(плита перекрытия)=42000, ЮСВ='32м+гаситель'=38000 (не различает уровни) -> среднее (39500+42000+38000)/3=39833.33, округлено до 39833. Новый registry_code, добавлен вместе с fixed 2026-08-08 разделением concrete_pump_32m_shift по разделам (было общее на все 3 раздела).",
    },
    {
        "section": "Ж/Б монолитная плита перекрытия 1-го этажа",
        "name": "Пиломатериал обрезной для устройства опалубки ГОСТ (для опалубки монолитных перемычек)",
        "unit": "м3",
        "min_quantity": 1,
        "price": 21500,
        "price_code": "lintel_formwork_timber_m3",
        "comment": "Добавлено 2026-08-06: та же цена, что и timber_m3 - отдельный price_code нужен калькулятору перемычек load_bearing_walls_lintels_calculator.py.",
    },
]

SORT_BLOCK_ORDER = {
    "general": 0,
    "earthworks": 1,
    "concrete_sand": 2,
    "formwork_timber": 3,
    "insulation_thermal": 4,
    "waterproofing": 5,
    "walls_lintels": 6,
    "flat_roof": 7,
    "schiedel": 8,
    "other": 9,
}

SECTION_SORT_BLOCK = {
    "бетон": "concrete_sand",
    "песок": "concrete_sand",
    "пиломатериал": "formwork_timber",
    "арматура": "walls_lintels",
    "газобетон": "walls_lintels",
    "поротерм": "walls_lintels",
    "несущие стены и перемычки": "walls_lintels",
    "плоская кровля": "flat_roof",
    "кровельное покрытие дома": "flat_roof",
    "schiedel": "schiedel",
    "земляные работы": "earthworks",
    "устройство фундаментной плиты": "formwork_timber",
    "ж/б монолитная плита перекрытия 1-го этажа": "formwork_timber",
    "фундамент": "insulation_thermal",
}


def row_sort_block(row: list[Any]) -> str:
    section = normalize_text(row[REGISTRY_HEADERS.index("Раздел")])
    name = normalize_text(row[REGISTRY_HEADERS.index("Наименование")])
    code = normalize_text(row[REGISTRY_HEADERS.index("price_code")])

    text = f"{section} {name} {code}"
    if section in SECTION_SORT_BLOCK and section != "работы":
        return SECTION_SORT_BLOCK[section]
    if code in {
        "axis_marking_shift",
        "construction_camp_setup_item",
        "site_cabin_connection_item",
        "waste_removal_container_truck",
        "waste_removal_loading_work_truck",
    } or "строительный городок" in text or "бытовк" in text or "вынос осей" in text:
        return "general"
    if any(term in text for term in ("excavator", "землян", "грунт", "котлован", "геотекст", "песок", "коммуникац")):
        return "earthworks"
    if any(term in text for term in ("газобетон", "поротерм", "кладк", "перемыч", "парапет", "блок", "арматур", "lintel", "rebar")):
        return "walls_lintels"
    if any(term in text for term in ("бетон", "concrete", "насос", "миксер", "перенос, подъем бетона")):
        return "concrete_sand"
    if any(term in text for term in ("опалуб", "фанер", "пиломатериал", "timber", "plywood", "formwork")):
        return "formwork_timber"
    if any(term in text for term in ("утепл", "эппс", "eps", "термов", "пеноплэкс", "пеноплекс")):
        return "insulation_thermal"
    if any(term in text for term in ("гидроизоляц", "мастик", "битум", "planter", "waterproof")):
        return "waterproofing"
    if any(term in text for term in ("кров", "roof", "мембран", "воронк", "аэратор", "водосток", "пвх")):
        return "flat_roof"
    if "schiedel" in text or "вентканал" in text:
        return "schiedel"
    return SECTION_SORT_BLOCK.get(section, "other")


def sort_registry_rows(rows: list[list[Any]]) -> list[list[Any]]:
    indexed_rows = list(enumerate(rows))
    return [
        row
        for original_index, row in sorted(
            indexed_rows,
            key=lambda item: (
                SORT_BLOCK_ORDER[row_sort_block(item[1])],
                clean_text(item[1][REGISTRY_HEADERS.index("Раздел")]).lower(),
                item[0],
            ),
        )
    ]


def headers(ws) -> dict[str, int]:
    return {
        clean_text(ws.cell(1, column).value): column
        for column in range(1, ws.max_column + 1)
        if clean_text(ws.cell(1, column).value)
    }


def previous_code_index(path: Path) -> dict[tuple[str, str], set[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["price_registry"]
    header_map = headers(ws)
    index: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, header_map["Наименование"]).value
        unit = ws.cell(row, header_map["Ед. изм."]).value
        code = clean_text(ws.cell(row, header_map["price_code"]).value)
        if name and unit and code:
            index[row_key(name, unit)].add(code)
    return index


def previous_registry_rows(path: Path) -> list[list[Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["price_registry"]
    header_map = headers(ws)
    rows: list[list[Any]] = []
    for row in range(2, ws.max_row + 1):
        if not clean_text(ws.cell(row, header_map["Наименование"]).value):
            continue
        rows.append([ws.cell(row, header_map[name]).value for name in REGISTRY_HEADERS])
    return rows


def deduplicate_rows_by_price_code(rows: list[list[Any]]) -> tuple[list[list[Any]], list[dict[str, Any]]]:
    code_index = REGISTRY_HEADERS.index("price_code")
    latest_by_code: dict[str, list[Any]] = {}
    first_position_by_code: dict[str, int] = {}
    uncoded_rows: list[list[Any]] = []
    duplicates: list[dict[str, Any]] = []

    for row in rows:
        code = clean_text(row[code_index])
        if not code:
            uncoded_rows.append(row)
            continue
        if code in latest_by_code:
            duplicates.append(
                {
                    "price_code": code,
                    "old_name": latest_by_code[code][REGISTRY_HEADERS.index("Наименование")],
                    "new_name": row[REGISTRY_HEADERS.index("Наименование")],
                }
            )
        else:
            first_position_by_code[code] = len(first_position_by_code)
        latest_by_code[code] = row

    coded_rows = sorted(
        latest_by_code.values(),
        key=lambda row: first_position_by_code[clean_text(row[code_index])],
    )
    return uncoded_rows + coded_rows, duplicates


def extract_work_price_rows(source_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(source_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[dict[str, Any]] = []
    for row in range(3, ws.max_row + 1):
        name = ws.cell(row, 2).value
        unit = ws.cell(row, 3).value
        quantity = ws.cell(row, 4).value
        price = ws.cell(row, 5).value
        if not clean_text(name):
            continue
        rows.append(
            {
                "source_row": row,
                "section": "Работы",
                "name": clean_text(name),
                "unit": clean_text(unit),
                "min_quantity": quantity,
                "price": price,
            }
        )
    return rows


def build_registry_workbook(
    source_path: Path,
    previous_registry_path: Path,
    output_path: Path,
    archive_source_path: Path,
) -> dict[str, Any]:
    archive_source_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, archive_source_path)

    previous_codes = previous_code_index(previous_registry_path)
    source_rows = extract_work_price_rows(source_path)
    registry_rows = [
        row
        for row in previous_registry_rows(previous_registry_path)
        if clean_text(row[REGISTRY_HEADERS.index("price_code")]) not in OBSOLETE_PREVIOUS_PRICE_CODES
    ]
    merged_rows, removed_existing_duplicates = deduplicate_rows_by_price_code(registry_rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "price_registry"
    ws.append(REGISTRY_HEADERS)

    matched = 0
    manual = 0
    ambiguous: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    ignored: list[dict[str, Any]] = []
    replaced_codes: set[str] = set()
    for item in source_rows:
        key = row_key(item["name"], item["unit"])
        if key in IGNORED_WORK_PRICE_ROWS:
            ignored.append(item)
            continue
        manual_code = MANUAL_WORK_PRICE_CODES.get(key)
        codes = sorted(previous_codes.get(key, set()))
        price_code = ""
        comment_parts = [f"Источник: первый лист нового прайса, строка {item['source_row']}."]
        if item["min_quantity"] not in (None, ""):
            comment_parts.append(f"Исходное кол-во: {item['min_quantity']}.")
        if manual_code:
            price_code = manual_code
            manual += 1
            comment_parts.append("price_code задан ручной картой normalize_elena_price_registry.py.")
        elif len(codes) == 1:
            price_code = codes[0]
            matched += 1
        elif len(codes) > 1:
            ambiguous.append({**item, "codes": codes})
            comment_parts.append(f"Неоднозначный price_code из v3: {', '.join(codes)}.")
        else:
            unmatched.append(item)
            comment_parts.append("price_code не найден автоматически: нужно сопоставить вручную.")

        new_row = [
            item["section"],
            PRICE_CODE_NAME_OVERRIDES.get(price_code, item["name"]),
            item["unit"],
            item["min_quantity"],
            item["price"],
            date.today().isoformat(),
            " ".join(comment_parts),
            price_code,
        ]
        if price_code:
            replaced = False
            for index, existing_row in enumerate(merged_rows):
                if clean_text(existing_row[REGISTRY_HEADERS.index("price_code")]) == price_code:
                    merged_rows[index] = new_row
                    replaced = True
                    replaced_codes.add(price_code)
                    break
            if not replaced:
                merged_rows.append(new_row)
                replaced_codes.add(price_code)
        else:
            merged_rows.append(new_row)

    for item in DERIVED_PRICE_ROWS:
        new_row = [
            item["section"],
            item["name"],
            item["unit"],
            item["min_quantity"],
            item["price"],
            date.today().isoformat(),
            item["comment"],
            item["price_code"],
        ]
        replaced = False
        for index, existing_row in enumerate(merged_rows):
            if clean_text(existing_row[REGISTRY_HEADERS.index("price_code")]) == item["price_code"]:
                merged_rows[index] = new_row
                replaced = True
                break
        if not replaced:
            merged_rows.append(new_row)
        replaced_codes.add(item["price_code"])

    final_rows, removed_final_duplicates = deduplicate_rows_by_price_code(merged_rows)
    final_rows = sort_registry_rows(final_rows)
    for row in final_rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    ws.freeze_panes = "A2"
    widths = [22, 78, 14, 24, 14, 18, 80, 36]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return {
        "source_path": source_path,
        "archive_source_path": archive_source_path,
        "output_path": output_path,
        "previous_registry_rows": len(registry_rows),
        "previous_unique_rows": len(merged_rows),
        "new_work_rows": len(source_rows),
        "matched": matched,
        "manual": manual,
        "replaced_codes": sorted(replaced_codes),
        "removed_existing_duplicates": removed_existing_duplicates,
        "removed_final_duplicates": removed_final_duplicates,
        "final_rows": len(final_rows),
        "ambiguous": ambiguous,
        "unmatched": unmatched,
        "ignored": ignored,
    }


def write_report(summary: dict[str, Any], report_path: Path) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Price Registry v4 Mapping Report",
        "",
        f"- Source file: `{summary['source_path']}`",
        f"- Archived source copy: `{summary['archive_source_path']}`",
        f"- Output registry: `{summary['output_path']}`",
        f"- Previous v3 registry rows copied: `{summary['previous_registry_rows']}`",
        f"- Previous duplicate price_code rows removed: `{len(summary['removed_existing_duplicates'])}`",
        f"- Work price rows from first sheet: `{summary['new_work_rows']}`",
        f"- Matched price_code from v3: `{summary['matched']}`",
        f"- Manual work price_code mappings: `{summary['manual']}`",
        f"- Existing price_code rows replaced by new price: `{len(summary['replaced_codes'])}`",
        f"- Final registry rows: `{summary['final_rows']}`",
        f"- Ambiguous matches: `{len(summary['ambiguous'])}`",
        f"- Unmatched rows: `{len(summary['unmatched'])}`",
        f"- Ignored source rows: `{len(summary['ignored'])}`",
        "",
    ]
    if summary["replaced_codes"]:
        lines.extend(["## Replaced price_code Rows", ", ".join(f"`{code}`" for code in summary["replaced_codes"]), ""])
    if summary["ambiguous"]:
        lines.extend(["## Ambiguous", "| source row | name | unit | candidate codes |", "|---:|---|---|---|"])
        for row in summary["ambiguous"]:
            lines.append(
                f"| {row['source_row']} | {row['name']} | {row['unit']} | {', '.join(row['codes'])} |"
            )
        lines.append("")
    if summary["unmatched"]:
        lines.extend(["## Unmatched", "| source row | name | unit | price |", "|---:|---|---|---:|"])
        for row in summary["unmatched"]:
            lines.append(f"| {row['source_row']} | {row['name']} | {row['unit']} | {row['price']} |")
        lines.append("")
    if summary["ignored"]:
        lines.extend(["## Ignored Source Rows", "| source row | name | unit | price | reason |", "|---:|---|---|---:|---|"])
        for row in summary["ignored"]:
            lines.append(
                f"| {row['source_row']} | {row['name']} | {row['unit']} | {row['price']} | "
                "Елена 2026-07-30: балки перекрытий считаем единой строкой по м.п. "
                "`beam_concrete_placing_work_m`; строку `м3 > 250мм` не используем. |"
            )
        lines.append("")
    report_path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--previous-registry", type=Path, default=DEFAULT_PREVIOUS_REGISTRY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--archive-source", type=Path, default=DEFAULT_ARCHIVE_SOURCE)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()

    summary = build_registry_workbook(
        args.source,
        args.previous_registry,
        args.output,
        args.archive_source,
    )
    write_report(summary, args.report)
    print(
        "price_registry v4:",
        f"previous_rows={summary['previous_registry_rows']}",
        f"final_rows={summary['final_rows']}",
        f"new_work_rows={summary['new_work_rows']}",
        f"matched={summary['matched']}",
        f"manual={summary['manual']}",
        f"replaced={len(summary['replaced_codes'])}",
        f"ambiguous={len(summary['ambiguous'])}",
        f"unmatched={len(summary['unmatched'])}",
        f"ignored={len(summary['ignored'])}",
        f"output={summary['output_path']}",
    )


if __name__ == "__main__":
    main()
