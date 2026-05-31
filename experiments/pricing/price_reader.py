from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_PRICE_REGISTRY_COLUMNS = {
    "Раздел",
    "Наименование",
    "Ед. изм.",
    "Цена",
    "Дата обновления",
    "Комментарий",
    "price_code",
}

PROJECT_OVERRIDES_COLUMNS = {
    "price_code",
    "Наименование",
    "Ед. изм.",
    "Цена для проекта",
    "Комментарий",
}


def _to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return None


def _headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, column).value).strip(): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(1, column).value is not None
    }


def _cell(ws, row: int, headers: dict[str, int], name: str) -> Any:
    column = headers.get(name)
    if column is None:
        return None
    return ws.cell(row, column).value


def load_price_registry(path: str | Path) -> dict[str, dict[str, Any]]:
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    if "price_registry" not in wb.sheetnames:
        raise ValueError(f"Sheet price_registry not found in {workbook_path}")

    ws = wb["price_registry"]
    headers = _headers(ws)
    missing = REQUIRED_PRICE_REGISTRY_COLUMNS - set(headers)
    if missing:
        raise ValueError(f"Missing required price_registry columns: {sorted(missing)}")

    registry: dict[str, dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        price_code = _cell(ws, row, headers, "price_code")
        if price_code is None or str(price_code).strip() == "":
            continue
        price = _to_decimal(_cell(ws, row, headers, "Цена"))
        if price is None:
            continue
        registry[str(price_code).strip()] = {
            "price": price,
            "name": _cell(ws, row, headers, "Наименование") or "",
            "unit": _cell(ws, row, headers, "Ед. изм.") or "",
            "section": _cell(ws, row, headers, "Раздел") or "",
            "source": "price_registry",
        }
    return registry


def load_project_price_overrides(path: str | Path) -> dict[str, dict[str, Any]]:
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    if "project_price_overrides" not in wb.sheetnames:
        return {}

    ws = wb["project_price_overrides"]
    headers = _headers(ws)
    missing = PROJECT_OVERRIDES_COLUMNS - set(headers)
    if missing:
        raise ValueError(f"Missing required project_price_overrides columns: {sorted(missing)}")

    overrides: dict[str, dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        price_code = _cell(ws, row, headers, "price_code")
        if price_code is None or str(price_code).strip() == "":
            continue
        price = _to_decimal(_cell(ws, row, headers, "Цена для проекта"))
        if price is None:
            continue
        overrides[str(price_code).strip()] = {
            "price": price,
            "name": _cell(ws, row, headers, "Наименование") or "",
            "unit": _cell(ws, row, headers, "Ед. изм.") or "",
            "source": "project_price_overrides",
        }
    return overrides


def resolve_price(
    price_code: str | None,
    fallback_price: Any,
    registry: dict[str, dict[str, Any]],
    overrides: dict[str, dict[str, Any]] | None = None,
) -> dict[str, Any]:
    fallback = _to_decimal(fallback_price)
    clean_code = str(price_code).strip() if price_code else ""
    overrides = overrides or {}

    if not clean_code:
        return {
            "price_code": clean_code,
            "price": fallback,
            "found": False,
            "source": "fallback_input",
            "warning": "price_code is empty, fallback input price used"
            if fallback is not None
            else "price_code is empty and fallback input price is missing",
        }

    if clean_code in overrides:
        return {
            "price_code": clean_code,
            "price": overrides[clean_code]["price"],
            "found": True,
            "source": "project_price_overrides",
            "warning": None,
        }

    if clean_code in registry:
        return {
            "price_code": clean_code,
            "price": registry[clean_code]["price"],
            "found": True,
            "source": "price_registry",
            "warning": None,
        }

    warning = "price_code not found in price_registry, fallback input price used"
    if fallback is None:
        warning = "price_code not found in price_registry and fallback input price is missing"

    return {
        "price_code": clean_code,
        "price": fallback,
        "found": False,
        "source": "fallback_input",
        "warning": warning,
    }
