from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_REGISTRY_PATH = ROOT / "output" / "price_registry_filled_v3.xlsx"
DEFAULT_REQUIRED_PATH = (
    ROOT / "output" / "required_price_codes_v3.csv"
    if (ROOT / "output" / "required_price_codes_v3.csv").exists()
    else ROOT / "output" / "required_price_codes_v2.csv"
)
DEFAULT_OUTPUT_PATH = Path(__file__).resolve().parent / "output" / "required_codes_coverage_report.md"


def headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, column).value).strip(): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(1, column).value is not None
    }


def collect_codes_from_sheet(wb, sheet_name: str) -> set[str]:
    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    header_map = headers(ws)
    if "price_code" not in header_map:
        return set()
    codes = set()
    col = header_map["price_code"]
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, col).value
        if value is not None and str(value).strip():
            codes.add(str(value).strip())
    return codes


def collect_required_codes(required_path: Path) -> dict[str, dict[str, Any]]:
    with required_path.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = row.get("price_code", "").strip()
        if code and code not in result:
            result[code] = row
    return result


def check_required_codes_against_registry(
    required_path: Path = DEFAULT_REQUIRED_PATH,
    registry_path: Path = DEFAULT_REGISTRY_PATH,
    output_path: Path = DEFAULT_OUTPUT_PATH,
) -> dict[str, Any]:
    required = collect_required_codes(required_path)
    wb = load_workbook(registry_path, read_only=True, data_only=True)

    registry_codes = collect_codes_from_sheet(wb, "price_registry")
    rows_to_add_codes = collect_codes_from_sheet(wb, "rows_to_add")
    required_codes = set(required)

    found_registry = required_codes & registry_codes
    found_rows_to_add = (required_codes - found_registry) & rows_to_add_codes
    missing_completely = required_codes - registry_codes - rows_to_add_codes

    result = {
        "required_path": str(required_path),
        "registry_path": str(registry_path),
        "required_unique_price_code_count": len(required_codes),
        "found_in_price_registry": sorted(found_registry),
        "found_in_rows_to_add": sorted(found_rows_to_add),
        "missing_completely": sorted(missing_completely),
        "fallback_needed": sorted(found_rows_to_add | missing_completely),
        "required_by_code": required,
    }
    write_report(result, output_path)
    return result


def write_report(result: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Required Codes Coverage Report",
        "",
        f"- Required file: `{result['required_path']}`",
        f"- Registry file: `{result['registry_path']}`",
        f"- Required unique price_code count: `{result['required_unique_price_code_count']}`",
        f"- Found in price_registry: `{len(result['found_in_price_registry'])}`",
        f"- Found in rows_to_add: `{len(result['found_in_rows_to_add'])}`",
        f"- Missing completely: `{len(result['missing_completely'])}`",
        f"- Fallback needed: `{len(result['fallback_needed'])}`",
        "",
        "## Found In Price Registry",
        ", ".join(f"`{code}`" for code in result["found_in_price_registry"]) or "none",
        "",
        "## Found In rows_to_add",
        ", ".join(f"`{code}`" for code in result["found_in_rows_to_add"]) or "none",
        "",
        "## Missing Completely",
        ", ".join(f"`{code}`" for code in result["missing_completely"]) or "none",
        "",
    ]
    output_path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    summary = check_required_codes_against_registry()
    print(
        "required code coverage:",
        f"required={summary['required_unique_price_code_count']}",
        f"registry={len(summary['found_in_price_registry'])}",
        f"rows_to_add={len(summary['found_in_rows_to_add'])}",
        f"missing={len(summary['missing_completely'])}",
        f"fallback={len(summary['fallback_needed'])}",
    )
