from __future__ import annotations

from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_REGISTRY_PATH = ROOT / "output" / "price_registry_filled_v3.xlsx"
DEFAULT_OUTPUT_PATH = Path(__file__).resolve().parent / "output" / "price_registry_validation_report.md"

REQUIRED_COLUMNS = [
    "Раздел",
    "Наименование",
    "Ед. изм.",
    "Цена",
    "Дата обновления",
    "Комментарий",
    "price_code",
]


def to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return None


def headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, column).value).strip(): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(1, column).value is not None
    }


def validate_price_registry(
    registry_path: Path = DEFAULT_REGISTRY_PATH,
    output_path: Path = DEFAULT_OUTPUT_PATH,
) -> dict[str, Any]:
    wb = load_workbook(registry_path, read_only=True, data_only=True)
    has_sheet = "price_registry" in wb.sheetnames

    result: dict[str, Any] = {
        "registry_path": str(registry_path),
        "has_price_registry_sheet": has_sheet,
        "missing_columns": [],
        "total_rows": 0,
        "filled_price_code_rows": 0,
        "empty_price_code_rows": 0,
        "duplicate_price_codes": {},
        "filled_code_empty_price_rows": [],
        "non_numeric_price_rows": [],
    }

    if not has_sheet:
        write_report(result, output_path)
        return result

    ws = wb["price_registry"]
    header_map = headers(ws)
    missing_columns = [name for name in REQUIRED_COLUMNS if name not in header_map]
    result["missing_columns"] = missing_columns
    if missing_columns:
        write_report(result, output_path)
        return result

    code_counter: Counter[str] = Counter()
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, header_map["Наименование"]).value
        if name is None or str(name).strip() == "":
            continue

        result["total_rows"] += 1
        price_code_value = ws.cell(row, header_map["price_code"]).value
        price_code = str(price_code_value).strip() if price_code_value else ""
        price_value = ws.cell(row, header_map["Цена"]).value
        price_decimal = to_decimal(price_value)

        if price_code:
            result["filled_price_code_rows"] += 1
            code_counter[price_code] += 1
            if price_value is None or price_value == "":
                result["filled_code_empty_price_rows"].append(row)
            elif price_decimal is None:
                result["non_numeric_price_rows"].append(row)
        else:
            result["empty_price_code_rows"] += 1

    result["duplicate_price_codes"] = {
        code: count for code, count in sorted(code_counter.items()) if count > 1
    }
    write_report(result, output_path)
    return result


def write_report(result: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Price Registry Validation Report",
        "",
        f"- Registry file: `{result['registry_path']}`",
        f"- Has `price_registry` sheet: `{result['has_price_registry_sheet']}`",
        f"- Total rows: `{result['total_rows']}`",
        f"- Filled price_code rows: `{result['filled_price_code_rows']}`",
        f"- Empty price_code rows: `{result['empty_price_code_rows']}`",
        f"- Missing columns: `{', '.join(result['missing_columns']) or 'none'}`",
        f"- Duplicate price_code count: `{len(result['duplicate_price_codes'])}`",
        f"- Filled price_code rows with empty price: `{len(result['filled_code_empty_price_rows'])}`",
        f"- Rows with non-numeric price: `{len(result['non_numeric_price_rows'])}`",
        "",
    ]
    if result["duplicate_price_codes"]:
        lines.extend(["## Duplicate price_code", "| price_code | count |", "|---|---:|"])
        for code, count in result["duplicate_price_codes"].items():
            lines.append(f"| `{code}` | {count} |")
        lines.append("")
    if result["filled_code_empty_price_rows"]:
        rows = ", ".join(str(row) for row in result["filled_code_empty_price_rows"])
        lines.extend(["## Empty Price Rows", rows, ""])
    if result["non_numeric_price_rows"]:
        rows = ", ".join(str(row) for row in result["non_numeric_price_rows"])
        lines.extend(["## Non-Numeric Price Rows", rows, ""])
    output_path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    summary = validate_price_registry()
    print(
        "price_registry validation:",
        f"rows={summary['total_rows']}",
        f"filled={summary['filled_price_code_rows']}",
        f"empty={summary['empty_price_code_rows']}",
        f"duplicates={len(summary['duplicate_price_codes'])}",
    )
