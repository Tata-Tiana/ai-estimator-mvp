from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re
import shutil
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_XLSX = Path("/Users/tatanamedzidova/Downloads/elena_parameter_review_pack.xlsx")
OUTPUT_DIR = REPO_ROOT / "output"
BACKUP_DIR = OUTPUT_DIR / "backups"
OUTPUT_XLSX = OUTPUT_DIR / "elena_parameter_review_pack_updated.xlsx"
AUDIT_MD = REPO_ROOT / "docs/elena_parameter_review_pack_audit.md"
DISCUSSION_SHEET = "К_обсуждению_с_Еленой"
MANUAL_SUMMARY_SHEET = "Итог_ручного_ввода"


@dataclass(frozen=True)
class Decision:
    status: str
    reason: str
    source_of_truth: str | None = None
    can_override: str | None = None
    visible_to_elena: str | None = None
    risk_level: str | None = None


PROFILE_DEFAULTS: dict[str, dict[str, str]] = {
    "AUTO_PROJECT": {
        "source_of_truth": "Project PDF/specification/review card",
        "can_override": "да",
        "visible_to_elena": "да",
        "risk_level": "high",
    },
    "AUTO_CALCULATED": {
        "source_of_truth": "deterministic calculator",
        "can_override": "нет",
        "visible_to_elena": "нет",
        "risk_level": "low",
    },
    "DEFAULT_VALUE": {
        "source_of_truth": "defaults_registry.py / system settings",
        "can_override": "да",
        "visible_to_elena": "нет",
        "risk_level": "low",
    },
    "MATERIAL_CATALOG": {
        "source_of_truth": "material_catalog.py / supplier catalog",
        "can_override": "да",
        "visible_to_elena": "нет",
        "risk_level": "low",
    },
    "PRICE_DATABASE": {
        "source_of_truth": "price_registry",
        "can_override": "да",
        "visible_to_elena": "нет",
        "risk_level": "medium",
    },
    "SUPPLIER_INPUT": {
        "source_of_truth": "supplier / Technonikol layout / commercial offer",
        "can_override": "да",
        "visible_to_elena": "да",
        "risk_level": "high",
    },
    "MANUAL_REQUIRED": {
        "source_of_truth": "estimator/manual project decision",
        "can_override": "да",
        "visible_to_elena": "да",
        "risk_level": "high",
    },
    "OPTIONAL_CONTROL": {
        "source_of_truth": "project geometry / control check",
        "can_override": "да",
        "visible_to_elena": "нет",
        "risk_level": "medium",
    },
    "DEPRECATED / LEGACY_ONLY": {
        "source_of_truth": "legacy locked cases only",
        "can_override": "нет",
        "visible_to_elena": "нет",
        "risk_level": "low",
    },
    "DEPRECATED / OPTIONAL_OVERRIDE": {
        "source_of_truth": "optional legacy override",
        "can_override": "да",
        "visible_to_elena": "нет",
        "risk_level": "medium",
    },
}


def norm(value: Any) -> str:
    return "" if value is None else str(value).strip()


def make_decision(
    status: str,
    reason: str,
    *,
    source_of_truth: str | None = None,
    can_override: str | None = None,
    visible_to_elena: str | None = None,
    risk_level: str | None = None,
) -> Decision:
    return Decision(status, reason, source_of_truth, can_override, visible_to_elena, risk_level)


def source_fields(decision: Decision) -> dict[str, str]:
    defaults = PROFILE_DEFAULTS[decision.status].copy()
    if decision.source_of_truth is not None:
        defaults["source_of_truth"] = decision.source_of_truth
    if decision.can_override is not None:
        defaults["can_override"] = decision.can_override
    if decision.visible_to_elena is not None:
        defaults["visible_to_elena"] = decision.visible_to_elena
    if decision.risk_level is not None:
        defaults["risk_level"] = decision.risk_level
    return defaults


def key_tuple(section: str, key: str, label: str = "") -> tuple[str, str, str]:
    return (section, key, label)


EXACT: dict[tuple[str, str], Decision] = {
    # Earthworks.
    ("Земляные работы", "assumptions.manual_excavation_override"): make_decision(
        "DEPRECATED / OPTIONAL_OVERRIDE",
        "Production-логика считает ручную разработку формулой из площади котлована, глубины доработки и траншей; override нужен только для legacy/аварийного вмешательства.",
    ),
    ("Земляные работы", "assumptions.sand_override"): make_decision(
        "DEPRECATED / OPTIONAL_OVERRIDE",
        "Production считает песок от проектного объема песка и объема траншей с коэффициентом уплотнения и округлением кратно машине; override только аварийный.",
    ),
    ("Земляные работы", "assumptions.geotextile_override"): make_decision(
        "DEPRECATED / OPTIONAL_OVERRIDE",
        "Production считает геотекстиль по площади, нахлесту и площади рулона; ручной override не является production-входом.",
    ),
    ("Земляные работы", "manual_refinement_depth_m"): make_decision(
        "DEFAULT_VALUE",
        "Глубина ручной доработки 0.08 м подтверждена как стандартное правило.",
    ),
    ("Земляные работы", "manual_excavation_quantity_for_estimate_m3"): make_decision(
        "AUTO_CALCULATED",
        "Считается как pit_area_m2 * manual_refinement_depth_m + trench_volume_total_m3.",
    ),
    ("Земляные работы", "trench_volume_m3"): make_decision(
        "AUTO_PROJECT",
        "Готовый объем траншей должен приходить из спецификации; если его нет, production может считать объем из routes отдельным расчетом.",
        visible_to_elena="да",
        risk_level="high",
    ),
    ("Земляные работы", "sand_truck_step_m3"): make_decision(
        "DEFAULT_VALUE",
        "Стандартный шаг заказа песка машиной 20 м3.",
    ),
    ("Земляные работы", "geotextile_roll_area_m2"): make_decision(
        "MATERIAL_CATALOG",
        "Площадь рулона геотекстиля 100 м2 является каталожным/стандартным значением.",
    ),
    ("Земляные работы", "excavator_shifts"): make_decision(
        "AUTO_CALCULATED",
        "Считается от объема машинной разработки и производительности 80 м3/смена.",
    ),
    # Foundation slab.
    ("Фундаментная плита", "slab_formwork_perimeter_m"): make_decision(
        "OPTIONAL_CONTROL",
        "Production использует готовую площадь боковой опалубки из спецификации; периметр остается контрольной геометрией.",
    ),
    ("Фундаментная плита", "slab_edge_height_m"): make_decision(
        "OPTIONAL_CONTROL",
        "Высота плиты может быть проектным контролем, но production-источник опалубки — готовая площадь боковой опалубки из спецификации.",
        source_of_truth="Project PDF/specification/review card",
    ),
    ("Фундаментная плита", "thermal_insert_length_m"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Одно общее поле длины термовставок заменено production-полями по 50/100 мм и материальными quantity из спецификации.",
    ),
    ("Фундаментная плита", "rebar_crane_shifts"): make_decision(
        "DEFAULT_VALUE",
        "Для текущей production-логики это стандартная смена крана, а не ручной ввод из проекта.",
    ),
    ("Фундаментная плита", "rebar_metal_delivery_trucks"): make_decision(
        "AUTO_CALCULATED",
        "Доставка арматуры должна считаться по общему весу арматуры коробки кратно 10 тоннам.",
    ),
    ("Фундаментная плита", "concrete_mixer_volume_m3"): make_decision(
        "DEFAULT_VALUE",
        "Объем автобетоносмесителя — стандартная настройка, например 9 м3.",
    ),
    ("Фундаментная плита", "concrete_pump_shifts"): make_decision(
        "DEFAULT_VALUE",
        "В подтвержденных кейсах количество смен бетононасоса является стандартным значением.",
    ),
    ("Фундаментная плита", "plywood_calc_method"): make_decision(
        "DEFAULT_VALUE",
        "Режим расчета фанеры является системной настройкой калькулятора; production использует подтвержденный стандарт, а не ручной ввод Елены.",
    ),
    ("Фундаментная плита", "slab_edge_height_strategy"): make_decision(
        "DEFAULT_VALUE",
        "Правило выбора высоты торца является системной стратегией калькулятора, не проектным ручным параметром.",
    ),
    # Waterproofing.
    ("Гидроизоляция", "slab_formwork_perimeter_m"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Production использует готовую площадь гидроизоляции из спецификации, а не периметр.",
    ),
    ("Гидроизоляция", "slab_edge_height_m"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Высота торца нужна только для старой геометрической проверки; production берет площадь гидроизоляции из спецификации.",
    ),
    ("Гидроизоляция", "glue_foam_min_units"): make_decision(
        "DEFAULT_VALUE",
        "Минимум клей-пены и расход считаются по стандартному правилу от площади утепления.",
    ),
    # Load-bearing walls and lintels.
    ("Несущие стены и перемычки", "scaffolding_setup_quantity"): make_decision(
        "AUTO_CALCULATED",
        "Считается от количества этажей: floors_count * scaffolding_setup_units_per_floor.",
    ),
    ("Несущие стены и перемычки", "scaffolding_timber_quantity_m3"): make_decision(
        "AUTO_CALCULATED",
        "Считается от количества этажей: floors_count * scaffolding_timber_m3_per_floor.",
    ),
    ("Несущие стены и перемычки", "gas_block_d400_pallet_volume_m3"): make_decision(
        "MATERIAL_CATALOG",
        "Объем газобетона в поддоне — каталожная характеристика.",
    ),
    ("Несущие стены и перемычки", "gas_block_d500_250_pallet_volume_m3"): make_decision(
        "MATERIAL_CATALOG",
        "Объем газобетона в поддоне — каталожная характеристика.",
    ),
    ("Несущие стены и перемычки", "gas_block_d500_150_pallet_volume_m3"): make_decision(
        "MATERIAL_CATALOG",
        "Объем газобетона в поддоне — каталожная характеристика.",
    ),
    ("Несущие стены и перемычки", "adhesive_consumption_bag_per_m3"): make_decision(
        "MATERIAL_CATALOG",
        "Расход клея на м3 кладки — каталожная/технологическая характеристика материала.",
    ),
    ("Несущие стены и перемычки", "sand_concrete_bag_weight_kg"): make_decision(
        "MATERIAL_CATALOG",
        "Вес мешка пескобетона — каталожное значение упаковки.",
    ),
    ("Несущие стены и перемычки", "gas_block_length_m"): make_decision(
        "MATERIAL_CATALOG",
        "Длина блока 0.6 м — стандартная характеристика блока.",
    ),
    ("Несущие стены и перемычки", "main_walls_crane_shifts"): make_decision(
        "AUTO_CALCULATED",
        "Считается от количества доставок блоков: до 3 доставок — 1 смена, 4 и более — 2 смены.",
    ),
    ("Несущие стены и перемычки", "lintel_section_width_m"): make_decision(
        "DEFAULT_VALUE",
        "Сечение бетонной части U-блока 0.125 м — подтвержденное стандартное значение.",
    ),
    ("Несущие стены и перемычки", "lintel_section_height_m"): make_decision(
        "DEFAULT_VALUE",
        "Сечение бетонной части U-блока 0.125 м — подтвержденное стандартное значение.",
    ),
    ("Несущие стены и перемычки", "lintel_concrete_min_order_volume_m3"): make_decision(
        "DEFAULT_VALUE",
        "Минимальный заказ 1 м3 — закупочное правило; заказанный объем затем округляется вверх до целого м3.",
    ),
    ("Несущие стены и перемычки", "concrete_delivery_trips"): make_decision(
        "AUTO_CALCULATED",
        "Рейсы доставки бетона считаются от заказанного объема бетона.",
    ),
    ("Несущие стены и перемычки", "parapet_enabled"): make_decision(
        "AUTO_CALCULATED",
        "В production включение парапета считается от flat_roof_enabled и объема парапета из спецификации кровли.",
    ),
    ("Несущие стены и перемычки", "parapet_masonry_volume_m3"): make_decision(
        "AUTO_PROJECT",
        "Объем кладки парапета должен приходить из спецификации кровли/парапета.",
    ),
    ("Несущие стены и перемычки", "second_light_masonry_enabled"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Термин second_light оставлен только для ЮСВ; production использует floors_count и блок floor_2_load_bearing_walls.",
    ),
    ("Несущие стены и перемычки", "second_light_masonry_case_specific"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Case-specific флаг ЮСВ не должен быть production-параметром.",
    ),
    ("Несущие стены и перемычки", "second_light_masonry_volume_m3"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Legacy-поле заменено production-параметром floor_2_masonry_volume_m3.",
    ),
    ("Несущие стены и перемычки", "vent_chimney_cladding_enabled"): make_decision(
        "AUTO_CALCULATED",
        "В production включение обкладки вентканалов считается от flat_roof_enabled и объема вентканалов из спецификации.",
    ),
    ("Несущие стены и перемычки", "vent_chimney_rows"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Production считает площадь обкладки вентканалов от vent_chimney_gas_block_spec_volume_m3 / 0.15, без рядов.",
    ),
    ("Несущие стены и перемычки", "block_height_m"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Высота блока нужна только для legacy geometry check вентканалов; production использует объем из спецификации и толщину 0.15 м.",
    ),
    ("Несущие стены и перемычки", "parapet_crane_shifts"): make_decision(
        "DEFAULT_VALUE",
        "Стандартно 1 смена крана.",
    ),
    # Floor slab 1.
    ("Плита перекрытия 1-го этажа", "geometry.total_concrete_volume_from_spec_m3"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Старый общий объем использовался для восстановления площади опалубки; production использует готовые площади опалубки из спецификации и отдельные объемы/контроли.",
    ),
    ("Плита перекрытия 1-го этажа", "geometry.slab_edge_perimeter_m"): make_decision(
        "OPTIONAL_CONTROL",
        "Для production-опалубки не главный источник; готовые площади опалубки приходят из спецификации. Может использоваться как контроль/связанный геометрический параметр.",
    ),
    ("Плита перекрытия 1-го этажа", "geometry.edge_formwork_height_m"): make_decision(
        "OPTIONAL_CONTROL",
        "Для production-опалубки не главный источник; edge_formwork_area_m2 приходит из спецификации.",
    ),
    ("Плита перекрытия 1-го этажа", "rates.formwork_supplier_quote_total"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Production использует прямую ставку/прайс для раздела, а не кросс-разделное предложение поставщика.",
    ),
    ("Плита перекрытия 1-го этажа", "rates.slab_2_formwork_area_for_rate_context_m2"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Площадь плиты 2-го этажа нужна только для старого расчета справочной ставки из предложения поставщика.",
    ),
    ("Плита перекрытия 1-го этажа", "rates.reserve_plywood_sheets"): make_decision(
        "DEFAULT_VALUE",
        "Резерв фанеры — стандартная настройка/запас, не проектный ручной ввод.",
    ),
    ("Плита перекрытия 1-го этажа", "rates.floor_slab_2_rebar_weight_for_delivery_context_kg"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Production плиты 1-го этажа отдает только section_rebar_delivery_weight_kg; доставка металла всей коробки считается на box-level.",
    ),
    ("Плита перекрытия 1-го этажа", "rates.max_rebar_delivery_weight_per_truck_kg"): make_decision(
        "DEFAULT_VALUE",
        "10 000 кг на машину — стандартное правило доставки.",
    ),
    ("Плита перекрытия 1-го этажа", "insulation.total_eps_volume_from_spec_m3"): make_decision(
        "AUTO_PROJECT",
        "Чистый объем ЭППС должен приходить из спецификации; материал считается с запасом и округлением до упаковок.",
    ),
    ("Плита перекрытия 1-го этажа", "manual_lines.formwork_delivery_trucks_override"): make_decision(
        "DEPRECATED / OPTIONAL_OVERRIDE",
        "Production считает доставку опалубки по площади: <= 180 м2 — 2 машины, > 180 м2 — 4 машины; override только для исключений.",
    ),
    ("Плита перекрытия 1-го этажа", "manual_lines.formwork_rebar_crane_shifts"): make_decision(
        "DEFAULT_VALUE",
        "Стандартно 2 смены.",
    ),
    ("Плита перекрытия 1-го этажа", "manual_lines.concrete_pump_shifts"): make_decision(
        "DEFAULT_VALUE",
        "Стандартно 1 смена.",
    ),
    ("Плита перекрытия 1-го этажа", "overheads.logistics_and_supply_percent"): make_decision(
        "DEFAULT_VALUE",
        "Стандартный процент раздела.",
    ),
    ("Плита перекрытия 1-го этажа", "overheads.consumables_and_tool_percent"): make_decision(
        "DEFAULT_VALUE",
        "Стандартный процент раздела.",
    ),
    # Floor slab 2.
    ("Плита перекрытия 2-го этажа", "slab_length_m"): make_decision(
        "OPTIONAL_CONTROL",
        "Длина нужна только для геометрического контроля; production-источник опалубки — готовые площади из спецификации.",
    ),
    ("Плита перекрытия 2-го этажа", "slab_width_m"): make_decision(
        "OPTIONAL_CONTROL",
        "Ширина нужна только для геометрического контроля; production-источник опалубки — готовые площади из спецификации.",
    ),
    ("Плита перекрытия 2-го этажа", "slab_area_m2"): make_decision(
        "OPTIONAL_CONTROL",
        "Площадь плиты нужна только для контроля; main_formwork_area_m2 приходит из спецификации.",
    ),
    ("Плита перекрытия 2-го этажа", "slab_edge_perimeter_m"): make_decision(
        "AUTO_PROJECT",
        "Нужен как длина работы по утеплению торца в м.п.; должен приходить из спецификации.",
    ),
    ("Плита перекрытия 2-го этажа", "main_formwork_area_m2"): make_decision(
        "AUTO_PROJECT",
        "Основная площадь опалубки должна приходить из спецификации.",
    ),
    ("Плита перекрытия 2-го этажа", "formwork_rental_supplier_quote_total"): make_decision(
        "DEPRECATED / LEGACY_ONLY",
        "Production использует площадь опалубки и ставку/прайс, а не ручную сумму предложения.",
    ),
    ("Плита перекрытия 2-го этажа", "formwork_delivery_trips"): make_decision(
        "AUTO_CALCULATED",
        "Production считает доставку опалубки по правилу 180 м2: до 180 включительно = 2 рейса, больше 180 = 4 рейса.",
    ),
    ("Плита перекрытия 2-го этажа", "crane_shifts"): make_decision(
        "DEFAULT_VALUE",
        "Стандартно 2 смены.",
    ),
    ("Плита перекрытия 2-го этажа", "edge_formwork_height_m"): make_decision(
        "OPTIONAL_CONTROL",
        "Для production-площади торцевой опалубки высота больше не главный вход: edge_formwork_area_m2 приходит из спецификации.",
    ),
    ("Плита перекрытия 2-го этажа", "plywood_sheet_working_area_m2"): make_decision(
        "MATERIAL_CATALOG",
        "Рабочая площадь листа фанеры — каталожная характеристика.",
    ),
    ("Плита перекрытия 2-го этажа", "plywood_reserve_sheets"): make_decision(
        "DEFAULT_VALUE",
        "Резерв фанеры — стандартная настройка/запас.",
    ),
    ("Плита перекрытия 2-го этажа", "concrete_mixer_volume_m3"): make_decision(
        "DEFAULT_VALUE",
        "9 м3 — стандартное значение автобетоносмесителя.",
    ),
    ("Плита перекрытия 2-го этажа", "concrete_pump_shifts"): make_decision(
        "DEFAULT_VALUE",
        "Стандартно 1 смена.",
    ),
    ("Плита перекрытия 2-го этажа", "edge_insulation_height_m"): make_decision(
        "AUTO_PROJECT",
        "Елена подтвердила высоту утепления торца 180 мм; значение должно быть в спецификации.",
    ),
    ("Плита перекрытия 2-го этажа", "foam_min_cans"): make_decision(
        "DEFAULT_VALUE",
        "Минимум клей-пены — стандартное закупочное правило.",
    ),
    ("Плита перекрытия 2-го этажа", "logistics_rate"): make_decision(
        "DEFAULT_VALUE",
        "Стандартный процент логистики.",
    ),
    ("Плита перекрытия 2-го этажа", "consumables_rate"): make_decision(
        "DEFAULT_VALUE",
        "Стандартный процент расходных материалов.",
    ),
    # Flat roof.
    ("Плоская кровля", "roof_area_level_1_m2"): make_decision(
        "AUTO_PROJECT",
        "Площадь кровли уровня 1 должна приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "roof_area_level_2_m2"): make_decision(
        "AUTO_PROJECT",
        "Площадь кровли уровня 2 должна приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "parapet_length_level_1_m"): make_decision(
        "AUTO_PROJECT",
        "Длина парапета уровня 1 должна приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "parapet_length_level_2_m"): make_decision(
        "AUTO_PROJECT",
        "Длина парапета уровня 2 должна приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "vent_wall_abutment_level_1_m"): make_decision(
        "AUTO_PROJECT",
        "Длина примыканий уровня 1 должна приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "vent_wall_abutment_level_2_m"): make_decision(
        "AUTO_PROJECT",
        "Длина примыканий уровня 2 должна приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "vapor_barrier_film_roll_area_m2"): make_decision(
        "MATERIAL_CATALOG",
        "150 м2 — площадь рулона пароизоляционной пленки.",
    ),
    ("Плоская кровля", "vent_shaft_abutment_count"): make_decision(
        "AUTO_PROJECT",
        "Количество примыканий к вентшахтам должно приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "rail_piece_length_m"): make_decision(
        "MATERIAL_CATALOG",
        "Алюминиевая рейка 3 м — каталожный размер.",
    ),
    ("Плоская кровля", "pvc_membrane_roll_width_m"): make_decision(
        "MATERIAL_CATALOG",
        "Ширина рулона ПВХ мембраны — каталожное значение.",
    ),
    ("Плоская кровля", "pvc_membrane_roll_length_m"): make_decision(
        "MATERIAL_CATALOG",
        "Длина рулона ПВХ мембраны — каталожное значение.",
    ),
    ("Плоская кровля", "roof_aerators_count"): make_decision(
        "AUTO_PROJECT",
        "Количество кровельных аэраторов должно приходить из проекта/спецификации.",
    ),
    ("Плоская кровля", "gas_block_wall_holes_count"): make_decision(
        "AUTO_PROJECT",
        "Количество отверстий должно приходить из проекта/спецификации/кровельных узлов.",
    ),
    ("Плоская кровля", "internal_drain_height_per_drain_m"): make_decision(
        "AUTO_PROJECT",
        "Высота внутреннего водостока должна приходить из проекта/спецификации/кровельных узлов.",
    ),
    ("Плоская кровля", "roof_crane_lifting_shifts"): make_decision(
        "DEFAULT_VALUE",
        "Стандартно 1 смена автокрана.",
    ),
    # Schiedel.
    ("Вентиляционные каналы Schiedel", "vent_channel_1_height_m"): make_decision(
        "AUTO_PROJECT",
        "Берется из спецификации вентканалов.",
    ),
    ("Вентиляционные каналы Schiedel", "vent_channel_2_height_m"): make_decision(
        "AUTO_PROJECT",
        "Берется из спецификации вентканалов.",
    ),
    ("Вентиляционные каналы Schiedel", "vent_channel_2_count"): make_decision(
        "AUTO_PROJECT",
        "Берется из спецификации вентканалов.",
    ),
    ("Вентиляционные каналы Schiedel", "schiedel_delivery_trips"): make_decision(
        "DEFAULT_VALUE",
        "Стандартно 1 доставка/смена.",
    ),
    ("Вентиляционные каналы Schiedel", "consumables_rate"): make_decision(
        "DEFAULT_VALUE",
        "Стандартный процент/фиксированная настройка раздела.",
    ),
}


ADDED_ROWS: list[dict[str, str]] = [
    {
        "section_name": "Фундаментная плита",
        "label": "Термовставка 50 мм: длина",
        "calculator_input_key": "thermal_insert_50_length_m",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production разделяет термовставки 50 и 100 мм; длина 50 мм должна приходить из спецификации.",
    },
    {
        "section_name": "Фундаментная плита",
        "label": "Термовставка 100 мм: длина",
        "calculator_input_key": "thermal_insert_100_length_m",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production разделяет термовставки 50 и 100 мм; длина 100 мм должна приходить из спецификации.",
    },
    {
        "section_name": "Фундаментная плита",
        "label": "Термовставка 50 мм: количество материала",
        "calculator_input_key": "thermal_insert_50_material_spec_qty",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Материальное количество термовставки 50 мм должно приходить из спецификации.",
    },
    {
        "section_name": "Фундаментная плита",
        "label": "Термовставка 100 мм: количество материала",
        "calculator_input_key": "thermal_insert_100_material_spec_qty",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Материальное количество термовставки 100 мм должно приходить из спецификации.",
    },
    {
        "section_name": "Несущие стены и перемычки",
        "label": "Площадь отсечной гидроизоляции под несущие стены",
        "calculator_input_key": "cutoff_waterproofing_load_bearing_walls_area_m2",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production берет готовую площадь отсечной гидроизоляции под несущие стены из спецификации.",
    },
    {
        "section_name": "Несущие стены и перемычки",
        "label": "Площадь отсечной гидроизоляции под перегородки",
        "calculator_input_key": "cutoff_waterproofing_partitions_area_m2",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Площадь отсечной гидроизоляции под перегородки должна жить в разделе перегородок, не в несущих стенах.",
    },
    {
        "section_name": "Несущие стены и перемычки",
        "label": "Общая длина перемычек в U-блоке",
        "calculator_input_key": "lintel_total_length_m",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production берет общую длину перемычек из спецификации, а не список length/count.",
    },
    {
        "section_name": "Несущие стены и перемычки",
        "label": "Объем газоблока для обкладки вентканалов",
        "calculator_input_key": "vent_chimney_gas_block_spec_volume_m3",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production считает площадь обкладки вентканалов как volume / 0.15.",
    },
    {
        "section_name": "Несущие стены и перемычки",
        "label": "Объем кладки несущих стен 2-го этажа",
        "calculator_input_key": "floor_2_masonry_volume_m3",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production заменяет legacy second_light на блок floor_2_load_bearing_walls от floors_count.",
    },
    {
        "section_name": "Плита перекрытия 1-го этажа",
        "label": "Площадь опалубки под плиту",
        "calculator_input_key": "main_formwork_area_m2",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production берет готовую площадь опалубки под плиту из спецификации.",
    },
    {
        "section_name": "Плита перекрытия 1-го этажа",
        "label": "Площадь торцевой опалубки плиты",
        "calculator_input_key": "edge_formwork_area_m2",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production берет готовую площадь торцевой опалубки из спецификации.",
    },
    {
        "section_name": "Плита перекрытия 1-го этажа",
        "label": "Площадь опалубки балок",
        "calculator_input_key": "beams_formwork_area_m2",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production берет готовую площадь опалубки балок из спецификации.",
    },
    {
        "section_name": "Плита перекрытия 1-го этажа",
        "label": "Площадь торцевой опалубки и балок",
        "calculator_input_key": "edge_and_beam_formwork_area_m2",
        "recommended_source_status": "AUTO_CALCULATED",
        "reason": "Считается как edge_formwork_area_m2 + beams_formwork_area_m2.",
    },
    {
        "section_name": "Плита перекрытия 1-го этажа",
        "label": "Высота утепления торца",
        "calculator_input_key": "edge_insulation_height_m",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Елена подтвердила 180 мм; высота должна приходить из спецификации, fallback — slab_thickness_m.",
    },
    {
        "section_name": "Плита перекрытия 2-го этажа",
        "label": "Площадь торцевой опалубки плиты",
        "calculator_input_key": "edge_formwork_area_m2",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production берет готовую площадь торцевой опалубки из спецификации.",
    },
    {
        "section_name": "Плита перекрытия 2-го этажа",
        "label": "Площадь опалубки балок",
        "calculator_input_key": "beams_formwork_area_m2",
        "recommended_source_status": "AUTO_PROJECT",
        "reason": "Production поддерживает готовую площадь опалубки балок из спецификации; для текущего ЮСВ без балок значение 0.",
    },
    {
        "section_name": "Плита перекрытия 2-го этажа",
        "label": "Площадь торцевой опалубки и балок",
        "calculator_input_key": "edge_and_beam_formwork_area_m2",
        "recommended_source_status": "AUTO_CALCULATED",
        "reason": "Считается как edge_formwork_area_m2 + beams_formwork_area_m2.",
    },
    {
        "section_name": "Плоская кровля",
        "label": "Общая площадь кровли",
        "calculator_input_key": "roof_area_total_m2",
        "recommended_source_status": "AUTO_CALCULATED",
        "reason": "Production detailed_project_geometry считает roof_area_total_m2 как сумма площадей уровней.",
    },
    {
        "section_name": "Плоская кровля",
        "label": "Общая длина парапетов и примыканий",
        "calculator_input_key": "parapet_and_abutment_total_length_m",
        "recommended_source_status": "AUTO_CALCULATED",
        "reason": "Production detailed_project_geometry считает сумму длин парапетов и примыканий по уровням.",
    },
]


def pattern_decision(section: str, key: str, label: str, current_recommended: str) -> Decision:
    lowered = key.lower()

    if "unit_price" in lowered or lowered.endswith("_rate") or "work_rate" in lowered or "price" in lowered:
        return make_decision("PRICE_DATABASE", "Цена или ставка должна приходить из price_registry / live pricing.")

    if "supplier_required_volume" in lowered:
        return make_decision(
            "SUPPLIER_INPUT",
            "Объем зависит от раскладки поставщика/Технониколь и пока не считается геометрией в MVP.",
        )

    if "override" in lowered:
        return make_decision(
            "DEPRECATED / OPTIONAL_OVERRIDE",
            "Production считает значение формулой; override оставлен только для исключений.",
        )

    if "cutoff_waterproofing_wall_" in lowered:
        return make_decision(
            "DEPRECATED / LEGACY_ONLY",
            "Production использует готовые площади отсечной гидроизоляции, а не списки длин стен 400/250 мм.",
        )

    if "lintel_lengths_m" in lowered:
        return make_decision(
            "DEPRECATED / LEGACY_ONLY",
            "Production использует lintel_total_length_m из спецификации; список length/count оставлен только для legacy.",
        )

    if "vent_chimney_segment_lengths_m" in lowered:
        return make_decision(
            "DEPRECATED / LEGACY_ONLY",
            "Production считает обкладку вентканалов от vent_chimney_gas_block_spec_volume_m3 / vent_chimney_block_thickness_m.",
        )

    if "non_insulated_edge_lengths_m" in lowered:
        return make_decision(
            "DEPRECATED / LEGACY_ONLY",
            "Production считает ЭППС 100 мм торец от объема из спецификации: volume / thickness.",
        )

    if "source_weight_parts_kg" in lowered:
        return make_decision(
            "DEPRECATED / LEGACY_ONLY",
            "Production-режим арматуры перешел на spec_length_m; веса оставлены для legacy/контроля.",
        )

    if re.search(r"rebar_items\[\d+\]\.code|lintel_rebar_items\[\d+\]\.code", lowered):
        return make_decision("AUTO_CALCULATED", "Код арматуры формируется по steel_class + diameter_mm / каталогу.")
    if re.search(r"rebar_items\[\d+\]\.name|lintel_rebar_items\[\d+\]\.name", lowered):
        return make_decision("AUTO_CALCULATED", "Наименование арматуры формируется по steel_class + diameter_mm / каталогу.")
    if re.search(r"rebar_items\[\d+\]\.steel_class|lintel_rebar_items\[\d+\]\.steel_class", lowered):
        return make_decision("AUTO_PROJECT", "Класс стали должен приходить из спецификации/ведомости арматуры.")
    if re.search(r"rebar_items\[\d+\]\.diameter_mm|lintel_rebar_items\[\d+\]\.diameter_mm", lowered):
        return make_decision("AUTO_PROJECT", "Диаметр арматуры должен приходить из спецификации/ведомости арматуры.")
    if "spec_length_m" in lowered and "rebar_items" in lowered:
        return make_decision("AUTO_PROJECT", "Длина арматуры в м.п. должна приходить из спецификации.")
    if "weight_parts_kg" in lowered and section == "Фундаментная плита":
        return make_decision("AUTO_PROJECT", "Вес арматуры фундаментной плиты пока приходит из спецификации; будущий режим может перейти на м.п.")

    catalog_tokens = [
        "pack_volume",
        "roll_area",
        "sheet_width",
        "sheet_height",
        "sheet_working_area",
        "working_area",
        "kg_per_meter",
        "kg_per_m",
        "rod_length",
        "pallet_volume",
        "piece_length",
        "piece_width",
        "piece_height",
        "bucket_weight",
        "canister_volume",
        "bag_weight",
        "block_thickness",
        "roll_width",
        "roll_length",
        "rail_piece_length",
        "gas_block_length",
        "primer_consumption",
        "mastic_bucket_weight",
        "adhesive_consumption",
    ]
    if any(token in lowered for token in catalog_tokens):
        return make_decision("MATERIAL_CATALOG", "Размер/расход/упаковка является каталожной характеристикой материала.")

    default_tokens = [
        "mastic_layers",
        "min_units",
        "min_cans",
        "shifts",
        "mixer_volume",
        "max_rebar_delivery_weight",
        "logistics_rate",
        "consumables_rate",
        "work_coeff",
        "overlap_coeff",
        "waste_coeff",
        "reserve",
    ]
    if any(token in lowered for token in default_tokens):
        return make_decision("DEFAULT_VALUE", "Стандартное правило/настройка раздела, не проектный ручной ввод.")

    if "beams.items" in lowered and (lowered.endswith(".code") or lowered.endswith(".name")):
        return make_decision("AUTO_CALCULATED", "Код/название балки формирует калькулятор.")
    if "beams.items" in lowered and any(lowered.endswith(suffix) for suffix in [".length_m", ".width_m", ".height_m", ".count"]):
        return make_decision("AUTO_PROJECT", "Геометрия балок должна приходить из спецификации.")

    if lowered.endswith("_count") or ".count" in lowered:
        return make_decision("AUTO_PROJECT", "Количество должно приходить из проекта/спецификации, если нет отдельного подтвержденного правила.")

    if current_recommended in PROFILE_DEFAULTS:
        return make_decision(current_recommended, "Сохранено из существующей рекомендации audit; явного override после разбора не найдено.")

    if any(token in lowered for token in ["area", "volume", "length", "height", "width", "perimeter"]):
        return make_decision("AUTO_PROJECT", "Проектная геометрия/объем/длина должны приходить из PDF/спецификации или review card.")

    return make_decision("AUTO_PROJECT", "Не найдено более точное правило; оставить как проектный параметр для review card.")


def final_decision(section: str, key: str, label: str, recommended: str) -> Decision:
    exact = EXACT.get((section, key))
    if exact:
        return exact
    return pattern_decision(section, key, label, recommended)


def get_headers(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(1, col).value
        if value is not None:
            headers[str(value)] = col
    return headers


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def normalize_headers(ws) -> dict[str, int]:
    headers = get_headers(ws)
    if "user_comment" not in headers:
        last_col = ws.max_column
        ws.cell(1, last_col).value = "user_comment"
    headers = get_headers(ws)
    required = [
        "section_name",
        "label",
        "calculator_input_key",
        "recommended_source_status",
        "reason",
        "elena_decision",
        "source_of_truth",
        "can_override",
        "visible_to_elena",
        "risk_level",
        "user_comment",
    ]
    missing = [header for header in required if header not in headers]
    if missing:
        raise ValueError(f"Missing required headers: {missing}")
    return headers


def row_identity(ws, row: int, headers: dict[str, int]) -> tuple[str, str, str]:
    section = norm(ws.cell(row, headers["section_name"]).value)
    key = norm(ws.cell(row, headers["calculator_input_key"]).value)
    label = norm(ws.cell(row, headers["label"]).value)
    return section, key, label


def find_last_section_row(ws, headers: dict[str, int], section: str) -> int:
    result = ws.max_row
    for row in range(2, ws.max_row + 1):
        if norm(ws.cell(row, headers["section_name"]).value) == section:
            result = row
    return result


def add_missing_rows(ws, headers: dict[str, int]) -> list[str]:
    existing = {
        (norm(ws.cell(row, headers["section_name"]).value), norm(ws.cell(row, headers["calculator_input_key"]).value))
        for row in range(2, ws.max_row + 1)
    }
    added: list[str] = []
    for payload in ADDED_ROWS:
        identity = (payload["section_name"], payload["calculator_input_key"])
        if identity in existing:
            continue
        insert_at = find_last_section_row(ws, headers, payload["section_name"]) + 1
        ws.insert_rows(insert_at)
        style_source = max(2, insert_at - 1)
        copy_row_style(ws, style_source, insert_at)
        for header, value in payload.items():
            if header in headers:
                ws.cell(insert_at, headers[header]).value = value
        for header in ["current_status", "current_input_type"]:
            if header in headers:
                ws.cell(insert_at, headers[header]).value = "new_production_field"
        if "user_comment" in headers:
            ws.cell(insert_at, headers["user_comment"]).value = None
        added.append(f"{payload['section_name']} :: {payload['calculator_input_key']}")
        existing.add(identity)
    return added


def apply_decision(ws, row: int, headers: dict[str, int], decision: Decision) -> None:
    fields = source_fields(decision)
    ws.cell(row, headers["recommended_source_status"]).value = decision.status
    ws.cell(row, headers["elena_decision"]).value = decision.status
    ws.cell(row, headers["reason"]).value = decision.reason
    ws.cell(row, headers["source_of_truth"]).value = fields["source_of_truth"]
    ws.cell(row, headers["can_override"]).value = fields["can_override"]
    ws.cell(row, headers["visible_to_elena"]).value = fields["visible_to_elena"]
    ws.cell(row, headers["risk_level"]).value = fields["risk_level"]


def action_needed(decision: str) -> str:
    if decision == "MANUAL_REQUIRED":
        return "Нужно ручное решение сметчика."
    if decision == "SUPPLIER_INPUT":
        return "Получить данные от поставщика/раскладки."
    if decision == "AUTO_PROJECT":
        return "Найти в проекте/спецификации или показать Елене в review form."
    if decision == "OPTIONAL_CONTROL":
        return "Показывать только как контроль, если требуется."
    return ""


def rebuild_manual_summary(wb, rows: list[dict[str, Any]]) -> None:
    if MANUAL_SUMMARY_SHEET in wb.sheetnames:
        del wb[MANUAL_SUMMARY_SHEET]
    ws = wb.create_sheet(MANUAL_SUMMARY_SHEET)
    headers = [
        "section_name",
        "label",
        "calculator_input_key",
        "elena_decision",
        "source_of_truth",
        "reason",
        "user_comment",
        "action_needed",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")

    allowed = {"MANUAL_REQUIRED", "SUPPLIER_INPUT", "AUTO_PROJECT", "OPTIONAL_CONTROL"}
    for row in rows:
        if row["elena_decision"] not in allowed:
            continue
        if row["elena_decision"] in {"AUTO_PROJECT", "OPTIONAL_CONTROL"} and row["visible_to_elena"] != "да":
            continue
        ws.append(
            [
                row["section_name"],
                row["label"],
                row["calculator_input_key"],
                row["elena_decision"],
                row["source_of_truth"],
                row["reason"],
                row["user_comment"],
                action_needed(row["elena_decision"]),
            ]
        )
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 28


def build_audit(rows_before_blank: int, rows: list[dict[str, Any]], added_rows: list[str], backup_path: Path) -> str:
    counts = Counter(row["elena_decision"] for row in rows)
    sections = Counter(row["section_name"] for row in rows)

    def list_rows(title: str, predicate) -> list[str]:
        result = [f"## {title}", ""]
        selected = [row for row in rows if predicate(row)]
        if not selected:
            result.append("- Нет.")
        else:
            for row in selected:
                result.append(
                    f"- {row['section_name']} :: `{row['calculator_input_key']}` — {row['label']} ({row['reason']})"
                )
        result.append("")
        return result

    lines = [
        "# Audit: Elena parameter review pack",
        "",
        f"Дата: {datetime.now().date().isoformat()}",
        "",
        f"- Исходный файл: `{SOURCE_XLSX}`",
        f"- Backup: `{backup_path}`",
        f"- Обновленный файл: `{OUTPUT_XLSX}`",
        "",
        "## Summary",
        "",
        f"- Строк на первом листе после обновления: {len(rows)}",
        f"- Пустых `elena_decision` до обновления: {rows_before_blank}",
        f"- Пустых `elena_decision` после обновления: {sum(1 for row in rows if not row['elena_decision'])}",
        "",
        "### По разделам",
        "",
    ]
    for section, count in sections.most_common():
        lines.append(f"- {section}: {count}")
    lines.extend(["", "### По elena_decision", ""])
    for decision, count in counts.most_common():
        lines.append(f"- {decision}: {count}")
    lines.extend(["", "## Добавленные строки", ""])
    if added_rows:
        lines.extend([f"- `{row}`" for row in added_rows])
    else:
        lines.append("- Нет.")
    lines.append("")

    lines.extend(list_rows("MANUAL_REQUIRED", lambda row: row["elena_decision"] == "MANUAL_REQUIRED"))
    lines.extend(list_rows("SUPPLIER_INPUT", lambda row: row["elena_decision"] == "SUPPLIER_INPUT"))
    lines.extend(list_rows("DEPRECATED / LEGACY_ONLY", lambda row: row["elena_decision"] == "DEPRECATED / LEGACY_ONLY"))
    lines.extend(
        list_rows("DEPRECATED / OPTIONAL_OVERRIDE", lambda row: row["elena_decision"] == "DEPRECATED / OPTIONAL_OVERRIDE")
    )
    lines.extend(
        list_rows(
            "AUTO_PROJECT high risk",
            lambda row: row["elena_decision"] == "AUTO_PROJECT" and row["risk_level"] == "high",
        )
    )
    lines.extend(
        [
            "## Notes",
            "",
            "- `SUPPLIER_INPUT` не считается ручным решением сметчика: это внешний ввод поставщика/раскладки.",
            "- `DEPRECATED / LEGACY_ONLY` и `DEPRECATED / OPTIONAL_OVERRIDE` не являются production input.",
            "- `DEFAULT_VALUE`, `MATERIAL_CATALOG`, `PRICE_DATABASE`, `AUTO_CALCULATED` не нужно спрашивать у Елены в каждом проекте.",
            "- `AUTO_PROJECT` нужно искать в проекте/спецификации; если parser не нашел, показывать в review form.",
            "",
        ]
    )
    return "\n".join(lines)


def workbook_rows(ws, headers: dict[str, int]) -> list[dict[str, Any]]:
    result = []
    for row in range(2, ws.max_row + 1):
        if not norm(ws.cell(row, headers["calculator_input_key"]).value):
            continue
        result.append(
            {
                "section_name": norm(ws.cell(row, headers["section_name"]).value),
                "label": norm(ws.cell(row, headers["label"]).value),
                "calculator_input_key": norm(ws.cell(row, headers["calculator_input_key"]).value),
                "elena_decision": norm(ws.cell(row, headers["elena_decision"]).value),
                "source_of_truth": norm(ws.cell(row, headers["source_of_truth"]).value),
                "reason": norm(ws.cell(row, headers["reason"]).value),
                "can_override": norm(ws.cell(row, headers["can_override"]).value),
                "visible_to_elena": norm(ws.cell(row, headers["visible_to_elena"]).value),
                "risk_level": norm(ws.cell(row, headers["risk_level"]).value),
                "user_comment": norm(ws.cell(row, headers["user_comment"]).value),
            }
        )
    return result


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backup_path = BACKUP_DIR / f"elena_parameter_review_pack_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    shutil.copy2(SOURCE_XLSX, backup_path)

    wb = load_workbook(SOURCE_XLSX)
    ws = wb[DISCUSSION_SHEET]
    headers = normalize_headers(ws)
    rows_before_blank = sum(
        1
        for row in range(2, ws.max_row + 1)
        if norm(ws.cell(row, headers["calculator_input_key"]).value)
        and not norm(ws.cell(row, headers["elena_decision"]).value)
    )

    added_rows = add_missing_rows(ws, headers)
    headers = normalize_headers(ws)

    for row in range(2, ws.max_row + 1):
        section, key, label = row_identity(ws, row, headers)
        if not key:
            continue
        recommended = norm(ws.cell(row, headers["recommended_source_status"]).value)
        decision = final_decision(section, key, label, recommended)
        apply_decision(ws, row, headers, decision)

    rows = workbook_rows(ws, headers)
    rebuild_manual_summary(wb, rows)
    AUDIT_MD.write_text(build_audit(rows_before_blank, rows, added_rows, backup_path), encoding="utf-8")
    wb.save(OUTPUT_XLSX)

    # Re-open as a corruption smoke test.
    load_workbook(OUTPUT_XLSX, read_only=True).close()

    print(f"backup: {backup_path}")
    print(f"output: {OUTPUT_XLSX}")
    print(f"audit: {AUDIT_MD}")
    print(f"rows: {len(rows)}")
    print(f"blank_elena_decision_before: {rows_before_blank}")
    print(f"blank_elena_decision_after: {sum(1 for row in rows if not row['elena_decision'])}")
    print("decisions:")
    for decision, count in Counter(row["elena_decision"] for row in rows).most_common():
        print(f"  {decision}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
